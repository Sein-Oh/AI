"""테스트용 실제 형식 샘플 파일 (docx, pptx, xlsx, pdf, hwp, hwpx, eml, doc).

HWP 5.0과 옛 Word(.doc)는 OLE 복합 파일이라 만들 도구가 없어서, 최소한의 OLE 작성기(write_ole)를 둡니다.
스트림은 모두 4096바이트 이상으로 채워 미니 스트림 없이 일반 섹터만 씁니다 (olefile이 그대로 읽음).
"""
import struct
import zipfile

ENDOFCHAIN, FREESECT, FATSECT, NOSTREAM = 0xFFFFFFFE, 0xFFFFFFFF, 0xFFFFFFFD, 0xFFFFFFFF


def write_ole(path, streams):
    """streams: {"이름": bytes, "저장소/이름": bytes}. 한 단계 저장소까지 지원."""
    sec = 512
    data_secs = []            # 섹터 순서대로의 바이트
    fat = []
    entries = [{"name": "Root Entry", "type": 5, "children": []}]
    storages = {}

    def add_entry(name, typ, parent):
        e = {"name": name, "type": typ, "children": []}
        entries.append(e)
        entries[parent]["children"].append(len(entries) - 1)
        return len(entries) - 1

    for full, blob in streams.items():
        parent = 0
        if "/" in full:
            st_name, name = full.split("/", 1)
            if st_name not in storages:
                storages[st_name] = add_entry(st_name, 1, 0)
            parent = storages[st_name]
        else:
            name = full
        idx = add_entry(name, 2, parent)
        blob = blob + b"\x00" * max(0, 4096 - len(blob))
        blob += b"\x00" * (-len(blob) % sec)
        start = len(data_secs)
        n = len(blob) // sec
        for i in range(n):
            data_secs.append(blob[i * sec:(i + 1) * sec])
            fat.append(start + i + 1 if i < n - 1 else ENDOFCHAIN)
        entries[idx]["start"], entries[idx]["size"] = start, len(blob)

    # 디렉터리: 형제는 오른쪽 형제로만 잇습니다 (모두 검은색)
    for e in entries:
        kids = e["children"]
        e["child"] = kids[0] if kids else NOSTREAM
        for a, b in zip(kids, kids[1:] + [None]):
            entries[a]["right"] = b if b is not None else NOSTREAM
    dir_bytes = b""
    for e in entries:
        name = e["name"].encode("utf-16-le") + b"\x00\x00"
        rec = name.ljust(64, b"\x00") + struct.pack("<HBB", len(name), e["type"], 1)
        rec += struct.pack("<III", NOSTREAM, e.get("right", NOSTREAM), e["child"])
        rec += b"\x00" * 16 + b"\x00" * 4 + b"\x00" * 16
        rec += struct.pack("<IQ", e.get("start", ENDOFCHAIN), e.get("size", 0))
        dir_bytes += rec
    dir_bytes += b"\x00" * (-len(dir_bytes) % sec)
    dir_start = len(data_secs)
    nd = len(dir_bytes) // sec
    for i in range(nd):
        data_secs.append(dir_bytes[i * sec:(i + 1) * sec])
        fat.append(dir_start + i + 1 if i < nd - 1 else ENDOFCHAIN)
    fat_sector = len(data_secs)
    fat.append(FATSECT)
    assert len(fat) <= 128, "테스트용 작성기는 FAT 섹터 1개까지만"
    fat += [FREESECT] * (128 - len(fat))
    data_secs.append(struct.pack("<128I", *fat))

    header = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 16
    header += struct.pack("<HHHHH", 0x3E, 3, 0xFFFE, 9, 6) + b"\x00" * 6
    header += struct.pack("<IIIIIIII", 0, 1, dir_start, 0, 4096, ENDOFCHAIN, 0, ENDOFCHAIN)
    header += struct.pack("<I", 0)                        # DIFAT 섹터 수
    header += struct.pack("<109I", fat_sector, *([FREESECT] * 108))
    assert len(header) == 512
    with open(path, "wb") as f:
        f.write(header + b"".join(data_secs))


def make_hwp(path, body, preview=None, props=0):
    """HWP 5.0: FileHeader + (선택) PrvText + BodyText/Section0 (압축 안 함)."""
    header = b"HWP Document File".ljust(32, b"\x00") + struct.pack("<II", 0x05000300, props)
    streams = {"FileHeader": header}
    if preview is not None:
        streams["PrvText"] = preview.encode("utf-16-le")
    sec = b""
    for para in body.split("\n"):
        t = para.encode("utf-16-le")
        sec += struct.pack("<I", 67 | (len(t) << 20)) + t
    streams["BodyText/Section0"] = sec
    write_ole(path, streams)


def make_doc(path, text):
    """옛 Word(.doc) 흉내: WordDocument 스트림에 UTF-16 본문."""
    write_ole(path, {"WordDocument": b"\xec\xa5" + b"\x00" * 30 + text.encode("utf-16-le")})


def make_docx(path, *paras):
    import docx
    d = docx.Document()
    for p in paras:
        d.add_paragraph(p)
    d.save(path)


def make_pptx(path, *titles):
    import pptx
    prs = pptx.Presentation()
    for t in titles:
        s = prs.slides.add_slide(prs.slide_layouts[5])
        s.shapes.title.text = t
    prs.save(path)


def make_xlsx(path, sheet, *rows):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    for r in rows:
        ws.append(list(r))
    wb.save(path)


def make_pdf(path, *pages):
    import pymupdf
    doc = pymupdf.open()
    for text in pages:
        doc.new_page().insert_text((72, 72), text, fontname="korea")
    doc.save(path)
    doc.close()


def make_hwpx(path, preview=None, body="본문"):
    sec = ('<?xml version="1.0" encoding="UTF-8"?><hs:sec xmlns:hs="x" xmlns:hp="y">'
           f'<hp:p><hp:run><hp:t>{body}</hp:t></hp:run></hp:p></hs:sec>')
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("mimetype", "application/hwp+zip")
        z.writestr("Contents/section0.xml", sec)
        if preview is not None:
            z.writestr("Preview/PrvText.txt", preview.encode("utf-8"))


def make_eml(path, subject, body):
    raw = (f"Subject: {subject}\r\nFrom: a@example.com\r\nTo: b@example.com\r\nMIME-Version: 1.0\r\n"
           f"Content-Type: text/plain; charset=utf-8\r\nContent-Transfer-Encoding: 8bit\r\n\r\n{body}\r\n").encode("utf-8")
    with open(path, "wb") as f:
        f.write(raw)
