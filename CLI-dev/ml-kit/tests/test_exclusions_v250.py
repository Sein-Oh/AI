# -*- coding: utf-8 -*-
"""
2.5.0: rows left out of training ('제외' column + automatic outliers), all-rows score, report terms.
The data file is never modified; the excluded rows are listed in <version>/excluded_rows.xlsx.
"""
import json
import os
import sys

import numpy as np
import pandas as pd
import pytest

HERE = os.path.dirname(os.path.abspath(__file__))
KIT = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(HERE))), "CLI", "ml-kit")  # 실행 파일은 CLI\ml-kit
sys.path.insert(0, HERE)
sys.path.insert(0, KIT)
from test_cli_smoke import run_tool, run_fail  # noqa: E402
import ml_train as mt  # noqa: E402
import ml_predict as mc  # noqa: E402

N = 200


def _data(marks=None, spikes=True):
    r = np.random.default_rng(0)
    t = r.uniform(170, 195, N)
    p = r.uniform(2.5, 3.5, N)
    c = r.uniform(2, 3, N)
    y = 0.5 * t + 8 * p - 3 * c + r.normal(0, 1, N)
    if spikes:
        y[5] += 60          # target outlier -> automatic
        y[17] -= 50         # target outlier, but marked '포함'
        t[30] = 400         # feature outlier -> automatic
    df = pd.DataFrame({"반응온도": t, "압축압력": p, "촉매비율": c, "생산효율": y,
                       "등급": np.where(y > np.median(y), "양품", "불량")})
    if marks is not None:
        df["제외"] = marks
    return df


def _marks():
    m = [""] * N
    m[50] = "센서오류"
    m[17] = "포함"
    return m


def test_marks_auto_outliers_and_all_rows_score(tmp_path):
    f = tmp_path / "d.xlsx"
    _data(_marks()).to_excel(f, index=False)
    before = f.read_bytes()
    res = run_tool("ml_train.py", "-f", str(f), "-x", "반응온도,압축압력,촉매비율,제외", "-y", "생산효율")
    ex = res["exclusions"]["생산효율"]
    assert (ex["manual"], ex["auto"], ex["forced_include"]) == (1, 2, 1)
    assert isinstance(ex["all_rows_score"], float) and ex["all_rows_score"] < res["mean_val_score"]
    assert "제외" not in res["feature_types"]["numeric"]                 # never a feature
    assert "전체 R2(제외 행 포함)" in res["preview"] and "검증 R2" in res["preview"]
    assert f.read_bytes() == before                                       # data file untouched

    rows = pd.read_excel(res["excluded_rows_file"])
    by_row = {int(r["원본 행"]): r for _, r in rows.iterrows()}
    assert set(by_row) == {7, 32, 52}                                     # Excel rows (header is row 1)
    assert by_row[52]["구분"] == "직접 표시" and by_row[52]["이유"] == "센서오류"
    assert by_row[7]["구분"] == "자동 이상치" and "생산효율" in by_row[7]["이유"]
    assert "반응온도 400" in by_row[32]["이유"]

    meta = json.load(open(os.path.join(res["version_dir"], "model_생산효율.json"), encoding="utf-8"))
    assert meta["feature_stats"]["반응온도"]["max"] < 200                 # training range ignores the outlier
    assert meta["excluded"]["auto"] == 2
    assert meta["target_summary"]["max"] < 150
    hist = [json.loads(ln) for ln in open(os.path.join(res["data_root"], "history.jsonl"), encoding="utf-8")]
    assert hist[-1]["exclusions"]["생산효율"]["manual"] == 1
    assert hist[-1]["data_profile"]["rows_used"]["생산효율"] == N - 3


def test_no_exclude_column_still_finds_outliers_and_clean_data_has_no_key(tmp_path):
    f = tmp_path / "d.csv"
    _data().to_csv(f, index=False)
    res = run_tool("ml_train.py", "-f", str(f), "-x", "반응온도,압축압력,촉매비율", "-y", "생산효율")
    assert res["exclusions"]["생산효율"]["auto"] == 3 and res["exclusions"]["생산효율"]["column"] is None

    g = tmp_path / "clean.csv"
    _data(spikes=False).to_csv(g, index=False)
    res = run_tool("ml_train.py", "-f", str(g), "-x", "반응온도,압축압력,촉매비율", "-y", "생산효율")
    assert "exclusions" not in res and "excluded_rows_file" not in res


def test_skewed_distribution_is_not_trimmed():
    """More candidates than 10% means a wide distribution, not outliers: nothing is excluded."""
    r = np.random.default_rng(1)
    y = pd.Series(np.concatenate([r.normal(0, 1, 160), r.normal(30, 1, 40)]))
    X = pd.DataFrame({"a": r.normal(0, 1, 200)})
    plan = mt.plan_exclusions(None, y, X, {"numeric": ["a"]}, [], mt.TASK_REGRESSION, "y")
    assert not plan["mask"].any() and "자동 제외하지 않았습니다" in plan["note"]


def test_classification_keeps_every_class(tmp_path):
    """An automatic outlier never removes a class from training."""
    y = pd.Series(["A"] * 98 + ["B", "B"])
    x = np.r_[np.random.default_rng(2).normal(0, 1, 98), 50.0, 60.0]   # both 'B' rows look like outliers
    plan = mt.plan_exclusions(None, y, pd.DataFrame({"a": x}), {"numeric": ["a"]}, [], mt.TASK_CLASSIFICATION, "y")
    assert not plan["mask"][98:].any()

    f = tmp_path / "c.xlsx"
    _data(_marks()).to_excel(f, index=False)
    res = run_tool("ml_train.py", "-f", str(f), "-x", "반응온도,압축압력,촉매비율", "-y", "등급",
                   "--task", "classification", "--positive-class", "불량")
    ex = res["exclusions"]["등급"]
    assert ex["manual"] == 1 and ex["auto"] == 1                          # target is a label: features only
    assert 0 <= ex["all_rows_score"] <= 1 and "전체 Macro-F1" in res["preview"]


def test_exclude_column_as_target_is_input_error(tmp_path):
    f = tmp_path / "d.csv"
    _data(_marks()).to_csv(f, index=False)
    payload = run_fail("ml_train.py", "-f", str(f), "-x", "반응온도", "-y", "제외")
    assert payload["agent_directive"].startswith("STOP_RETRY") and "제외" in payload["message"]


def test_loader_keeps_source_row_numbers(tmp_path):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "공정 기록"
    ws.append([])
    ws.append(["a", "b"])           # header on row 3
    ws.append([1, 2])               # row 4
    ws.append([None, None])         # blank row 5 is dropped
    ws.append([3, 4])               # row 6
    path = tmp_path / "s.xlsx"
    wb.save(path)
    df = mc.load_data_file(str(path))
    assert list(df["a"]) == [1, 3] and df.attrs["source_rows"] == [4, 6]


def test_report_shows_terms_and_excluded_rows(tmp_path):
    from docx import Document
    f = tmp_path / "d.xlsx"
    _data(_marks()).to_excel(f, index=False)
    res = run_tool("ml_train.py", "-f", str(f), "-x", "반응온도,압축압력,촉매비율", "-y", "생산효율")
    rep = run_tool("ml_report.py", "-m", res["version_dir"])
    doc = Document(rep["report_file"])
    text = "\n".join(p.text for p in doc.paragraphs)
    cells = " ".join(c.text for tb in doc.tables for row in tb.rows for c in row.cells)
    assert "검증 R²" in cells and "전체 R²" in cells
    assert "학습 제외 3행 (직접 표시 1, 자동 이상치 2)" in text and "excluded_rows.xlsx" in text
    assert "학습 R² = 모델이 배운 행" in text
    assert "제외 행 포함 전체" not in text + cells and "교차검증 R²는" not in text
