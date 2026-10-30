# -*- coding: utf-8 -*-
"""
2.9.6 (전체 검토): calculator range from the rows actually used for training, constant inputs, 8+ inputs
(Excel's 255-argument limit), derived-feature sources matched by whole name, '제외' 0/FALSE marks,
numbers in nearest rows, input errors from --find.
"""
import os
import sys

import numpy as np
import openpyxl
import pandas as pd
import pytest

HERE = os.path.dirname(os.path.abspath(__file__))
KIT = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(HERE))), "CLI", "ml-kit")
sys.path.insert(0, HERE)
sys.path.insert(0, KIT)
from test_cli_smoke import run_tool, run_fail  # noqa: E402
from test_export_v290 import SOFFICE, _grid, _recalc  # noqa: E402
import ml_export as mx  # noqa: E402


def _strings(wb, sheet):
    return " ".join(str(c.value) for row in wb[sheet].iter_rows() for c in row if c.value is not None)


def test_calculator_range_ignores_excluded_rows_and_constant_input(tmp_path):
    r = np.random.default_rng(5)
    n = 300
    df = pd.DataFrame({"반응온도": r.uniform(150, 200, n), "압력": r.uniform(1, 3, n), "pH": 7.0})
    df["효율"] = 0.4 * df.반응온도 + 6 * df.압력 + r.normal(0, 0.3, n)
    df.loc[10, "반응온도"] = 2000                                # typo, excluded as an automatic outlier
    p = tmp_path / "c.csv"
    df.to_csv(p, index=False)
    v = run_tool("ml_train.py", "-f", str(p), "-x", "반응온도,압력,pH", "-y", "효율")["version_dir"]
    res = run_tool("ml_export.py", "-m", v, "--save", str(tmp_path / "c.xlsx"))
    t = res["targets"]["효율"]
    assert "pH" not in t["inputs"] and any("pH" in w for w in res["warnings"])
    wb = openpyxl.load_workbook(res["output_file"])
    text = _strings(wb, "계산기") + _strings(wb, "일괄계산") + _strings(wb, "정보")
    assert "2000" not in text and "1999" not in text                       # range / out-of-range checks use 150~200


def test_inputs_and_keep_drop_constant_with_hint(tmp_path):
    r = np.random.default_rng(6)
    df = pd.DataFrame({"a": r.uniform(0, 1, 200), "k": 1.0})
    df["y"] = 3 * df.a + r.normal(0, 0.05, 200)
    p = tmp_path / "k.csv"
    df.to_csv(p, index=False)
    v = run_tool("ml_train.py", "-f", str(p), "-x", "a,k", "-y", "y")["version_dir"]
    bad = run_fail("ml_export.py", "-m", v, "--inputs", "k")
    assert bad["agent_directive"].startswith("STOP_RETRY") and "[추천] a" in bad["message"]
    ok = run_tool("ml_export.py", "-m", v, "--keep", "k", "--save", str(tmp_path / "k.xlsx"))
    assert ok["targets"]["y"]["inputs"] == ["a"]


@pytest.mark.skipif(SOFFICE is None, reason="LibreOffice not installed")
def test_eight_inputs_use_range_sums_and_recalculate(tmp_path):
    r = np.random.default_rng(8)
    n = 400
    cols = [f"x{i}" for i in range(8)]
    df = pd.DataFrame({c: r.uniform(0, 1, n) for c in cols})
    df["y"] = sum((i + 1) * df[c] for i, c in enumerate(cols)) + r.normal(0, 0.05, n)
    p = tmp_path / "e.csv"
    df.to_csv(p, index=False)
    v = run_tool("ml_train.py", "-f", str(p), "-x", ",".join(cols), "-y", "y")["version_dir"]
    res = run_tool("ml_export.py", "-m", v, "--inputs", ",".join(cols), "--save", str(tmp_path / "e.xlsx"))
    wb = openpyxl.load_workbook(res["output_file"])
    sums = [c.value for row in wb["계산기"].iter_rows() for c in row if isinstance(c.value, str) and "SUM(" in c.value]
    assert sums and all(s.count(",") < 255 for s in sums)
    out = _recalc(res["output_file"], tmp_path)
    x = [out["계산기"][f"C{6 + i}"].value for i in range(8)]
    want = float(_grid(openpyxl.load_workbook(res["output_file"]), "표_y", 8)([x])[0])
    assert out["계산기"][f"C{5 + 8 + 3}"].value == pytest.approx(want, abs=1e-9)


def test_derived_sources_are_whole_names_and_chains_drop_together():
    raw = ["온도", "반응온도", "압력"]
    evals = ["t2 = 반응온도 * 반응온도", "tp = t2 * 압력"]
    assert mx.evals_for(["반응온도", "압력"], evals, raw) == evals            # '온도' is not inside '반응온도'
    assert mx.evals_for(["온도", "압력"], evals, raw) == []                    # t2 dropped -> tp dropped


def test_exclude_column_zero_and_false_are_not_marks():
    import ml_train as mt
    assert [mt._exclude_mark(v) for v in (0, 0.0, "0", "FALSE", False, None, float("nan"), "", " 센서오류 ", 1, "포함")] == \
        ["", "", "", "", "", "", "", "", "센서오류", "1", "포함"]


def test_nearest_rows_are_numbers_and_find_rejects_text_input(tmp_path):
    r = np.random.default_rng(3)
    df = pd.DataFrame({"온도": r.uniform(150, 200, 200), "압력": r.uniform(1, 3, 200)})
    df["효율"] = 0.4 * df.온도 + 5 * df.압력 + r.normal(0, 0.2, 200)
    p = tmp_path / "n.csv"
    df.to_csv(p, index=False)
    m = run_tool("ml_map.py", "-f", str(p), "--find", "효율=80", "--vary", "온도", "-i", "압력=2")
    assert m["nearest_rows"] and all(isinstance(nr["온도"], float) for nr in m["nearest_rows"])
    v = run_tool("ml_train.py", "-f", str(p), "-x", "온도,압력", "-y", "효율")["version_dir"]
    f = run_tool("ml_predict.py", "-m", v, "--find", "효율=80", "--vary", "온도", "-i", "압력=2")
    assert all(isinstance(nr["온도"], float) for nr in f["nearest_rows"])
    bad = run_fail("ml_predict.py", "-m", v, "--find", "효율=80", "--vary", "온도", "-i", "압력=x")
    assert bad["agent_directive"].startswith("STOP_RETRY") and "압력" in bad["message"] and "숫자" in bad["message"]


def test_model_from_before_291_asks_for_retraining(tmp_path):
    """2.9.7 (beta, no compatibility paths): an older model is refused with a clear message, not half-supported."""
    import json
    import shutil
    r = np.random.default_rng(2)
    df = pd.DataFrame({"a": r.uniform(0, 1, 120), "b": r.uniform(0, 1, 120)})
    df["y"] = 2 * df.a + df.b + r.normal(0, 0.05, 120)
    p = tmp_path / "o.csv"
    df.to_csv(p, index=False)
    v = run_tool("ml_train.py", "-f", str(p), "-x", "a,b", "-y", "y")["version_dir"]
    old = tmp_path / "v_old"
    shutil.copytree(v, old)
    meta_p = old / "model_y.json"
    meta = json.loads(meta_p.read_text(encoding="utf-8"))
    del meta["input_columns_required"]
    meta_p.write_text(json.dumps(meta, ensure_ascii=False), encoding="utf-8")
    bad = run_fail("ml_predict.py", "-m", str(old), "-i", "a=0.5,b=0.5")
    assert bad["agent_directive"].startswith("STOP_RETRY") and "다시 학습" in bad["message"]
    res = run_tool("ml_train.py", "-f", str(p), "--history")
    assert "mean_val_r2" not in res["records"][0] and "best_mean_val_r2" not in res
