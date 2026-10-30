# -*- coding: utf-8 -*-
"""
2.9.2 (실기 2026-10-30): the report and the performance calculator are a set. ml_report.py also makes
성능예측표_<data>_<version>.xlsx in the same <data>_ml_data folder as the report and history.jsonl
(regression only). A calculator failure never fails the report.
"""
import os
import sys

import numpy as np
import pandas as pd
import pytest

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from test_cli_smoke import run_tool  # noqa: E402


@pytest.fixture
def with_export(monkeypatch):
    monkeypatch.delenv("ML_KIT_SKIP_EXPORT", raising=False)


def _data(tmp_path):
    r = np.random.default_rng(12)
    n = 250
    df = pd.DataFrame({"온도": r.uniform(0, 1, n), "압력": r.uniform(0, 1, n), "잡음": r.uniform(0, 1, n)})
    df["효율"] = 8 * df.온도 + 3 * df.압력 + r.normal(0, 0.05, n)
    df["등급"] = np.where(df.효율 > df.효율.median(), "상", "하")
    path = tmp_path / "p.csv"
    df.to_csv(path, index=False)
    return path


def test_report_makes_calculator_in_the_same_folder(tmp_path, with_export):
    path = _data(tmp_path)
    tr = run_tool("ml_train.py", "-f", str(path), "-x", "온도,압력,잡음", "-y", "효율")
    rep = run_tool("ml_report.py", "-m", tr["version_dir"], "--title", "시험 효율 예측")
    root = os.path.dirname(tr["version_dir"])
    assert os.path.dirname(rep["report_file"]) == root and os.path.isfile(os.path.join(root, "history.jsonl"))
    assert rep["export_file"] == os.path.join(root, "성능예측표_p_v1.xlsx") and os.path.isfile(rep["export_file"])
    assert rep["export"]["효율"]["inputs"] == ["온도", "압력"] and "성능예측표" in rep["preview"]
    assert "export_file" in rep["agent_directive"]
    a = rep["export"]["효율"]["agreement"]                                       # 2.9.4: agreement with the model
    assert 95 < a["agreement_pct"] <= 100 and a["p95_diff_pct"] >= 0
    import openpyxl
    wb = openpyxl.load_workbook(rep["export_file"])
    assert wb["계산기"]["B2"].value == "시험 효율 예측"
    notes = " ".join(str(c.value) for row in wb["계산기"].iter_rows() for c in row if isinstance(c.value, str))
    assert f"평균 {a['agreement_pct']:.1f}% 일치" in notes
    from docx import Document
    text = "\n".join(p.text for p in Document(rep["report_file"]).paragraphs)
    assert "성능예측표 (엑셀 계산기): 성능예측표_p_v1.xlsx" in text and f"평균 {a['agreement_pct']:.1f}% 일치" in text


def test_classification_report_has_no_calculator(tmp_path, with_export):
    path = _data(tmp_path)
    tr = run_tool("ml_train.py", "-f", str(path), "-x", "온도,압력", "-y", "등급", "--task", "classification")
    rep = run_tool("ml_report.py", "-m", tr["version_dir"])
    assert "export_file" not in rep and "export_note" not in rep


def test_calculator_failure_keeps_the_report(tmp_path, with_export):
    path = _data(tmp_path)
    tr = run_tool("ml_train.py", "-f", str(path), "-x", "온도,압력", "-y", "효율")
    os.remove(os.path.join(tr["version_dir"], "model_dataset.csv"))          # calculator needs the training snapshot
    rep = run_tool("ml_report.py", "-m", tr["version_dir"])
    assert os.path.isfile(rep["report_file"]) and "export_file" not in rep
    assert rep["export_note"].startswith("성능예측표를 만들지 못했습니다") and "export_note" in rep["agent_directive"]
