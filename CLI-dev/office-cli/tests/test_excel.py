"""excel.py 테스트: 가짜 Excel 격자로 명령 로직과 응답 계약을 검증하고, 닫힌 파일은 openpyxl로 실제 파일을 읽습니다."""
import json
import os

import pytest

import _office as off
import excel
from fake_excel import FakeExcelWorkbook, FakeXlApp
from fake_office import make_win32com


@pytest.fixture
def xl(monkeypatch):
    """활성 통합문서가 가짜 Excel인 상태. (app, wb) 반환."""
    app = FakeXlApp()
    wb = FakeExcelWorkbook(app, "보고.xlsx", sheets=("Sheet1", "요약"))
    fake, _ = make_win32com()
    monkeypatch.setattr(off, "win32com", fake)
    monkeypatch.setattr(off, "pythoncom", None)
    monkeypatch.setattr(off, "list_open", lambda kind: [])
    monkeypatch.setattr(off, "get_active", lambda kind: {"name": wb.Name, "path": wb.FullName, "app": app, "doc": wb})
    return app, wb


def run(capsys, *argv):
    with pytest.raises(SystemExit) as ei:
        excel.main(list(argv))
    out = [ln for ln in capsys.readouterr().out.splitlines() if ln.strip()]
    assert len(out) == 1, out
    return ei.value.code, json.loads(out[0])


def ws_of(wb, name="Sheet1"):
    return wb.Worksheets(name)


# ---------------------------------------------------------------------------
# 도우미
# ---------------------------------------------------------------------------

def test_coerce_and_address_helpers():
    assert excel.coerce("1,200") == 1200 and excel.coerce("3.5") == 3.5
    assert excel.coerce("001") == "001" and excel.coerce("0.5") == 0.5 and excel.coerce("") is None
    assert excel.split_sheet("'My Sheet'!A1:B2") == ("My Sheet", "A1:B2")
    assert excel.addr(2, 1, 101, 2) == "A2:B101" and excel.col_letter(28) == "AB"


def test_load_table_file_markdown_and_csv(tmp_path):
    md = tmp_path / "t.md"
    md.write_text("# 제목\n\n| 공정 | 수율 |\n|---|---|\n| 가열 | 98.2 |\n", encoding="utf-8")
    assert excel.load_table_file(str(md)) == [["공정", "수율"], ["가열", 98.2]]
    csv = tmp_path / "t.csv"
    csv.write_bytes("코드,값\n001,7\n".encode("cp949"))
    assert excel.load_table_file(str(csv)) == [["코드", "값"], ["001", 7]]


# ---------------------------------------------------------------------------
# write
# ---------------------------------------------------------------------------

def test_write_sine_header_and_column_formulas(xl, capsys):
    app, wb = xl
    code, out = run(capsys, "write", "--range", "A1:B101", "--header", "x,sin(x)",
                    "--formula", "=(ROW()-2)*4*PI()/99", "=SIN(A2)")
    assert code == 0 and out["range_written"] == "A1:B101" and out["rows"] == 101
    ws = ws_of(wb)
    assert ws.values[(1, 1)] == "x" and ws.values[(1, 2)] == "sin(x)"
    assert ("Formula", "$A$2:$A$101", "=(ROW()-2)*4*PI()/99") in ws.log
    assert ("Formula", "$B$2:$B$101", "=SIN(A2)") in ws.log
    assert app.ScreenUpdating is True and app.Calculation == -4105       # 대량 쓰기 설정 복원
    assert "chart" in out["next_commands"][0]["command"]
    assert out["target_resolved"]["sheet"] == "Sheet1"


def test_write_formula_count_mismatch_is_input_error(xl, capsys):
    code, out = run(capsys, "write", "--range", "A2:C10", "--formula", "=1", "=2")
    assert code == 1 and out["error_type"] == "input"


def test_write_checks_everything_before_writing_the_header(xl, capsys):
    # 회사 PC 오류 로그: 헤더를 먼저 쓰고 수식 검사에서 실패해 A1:B1만 남았음. 추천 범위는 머리글 행부터
    app, wb = xl
    code, out = run(capsys, "write", "--range", "A1", "--header", "x,sin(x)", "--formula", "=(ROW()-2)*0.1", "=SIN(A2)")
    assert code == 1 and out["message"].endswith("[추천] --range A1:B101")
    assert (1, 1) not in ws_of(wb).values and (1, 2) not in ws_of(wb).values


def test_write_numbers_given_as_formulas_points_to_values(xl, capsys):
    code, out = run(capsys, "write", "--range", "C6:C8", "--formula", "73.42", "49.97", "14.7")
    assert code == 1 and "[추천] --values '[[73.42], [49.97], [14.7]]'" in out["message"]


def test_write_over_existing_values_needs_overwrite(xl, capsys):
    app, wb = xl
    ws = ws_of(wb)
    ws.set(2, 2, "keep")
    code, out = run(capsys, "write", "--range", "A1", "--values", '[["a","b"],["c","new"]]')
    assert code == 1 and out["error_type"] == "needs_user_decision" and "B2=keep" in out["message"]
    assert ws.values[(2, 2)] == "keep" and (1, 1) not in ws.values
    code, out = run(capsys, "write", "--range", "A1", "--values", '[["a","b"],["c","new"]]', "--overwrite")
    assert code == 0 and ws.values[(2, 2)] == "new"


def test_write_reads_back_and_reports_mismatch(xl, capsys, monkeypatch):
    # R1: VERIFIED는 쓴 칸을 다시 읽어 같을 때만. Excel이 값을 바꿨으면 UNVERIFIED와 그 칸
    app, wb = xl
    code, out = run(capsys, "write", "--range", "A1", "--values", '[["코드","값"],["001",7]]')
    assert code == 0 and out["agent_directive"].startswith("VERIFIED:") and out["verified"] is True
    ws = ws_of(wb)
    real_set = ws.set
    monkeypatch.setattr(ws, "set", lambda r, c, v: real_set(r, c, 999 if (r, c) == (2, 5) else v))
    code, out = run(capsys, "write", "--range", "D1", "--values", '[["x","y"],[1,2]]')
    assert code == 0 and out["agent_directive"].startswith("UNVERIFIED:") and "E2" in out["unverified"][0]


def test_write_csv_new_sheet_and_append_skips_header(xl, capsys, tmp_path):
    app, wb = xl
    f = tmp_path / "d.csv"
    f.write_text("코드,금액\n001,1200\n002,800\n", encoding="utf-8")
    code, out = run(capsys, "write", "--file", str(f), "--sheet", "검색결과", "--new-sheet")
    assert code == 0 and out["sheet"] == "검색결과" and out["range_written"] == "A1:B3"
    ws = ws_of(wb, "검색결과")
    assert ws.values[(2, 1)] == "001" and ws.values[(2, 2)] == 1200
    f2 = tmp_path / "e.csv"
    f2.write_text("코드,금액\n003,500\n", encoding="utf-8")
    code, out = run(capsys, "write", "--file", str(f2), "--sheet", "검색결과", "--append")
    assert out["range_written"] == "A4:B4" and ws.values[(4, 1)] == "003"


def _sheet_names(wb):
    return [wb.Worksheets(i).Name for i in range(1, wb.Worksheets.Count + 1)]


def _no_create_imperative(out):
    # 시나리오 10: message나 [후보]에 --new-sheet / '만드세요'가 있으면 경량 LLM이 묻지 않고 새 시트를 만든다
    texts = [out["message"]] + list(out.get("candidates", []))
    assert not any("--new-sheet" in x or "만드세요" in x or "붙이세요" in x for x in texts), texts


def test_write_missing_sheet_is_not_created_without_flag(xl, capsys):
    app, wb = xl
    code, out = run(capsys, "write", "--sheet", "결과표", "--value", "1")
    assert code == 1 and out["error_type"] == "needs_user_decision"
    assert out["agent_directive"].startswith("STOP_RETRY: 다시 실행하지 마세요.")
    assert "만들지 물어보세요" in out["agent_directive"]
    assert "만들라고 답한 경우에만 같은 명령에 --new-sheet" in out["agent_directive"]   # 플래그는 사용자 동의 조건으로만
    _no_create_imperative(out)
    assert "[추천]" not in out["message"] and "[후보]" not in out["message"] and "candidates" not in out
    assert "있는 시트: Sheet1, 요약" in out["message"] and "사용자에게 물어보세요" in out["message"]
    assert _sheet_names(wb) == ["Sheet1", "요약"]                       # 새 시트 없음
    code, out = run(capsys, "write", "--range", "'결과표'!A1", "--value", "1")
    assert code == 1 and _sheet_names(wb) == ["Sheet1", "요약"]


def test_write_sheet_differs_only_in_spacing_recommends_existing(xl, capsys):
    app, wb = xl
    code, out = run(capsys, "write", "--sheet", "Sheet 1", "--value", "1")
    assert code == 1 and out["error_type"] == "input" and out["message"].endswith("[추천] Sheet1")
    _no_create_imperative(out)
    assert _sheet_names(wb) == ["Sheet1", "요약"]


def test_write_similar_but_different_sheet_asks_user(xl, capsys):
    # Sheet2를 Sheet1의 오타로 단정해 덮어쓰지 않음: [후보]는 실제 기존 시트만, 새 시트 여부는 사용자에게 물음
    app, wb = xl
    code, out = run(capsys, "write", "--sheet", "Sheet2", "--value", "1")
    assert code == 1 and out["error_type"] == "needs_user_decision" and out["message"].endswith("[후보] Sheet1")
    assert out["candidates"] == ["Sheet1"]
    assert out["agent_directive"].startswith("STOP_RETRY: 다시 실행하지 마세요.")
    assert "만들라고 답한 경우에만" in out["agent_directive"]
    _no_create_imperative(out)
    assert _sheet_names(wb) == ["Sheet1", "요약"] and (1, 1) not in ws_of(wb, "Sheet1").values


def test_scenario10_blank_workbook_sheet9_asks_user(capsys, monkeypatch):
    # 시나리오 10: Sheet1만 있는 빈 통합문서에 "Sheet9에 써줘" -> 만들지 않고 사용자에게 물음
    app = FakeXlApp()
    wb = FakeExcelWorkbook(app, "Book1", sheets=("Sheet1",))
    wb.FullName = "Book1"                                                # 저장한 적 없는 새 통합문서
    fake, _ = make_win32com()
    monkeypatch.setattr(off, "win32com", fake)
    monkeypatch.setattr(off, "pythoncom", None)
    monkeypatch.setattr(off, "list_open", lambda kind: [])
    monkeypatch.setattr(off, "get_active", lambda kind: {"name": wb.Name, "path": "", "app": app, "doc": wb})
    code, out = run(capsys, "write", "--sheet", "Sheet9", "--value", "테스트")
    assert code == 1 and out["error_type"] == "needs_user_decision"
    assert out.get("candidates", []) == ["Sheet1"] and "[추천]" not in out["message"]
    _no_create_imperative(out)
    assert "Sheet9" in out["agent_directive"] and "물어보세요" in out["agent_directive"]
    assert _sheet_names(wb) == ["Sheet1"]


def test_write_new_sheet_flag_rules(xl, capsys):
    app, wb = xl
    code, out = run(capsys, "write", "--sheet", "요약", "--new-sheet", "--value", "1")
    assert code == 1 and "이미 있어" in out["message"] and (1, 1) not in ws_of(wb, "요약").values
    code, out = run(capsys, "write", "--new-sheet", "--value", "1")
    assert code == 1 and "새 시트 이름" in out["message"]
    code, out = run(capsys, "write", "--range", "'새시트'!B2", "--new-sheet", "--value", "7")
    assert code == 0 and out["sheet"] == "새시트" and ws_of(wb, "새시트").values[(2, 2)] == 7


def test_write_requires_content(xl, capsys):
    code, out = run(capsys, "write", "--range", "A1")
    assert code == 1 and "--value" in out["message"]


# ---------------------------------------------------------------------------
# read
# ---------------------------------------------------------------------------

def test_read_values_search_and_truncation(xl, capsys):
    app, wb = xl
    ws = ws_of(wb)
    for r in range(1, 31):
        ws.set(r, 1, f"행{r}")
        ws.set(r, 2, float(r))
    code, out = run(capsys, "read", "--range", "A1:B30", "--max-rows", "10")
    assert code == 0 and out["truncated"] and len(out["values"]) == 10
    assert out["values"][1] == ["행2", 2]                    # 2.0 -> 2
    assert "A11:B30" in out["next_commands"][0]["command"]
    code, out = run(capsys, "read", "--search", "행2")
    assert out["total_matches"] == 11 and out["matches"][0]["cell"] == "A2"


def test_read_unknown_sheet_suggests(xl, capsys):
    code, out = run(capsys, "read", "--sheet", "요약표")
    assert code == 1 and "[후보] 요약" in out["message"] and "[추천]" not in out["message"]   # 부분 일치는 추측하지 않음


def test_read_closed_xlsx_with_openpyxl(tmp_path, capsys, monkeypatch):
    import openpyxl
    p = tmp_path / "매출.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "매출"
    wb.active.append(["지역", "금액"])
    wb.active.append(["서울", 1200])
    wb.save(p)
    fake, calls = make_win32com()
    monkeypatch.setattr(off, "win32com", fake)
    monkeypatch.setattr(off, "list_open", lambda kind: [])
    code, out = run(capsys, "read", "--target", str(p))
    assert code == 0 and out["values"] == [["지역", "금액"], ["서울", 1200]]
    assert out["target_resolved"]["source"] == "file(openpyxl)"
    assert not any(c[0] == "DispatchEx" for c in calls)          # Excel을 띄우지 않음
    code, out = run(capsys, "read", "--target", str(p), "--range", "B:B")       # 열 전체 (열린 문서와 같게)
    assert code == 0 and out["values"] == [["금액"], [1200]]
    code, out = run(capsys, "read", "--target", str(p), "--range", "2:2")
    assert code == 0 and out["values"] == [["서울", 1200]]


# ---------------------------------------------------------------------------
# replace / style / range / sheet
# ---------------------------------------------------------------------------

def test_replace_keeps_formula_cells(xl, capsys):
    app, wb = xl
    ws = ws_of(wb)
    ws.set(1, 1, "초안 보고")
    ws.set(1, 2, 3.0)
    ws.formulas[(1, 2)] = "=1+2"
    ws.set(1, 3, 3.0)
    code, out = run(capsys, "replace", "--find", "3", "--replace", "9", "--whole", "--bg", "yellow")
    assert code == 0 and out["changed_cells"] == 2
    assert ws.formulas[(1, 2)] == "=1+2" and ws.values[(1, 3)] == 9
    assert out["formula_cells_kept"] == ["Sheet1!B1"]
    code, out = run(capsys, "replace", "--map", '{"초안": "최종"}')
    assert ws.values[(1, 1)] == "최종 보고"


def test_replace_no_match_is_ok_with_stop_hint(xl, capsys):
    code, out = run(capsys, "replace", "--find", "없는말", "--replace", "x")
    assert code == 0 and out["changed_cells"] == 0 and "재시도하지" in out["agent_directive"]


def test_style_theme_zebra_and_highlight(xl, capsys):
    app, wb = xl
    ws = ws_of(wb)
    rows = [["지역", "금액"]] + [[f"지역{i}", float(i * 100)] for i in range(1, 8)] + [["합계", 2800.0]]
    for r, row in enumerate(rows, 1):
        for c, v in enumerate(row, 1):
            ws.set(r, c, v)
    code, out = run(capsys, "style", "--theme", "navy", "--highlight", ">=500:pink")
    assert code == 0
    applied = out["applied"]
    assert applied["summary_row"] is True and applied["number_columns"] == ["B"]
    assert applied["zebra_rows"] == 3
    assert any(k[0].startswith("$A$3:$B$3,$A$5:$B$5") for k in ws.styles)
    fc = ws.fcs["$A$1:$B$9"].items[0]
    assert fc.kw["Type"] == 1 and fc.kw["Operator"] == 7


def test_style_unknown_theme_and_color(xl, capsys):
    code, out = run(capsys, "style", "--theme", "nvy")
    assert code == 1 and "[추천] navy" in out["message"]
    code, out = run(capsys, "style", "--bg", "노랑")
    assert code == 0


def test_range_move_uses_cut_and_merge_rules(xl, capsys):
    app, wb = xl
    ws = ws_of(wb)
    ws.set(1, 1, "가")
    code, out = run(capsys, "range", "--action", "move", "--range", "A1:B3", "--dest", "B2")
    assert code == 0 and ("Cut", "$A$1:$B$3", "$B$2") in ws.log
    code, out = run(capsys, "range", "--action", "merge")
    assert code == 1
    run(capsys, "range", "--action", "merge", "--range", "A1:B1")
    assert ("Merge", "$A$1:$B$1", False) in ws.log and app.DisplayAlerts is True


def test_sheet_actions(xl, capsys):
    app, wb = xl
    code, out = run(capsys, "sheet", "--action", "add", "--name", "sheet1")
    assert code == 0 and "이미 있어" in out["preview"] and out["sheets"] == ["Sheet1", "요약"]
    code, out = run(capsys, "sheet", "--action", "rename", "--name", "요약", "--new-name", "Sheet1")
    assert code == 1
    code, out = run(capsys, "sheet", "--action", "delete", "--name", "요약")
    assert code == 0 and out["sheets"] == ["Sheet1"]
    code, out = run(capsys, "sheet", "--action", "delete", "--name", "Sheet1")
    assert code == 1 and "마지막" in out["message"]


# ---------------------------------------------------------------------------
# chart
# ---------------------------------------------------------------------------

def test_chart_numeric_x_becomes_scatter_and_offers_transfer(xl, capsys):
    app, wb = xl
    ws = ws_of(wb)
    ws.set(1, 1, "x")
    ws.set(1, 2, "sin(x)")
    for r in range(2, 12):
        ws.set(r, 1, float(r))
        ws.set(r, 2, 0.5)
    code, out = run(capsys, "chart", "--range", "A1:B11", "--type", "line", "--title", "사인 2주기")
    assert code == 0 and out["chart_type"] == "scatter_lines" and out["chart_name"] == "사인 2주기"
    cmd = out["next_commands"][0]["command"]
    assert "transfer" in cmd and '"사인 2주기"' in cmd
    code, out = run(capsys, "chart", "--range", "A1:B11", "--title", "사인 2주기")
    assert out["chart_name"] == "사인 2주기 2"                    # 이름 중복 방지
    code, out = run(capsys, "chart", "--action", "list")
    assert [c["name"] for c in out["charts"]] == ["사인 2주기", "사인 2주기 2"]


@pytest.mark.parametrize("kind,expected", [("scatter", -4169), ("scatter_lines", 75), ("line", 75), ("column", 51)])
def test_chart_type_survives_set_source_data(xl, capsys, kind, expected):
    # 실기 I5: 종류를 SetSourceData 앞에서만 정하면 Excel 2007이 분산형을 74(선+표식)로 바꾸고,
    # 뒤에서만 정하면 첫 열(x)이 계열로 그려짐. 앞뒤로 모두 정해야 함. 'scatter' 상수도 74로 잘못돼 있었음
    app, wb = xl
    ws = ws_of(wb)
    ws.set(1, 1, "x")
    ws.set(1, 2, "y")
    for r in range(2, 6):
        ws.set(r, 1, float(r))
        ws.set(r, 2, float(r * r))
    code, out = run(capsys, "chart", "--range", "A1:B5", "--type", kind, "--title", "c")
    assert code == 0, out
    chart = ws.charts[-1]
    assert chart.type_at_source == expected and chart.Chart.ChartType == expected


def test_info_detail_reports_chart_type_series_and_covered_data(xl, capsys):
    app, wb = xl
    ws = ws_of(wb)
    ws.set(1, 1, "x")
    ws.set(1, 2, "y")
    for r in range(2, 12):
        ws.set(r, 1, float(r))
        ws.set(r, 2, float(r))
    run(capsys, "chart", "--range", "A1:B11", "--type", "scatter", "--title", "옆 차트")            # 데이터 오른쪽
    run(capsys, "chart", "--range", "A1:B11", "--type", "column", "--title", "덮는 차트", "--at", "A3")  # 데이터 위
    code, out = run(capsys, "info", "--detail")
    assert code == 0
    sheet = [s for s in out["sheets"] if s["charts"]][0]
    charts = {c["name"]: c for c in sheet["charts"]}
    assert charts["옆 차트"]["chart_type"] == "scatter" and charts["옆 차트"]["series"] == 1
    assert not charts["옆 차트"].get("covers_data") and charts["덮는 차트"]["covers_data"] is True
    assert out["layout_warnings"] == 1 and "덮는 차트" in sheet["warnings"][0] and "배치 경고 1건" in out["preview"]
    code, plain = run(capsys, "info")
    assert "layout_warnings" not in plain and "chart_type" not in plain["sheets"][0].get("charts", [{}])[0]


def test_chart_default_position_does_not_cover_other_data(xl, capsys):
    # 실기 J2: A1:C3 차트를 바로 오른쪽(D열)에 두어 F1:G5의 다른 데이터를 가렸음
    app, wb = xl
    ws = ws_of(wb)
    for r in range(1, 4):
        for c in range(1, 4):
            ws.set(r, c, "머리" if r == 1 else float(r * c))
    for r in range(1, 6):
        ws.set(r, 6, float(r))
        ws.set(r, 7, float(r * 2))
    code, out = run(capsys, "chart", "--range", "A1:C3", "--type", "column", "--title", "막대")
    assert code == 0
    chart = ws.charts[-1]
    assert chart.Left >= 7 * 48.0 and chart.Top == 0           # G열(7) 오른쪽, 범위의 첫 행
    code, out = run(capsys, "info", "--detail")
    assert out["layout_warnings"] == 0


def test_chart_delete_rules(xl, capsys):
    app, wb = xl
    ws = ws_of(wb)
    ws.set(1, 1, "a")
    ws.set(2, 1, 1.0)
    run(capsys, "chart", "--range", "A1:A2", "--title", "매출")
    code, out = run(capsys, "chart", "--action", "delete")
    assert code == 1
    code, out = run(capsys, "chart", "--action", "delete", "--name", "매출차트")
    assert code == 1 and "[후보] 매출" in out["message"] and "[추천]" not in out["message"]
    code, out = run(capsys, "chart", "--action", "delete", "--name", "매출")
    assert code == 0 and ws.charts == []


def test_argument_error_is_json(capsys):
    code, out = run(capsys, "chart", "--type", "donut")
    assert code == 2 and out["tool"] == "excel"


# ---------------------------------------------------------------------------
# 점검 반영: 데이터 손상 방지
# ---------------------------------------------------------------------------

def test_text_that_excel_would_reinterpret_is_kept_as_text(xl, capsys, tmp_path):
    app, wb = xl
    f = tmp_path / "코드.csv"
    f.write_text("코드,품번,금액,큰수,지수,날짜\n01234,1-2,(100),12345678901234567890,1E5,2026-10-28\n", encoding="utf-8")
    code, out = run(capsys, "write", "--file", str(f))
    assert code == 0, out
    ws = ws_of(wb)
    assert [ws.values[(2, c)] for c in range(1, 7)] == ["01234", "1-2", "(100)", "12345678901234567890", "1E5", "2026-10-28"]
    assert {(2, c) for c in range(1, 6)} <= ws.prefixed and (2, 6) not in ws.prefixed     # 연도가 있는 날짜는 날짜로 변환되게 둠
    assert excel.coerce("12345678901234567890") == "12345678901234567890" and excel.coerce("1200") == 1200


def test_replace_result_that_looks_numeric_stays_text(xl, capsys):
    app, wb = xl
    ws = ws_of(wb)
    ws.set(1, 1, "AB-001")
    code, out = run(capsys, "replace", "--find", "AB-", "--replace", "")
    assert code == 0 and ws.values[(1, 1)] == "001" and (1, 1) in ws.prefixed


def test_range_clear_requires_range(xl, capsys):
    app, wb = xl
    ws = ws_of(wb)
    ws.set(1, 1, "중요")
    code, out = run(capsys, "range", "--action", "clear")
    assert code == 1 and "--range" in out["message"] and ws.values[(1, 1)] == "중요"


def test_merge_with_several_values_needs_overwrite(xl, capsys):
    app, wb = xl
    ws = ws_of(wb)
    for c, v in enumerate(["이름", "부서", "직급"], 1):
        ws.set(1, c, v)
    code, out = run(capsys, "range", "--action", "merge", "--range", "A1:C1")
    assert code == 1 and out["error_type"] == "needs_user_decision" and "B1=부서" in out["message"]
    assert not [x for x in ws.log if x[0] == "Merge"]
    code, out = run(capsys, "range", "--action", "merge", "--range", "A1:C1", "--overwrite")
    assert code == 0 and ("Merge", "$A$1:$C$1", False) in ws.log


def test_copy_over_existing_values_needs_overwrite(xl, capsys):
    app, wb = xl
    src, dst = ws_of(wb), ws_of(wb, "요약")
    src.set(1, 1, "원본")
    dst.set(2, 2, "기존 값")
    code, out = run(capsys, "range", "--action", "copy", "--range", "A1:B2", "--dest", "요약!A1")
    assert code == 1 and out["error_type"] == "needs_user_decision" and "B2=기존 값" in out["message"]
    assert not [x for x in src.log if x[0] == "Copy"]
    code, out = run(capsys, "range", "--action", "copy", "--range", "A1:B2", "--dest", "요약!D1")
    assert code == 0 and ("Copy", "$A$1:$B$2", "$D$1") in src.log


def test_hide_row_address_hides_rows_not_columns(xl, capsys):
    import types
    assert excel._axis(types.SimpleNamespace(axis=None, range="5:7")) == "rows"
    assert excel._axis(types.SimpleNamespace(axis=None, range="'시트 1'!$5:$7")) == "rows"
    assert excel._axis(types.SimpleNamespace(axis=None, range="C:E")) == "cols"
    assert excel._axis(types.SimpleNamespace(axis="cols", range="5:7")) == "cols"


def test_append_uses_lowest_row_of_all_columns_and_start_row_header(xl, capsys, tmp_path):
    app, wb = xl
    ws = ws_of(wb)
    # C5부터 머리글, A열은 비어 있고 D열이 더 길게 차 있는 표
    for r, row in enumerate([["품목", "수량"], ["사과", 3], [None, 5]], 5):
        for c, v in enumerate(row, 3):
            if v is not None:
                ws.set(r, c, v)
    f = tmp_path / "추가.csv"
    f.write_text("품목,수량\n배,2\n", encoding="utf-8")
    code, out = run(capsys, "write", "--file", str(f), "--range", "C5", "--append")
    assert code == 0, out
    assert ws.values[(8, 3)] == "배" and ws.values[(8, 4)] == 2 and ws.values[(7, 4)] == 5     # D7을 덮어쓰지 않음
    assert out["range_written"] == "C8:D8"                                                   # 머리글(5행)을 다시 쓰지 않음
    code, out = run(capsys, "write", "--file", str(f), "--range", "C5", "--append", "--header", "품목,수량")
    assert code == 0 and out["range_written"] == "C9:D9" and ws.values[(9, 3)] == "배"          # --header도 같은 머리글이면 생략


def test_read_closed_xlsx_without_dimension_reads_all_rows(tmp_path, capsys, monkeypatch):
    import re as _re
    import zipfile
    import openpyxl
    src, dst = tmp_path / "src.xlsx", tmp_path / "nodim.xlsx"
    wb = openpyxl.Workbook()
    for r in [["a", "b", "c"], [1, 2, 3], [4, 5, 6], [7, 8, 9], [10, 11, 12]]:
        wb.active.append(r)
    wb.save(str(src))
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dst, "w") as zout:     # 다른 프로그램처럼 <dimension> 없이
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename.startswith("xl/worksheets/sheet"):
                data = _re.sub(rb"<dimension[^>]*/>", b"", data)
            zout.writestr(item, data)
    monkeypatch.setattr(off, "win32com", None)
    code, out = run(capsys, "read", "--target", str(dst))
    assert code == 0 and out["values"][-1] == [10, 11, 12] and len(out["values"]) == 5


def test_read_csv_target_without_excel(capsys, monkeypatch, tmp_path):
    """CSV도 excel.py read로 읽음 (2026-10-29 Gemma: 확장자 오류 뒤 pandas 코드를 직접 짬). CP949, 숫자 변환, --range."""
    p = tmp_path / "공정.csv"
    p.write_bytes("반응온도(°C),생산효율(%)\r\n200.5,61\r\n210,63\r\n".encode("cp949"))
    monkeypatch.setattr(off, "win32com", None)
    code, out = run(capsys, "read", "--target", str(p))
    assert code == 0 and out["rows"] == 3 and out["values"][1] == [200.5, 61]
    assert out["target_resolved"]["source"] == "file(csv)"
    code, out = run(capsys, "read", "--target", str(p), "--range", "B2:B3")
    assert code == 0 and out["values"] == [[61], [63]] and out["range"] == "B2:B3"
    code, out = run(capsys, "read", "--target", str(p), "--range", "A:A")
    assert code == 0 and out["values"] == [["반응온도(°C)"], [200.5], [210]]
