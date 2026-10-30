"""office-cli 1.0 Windows 실기 테스트 러너

사용법 (CLI/office-cli 폴더에서 실행):
    python tools/run_win_tests.py                     # 전체 진행 (눈으로 확인할 항목은 y/n 질문)
    python tools/run_win_tests.py --only A,C          # 특정 그룹만
    python tools/run_win_tests.py --self-email me@company.com   # G3 실제 발송 테스트 포함
    python tools/run_win_tests.py --auto              # 질문 없이 자동 판정 항목만
    python tools/run_win_tests.py --dir D:\\office_test --skip-outlook
    python tools/run_win_tests.py --agent --skip-outlook   # 에이전트(OpenCode 스킬)용: 질문 없음, stdout은 JSON 1줄

진행 방식:
    1) 테스트 폴더(기본 C:\\office_test)에 A/B 테스트 파일(xlsx, docx, pptx)을 새로 만듭니다.
    2) A.xlsx, A.docx, A.pptx를 더블클릭한 것처럼 화면에 엽니다.
    3) 체크리스트 순서대로 도구를 실행하고, COM으로 결과(셀 값, 파일 수정 시각, 슬라이드 수 등)를 자동 검증합니다.
    4) 화면 깜빡임처럼 눈으로만 확인 가능한 항목은 y/n으로 묻습니다.
    5) 결과를 테스트 폴더의 report_*.txt(요약)와 errors_*.txt(오류 분석용 상세 로그)로 저장합니다.
       문제가 있으면 errors_*.txt 파일을 그대로 공유해 주세요.

주의:
    - 실행 전에 다른 Excel/Word/PowerPoint 문서는 저장 후 모두 닫아 주세요 (활성 문서 판정이 섞이지 않도록).
    - 테스트가 진행되는 동안 마우스/키보드로 Office 창을 조작하지 마세요 (질문에 답할 때 제외).
"""
import argparse
import json
import os
import platform
import shutil
import subprocess
import sys
import time
import traceback
from datetime import datetime

DEV = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))          # CLI-dev\office-cli (tests, tools)
ROOT = os.path.join(os.path.dirname(os.path.dirname(DEV)), "CLI", "office-cli")  # 실행 파일: CLI-dev 옆의 CLI\office-cli
TOOLS_DIR = os.path.join(DEV, "tools")
if TOOLS_DIR not in sys.path:
    sys.path.insert(0, TOOLS_DIR)
import win_test_rules as rules   # noqa: E402  (판단 규칙: 끝내도 되는 프로세스, 문서 없음 오류)

PY = sys.executable
CMD_TIMEOUT = 120
RUN_ENV = {}   # 도구에 넘길 환경변수 (테스트 폴더 전용 캐시/데이터 폴더). main()에서 채움
AGENT = "--agent" in sys.argv   # 에이전트 모드: 로그는 stderr, stdout은 마지막 JSON 1줄만

# 오류 로그(errors_*.txt)용 기록: 모든 로그 줄과 도구 실행(명령, 종료 코드, stdout/stderr 전체)
LOG_LINES = []
CALLS = []
CALL_OUTPUT_LIMIT = 20000   # 실행 1건의 stdout/stderr 보관 상한(글자)

if AGENT:
    for _stream in (sys.stdout, sys.stderr):
        try:
            _stream.reconfigure(encoding="utf-8", errors="replace")
        except Exception:
            pass


def emit_agent_json(payload, code):
    """에이전트 모드의 마지막 출력: stdout에 JSON 1줄을 쓰고 종료합니다."""
    sys.stdout.write(json.dumps(payload, ensure_ascii=False) + "\n")
    sys.stdout.flush()
    sys.exit(code)


try:
    import pythoncom
    import win32com.client
except ImportError:
    if AGENT:
        emit_agent_json({"status": "error", "error_type": "dependency", "message": "pywin32가 필요합니다.",
                         "agent_directive": "STOP_RETRY: pywin32가 설치되어 있지 않습니다. "
                                            "사용자에게 'python -m pip install pywin32'가 필요하다고 알리세요."}, 1)
    print("[ERROR] pywin32가 필요합니다. (python -m pip install pywin32)")
    sys.exit(1)


class TestDocMissing(Exception):
    """테스트 문서(A.xlsx 등)가 닫혀 있고 다시 열지도 못함. 그 문서를 쓰는 항목은 바로 ERROR로 끝냅니다."""
    def __init__(self, name):
        super().__init__(f"테스트 문서 {name}가 닫혀 다시 열지 못함")
        self.name = name


class SetupError(Exception):
    """테스트 준비(파일 만들기, 창 열기)가 실패해 테스트를 진행할 수 없음."""
    def __init__(self, label, detail="", trace=""):
        super().__init__(f"{label}: {detail}" if detail else label)
        self.label, self.detail, self.trace = label, detail, trace


# ---------------------------------------------------------------------------
# 공통 유틸
# ---------------------------------------------------------------------------
class Ctx:
    def __init__(self, args):
        self.args = args
        self.dir = os.path.abspath(args.dir)
        self.results = []
        self.auto = args.auto
        self.untitled_before = []
        self.start_time = time.time()           # 이 시각 뒤에 생긴 Office 프로세스만 정리 대상
        self.stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        self.mark = (0, 0)                      # 현재 항목이 시작된 CALLS, LOG_LINES 위치
        self.docs_ready = False                 # A 문서를 화면에 연 뒤부터 '닫힌 테스트 문서 다시 열기'를 함
        self.unrecoverable = set()              # 이번 그룹에서 다시 열지 못한 테스트 문서 (그룹마다 다시 시도)
        self.recoveries = []                    # 러너 자동 복구 기록 (errors_*.txt)
        self.cleanup_log = []                   # 창/프로세스 정리 기록 (errors_*.txt)

    def path(self, name):
        return os.path.join(self.dir, name)


def log(msg):
    LOG_LINES.append(str(msg))
    print(msg, flush=True, file=sys.stderr if AGENT else sys.stdout)


def run_cli(module, *cli_args):
    """도구를 별도 프로세스로 실행하고 (JSON dict 또는 None, 종료 코드, stdout, stderr 마지막 부분)을 반환합니다."""
    cmd = [PY, os.path.join(ROOT, module)] + [str(a) for a in cli_args]
    shown = f"python {module} {subprocess.list2cmdline([str(a) for a in cli_args])}"
    log(f"    $ {shown}")
    env = dict(os.environ, PYTHONIOENCODING="utf-8", **RUN_ENV)
    call = {"command": shown, "started": datetime.now().strftime("%H:%M:%S")}
    CALLS.append(call)
    t0 = time.time()
    try:
        proc = subprocess.run(cmd, cwd=ROOT, capture_output=True, timeout=CMD_TIMEOUT, env=env)
    except subprocess.TimeoutExpired as e:
        call.update(exit_code="TIMEOUT", elapsed=round(time.time() - t0, 2),
                    stdout=(e.stdout or b"").decode("utf-8", errors="replace")[-CALL_OUTPUT_LIMIT:],
                    stderr=(e.stderr or b"").decode("utf-8", errors="replace")[-CALL_OUTPUT_LIMIT:])
        return None, "TIMEOUT", "", f"{CMD_TIMEOUT}초 안에 끝나지 않음 (대화상자가 떠서 멈췄을 가능성)"
    out = proc.stdout.decode("utf-8", errors="replace")
    err = proc.stderr.decode("utf-8", errors="replace")
    call.update(exit_code=proc.returncode, elapsed=round(time.time() - t0, 2),
                stdout=out[-CALL_OUTPUT_LIMIT:], stderr=err[-CALL_OUTPUT_LIMIT:])
    data = None
    for line in reversed(out.strip().splitlines()):
        line = line.strip()
        if line.startswith("{"):
            try:
                data = json.loads(line)
            except Exception:
                pass
            break
    return data, proc.returncode, out, err[-1500:]


def ask(ctx, question):
    """눈으로 확인할 항목. y=정상(PASS), n=문제(FAIL), s=건너뜀."""
    if ctx.auto:
        return None
    while True:
        ans = input(f"    [확인] {question} (y=예 / n=아니오 / s=건너뜀): ").strip().lower()
        if ans in ("y", "n", "s"):
            return {"y": True, "n": False, "s": None}[ans]


def com_retry(fn, tries=10, delay=0.5):
    """Office가 바쁠 때(RPC_E_CALL_REJECTED 등) 잠시 기다렸다가 재시도합니다.
    '그 이름의 문서/시트 없음'(DISP_E_BADINDEX)과 테스트 문서 없음은 바쁜 것이 아니므로 바로 올립니다
    (중첩된 com_retry가 10x10회씩 기다리며 멈춘 것처럼 보이던 문제)."""
    last = None
    for _ in range(tries):
        try:
            return fn()
        except TestDocMissing:
            raise
        except Exception as e:
            if rules.is_missing_doc_error(e):
                raise
            last = e
            time.sleep(delay)
    raise last


_WAITING = [0]   # wait_for 안에서는 닫힌 테스트 문서를 다시 열지 않음 (undo가 다시 여는 중일 수 있음)


def wait_for(fn, timeout=40, delay=1.0):
    end = time.time() + timeout
    _WAITING[0] += 1
    try:
        while time.time() < end:
            try:
                v = fn()
                if v:
                    return v
            except Exception:
                pass
            time.sleep(delay)
        return None
    finally:
        _WAITING[0] -= 1


def process_pids(image):
    """실행 중인 프로세스 PID 목록. 확인 실패는 None."""
    try:
        out = subprocess.run(["tasklist", "/FI", f"IMAGENAME eq {image}", "/FO", "CSV", "/NH"],
                             capture_output=True).stdout.decode("cp949", errors="replace")
        return rules.parse_tasklist_pids(out, image)
    except Exception:
        return None


def process_count(image):
    pids = process_pids(image)
    return -1 if pids is None else len(pids)


def describe_com_error(e):
    """pywintypes.com_error를 사람이 읽을 수 있게 풀어 씁니다 (HRESULT, Office가 준 설명 문구)."""
    args = getattr(e, "args", ())
    if len(args) >= 3 and isinstance(args[0], int):
        hr, msg, exc = args[0], args[1], args[2]
        text = f"HRESULT={hr} (0x{hr & 0xFFFFFFFF:08X}) {msg}"
        if exc and len(exc) >= 6:
            src, desc, scode = exc[1], exc[2], exc[5]
            text += f" | 원본={src} | 설명={desc}"
            if isinstance(scode, int):
                text += f" | scode=0x{scode & 0xFFFFFFFF:08X}"
        return text
    return f"{type(e).__name__}: {e}"


def step(label, fn):
    """준비 단계 하나를 실행하고, 실패하면 어느 단계에서 왜 실패했는지 출력한 뒤 종료합니다."""
    try:
        return fn()
    except Exception as e:
        log(f"[ERROR] 준비 단계 실패: {label}")
        log(f"        {describe_com_error(e)}")
        log(traceback.format_exc())
        raise SetupError(label, describe_com_error(e), traceback.format_exc())


def mtime(path):
    return os.path.getmtime(path) if os.path.exists(path) else None


def short(data, raw_out=""):
    if data is None:
        return (raw_out or "").strip()[-300:] or "(JSON 출력 없음)"
    s = json.dumps(data, ensure_ascii=False)
    return s if len(s) <= 400 else s[:400] + "..."


# ---------------------------------------------------------------------------
# 테스트 실행/기록
# ---------------------------------------------------------------------------
ENV_SKIP = "env-skip"   # 이 PC에 기능이 없어 시험할 수 없음 (예: Office 2007에 PDF 추가 기능 없음)


def pdf_unavailable(data):
    """도구가 'PDF 내보내기 기능 없음'을 STOP_RETRY 입력 오류로 정상 보고했는지."""
    data = data or {}
    return data.get("error_type") == "input" and "PDF 내보내기 기능이 없습니다" in data.get("message", "")


def _evidence(ctx):
    """지금 항목이 시작된 뒤의 도구 실행 기록과 로그 줄을 꺼내고 다음 항목의 시작 위치를 옮깁니다."""
    c0, l0 = ctx.mark
    ctx.mark = (len(CALLS), len(LOG_LINES))
    return CALLS[c0:], LOG_LINES[l0:]


def record(ctx, tid, title, ok, detail="", data=None, code=None, raw="", err=""):
    if ok is None:
        status = "MANUAL-SKIP"
    elif ok == ENV_SKIP:
        status = "ENV-SKIP"
    else:
        status = "PASS" if ok else "FAIL"
    entry = {"id": tid, "title": title, "status": status, "detail": detail,
             "exit_code": code, "json": short(data, raw)}
    if status == "FAIL" and err:
        entry["stderr_tail"] = err.strip()[-600:]
    log(f"  => [{status}] {tid} {title} {('- ' + detail) if detail else ''}")
    calls, lines = _evidence(ctx)
    if status == "FAIL":
        entry["calls"], entry["log"] = calls, lines
    ctx.results.append(entry)


def run_test(ctx, tid, title, fn):
    ctx.mark = (len(CALLS), len(LOG_LINES))
    log(f"\n[{tid}] {title}")
    try:
        fn()
    except TestDocMissing as e:
        # 테스트 문서가 닫혀 다시 열지 못함: 이유 하나로 바로 끝냅니다 (항목마다 오래 재시도하지 않음)
        log(f"  => [ERROR] {tid} {e}")
        calls, lines = _evidence(ctx)
        ctx.results.append({"id": tid, "title": title, "status": "ERROR", "detail": str(e),
                            "calls": calls, "log": lines})
    except Exception as e:
        trace = traceback.format_exc()
        log(f"  => [ERROR] {tid} {e}")
        calls, lines = _evidence(ctx)
        ctx.results.append({"id": tid, "title": title, "status": "ERROR",
                            "detail": f"테스트 스크립트 예외: {describe_com_error(e)}", "trace": trace[-800:],
                            "full_trace": trace, "calls": calls, "log": lines})


# ---------------------------------------------------------------------------
# Office 객체 조회 (테스트 검증용, office-cli 코드와 독립)
# ---------------------------------------------------------------------------
CTX = None   # main()에서 채움 (닫힌 테스트 문서 다시 열기와 기록에 사용)
_COLL = {"excel": "Workbooks", "word": "Documents", "ppt": "Presentations"}


def _office_mod():
    if ROOT not in sys.path:
        sys.path.insert(0, ROOT)
    import _office
    return _office


def office_app(kind):
    """사용자가 띄운 Office 앱. 더블클릭으로 연 Excel은 포커스를 잃기 전까지 ROT에 등록되지 않아
    GetActiveObject가 실패하므로(MS KB 238610), 창 핸들 바인딩(_office.find_apps)을 먼저 씁니다."""
    _office = _office_mod()
    apps = _office.find_apps(kind)
    if apps:
        return apps[0]
    return win32com.client.GetActiveObject(_office.PROG_ID[kind])


def _lookup_doc(kind, name):
    """모든 인스턴스에서 이름으로 문서를 찾습니다 (Excel이 둘이면 앞 인스턴스에 없을 수 있음). 1회만 시도."""
    _office = _office_mod()
    apps = _office.find_apps(kind)
    if not apps:
        apps = [win32com.client.GetActiveObject(_office.PROG_ID[kind])]
    last = None
    for app in apps:
        try:
            return getattr(app, _COLL[kind])(name)
        except Exception as e:
            last = e
    raise last


def _note_recovery(msg):
    log(f"[WARN] {msg}")
    if CTX is not None:
        CTX.recoveries.append(f"{datetime.now().strftime('%H:%M:%S')} {msg}")


def reopen_test_doc(kind, name):
    """닫혀 있는 테스트 문서를 준비 단계와 같은 방법(더블클릭과 같은 os.startfile)으로 다시 엽니다.
    다시 열면 True. 못 열면 이번 그룹에서는 다시 시도하지 않도록 표시하고 False."""
    ctx = CTX
    if ctx is None or name in ctx.unrecoverable:
        return False
    try:
        os.startfile(ctx.path(name))
    except Exception as e:
        ctx.unrecoverable.add(name)
        _note_recovery(f"{name}가 닫혀 있어 다시 열려고 했으나 실패했습니다: {describe_com_error(e)}")
        return False
    _WAITING[0] += 1          # 아래 확인 중에는 다시 열기를 중첩하지 않음
    try:
        end = time.time() + 40
        while time.time() < end:
            try:
                _lookup_doc(kind, name)
                _note_recovery(f"{name}가 닫혀 있어 다시 열었습니다")
                time.sleep(2)
                return True
            except Exception:
                time.sleep(1)
    finally:
        _WAITING[0] -= 1
    ctx.unrecoverable.add(name)
    _note_recovery(f"{name}가 닫혀 있어 다시 열려고 했으나 40초 안에 열리지 않았습니다")
    return False


def _get_doc(kind, name):
    """이름으로 열린 문서를 얻습니다. 테스트 문서(A.*)가 닫혀 있으면 한 번 다시 열고(JIT 복구),
    다시 열지 못하면 TestDocMissing으로 바로 끝냅니다 (com_retry로 오래 기다리지 않음)."""
    is_test_doc = rules.TEST_DOCS.get(kind) == name
    if is_test_doc and CTX is not None and name in CTX.unrecoverable:
        raise TestDocMissing(name)
    try:
        return com_retry(lambda: _lookup_doc(kind, name))
    except Exception as e:
        if not (is_test_doc and rules.is_doc_gone_error(e) and CTX is not None and CTX.docs_ready
                and _WAITING[0] == 0):
            raise
        # undo가 문서를 닫았다 다시 여는 중일 수 있으므로 잠시 기다려 본 뒤 다시 엽니다
        end = time.time() + 5
        while time.time() < end:
            time.sleep(1)
            try:
                return _lookup_doc(kind, name)
            except Exception:
                pass
        if not reopen_test_doc(kind, name):
            raise TestDocMissing(name)
        return com_retry(lambda: _lookup_doc(kind, name))


def ensure_test_docs(ctx):
    """그룹 시작 전: A.xlsx / A.docx / A.pptx가 열려 있는지 확인하고, 닫혀 있으면 다시 엽니다."""
    ctx.unrecoverable.clear()
    for kind, name in rules.TEST_DOCS.items():
        try:
            _lookup_doc(kind, name)
        except Exception as e:
            if rules.is_doc_gone_error(e):           # 바쁜 경우 등은 건드리지 않음 (열려 있는 문서를 다시 열지 않게)
                reopen_test_doc(kind, name)
            else:
                log(f"[WARN] {name} 열림 확인 실패 (다시 열지 않음): {describe_com_error(e)}")


def xl_app():
    return office_app("excel")


def xl_wb(name):
    return _get_doc("excel", name)


def wd_doc(name):
    return _get_doc("word", name)


def pp_prs(name):
    return _get_doc("ppt", name)


def pp_app():
    return office_app("ppt")


def activate_xl(name, sheet=None):
    def _do():
        wb = xl_wb(name)
        wb.Activate()
        if sheet:
            wb.Worksheets(sheet).Activate()
        return wb
    return com_retry(_do)


# ---------------------------------------------------------------------------
# 준비: 테스트 파일 생성 및 열기
# ---------------------------------------------------------------------------
FILES = ["A.xlsx", "B.xlsx", "A.docx", "B.docx", "A.pptx", "B.pptx"]


def _set_sheet_name(wb, idx, name):
    """이미 같은 이름이면 건너뛰고, 다른 시트가 그 이름을 쓰고 있으면 먼저 비켜 줍니다."""
    ws = wb.Worksheets(idx)
    if ws.Name == name:
        return
    for other in wb.Worksheets:
        if other.Name == name:
            other.Name = f"{name}_old"
    ws.Name = name


def _save_workbook(wb, path):
    try:
        wb.SaveAs(path, FileFormat=51)  # xlOpenXMLWorkbook
    except Exception as e:
        log(f"        SaveAs(FileFormat=51) 실패 -> 형식 지정 없이 재시도: {describe_com_error(e)}")
        wb.SaveAs(path)


def _save_document(doc, path):
    """SaveAs2는 Word 2010+에만 있으므로 Word 2007은 SaveAs로 저장합니다."""
    try:
        save = doc.SaveAs2
    except AttributeError:
        save = doc.SaveAs
    save(path, FileFormat=16)  # wdFormatDocumentDefault


def isolate_tool_dirs(ctx):
    """캐시(스냅샷, 색인)와 data 폴더를 테스트 폴더 아래로 분리하고 캐시는 매번 비웁니다.
    이전 실행의 스냅샷이 남아 있으면 같은 경로로 다시 만든 A 파일에 undo가 이어져 C7이 틀어집니다."""
    import shutil
    cache, data = ctx.path("_cache"), ctx.path("_data")
    shutil.rmtree(cache, ignore_errors=True)
    os.makedirs(cache, exist_ok=True)
    os.makedirs(data, exist_ok=True)
    RUN_ENV.update(OFFICE_CLI_CACHE_DIR=cache, OFFICE_CLI_DATA_DIR=data)
    try:   # 이전 실행의 도구 오류 기록은 빼고 이번 실행분만 errors_*.txt에 옮겨 담음
        TOOL_LOG_OFFSET[0] = os.path.getsize(os.path.join(data, "logs", "errors.txt"))
    except OSError:
        TOOL_LOG_OFFSET[0] = 0


def setup_files(ctx):
    step(f"테스트 폴더 생성/쓰기 확인 ({ctx.dir})", lambda: _check_writable(ctx.dir))
    step("도구 캐시/데이터 폴더 분리", lambda: isolate_tool_dirs(ctx))
    for f in FILES:
        p = ctx.path(f)
        if os.path.exists(p):
            try:
                os.remove(p)
            except PermissionError:
                log(f"[ERROR] {p} 파일이 열려 있어 다시 만들 수 없습니다. 닫고 다시 실행하세요.")
                raise SetupError(f"{f} 다시 만들기", "파일이 열려 있음 (닫고 다시 실행)")

    log("[준비] Excel 테스트 파일 생성")
    xl = step("Excel 백그라운드 실행 (DispatchEx)", lambda: win32com.client.DispatchEx("Excel.Application"))
    try:
        step("Excel 표시/경고 설정", lambda: (setattr(xl, "Visible", False), setattr(xl, "DisplayAlerts", False)))
        for name, a1 in (("A.xlsx", "원래값"), ("B.xlsx", "B파일")):
            wb = step(f"{name}: 새 통합문서 만들기", lambda: xl.Workbooks.Add())

            def _sheets():
                while wb.Worksheets.Count < 2:
                    wb.Worksheets.Add(After=wb.Worksheets(wb.Worksheets.Count))
                _set_sheet_name(wb, 1, "Sheet1")
                _set_sheet_name(wb, 2, "Sheet2")
            step(f"{name}: 시트 2개 구성", _sheets)

            def _values():
                wb.Worksheets(1).Range("A1").Value = a1
                wb.Worksheets(2).Range("A1").Value = "시트2 데이터"
            step(f"{name}: 셀 값 입력", _values)
            step(f"{name}: 저장 ({ctx.path(name)})", lambda: _save_workbook(wb, ctx.path(name)))
            step(f"{name}: 닫기", lambda: wb.Close(SaveChanges=False))
    finally:
        try:
            xl.Quit()
        except Exception:
            pass

    log("[준비] Word 테스트 파일 생성")
    wd = step("Word 백그라운드 실행 (DispatchEx)", lambda: win32com.client.DispatchEx("Word.Application"))
    try:
        step("Word 표시/경고 설정", lambda: (setattr(wd, "Visible", False), setattr(wd, "DisplayAlerts", 0)))
        for name in ("A.docx", "B.docx"):
            doc = step(f"{name}: 새 문서 만들기", lambda: wd.Documents.Add())
            step(f"{name}: 본문 입력",
                 lambda: setattr(doc.Content, "Text", f"{name} 첫 문단입니다.\r두 번째 문단입니다."))
            step(f"{name}: 저장 ({ctx.path(name)})", lambda: _save_document(doc, ctx.path(name)))
            step(f"{name}: 닫기", lambda: doc.Close(SaveChanges=0))
    finally:
        try:
            wd.Quit(SaveChanges=0)
        except Exception:
            pass

    log("[준비] PowerPoint 테스트 파일 생성")
    running = True
    try:
        win32com.client.GetActiveObject("PowerPoint.Application")
    except Exception:
        running = False
    pp = step("PowerPoint 실행/연결", lambda: win32com.client.Dispatch("PowerPoint.Application"))
    for name in ("A.pptx", "B.pptx"):
        prs = step(f"{name}: 새 프레젠테이션 만들기 (창 없이)", lambda: pp.Presentations.Add(0))

        def _slides():
            for i in range(1, 4):
                s = prs.Slides.Add(i, 12)  # ppLayoutBlank
                s.Shapes.AddTextbox(1, 50, 50, 400, 50).TextFrame.TextRange.Text = f"{name} 슬라이드 {i}"
                s.Shapes.AddTextbox(1, 50, 150, 400, 50).TextFrame.TextRange.Text = f"본문 {i}"
        step(f"{name}: 슬라이드 3장 구성", _slides)
        step(f"{name}: 저장 ({ctx.path(name)})", lambda: prs.SaveAs(ctx.path(name), 24))
        step(f"{name}: 닫기", lambda: prs.Close())
    if not running:
        try:
            if pp.Presentations.Count == 0:
                pp.Quit()
        except Exception:
            pass


def _check_writable(folder):
    os.makedirs(folder, exist_ok=True)
    probe = os.path.join(folder, "_write_test.tmp")
    with open(probe, "w", encoding="utf-8") as f:
        f.write("ok")
    os.remove(probe)


def open_user_windows(ctx):
    log("[준비] A.xlsx / A.docx / A.pptx를 화면에 엽니다 (더블클릭과 동일)")
    os.startfile(ctx.path("A.xlsx"))
    if not wait_for(lambda: xl_wb("A.xlsx")):
        log("[ERROR] A.xlsx가 열리지 않았습니다.")
        raise SetupError("A.xlsx 화면에 열기", "제한 시간 안에 열린 문서를 찾지 못함")
    os.startfile(ctx.path("A.docx"))
    if not wait_for(lambda: wd_doc("A.docx")):
        log("[ERROR] A.docx가 열리지 않았습니다.")
        raise SetupError("A.docx 화면에 열기", "제한 시간 안에 열린 문서를 찾지 못함")
    os.startfile(ctx.path("A.pptx"))
    if not wait_for(lambda: pp_prs("A.pptx")):
        log("[ERROR] A.pptx가 열리지 않았습니다.")
        raise SetupError("A.pptx 화면에 열기", "제한 시간 안에 열린 문서를 찾지 못함")
    time.sleep(2)
    ctx.docs_ready = True


# ---------------------------------------------------------------------------
# A. 사용자 창 간섭 / 닫힌 파일 처리
# ---------------------------------------------------------------------------
def test_A(ctx):
    def a1():
        activate_xl("A.xlsx")
        before = process_count("EXCEL.EXE")
        data, code, raw, err = run_cli("excel.py", "read", "--target", ctx.path("B.xlsx"))
        after = process_count("EXCEL.EXE")
        src = ((data or {}).get("target_resolved") or {}).get("source", "")
        ok = (data or {}).get("status") == "ok" and src.startswith("file(openpyxl)") and after <= before
        record(ctx, "A1", "닫힌 엑셀 읽기는 Excel을 띄우지 않음 (openpyxl)", ok,
               f"EXCEL.EXE {before}->{after}, source={src}", data, code, raw, err)

    def a2():
        activate_xl("A.xlsx")
        before = process_count("EXCEL.EXE")
        b_mtime = mtime(ctx.path("B.xlsx"))
        data, code, raw, err = run_cli("excel.py", "write", "--target", ctx.path("B.xlsx"), "--range", "C1", "--value", "닫힌파일쓰기")
        time.sleep(2)
        after = process_count("EXCEL.EXE")
        still_open = bool(wait_for(lambda: xl_wb("A.xlsx"), timeout=5))
        visible = com_retry(lambda: bool(xl_app().Visible))
        ok = (data or {}).get("status") == "ok" and still_open and visible and after <= before \
            and mtime(ctx.path("B.xlsx")) != b_mtime
        manual = ask(ctx, "실행 중 A.xlsx 창이 닫히거나 깜빡이지 않고 그대로 있었나요?")
        record(ctx, "A2", "닫힌 엑셀 파일 쓰기 시 사용자 창 유지", ok and manual is not False,
               f"EXCEL.EXE {before}->{after}, A.xlsx 열림={still_open}, Visible={visible}", data, code, raw, err)

    def a3():
        before = process_count("WINWORD.EXE")
        data, code, raw, err = run_cli("word.py", "write", "--target", ctx.path("B.docx"), "--text", "닫힌 문서에 추가")
        time.sleep(2)
        after = process_count("WINWORD.EXE")
        still_open = bool(wait_for(lambda: wd_doc("A.docx"), timeout=5))
        ok = (data or {}).get("status") == "ok" and still_open and after <= before
        manual = ask(ctx, "실행 중 A.docx 창이 그대로 있었나요?")
        record(ctx, "A3", "닫힌 워드 파일 쓰기 시 사용자 창 유지", ok and manual is not False,
               f"WINWORD.EXE {before}->{after}, A.docx 열림={still_open}", data, code, raw, err)

    def a4():
        b_mtime = mtime(ctx.path("B.pptx"))
        data, code, raw, err = run_cli("ppt.py", "info", "--target", ctx.path("B.pptx"))
        time.sleep(1)
        state = com_retry(lambda: pp_app().WindowState)
        unchanged = mtime(ctx.path("B.pptx")) == b_mtime
        ok = (data or {}).get("status") == "ok" and unchanged and state != 2
        manual = ask(ctx, "PowerPoint 창이 최소화되거나 깜빡이지 않았나요?")
        record(ctx, "A4", "닫힌 PPT 파일 조회 시 창 유지 / 파일 미변경", ok and manual is not False,
               f"B.pptx 수정시각 불변={unchanged}, WindowState={state}", data, code, raw, err)

    def a5():
        c_path = ctx.path("C_second_instance.xlsx")
        if os.path.exists(c_path):
            os.remove(c_path)
        import shutil
        shutil.copy2(ctx.path("B.xlsx"), c_path)
        pids0 = process_pids("EXCEL.EXE")
        if pids0 is None:      # 기존 PID를 모르면 '이 단계가 띄운 프로세스'를 가릴 수 없음
            record(ctx, "A5", "다중 Excel 프로세스 목록", None, "tasklist로 Excel 프로세스를 확인하지 못해 건너뜀")
            return
        before_pids = set(pids0)
        before = len(before_pids)
        subprocess.Popen(["cmd", "/c", "start", "", "excel.exe", "/x", c_path])
        if not wait_for(lambda: set(process_pids("EXCEL.EXE") or []) - before_pids, timeout=30):
            record(ctx, "A5", "다중 Excel 프로세스 목록", None, "두 번째 Excel 프로세스를 띄우지 못해 건너뜀")
            return
        time.sleep(5)
        # 이 단계가 띄운 프로세스만 정리 대상 (A.xlsx를 연 사용자 Excel은 실행 중에 생겼어도 대상이 아님)
        spawned = sorted(set(process_pids("EXCEL.EXE") or []) - before_pids)
        log(f"    두 번째 Excel PID: {spawned} (기존 PID: {sorted(before_pids)})")
        data, code, raw, err = run_cli("excel.py", "info", "--all")
        text = json.dumps(data, ensure_ascii=False) if data else raw
        ok = "A.xlsx" in text and "C_second_instance.xlsx" in text
        record(ctx, "A5", "다중 Excel 프로세스 목록", ok, f"EXCEL.EXE 프로세스 {process_count('EXCEL.EXE')}개", data, code, raw, err)
        # 두 번째 Excel도 문서가 없어지면 종료. 끝내는 대상은 spawned PID뿐이고 문서가 남은 프로세스는 건드리지 않음
        close_test_documents(ctx, only_name="C_second_instance.xlsx", allowed_pids=spawned, label="A5")
        wait_for(lambda: process_count("EXCEL.EXE") <= before, timeout=15)
        if not wait_for(lambda: _lookup_doc("excel", "A.xlsx"), timeout=5):
            log("[WARN] A5 정리 뒤 A.xlsx가 보이지 않습니다. 다음 그룹 시작 전에 다시 엽니다.")
        ask(ctx, "두 번째 Excel 창(C_second_instance.xlsx)이 닫혔나요? 남아 있으면 저장하지 않고 직접 닫은 뒤 y를 누르세요")

    run_test(ctx, "A1", "닫힌 엑셀 읽기", a1)
    run_test(ctx, "A2", "닫힌 엑셀 쓰기와 사용자 창", a2)
    run_test(ctx, "A3", "닫힌 워드 쓰기와 사용자 창", a3)
    run_test(ctx, "A4", "닫힌 PPT 조회와 사용자 창", a4)
    run_test(ctx, "A5", "다중 Excel 프로세스 목록", a5)


# ---------------------------------------------------------------------------
# B. 읽기 전용 동작
# ---------------------------------------------------------------------------
def test_B(ctx):
    def b1():
        b_mtime = mtime(ctx.path("B.xlsx"))
        data, code, raw, err = run_cli("excel.py", "read", "--target", ctx.path("B.xlsx"), "--sheet", "없는시트")
        msg = (data or {}).get("message", "")
        ok = (data or {}).get("status") == "error" and ("Sheet1" in msg or "[추천]" in msg) \
            and mtime(ctx.path("B.xlsx")) == b_mtime and code == 1
        record(ctx, "B1", "없는 시트 읽기 -> error, 파일 미변경", ok, msg[:120], data, code, raw, err)

    def b2():
        activate_xl("A.xlsx", "Sheet1")
        data, code, raw, err = run_cli("excel.py", "read", "--target", "active", "--sheet", "Sheet2")
        active = com_retry(lambda: xl_wb("A.xlsx").ActiveSheet.Name)
        got = json.dumps(data, ensure_ascii=False) if data else ""
        ok = (data or {}).get("status") == "ok" and active == "Sheet1" and "시트2 데이터" in got
        record(ctx, "B2", "열린 문서 다른 시트 읽기 시 화면 시트 유지", ok, f"읽은 뒤 활성 시트={active}", data, code, raw, err)

    def b3():
        data, code, raw, err = run_cli("excel.py", "read", "--target", "active", "--sheet", "sheet2")
        ok = (data or {}).get("status") == "ok" and "시트2 데이터" in json.dumps(data or {}, ensure_ascii=False)
        record(ctx, "B3", "시트명 대소문자 무시", ok, "", data, code, raw, err)

    def b4():
        missing = ctx.path("없는파일.docx")
        if os.path.exists(missing):
            os.remove(missing)
        data, code, raw, err = run_cli("word.py", "read", "--target", missing)
        ok = (data or {}).get("status") == "error" and not os.path.exists(missing)
        record(ctx, "B4", "없는 워드 파일 읽기 -> error, 파일 미생성", ok, f"파일 생성됨={os.path.exists(missing)}", data, code, raw, err)

    def b5():
        prs = pp_prs("A.pptx")
        com_retry(lambda: pp_app().ActiveWindow.View.GotoSlide(1))
        data, code, raw, err = run_cli("ppt.py", "read", "--target", "active")
        cur = com_retry(lambda: pp_app().ActiveWindow.View.Slide.SlideIndex)
        ok = (data or {}).get("status") == "ok" and cur == 1
        record(ctx, "B5", "PPT 읽기 시 화면이 마지막 슬라이드로 이동하지 않음", ok, f"읽은 뒤 화면 슬라이드={cur}", data, code, raw, err)
        del prs

    run_test(ctx, "B1", "없는 시트 읽기", b1)
    run_test(ctx, "B2", "활성 시트 유지", b2)
    run_test(ctx, "B3", "시트명 대소문자 무시", b3)
    run_test(ctx, "B4", "없는 파일 읽기", b4)
    run_test(ctx, "B5", "PPT 읽기 화면 유지", b5)


# ---------------------------------------------------------------------------
# C. 되돌리기
# ---------------------------------------------------------------------------
def a1_value(book="A.xlsx", sheet="Sheet1"):
    return com_retry(lambda: xl_wb(book).Worksheets(sheet).Range("A1").Value)


def test_C(ctx):
    def c1():
        activate_xl("A.xlsx", "Sheet1")
        data, code, raw, err = run_cli("excel.py", "write", "--range", "A1", "--value", "안녕")
        filled = a1_value()
        if filled != "안녕":
            record(ctx, "C1", "Excel 쓰기 후 undo", False, f"쓰기 후 A1={filled!r}", data, code, raw, err)
            return
        data, code, raw, err = run_cli("office.py", "undo", "--app", "excel")
        wait_for(lambda: xl_wb("A.xlsx"), timeout=20)
        restored = a1_value()
        ok = (data or {}).get("status") == "ok" and restored == "원래값" and (data or {}).get("restored_as") == "reopened_in_office"
        record(ctx, "C1", "Excel 쓰기 후 undo", ok,
               f"undo 후 A1={restored!r}, restored_as={(data or {}).get('restored_as')}", data, code, raw, err)

    def c2():
        activate_xl("A.xlsx", "Sheet1")
        run_cli("excel.py", "write", "--range", "A1", "--value", "값1")
        activate_xl("A.xlsx", "Sheet1")
        run_cli("excel.py", "write", "--range", "A1", "--value", "값2")
        run_cli("office.py", "undo", "--app", "excel")
        wait_for(lambda: xl_wb("A.xlsx"), timeout=20)
        v1 = a1_value()
        d2, c2_, r2, e2 = run_cli("office.py", "undo", "--app", "excel")
        wait_for(lambda: xl_wb("A.xlsx"), timeout=20)
        v2 = a1_value()
        record(ctx, "C2", "연속 undo 두 번", v1 == "값1" and v2 == "원래값", f"1차 후 A1={v1!r}, 2차 후 A1={v2!r}", d2, c2_, r2, e2)

    def c3():
        # 1.2.2: 저장 안 한 문서 undo는 되돌린 내용을 새 창으로 열고 바뀐 통합문서는 닫음 -> 통합문서 수 그대로
        app = xl_app()
        new_wb = com_retry(lambda: app.Workbooks.Add())
        new_name = new_wb.Name

        def _seed():
            new_wb.Worksheets(1).Range("A1").Value = "원래새값"
        com_retry(_seed)
        new_wb.Activate()
        before = app.Workbooks.Count
        run_cli("excel.py", "write", "--range", "A1", "--value", "새문서값")
        data, code, raw, err = run_cli("office.py", "undo", "--app", "excel")
        time.sleep(2)
        app = xl_app()
        after = com_retry(lambda: app.Workbooks.Count)
        names = [com_retry(lambda i=i: app.Workbooks(i).Name) for i in range(1, after + 1)]
        restored = (data or {}).get("restored_name", "")
        value = com_retry(lambda: xl_wb(restored).Worksheets(1).Range("A1").Value) if restored in names else None
        ok = (data or {}).get("status") == "ok" and (data or {}).get("restored_as") == "replaced_untitled_document" \
            and after == before and new_name not in names and value == "원래새값"
        record(ctx, "C3", "저장 안 한 새 통합문서 undo -> 바뀐 창 닫고 되돌린 창 하나", ok,
               f"통합문서 수 {before}->{after}, '{new_name}' 남음={new_name in names}, 되돌린 A1={value!r}", data, code, raw, err)
        try:
            app = xl_app()
            for i in range(app.Workbooks.Count, 0, -1):
                wb = app.Workbooks(i)
                if wb.Name != "A.xlsx" and (wb.Name in (new_name, restored) or not wb.Path):
                    wb.Close(SaveChanges=False)
        except Exception:
            pass

    def c4():
        doc = wd_doc("A.docx")
        com_retry(lambda: doc.Activate())
        com_retry(lambda: doc.Range(0, 0).InsertBefore("가"))
        m_before = mtime(ctx.path("A.docx"))
        data, code, raw, err = run_cli("word.py", "write", "--text", "테스트 문장")
        doc = wd_doc("A.docx")
        text = com_retry(lambda: doc.Content.Text)
        ok = (data or {}).get("status") == "ok" and "테스트 문장" in text and mtime(ctx.path("A.docx")) == m_before \
            and not com_retry(lambda: bool(doc.Saved))
        record(ctx, "C4", "워드 쓰기 성공 + 원본 자동 저장 안 됨", ok,
               f"수정시각 불변={mtime(ctx.path('A.docx')) == m_before}", data, code, raw, err)

    def c5():
        data, code, raw, err = run_cli("office.py", "undo", "--app", "word")
        wait_for(lambda: wd_doc("A.docx"), timeout=20)
        text = com_retry(lambda: wd_doc("A.docx").Content.Text)
        ok = (data or {}).get("status") == "ok" and "테스트 문장" not in text and text.startswith("가")
        record(ctx, "C5", "워드 undo -> 도구가 쓴 문장만 제거", ok,
               f"'테스트 문장' 남음={'테스트 문장' in text}, '가' 유지={text.startswith('가')}", data, code, raw, err)

    def c6():
        n0 = com_retry(lambda: pp_prs("A.pptx").Slides.Count)
        run_cli("ppt.py", "slide", "--action", "add")
        n1 = com_retry(lambda: pp_prs("A.pptx").Slides.Count)
        data, code, raw, err = run_cli("office.py", "undo", "--app", "ppt")
        wait_for(lambda: pp_prs("A.pptx"), timeout=20)
        n2 = com_retry(lambda: pp_prs("A.pptx").Slides.Count)
        ok = n1 == n0 + 1 and n2 == n0 and (data or {}).get("status") == "ok"
        record(ctx, "C6", "PPT 슬라이드 추가 후 undo", ok, f"슬라이드 수 {n0}->{n1}->{n2}", data, code, raw, err)

    def c7():
        last = None
        for _ in range(4):
            last = run_cli("office.py", "undo", "--app", "ppt")
            if (last[0] or {}).get("restored") is False:
                break
            wait_for(lambda: pp_prs("A.pptx"), timeout=20)
        data, code, raw, err = last
        ok = (data or {}).get("restored") is False and code == 0
        record(ctx, "C7", "스냅샷 소진 시 '되돌릴 것 없음'", ok, "", data, code, raw, err)

    def c8():
        # 시나리오 12: 저장 안 한 새 워드 문서에 쓴 뒤 undo -> 같은 창에서 되돌림, 새 창 없음
        wd = office_app("word")
        doc = com_retry(lambda: wd.Documents.Add())
        name = doc.Name
        com_retry(lambda: doc.Activate())
        before = com_retry(lambda: wd.Documents.Count)
        run_cli("word.py", "write", "--text", "새문서문장")
        written = "새문서문장" in com_retry(lambda: wd_doc(name).Content.Text)
        data, code, raw, err = run_cli("office.py", "undo", "--app", "word")
        time.sleep(2)
        after = com_retry(lambda: wd.Documents.Count)
        try:
            text = com_retry(lambda: wd_doc(name).Content.Text)
        except Exception:
            text = None
        ok = written and (data or {}).get("status") == "ok" and (data or {}).get("restored_as") == "restored_in_place" \
            and after == before and text is not None and "새문서문장" not in text
        record(ctx, "C8", "저장 안 한 새 워드 문서 undo -> 같은 창에서 되돌림", ok,
               f"쓰기 확인={written}, 문서 수 {before}->{after}, 같은 창 남음={text is not None}, "
               f"문장 남음={text is not None and '새문서문장' in text}", data, code, raw, err)
        try:
            for i in range(wd.Documents.Count, 0, -1):
                d = wd.Documents(i)
                if not d.Path:
                    d.Close(SaveChanges=0)
        except Exception:
            pass

    run_test(ctx, "C1", "Excel undo", c1)
    run_test(ctx, "C2", "연속 undo", c2)
    run_test(ctx, "C3", "새 통합문서 undo", c3)
    run_test(ctx, "C4", "워드 쓰기", c4)
    run_test(ctx, "C5", "워드 undo", c5)
    run_test(ctx, "C6", "PPT undo", c6)
    run_test(ctx, "C7", "스냅샷 소진", c7)
    run_test(ctx, "C8", "새 워드 문서 undo", c8)


# ---------------------------------------------------------------------------
# D. 내보내기
# ---------------------------------------------------------------------------
def test_D(ctx):
    def d1():
        doc = wd_doc("A.docx")
        com_retry(lambda: doc.Activate())
        footer_before = com_retry(lambda: doc.Sections(1).Footers(1).PageNumbers.Count)
        saved_before = com_retry(lambda: bool(doc.Saved))
        data, code, raw, err = run_cli("word.py", "export", "--save", "office_test_pagenum.pdf", "--page-number")
        doc = wd_doc("A.docx")
        footer_after = com_retry(lambda: doc.Sections(1).Footers(1).PageNumbers.Count)
        saved_after = com_retry(lambda: bool(doc.Saved))
        pdf = (data or {}).get("output_file", "")
        if pdf_unavailable(data) and footer_before == footer_after and saved_before == saved_after:
            record(ctx, "D1", "PDF 쪽 번호, 원본 바닥글 미변경", ENV_SKIP,
                   "이 PC의 Word에 PDF 내보내기 기능이 없음 (원본은 그대로)", data, code, raw, err)
            return
        ok = (data or {}).get("status") == "ok" and os.path.exists(pdf) and footer_before == footer_after and saved_before == saved_after
        manual = None
        if ok and not ctx.auto:
            os.startfile(pdf)
            manual = ask(ctx, "방금 열린 PDF 아래쪽 가운데에 쪽 번호가 보이나요? (확인 후 PDF 창은 닫아 주세요)")
        record(ctx, "D1", "PDF 쪽 번호, 원본 바닥글 미변경", ok and manual is not False,
               f"원본 쪽번호 수 {footer_before}->{footer_after}, 저장상태 {saved_before}->{saved_after}", data, code, raw, err)

    def d2():
        doc = wd_doc("A.docx")
        full_before = com_retry(lambda: doc.FullName)
        data, code, raw, err = run_cli("word.py", "export", "--save", "office_test_copy.docx")
        fp = (data or {}).get("output_file", "")
        full_after = com_retry(lambda: wd_doc("A.docx").FullName)
        ok = (data or {}).get("status") == "ok" and os.path.exists(fp) and full_before == full_after
        record(ctx, "D2", "DOCX 사본 저장, 원본 경로 그대로", ok, fp, data, code, raw, err)

    def d3():
        activate_xl("A.xlsx", "Sheet1")
        before = com_retry(lambda: xl_app().Workbooks.Count)
        data, code, raw, err = run_cli("excel.py", "export", "--save", "office_test_t.xlsx")
        after = com_retry(lambda: xl_app().Workbooks.Count)
        fp = (data or {}).get("output_file", "")
        ok = (data or {}).get("status") == "ok" and os.path.exists(fp) and before == after
        record(ctx, "D3", "XLSX 내보내기 후 임시 통합문서 미잔류", ok, f"통합문서 수 {before}->{after}", data, code, raw, err)

    def d4():
        data, code, raw, err = run_cli("ppt.py", "export", "--save", "office_test_deck.pdf")
        fp = (data or {}).get("output_file", "")
        if pdf_unavailable(data):
            record(ctx, "D4", "PPT PDF 내보내기", ENV_SKIP, "이 PC의 PowerPoint에 PDF 내보내기 기능이 없음", data, code, raw, err)
            return
        record(ctx, "D4", "PPT PDF 내보내기", (data or {}).get("status") == "ok" and os.path.exists(fp), fp, data, code, raw, err)

    run_test(ctx, "D1", "워드 PDF 쪽 번호", d1)
    run_test(ctx, "D2", "워드 DOCX 사본", d2)
    run_test(ctx, "D3", "엑셀 XLSX 내보내기", d3)
    run_test(ctx, "D4", "PPT PDF 내보내기", d4)


# ---------------------------------------------------------------------------
# E. 엑셀 편집
# ---------------------------------------------------------------------------
def prepare_e_sheets():
    def _do():
        wb = xl_wb("A.xlsx")
        app = wb.Application
        prev = app.DisplayAlerts
        app.DisplayAlerts = False
        try:
            for name in ("E_TEST", "My Sheet"):
                for s in wb.Worksheets:
                    if s.Name == name:
                        s.Delete()
                        break
                ws = wb.Worksheets.Add(After=wb.Worksheets(wb.Worksheets.Count))
                ws.Name = name
        finally:
            app.DisplayAlerts = prev
        ws = wb.Worksheets("E_TEST")
        ws.Range("A1:B3").Value = (("가", "나"), ("다", "라"), ("마", "바"))
        ws.Range("A10").Value = "병합1"
        ws.Range("B10").Value = "병합2"
        ws.Range("A20").Value = 1
        ws.Range("B20").Value = 2
        ws.Range("C20").Formula = "=A20+B20"
        wb.Activate()
        ws.Activate()
    com_retry(_do)


def test_E(ctx):
    prepare_e_sheets()

    def e_ws():
        return xl_wb("A.xlsx").Worksheets("E_TEST")

    def e1():
        activate_xl("A.xlsx", "E_TEST")
        data, code, raw, err = run_cli("excel.py", "range", "--action", "move", "--range", "A1:B3", "--dest", "B2")
        vals = com_retry(lambda: e_ws().Range("B2:C4").Value)
        a1 = com_retry(lambda: e_ws().Range("A1").Value)
        ok = (data or {}).get("status") == "ok" and vals == (("가", "나"), ("다", "라"), ("마", "바")) and a1 is None
        record(ctx, "E1", "범위 이동 데이터 손실 없음", ok, f"B2:C4={vals}, A1={a1!r}", data, code, raw, err)

    def e2():
        activate_xl("A.xlsx", "E_TEST")
        data, code, raw, err = run_cli("excel.py", "range", "--action", "merge")
        record(ctx, "E2", "범위 없는 병합 -> error", (data or {}).get("status") == "error", "", data, code, raw, err)

    def e3():
        activate_xl("A.xlsx", "E_TEST")
        d0 = run_cli("excel.py", "range", "--action", "merge", "--range", "A10:B10")[0]       # 값 2개: 사용자 확인 필요
        held = not com_retry(lambda: bool(e_ws().Range("A10").MergeCells)) and com_retry(lambda: e_ws().Range("B10").Value) == "병합2"
        data, code, raw, err = run_cli("excel.py", "range", "--action", "merge", "--range", "A10:B10", "--overwrite")
        merged = com_retry(lambda: bool(e_ws().Range("A10").MergeCells))
        ok = (d0 or {}).get("error_type") == "needs_user_decision" and held \
            and (data or {}).get("status") == "ok" and merged and code != "TIMEOUT"
        record(ctx, "E3", "값이 있는 셀 병합: 확인 없이는 멈춤, --overwrite면 대화상자 없이 완료", ok,
               f"확인 요청={(d0 or {}).get('error_type')}, 값 유지={held}, 병합됨={merged}", data, code, raw, err)

    def e4():
        activate_xl("A.xlsx", "E_TEST")
        data, code, raw, err = run_cli("excel.py", "range", "--action", "width", "--range", "'My Sheet'!A1", "--size", "30")
        width = com_retry(lambda: xl_wb("A.xlsx").Worksheets("My Sheet").Columns(1).ColumnWidth)
        ok = (data or {}).get("status") == "ok" and abs(float(width) - 30) < 0.01
        record(ctx, "E4", "따옴표 시트명 주소 해석", ok, f"My Sheet A열 너비={width}", data, code, raw, err)

    def e5():
        activate_xl("A.xlsx", "E_TEST")
        data, code, raw, err = run_cli("excel.py", "replace", "--find", "3", "--replace", "9", "--whole", "--bg", "yellow")
        formula = com_retry(lambda: e_ws().Range("C20").Formula)
        color = com_retry(lambda: int(e_ws().Range("C20").Interior.Color))
        ok = (data or {}).get("status") == "ok" and formula == "=A20+B20" and color == 65535
        record(ctx, "E5", "수식 셀 치환 시 수식 보존 + 서식 적용", ok, f"C20 수식={formula}, 배경색={color}", data, code, raw, err)

    def e6():
        activate_xl("A.xlsx", "E_TEST")
        data, code, raw, err = run_cli("excel.py", "style", "--range", "B2:C4", "--theme", "navy")
        color = com_retry(lambda: int(e_ws().Range("B2").Interior.Color))
        ok = (data or {}).get("status") == "ok" and color == 0x5D361B
        record(ctx, "E6", "테마 적용 (머리글 색)", ok, f"B2 배경색={color:#x}", data, code, raw, err)

    run_test(ctx, "E1", "move", e1)
    run_test(ctx, "E2", "merge 범위 없음", e2)
    run_test(ctx, "E3", "merge 대화상자", e3)
    run_test(ctx, "E4", "따옴표 시트명", e4)
    run_test(ctx, "E5", "수식 보존", e5)
    run_test(ctx, "E6", "테마", e6)


# ---------------------------------------------------------------------------
# F. PowerPoint
# ---------------------------------------------------------------------------
def test_F(ctx):
    def counts():
        prs = pp_prs("A.pptx")
        return [prs.Slides(i).Shapes.Count for i in range(1, prs.Slides.Count + 1)]

    def f1():
        before = com_retry(counts)
        data, code, raw, err = run_cli("ppt.py", "slide", "--action", "clear")
        after = com_retry(counts)
        ok = (data or {}).get("status") == "error" and before == after
        record(ctx, "F1", "슬라이드 번호 없는 clear -> error, 삭제 없음", ok, f"도형 수 {before}->{after}", data, code, raw, err)

    def f2():
        before = com_retry(counts)
        data, code, raw, err = run_cli("ppt.py", "slide", "--action", "clear", "--slide", "2")
        after = com_retry(counts)
        expected = list(before)
        if len(expected) >= 2:
            expected[1] = 0
        ok = (data or {}).get("status") == "ok" and after == expected
        record(ctx, "F2", "2번 슬라이드만 clear", ok, f"도형 수 {before}->{after}", data, code, raw, err)

    def f3():
        n0 = com_retry(lambda: pp_prs("A.pptx").Slides.Count)
        data, code, raw, err = run_cli("ppt.py", "write", "--slide", "new", "--title", "추진 전략", "--text", "- 안정성\\n- 확장성")
        n1 = com_retry(lambda: pp_prs("A.pptx").Slides.Count)
        cur = com_retry(lambda: pp_app().ActiveWindow.View.Slide.SlideIndex)
        ok = (data or {}).get("status") == "ok" and n1 == n0 + 1 and cur == n1
        manual = ask(ctx, "새 슬라이드에 '추진 전략' 제목과 글머리 2줄이 보기 좋게 들어갔나요?")
        record(ctx, "F3", "새 슬라이드 쓰기 + 고친 슬라이드로 이동", ok and manual is not False,
               f"슬라이드 {n0}->{n1}, 화면={cur}", data, code, raw, err)

    run_test(ctx, "F1", "clear 보호", f1)
    run_test(ctx, "F2", "clear --slide 2", f2)
    run_test(ctx, "F3", "새 슬라이드 쓰기", f3)


# ---------------------------------------------------------------------------
# G. Outlook
# ---------------------------------------------------------------------------
def test_G(ctx):
    me = ctx.args.self_email

    def g1():
        to = me or "office-test@example.com"
        data, code, raw, err = run_cli("outlook.py", "new", "--to", to, "--subject", "[office-cli 테스트] 표시 전용",
                                       "--body", "본문 테스트")
        ok = (data or {}).get("status") == "ok" and (data or {}).get("sent") is False
        manual = ask(ctx, "메일 작성 창이 떴나요? (서명을 등록해 두었다면 본문 아래 서명도 보여야 합니다. 확인 후 저장하지 않고 닫아 주세요)")
        record(ctx, "G1", "기본 동작: 작성 창만 표시, 미발송", ok and manual is not False, "", data, code, raw, err)

    def g2():
        data, code, raw, err = run_cli("outlook.py", "new", "--to", "zz_office_invalid_recipient_000",
                                       "--subject", "[office-cli 테스트] 잘못된 수신자", "--body", "발송되면 안 됨", "--send")
        ok = (data or {}).get("status") == "error" and (data or {}).get("sent") is False
        ask(ctx, "작성 창이 떴다면 저장하지 않고 닫아 주세요. 닫았나요?")
        record(ctx, "G2", "잘못된 수신자 + --send -> 발송 안 됨", ok,
               f"unresolved={(data or {}).get('unresolved_recipients')}", data, code, raw, err)

    def g3():
        if not me:
            record(ctx, "G3", "명시적 --send 실제 발송", None, "--self-email 미지정으로 건너뜀")
            return
        if not ctx.auto and ask(ctx, f"{me} 주소로 실제 테스트 메일을 1통 발송합니다. 진행할까요?") is not True:
            record(ctx, "G3", "명시적 --send 실제 발송", None, "사용자가 건너뜀")
            return
        data, code, raw, err = run_cli("outlook.py", "new", "--to", me, "--subject", "[office-cli 테스트] 발송 테스트",
                                       "--body", "자동 발송 테스트 메일입니다.", "--send")
        record(ctx, "G3", "명시적 --send 실제 발송", (data or {}).get("sent") is True, "", data, code, raw, err)

    def g4():
        data, code, raw, err = run_cli("outlook.py", "list", "--limit", "3")
        record(ctx, "G4", "받은편지함 최근 메일 목록", (data or {}).get("status") == "ok",
               f"{len((data or {}).get('mails', []))}통", data, code, raw, err)

    run_test(ctx, "G1", "표시 전용", g1)
    run_test(ctx, "G2", "잘못된 수신자", g2)
    run_test(ctx, "G3", "실제 발송", g3)
    run_test(ctx, "G4", "메일 목록", g4)


# ---------------------------------------------------------------------------
# H. 체이닝 (자주 쓰는 시나리오)
# ---------------------------------------------------------------------------
def test_H(ctx):
    def h1():
        activate_xl("A.xlsx")
        run_cli("excel.py", "sheet", "--action", "add", "--name", "사인")
        d1, c1, r1, e1 = run_cli("excel.py", "write", "--sheet", "사인", "--range", "A1:B101", "--header", "x,sin(x)",
                                 "--formula", "=(ROW()-2)*4*PI()/99", "=SIN(A2)")
        data, code, raw, err = run_cli("excel.py", "chart", "--sheet", "사인", "--range", "A1:B101", "--type", "line",
                                       "--title", "사인 2주기")
        ok = (d1 or {}).get("status") == "ok" and (data or {}).get("status") == "ok" and (data or {}).get("chart_name") == "사인 2주기"
        manual = ask(ctx, "엑셀 '사인' 시트에 사인 곡선 2주기 차트가 그려졌나요?")
        record(ctx, "H1", "사인 2주기 그래프 (쓰기 2회)", ok and manual is not False,
               f"chart_type={(data or {}).get('chart_type')}", data, code, raw, err)

    def h2():
        data, code, raw, err = run_cli("office.py", "transfer", "--probe")
        record(ctx, "H2", "차트 전달 방식 시험 (--probe)", (data or {}).get("status") == "ok",
               f"저장된 순서={(data or {}).get('saved_order')}", data, code, raw, err)

    def h3():
        wd_before = com_retry(lambda: wd_doc("A.docx").InlineShapes.Count)
        pp_before = com_retry(lambda: pp_prs("A.pptx").Slides.Count)
        data, code, raw, err = run_cli("office.py", "transfer", "--from", "excel", "--chart", "사인 2주기", "--to", "word,ppt")
        wd_after = com_retry(lambda: wd_doc("A.docx").InlineShapes.Count)
        pp_after = com_retry(lambda: pp_prs("A.pptx").Slides.Count)
        methods = {k: v.get("method") for k, v in ((data or {}).get("results") or {}).items()}
        ok = (data or {}).get("status") == "ok" and wd_after == wd_before + 1 and pp_after == pp_before + 1
        manual = ask(ctx, "워드 문서 끝과 PPT 새 슬라이드에 사인 그래프 그림이 들어갔나요?")
        record(ctx, "H3", "차트를 워드와 PPT에 그림으로", ok and manual is not False,
               f"사용 방식={methods}, 워드 그림 {wd_before}->{wd_after}, PPT 슬라이드 {pp_before}->{pp_after}",
               data, code, raw, err)

    def h4():
        d1, c1, r1, e1 = run_cli("office.py", "transfer", "--from", "word", "--what", "text", "--to", "excel")
        d2, c2, r2, e2 = run_cli("office.py", "transfer", "--from", "ppt", "--what", "text", "--to", "excel")
        ok = (d1 or {}).get("status") == "ok" and (d2 or {}).get("status") == "ok"
        record(ctx, "H4", "워드/PPT 내용을 엑셀 새 시트로", ok,
               f"시트: {(d1 or {}).get('sheet')}, {(d2 or {}).get('sheet')}", d2, c2, r2, e2)

    def h5():
        run_cli("reader.py", "read", "--file", ctx.path("B.docx"), "--format", "summary")
        data, code, raw, err = run_cli("reader.py", "search", "--query", "첫 문단", "--save")
        out = (data or {}).get("output_file", "")
        before = com_retry(lambda: len(wd_doc("A.docx").Content.Text))
        d2, c2, r2, e2 = run_cli("word.py", "write", "--file", out) if out else (None, None, "", "")
        after = com_retry(lambda: len(wd_doc("A.docx").Content.Text))
        ok = (data or {}).get("total_matches", 0) >= 1 and (d2 or {}).get("status") == "ok" and after > before
        record(ctx, "H5", "캐시 문서 검색 -> 워드에 넣기", ok,
               f"일치 {(data or {}).get('total_matches')}건, 워드 글자 {before}->{after}", d2, c2, r2, e2)

    run_test(ctx, "H1", "사인 그래프", h1)
    run_test(ctx, "H2", "probe", h2)
    run_test(ctx, "H3", "차트 전달", h3)
    run_test(ctx, "H4", "워드/PPT -> 엑셀", h4)
    run_test(ctx, "H5", "검색 -> 워드", h5)


# ---------------------------------------------------------------------------
# I. 기능 확대 (모두 자동 판정: 결과를 COM/파일로 다시 읽어 확인)
# ---------------------------------------------------------------------------
XL_TYPES = {"column": 51, "pie": 5, "scatter": -4169, "scatter_lines": 75}


def write_csv(path, rows):
    import csv
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        csv.writer(f).writerows(rows)
    return path


def make_png(path, w_px=600, h_px=300):
    """가로세로 2:1 PNG (그림 비율 유지 확인용)."""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    fig = plt.figure(figsize=(w_px / 100, h_px / 100), dpi=100)
    plt.plot([0, 1, 2, 3], [0, 1, 0, 1])
    fig.savefig(path, dpi=100)
    plt.close(fig)
    return path


def xl_sheet(name):
    return com_retry(lambda: xl_wb("A.xlsx").Worksheets(name))


def sheet_names():
    return com_retry(lambda: [s.Name for s in xl_wb("A.xlsx").Worksheets])


def chart_types(sheet):
    def _do():
        ws = xl_sheet(sheet)
        return {ws.ChartObjects(i).Name: int(ws.ChartObjects(i).Chart.ChartType)
                for i in range(1, ws.ChartObjects().Count + 1)}
    return com_retry(_do)


def ppt_last_slide():
    prs = pp_prs("A.pptx")
    return prs.Slides(prs.Slides.Count)


def slide_title(slide):
    for i in range(1, slide.Shapes.Count + 1):
        sp = slide.Shapes(i)
        if sp.Type == 14 and sp.HasTextFrame:
            return sp.TextFrame.TextRange.Text
    return ""


def inside_slide(sp):
    prs = pp_prs("A.pptx")
    sw, sh = prs.PageSetup.SlideWidth, prs.PageSetup.SlideHeight
    return sp.Left >= -0.5 and sp.Top >= -0.5 and sp.Left + sp.Width <= sw + 0.5 and sp.Top + sp.Height <= sh + 0.5


def test_I(ctx):
    csv_main = write_csv(ctx.path("i_data.csv"), [["품목", "수량", "금액"], ["사과", 3, 1500], ["배", 2, 4000]])
    csv_more = write_csv(ctx.path("i_more.csv"), [["품목", "수량", "금액"], ["포도", 5, 120]])
    png = make_png(ctx.path("i_img.png"))

    def ok_(d):
        return (d or {}).get("status") == "ok"

    # --- 엑셀 쓰기/읽기 ------------------------------------------------------
    def i1():
        activate_xl("A.xlsx", "Sheet1")
        d, c, r, e = run_cli("excel.py", "write", "--sheet", "I_DATA", "--new-sheet", "--range", "A1", "--file", csv_main)
        d2, c2, r2, e2 = run_cli("excel.py", "write", "--sheet", "I_DATA", "--range", "A1", "--file", csv_more, "--append")
        vals = com_retry(lambda: xl_sheet("I_DATA").Range("A1:C5").Value)
        expected = (("품목", "수량", "금액"), ("사과", 3.0, 1500.0), ("배", 2.0, 4000.0), ("포도", 5.0, 120.0), (None, None, None))
        record(ctx, "I1", "엑셀 CSV 쓰기 + --append(같은 머리글 생략), 숫자는 숫자로", ok_(d) and ok_(d2) and vals == expected,
               f"A1:C5={vals}", d2, c2, r2, e2)

    def i2():
        d, c, r, e = run_cli("excel.py", "write", "--sheet", "I_DATA", "--range", "F1",
                             "--values", '[["x","y"],[1,2],[2,4],[3,1],[4,3]]')
        vals = com_retry(lambda: xl_sheet("I_DATA").Range("F1:G5").Value)
        ok = ok_(d) and vals == (("x", "y"), (1.0, 2.0), (2.0, 4.0), (3.0, 1.0), (4.0, 3.0))
        record(ctx, "I2", "엑셀 --values JSON 쓰기", ok, f"F1:G5={vals}", d, c, r, e)

    def i3():
        d, c, r, e = run_cli("excel.py", "read", "--sheet", "I_DATA", "--search", "포도")
        text = json.dumps(d or {}, ensure_ascii=False)
        d2, c2, r2, e2 = run_cli("excel.py", "read", "--sheet", "I_DATA", "--range", "A1:C4", "--save")
        out = (d2 or {}).get("output_file", "")
        first = ""
        if out and os.path.exists(out):
            with open(out, encoding="utf-8-sig") as f:
                first = f.readline().strip()
        active = com_retry(lambda: xl_wb("A.xlsx").ActiveSheet.Name)
        ok = ok_(d) and (d or {}).get("total_matches", 0) >= 1 and "A4" in text and ok_(d2) and first == "품목,수량,금액"
        record(ctx, "I3", "엑셀 --search 셀 주소 + --save CSV", ok,
               f"일치={(d or {}).get('total_matches')}, CSV 첫 줄={first!r}, 읽은 뒤 활성 시트={active}", d2, c2, r2, e2)

    def i4():
        run_cli("excel.py", "write", "--sheet", "I_BIG", "--new-sheet", "--range", "A1:A300", "--formula", "=ROW()")
        d, c, r, e = run_cli("excel.py", "read", "--sheet", "I_BIG", "--range", "A1:A300")
        nc = (d or {}).get("next_commands") or []
        rows = len((d or {}).get("values") or [])
        ok = ok_(d) and d.get("truncated") is True and rows < 300 and nc and "A201" in nc[0].get("command", "")
        record(ctx, "I4", "엑셀 300행 읽기 -> 잘림 + 이어 읽기 명령", ok,
               f"표시 {rows}행, next={nc[0].get('command', '')[-60:] if nc else None}", d, c, r, e)

    # --- 엑셀 차트 -------------------------------------------------------------
    def i5():
        d1 = run_cli("excel.py", "chart", "--sheet", "I_DATA", "--range", "A1:C3", "--type", "column", "--title", "I 막대")[0]
        d2 = run_cli("excel.py", "chart", "--sheet", "I_DATA", "--range", "A1:B4", "--type", "pie", "--title", "I 원형")[0]
        d3, c, r, e = run_cli("excel.py", "chart", "--sheet", "I_DATA", "--range", "F1:G5", "--type", "scatter", "--title", "I 분산")
        types = chart_types("I_DATA")
        want = {"I 막대": XL_TYPES["column"], "I 원형": XL_TYPES["pie"], "I 분산": XL_TYPES["scatter"]}
        extra = ""
        if "사인" in sheet_names():
            sine = chart_types("사인")
            if "사인 2주기" in sine:
                types["사인 2주기"] = sine["사인 2주기"]
                want["사인 2주기"] = XL_TYPES["scatter_lines"]
        # 분산형은 첫 열이 x값이어야 함: 계열 1개 (x가 계열로 그려지면 2개)
        series = {"I 분산": com_retry(lambda: int(xl_sheet("I_DATA").ChartObjects("I 분산").Chart.SeriesCollection().Count))}
        if "사인 2주기" in want:
            series["사인 2주기"] = com_retry(lambda: int(xl_sheet("사인").ChartObjects("사인 2주기").Chart.SeriesCollection().Count))
        ok = all(ok_(x) for x in (d1, d2, d3)) and all(types.get(k) == v for k, v in want.items()) \
            and all(n == 1 for n in series.values())
        record(ctx, "I5", "차트 종류 (막대/원형/분산형, 분산형 선) + 분산형 x축", ok,
               f"실제 ChartType={types}, 기대={want}{extra}, 분산형 계열 수={series} (기대 1)", d3, c, r, e)

    def i6():
        d, c, r, e = run_cli("excel.py", "chart", "--action", "list", "--sheet", "I_DATA")
        listed = [x.get("name") for x in (d or {}).get("charts", [])]
        dd, dc, dr, de = run_cli("excel.py", "chart", "--action", "delete", "--sheet", "I_DATA", "--name", "I 원형")
        after = set(chart_types("I_DATA"))
        bad, bc, br, be = run_cli("excel.py", "chart", "--action", "delete", "--sheet", "I_DATA", "--name", "없는차트")
        after2 = set(chart_types("I_DATA"))
        ok = ok_(d) and {"I 막대", "I 원형", "I 분산"} <= set(listed) and ok_(dd) and after == {"I 막대", "I 분산"} \
            and (bad or {}).get("status") == "error" and after2 == after
        record(ctx, "I6", "차트 목록 / 이름으로 삭제 / 없는 이름은 삭제 안 함", ok,
               f"목록={listed}, 삭제 후={sorted(after)}, 없는이름 삭제={(bad or {}).get('status')} 후={sorted(after2)}",
               bad, bc, br, be)

    # --- 엑셀 서식/시트/범위 ---------------------------------------------------
    def i7():
        d1 = run_cli("excel.py", "style", "--sheet", "I_DATA", "--range", "C2:C4", "--highlight", ">=1000:pink")[0]
        d2 = run_cli("excel.py", "style", "--sheet", "I_DATA", "--range", "C2:C4", "--number-format", "#,##0")[0]
        rng = lambda: xl_sheet("I_DATA").Range("C2:C4")  # noqa: E731
        fc = com_retry(lambda: int(rng().FormatConditions.Count))
        nf = com_retry(lambda: str(xl_sheet("I_DATA").Range("C2").NumberFormat))
        d3, c, r, e = run_cli("excel.py", "style", "--sheet", "I_DATA", "--range", "C2:C4", "--clear")
        fc2 = com_retry(lambda: int(rng().FormatConditions.Count))
        ok = ok_(d1) and ok_(d2) and ok_(d3) and fc >= 1 and nf == "#,##0" and fc2 == 0
        record(ctx, "I7", "조건부 강조 / 숫자 서식 / 서식 지우기", ok,
               f"조건부 서식 {fc}개 -> 지운 뒤 {fc2}개, 숫자 서식={nf!r}", d3, c, r, e)

    def i8():
        run_cli("excel.py", "sheet", "--action", "add", "--name", "I_TMP")
        names0 = sheet_names()
        d1 = run_cli("excel.py", "sheet", "--action", "rename", "--name", "I_TMP", "--new-name", "I_TMP2")[0]
        d2 = run_cli("excel.py", "sheet", "--action", "copy", "--name", "I_TMP2", "--new-name", "I_TMP3")[0]
        names1 = sheet_names()
        d3 = run_cli("excel.py", "sheet", "--action", "delete", "--name", "I_TMP2")[0]
        d4, c, r, e = run_cli("excel.py", "sheet", "--action", "delete", "--name", "I_TMP3")
        names2 = sheet_names()
        ok = all(ok_(x) for x in (d1, d2, d3, d4)) and "I_TMP2" in names1 and "I_TMP3" in names1 \
            and "I_TMP" not in names1 and not {"I_TMP2", "I_TMP3"} & set(names2)
        pos = names0.index("I_TMP") + 1 if "I_TMP" in names0 else None
        record(ctx, "I8", "시트 이름변경 / 복사 / 삭제", ok,
               f"추가된 시트 위치 {pos}/{len(names0)}번째, 복사 후 {names1}", d4, c, r, e)

    def i9():
        d1 = run_cli("excel.py", "range", "--action", "copy", "--sheet", "I_DATA", "--range", "A1:C3", "--dest", "A20")[0]
        copied = com_retry(lambda: xl_sheet("I_DATA").Range("A20:C22").Value)
        src = com_retry(lambda: xl_sheet("I_DATA").Range("A1:C3").Value)
        d2 = run_cli("excel.py", "range", "--action", "hide", "--sheet", "I_DATA", "--range", "E1")[0]
        hidden = com_retry(lambda: bool(xl_sheet("I_DATA").Columns(5).Hidden))
        run_cli("excel.py", "range", "--action", "unhide", "--sheet", "I_DATA", "--range", "E1")
        shown = not com_retry(lambda: bool(xl_sheet("I_DATA").Columns(5).Hidden))
        d3 = run_cli("excel.py", "range", "--action", "freeze", "--sheet", "I_DATA", "--range", "A2")[0]
        frozen = com_retry(lambda: (bool(xl_app().ActiveWindow.FreezePanes), int(xl_app().ActiveWindow.SplitRow)))
        d4, c, r, e = run_cli("excel.py", "range", "--action", "unfreeze", "--sheet", "I_DATA")
        unfrozen = not com_retry(lambda: bool(xl_app().ActiveWindow.FreezePanes))
        ok = all(ok_(x) for x in (d1, d2, d3, d4)) and copied == src and hidden and shown and frozen == (True, 1) and unfrozen
        record(ctx, "I9", "범위 복사 / 열 숨김·해제 / 틀 고정·해제", ok,
               f"복사 일치={copied == src}, 숨김={hidden}, 해제={shown}, 고정(FreezePanes,SplitRow)={frozen}, 해제={unfrozen}",
               d4, c, r, e)

    # --- 워드 -------------------------------------------------------------------
    def i10():
        doc = wd_doc("A.docx")
        com_retry(lambda: doc.Activate())
        n0 = com_retry(lambda: doc.InlineShapes.Count)
        d, c, r, e = run_cli("word.py", "write", "--image", png)
        doc = wd_doc("A.docx")
        n1 = com_retry(lambda: doc.InlineShapes.Count)
        shp = com_retry(lambda: doc.InlineShapes(n1))
        ratio = com_retry(lambda: shp.Width / shp.Height)
        ok = ok_(d) and n1 == n0 + 1 and abs(ratio - 2.0) < 0.05
        record(ctx, "I10", "워드 그림 넣기 (비율 유지)", ok, f"그림 {n0}->{n1}, 가로/세로={ratio:.3f} (기대 2.0)", d, c, r, e)

    def i11():
        doc = wd_doc("A.docx")
        t0 = com_retry(lambda: doc.Tables.Count)
        d, c, r, e = run_cli("word.py", "write", "--file", csv_main)
        doc = wd_doc("A.docx")
        t1 = com_retry(lambda: doc.Tables.Count)
        shape = com_retry(lambda: (doc.Tables(t1).Rows.Count, doc.Tables(t1).Columns.Count)) if t1 else None
        cell = com_retry(lambda: doc.Tables(t1).Cell(2, 1).Range.Text.rstrip("\r\x07")) if t1 else None
        ok = ok_(d) and t1 == t0 + 1 and shape == (3, 3) and cell == "사과"
        record(ctx, "I11", "워드 CSV -> 표", ok, f"표 {t0}->{t1}, 크기={shape}, (2,1)={cell!r}", d, c, r, e)

    def i12():
        d1 = run_cli("word.py", "write", "--text", "앞에 넣은 문장", "--anchor", "두 번째 문단", "--pos", "before")[0]
        paras = com_retry(lambda: [p.Range.Text.strip() for p in wd_doc("A.docx").Paragraphs])
        idx_new = next((i for i, p in enumerate(paras) if p == "앞에 넣은 문장"), -1)
        idx_anchor = next((i for i, p in enumerate(paras) if p.startswith("두 번째 문단")), -1)
        d2, c, r, e = run_cli("word.py", "write", "--text", "바꾼 문장", "--anchor", "앞에 넣은 문장", "--pos", "replace")
        paras2 = com_retry(lambda: [p.Range.Text.strip() for p in wd_doc("A.docx").Paragraphs])
        ok = ok_(d1) and ok_(d2) and idx_new >= 0 and idx_new == idx_anchor - 1 \
            and "앞에 넣은 문장" not in paras2 and "바꾼 문장" in paras2
        record(ctx, "I12", "워드 앵커 앞에 넣기 / 앵커 단락 바꾸기", ok,
               f"새 단락 {idx_new}번, 앵커 {idx_anchor}번, 바꾼 뒤 '바꾼 문장' 있음={'바꾼 문장' in paras2}", d2, c, r, e)

    def i13():
        d1 = run_cli("word.py", "write", "--text", "소제목 테스트", "--heading", "2")[0]
        level = com_retry(lambda: next((int(p.OutlineLevel) for p in wd_doc("A.docx").Paragraphs
                                        if p.Range.Text.strip() == "소제목 테스트"), None))
        d2 = run_cli("word.py", "write", "--page-break")[0]
        has_break = "\x0c" in com_retry(lambda: wd_doc("A.docx").Content.Text)
        d3, c, r, e = run_cli("word.py", "write", "--toc")
        toc = com_retry(lambda: wd_doc("A.docx").TablesOfContents.Count)
        ok = ok_(d1) and ok_(d2) and ok_(d3) and level == 2 and has_break and toc >= 1
        record(ctx, "I13", "워드 제목 수준 / 쪽 나누기 / 목차", ok,
               f"개요 수준={level}, 쪽 나누기={has_break}, 목차 {toc}개", d3, c, r, e)

    def i14():
        doc = wd_doc("A.docx")
        com_retry(lambda: setattr(doc.Sections(1).Footers(1).Range, "Text", "바닥글 원본표시"))
        run_cli("word.py", "write", "--text", "굵게바꿀말 사과 포도")
        d1 = run_cli("word.py", "replace", "--find", "굵게바꿀말", "--replace", "굵은말", "--bold")[0]
        d2 = run_cli("word.py", "replace", "--map", '{"사과 포도":"배 귤"}')[0]
        d3, c, r, e = run_cli("word.py", "replace", "--find", "원본표시", "--replace", "변경됨")
        doc = wd_doc("A.docx")
        text = com_retry(lambda: doc.Content.Text)

        def bold_of(word):
            rng = doc.Content
            f = rng.Find
            f.ClearFormatting()
            return bool(f.Execute(FindText=word, MatchCase=True)) and int(rng.Bold) == -1
        bold = com_retry(lambda: bold_of("굵은말"))
        footer = com_retry(lambda: doc.Sections(1).Footers(1).Range.Text.strip())
        ok = ok_(d1) and ok_(d2) and ok_(d3) and "굵은말 배 귤" in text and bold and footer == "바닥글 변경됨"
        record(ctx, "I14", "워드 치환 (굵게 / --map / 바닥글 안)", ok,
               f"굵게={bold}, map 결과 있음={'굵은말 배 귤' in text}, 바닥글={footer!r}", d3, c, r, e)

    # --- PPT ----------------------------------------------------------------
    def i15():
        n0 = com_retry(lambda: pp_prs("A.pptx").Slides.Count)
        d, c, r, e = run_cli("ppt.py", "write", "--slide", "new", "--title", "그림 슬라이드", "--image", png)
        s = com_retry(ppt_last_slide)
        pics = com_retry(lambda: [s.Shapes(i) for i in range(1, s.Shapes.Count + 1) if s.Shapes(i).Type == 13])
        info = com_retry(lambda: [(p.Width / p.Height, inside_slide(p)) for p in pics])
        ok = ok_(d) and com_retry(lambda: pp_prs("A.pptx").Slides.Count) == n0 + 1 and len(info) == 1 \
            and abs(info[0][0] - 2.0) < 0.05 and info[0][1]
        record(ctx, "I15", "PPT 그림 슬라이드 (비율 유지, 슬라이드 안)", ok, f"그림(비율, 안쪽)={info}", d, c, r, e)

    def i16():
        d, c, r, e = run_cli("ppt.py", "write", "--slide", "new", "--title", "카드", "--cards", "속도:빠름|품질:높음|비용:낮음")
        s = com_retry(ppt_last_slide)
        cards = com_retry(lambda: [(s.Shapes(i).TextFrame.TextRange.Text.split("\r")[0], inside_slide(s.Shapes(i)))
                                   for i in range(1, s.Shapes.Count + 1) if s.Shapes(i).Type == 1])
        ok = ok_(d) and [x[0] for x in cards] == ["속도", "품질", "비용"] and all(x[1] for x in cards)
        record(ctx, "I16", "PPT 카드 3개 (슬라이드 안)", ok, f"카드={cards}", d, c, r, e)

    def i17():
        md = ctx.path("i_slides.md")
        with open(md, "w", encoding="utf-8") as f:
            f.write("# 첫째 장\n\n- 가\n- 나\n\n# 둘째 장\n\n본문 문장\n")
        n0 = com_retry(lambda: pp_prs("A.pptx").Slides.Count)
        d, c, r, e = run_cli("ppt.py", "write", "--file", md)
        prs = pp_prs("A.pptx")
        n1 = com_retry(lambda: prs.Slides.Count)
        titles = com_retry(lambda: [slide_title(prs.Slides(i)) for i in range(n0 + 1, n1 + 1)])
        ok = ok_(d) and titles == ["첫째 장", "둘째 장"]
        record(ctx, "I17", "PPT 마크다운 -> 제목마다 새 슬라이드", ok, f"새 슬라이드 제목={titles}", d, c, r, e)

    def i18():
        rows = [["품목", "값"]] + [[f"품목{i:02d}", i] for i in range(1, 31)]
        long_csv = write_csv(ctx.path("i_long.csv"), rows)
        n0 = com_retry(lambda: pp_prs("A.pptx").Slides.Count)
        d, c, r, e = run_cli("ppt.py", "write", "--slide", "new", "--file", long_csv)
        prs = pp_prs("A.pptx")
        n1 = com_retry(lambda: prs.Slides.Count)

        def tables():
            out = []
            for i in range(n0 + 1, n1 + 1):
                s = prs.Slides(i)
                for j in range(1, s.Shapes.Count + 1):
                    sp = s.Shapes(j)
                    if sp.HasTable:
                        t = sp.Table
                        out.append((t.Cell(1, 1).Shape.TextFrame.TextRange.Text, t.Rows.Count - 1, inside_slide(sp)))
            return out
        tbls = com_retry(tables)
        ok = ok_(d) and len(tbls) >= 2 and all(t[0] == "품목" and t[2] for t in tbls) and sum(t[1] for t in tbls) == 30
        record(ctx, "I18", "PPT 긴 CSV 표 -> 여러 슬라이드, 머리글 반복", ok,
               f"(머리글, 데이터 행, 안쪽)={tbls}", d, c, r, e)

    def i19():
        prs = pp_prs("A.pptx")
        n0 = com_retry(lambda: prs.Slides.Count)
        t1 = com_retry(lambda: slide_title(prs.Slides(1)))
        d1 = run_cli("ppt.py", "slide", "--action", "duplicate", "--slide", "1")[0]
        n1 = com_retry(lambda: pp_prs("A.pptx").Slides.Count)
        t2 = com_retry(lambda: slide_title(pp_prs("A.pptx").Slides(2)))
        last_title = com_retry(lambda: slide_title(ppt_last_slide()))
        d2 = run_cli("ppt.py", "slide", "--action", "move", "--slide", str(n1), "--to", "1")[0]
        moved = com_retry(lambda: slide_title(pp_prs("A.pptx").Slides(1)))
        run_cli("ppt.py", "slide", "--action", "move", "--slide", "1", "--to", str(n1))
        d3 = run_cli("ppt.py", "slide", "--action", "delete", "--slide", "2")[0]
        n2 = com_retry(lambda: pp_prs("A.pptx").Slides.Count)
        d4, c, r, e = run_cli("ppt.py", "slide", "--action", "list")
        listed = len((d4 or {}).get("slides", []))
        ok = all(ok_(x) for x in (d1, d2, d3, d4)) and n1 == n0 + 1 and t2 == t1 and moved == last_title \
            and n2 == n0 and listed == n0
        record(ctx, "I19", "PPT 슬라이드 복제 / 이동 / 삭제 / 목록", ok,
               f"장수 {n0}->{n1}->{n2}, 복제 제목 일치={t2 == t1}, 이동 후 1번={moved!r}, 목록 {listed}장", d4, c, r, e)

    def i20():
        prs = pp_prs("A.pptx")

        def make_group():
            s = prs.Slides.Add(prs.Slides.Count + 1, 12)
            a = s.Shapes.AddTextbox(1, 50, 50, 200, 40)
            a.TextFrame.TextRange.Text = "그룹안글자"
            b = s.Shapes.AddTextbox(1, 50, 120, 200, 40)
            b.TextFrame.TextRange.Text = "옆 상자"
            s.Shapes.Range([a.Name, b.Name]).Group()
            return prs.Slides.Count
        gslide = com_retry(make_group)
        d, c, r, e = run_cli("ppt.py", "replace", "--map", '{"품목07":"변경07","그룹안글자":"그룹바뀜"}')
        group_text = com_retry(lambda: [prs.Slides(gslide).Shapes(1).GroupItems(i).TextFrame.TextRange.Text
                                        for i in range(1, prs.Slides(gslide).Shapes(1).GroupItems.Count + 1)])

        def table_hit():
            for i in range(1, prs.Slides.Count + 1):
                s = prs.Slides(i)
                for j in range(1, s.Shapes.Count + 1):
                    sp = s.Shapes(j)
                    if sp.HasTable:
                        t = sp.Table
                        for rr in range(1, t.Rows.Count + 1):
                            if t.Cell(rr, 1).Shape.TextFrame.TextRange.Text == "변경07":
                                return True
            return False
        hit = com_retry(table_hit)
        ok = ok_(d) and hit and "그룹바뀜" in group_text
        record(ctx, "I20", "PPT 치환 (표 셀, 그룹 도형 안)", ok, f"표 셀 변경={hit}, 그룹 글자={group_text}", d, c, r, e)

    def i21():
        prs = pp_prs("A.pptx")
        full = com_retry(lambda: prs.FullName)
        n = com_retry(lambda: prs.Slides.Count)
        d, c, r, e = run_cli("ppt.py", "export", "--save", "i_copy.pptx")
        out = (d or {}).get("output_file", "")
        count = None
        if out and os.path.exists(out):
            import pptx
            count = len(pptx.Presentation(out).slides)
        full_after = com_retry(lambda: pp_prs("A.pptx").FullName)
        ok = ok_(d) and count == n and full_after == full
        record(ctx, "I21", "PPT .pptx 사본 저장 (원본 경로 그대로)", ok,
               f"사본 장수={count} (원본 {n}), 원본 경로 유지={full_after == full}", d, c, r, e)

    # --- 앱 간 전달 -------------------------------------------------------------
    def i22():
        activate_xl("A.xlsx", "I_DATA")
        doc = wd_doc("A.docx")
        t0 = com_retry(lambda: doc.Tables.Count)
        n0 = com_retry(lambda: pp_prs("A.pptx").Slides.Count)
        d, c, r, e = run_cli("office.py", "transfer", "--from", "excel", "--sheet", "I_DATA", "--range", "A1:C4", "--to", "word,ppt")
        doc = wd_doc("A.docx")
        t1 = com_retry(lambda: doc.Tables.Count)
        wcell = com_retry(lambda: doc.Tables(t1).Cell(4, 1).Range.Text.rstrip("\r\x07")) if t1 > t0 else None

        def ppt_cell():
            prs = pp_prs("A.pptx")
            for i in range(n0, prs.Slides.Count + 1):
                s = prs.Slides(i)
                for j in range(1, s.Shapes.Count + 1):
                    if s.Shapes(j).HasTable:
                        t = s.Shapes(j).Table
                        if t.Rows.Count >= 4 and t.Cell(1, 1).Shape.TextFrame.TextRange.Text == "품목":
                            return t.Cell(4, 1).Shape.TextFrame.TextRange.Text
            return None
        pcell = com_retry(ppt_cell)
        ok = ok_(d) and t1 == t0 + 1 and wcell == "포도" and pcell == "포도"
        record(ctx, "I22", "엑셀 범위 -> 워드/PPT 표", ok, f"워드 표 {t0}->{t1} (4,1)={wcell!r}, PPT 표 (4,1)={pcell!r}",
               d, c, r, e)

    def i23():
        before = set(sheet_names())
        d, c, r, e = run_cli("office.py", "transfer", "--from", "word", "--what", "table", "--table", "1", "--to", "excel")
        new = [s for s in sheet_names() if s not in before]
        vals = com_retry(lambda: xl_sheet(new[0]).Range("A1:C2").Value) if new else None
        ok = ok_(d) and len(new) == 1 and vals == (("품목", "수량", "금액"), ("사과", 3.0, 1500.0))
        record(ctx, "I23", "워드 표 -> 엑셀 새 시트", ok, f"새 시트={new}, A1:C2={vals}", d, c, r, e)

    def i24():
        doc = wd_doc("A.docx")
        n0 = com_retry(lambda: doc.InlineShapes.Count)
        d, c, r, e = run_cli("office.py", "transfer", "--from", "excel", "--chart", "I 막대", "--to", "word", "--method", "redraw")
        n1 = com_retry(lambda: wd_doc("A.docx").InlineShapes.Count)
        method = (((d or {}).get("results") or {}).get("word") or {}).get("method")
        ok = ok_(d) and method == "redraw" and n1 == n0 + 1
        record(ctx, "I24", "차트 전달 redraw 방식 강제 (DRM 대비 경로)", ok, f"방식={method}, 워드 그림 {n0}->{n1}", d, c, r, e)

    # --- 되돌리기 / 창 -----------------------------------------------------------
    def i25():
        activate_xl("A.xlsx", "Sheet1")
        run_cli("excel.py", "sheet", "--action", "add", "--name", "I_UNDO")
        run_cli("excel.py", "write", "--sheet", "I_UNDO", "--range", "A1", "--value", "v1")
        run_cli("excel.py", "write", "--sheet", "I_UNDO", "--range", "A1", "--value", "v2")
        dl = run_cli("office.py", "undo", "--app", "excel", "--list")[0]
        snaps = (dl or {}).get("snapshots") or []
        d, c, r, e = run_cli("office.py", "undo", "--app", "excel", "--pick", "2")
        wait_for(lambda: xl_wb("A.xlsx"), timeout=20)
        names = sheet_names()
        a1 = com_retry(lambda: xl_sheet("I_UNDO").Range("A1").Value) if "I_UNDO" in names else "(시트 없음)"
        ok = ok_(dl) and len(snaps) >= 2 and snaps[0].get("pick") == 1 and ok_(d) and a1 is None
        record(ctx, "I25", "undo --list / --pick 2 (두 단계 전으로)", ok,
               f"지점 {len(snaps)}개, pick 2 후 I_UNDO!A1={a1!r}", d, c, r, e)

    def i26():
        b = ctx.path("B.xlsx")
        d1 = run_cli("excel.py", "write", "--target", b, "--range", "D5", "--value", "닫힌값")[0]
        import openpyxl
        v1 = openpyxl.load_workbook(b).active["D5"].value
        d, c, r, e = run_cli("office.py", "undo", "--app", "excel")
        v2 = openpyxl.load_workbook(b).active["D5"].value
        ok = ok_(d1) and v1 == "닫힌값" and ok_(d) and v2 is None
        record(ctx, "I26", "닫힌 파일 쓰기 후 undo", ok,
               f"쓰기 후 D5={v1!r}, undo 후 D5={v2!r}, restored_as={(d or {}).get('restored_as')}", d, c, r, e)

    def i27():
        app = xl_app()
        com_retry(lambda: setattr(app, "WindowState", -4137))       # xlMaximized
        com_retry(lambda: wd_doc("A.docx").Activate())
        d, c, r, e = run_cli("office.py", "focus", "--app", "excel")
        state = com_retry(lambda: int(xl_app().WindowState))
        fg = ""
        try:
            import win32gui
            fg = win32gui.GetClassName(win32gui.GetForegroundWindow())
        except Exception:
            pass
        ok = ok_(d) and state == -4137
        record(ctx, "I27", "창 전환 시 최대화 유지", ok, f"WindowState={state} (최대화 -4137), 앞 창 클래스={fg}", d, c, r, e)

    for tid, title, fn in [
        ("I1", "엑셀 CSV/append", i1), ("I2", "엑셀 --values", i2), ("I3", "엑셀 search/save", i3), ("I4", "엑셀 잘림", i4),
        ("I5", "차트 종류", i5), ("I6", "차트 목록/삭제", i6), ("I7", "엑셀 서식", i7), ("I8", "시트 관리", i8),
        ("I9", "범위 작업", i9), ("I10", "워드 그림", i10), ("I11", "워드 표", i11), ("I12", "워드 앵커", i12),
        ("I13", "워드 제목/쪽/목차", i13), ("I14", "워드 치환", i14), ("I15", "PPT 그림", i15), ("I16", "PPT 카드", i16),
        ("I17", "PPT 마크다운", i17), ("I18", "PPT 긴 표", i18), ("I19", "PPT 슬라이드 관리", i19), ("I20", "PPT 치환", i20),
        ("I21", "PPT 사본", i21), ("I22", "범위 -> 표", i22), ("I23", "워드 표 -> 엑셀", i23), ("I24", "redraw", i24),
        ("I25", "undo pick", i25), ("I26", "닫힌 파일 undo", i26), ("I27", "창 최대화", i27),
    ]:
        run_test(ctx, tid, title, fn)


# ---------------------------------------------------------------------------
# J. 배치 점검 (info --detail): 도구가 만든 결과가 슬라이드/쪽 밖으로 나가거나 겹치거나 데이터를 가리지 않는지
# ---------------------------------------------------------------------------
def test_J(ctx):
    def check(tid, title, module):
        d, c, r, e = run_cli(module, "info", "--detail")
        if module == "ppt.py":
            warns = [f"슬라이드 {s['no']}: {w}" for s in (d or {}).get("slides", []) for w in s.get("warnings", [])]
        elif module == "excel.py":
            warns = [f"{s['name']}: {w}" for s in (d or {}).get("sheets", []) for w in s.get("warnings", [])]
        else:
            warns = list((d or {}).get("warnings", []))
        ok = (d or {}).get("status") == "ok" and (d or {}).get("layout_warnings") == len(warns) and not warns
        record(ctx, tid, title, ok, f"경고 {len(warns)}건" + (": " + "; ".join(warns[:8]) if warns else ""), d, c, r, e)

    run_test(ctx, "J1", "PPT 배치", lambda: check("J1", "PPT 전체 슬라이드 배치 경고 없음 (info --detail)", "ppt.py"))
    run_test(ctx, "J2", "엑셀 차트 배치", lambda: check("J2", "엑셀 차트가 데이터를 가리지 않음 (info --detail)", "excel.py"))
    run_test(ctx, "J3", "워드 그림 크기", lambda: check("J3", "워드 그림이 쪽 너비 안 (info --detail)", "word.py"))


# ---------------------------------------------------------------------------
# K. 검토 기능: 워드 목차·장절·쪽·단락 읽기, 검토 메모, 변경 추적, PPT 노트
# ---------------------------------------------------------------------------
def test_K(ctx):
    tag = time.strftime("%H%M%S")           # 다시 실행해도 기준 글이 겹치지 않게
    docx = ctx.path("A.docx")

    def ok_(d):
        return (d or {}).get("status") == "ok"

    def k1():
        md = f"# K장 {tag}\\nK 본문 {tag}\\n## K절 {tag}\\nK 세부 {tag}\\n# K다음장 {tag}\\nK 다음 본문 {tag}"
        run_cli("word.py", "write", "--target", docx, "--text", md)
        d, c, r, e = run_cli("word.py", "outline", "--target", docx, "--start", "1")
        heads = {h["text"]: h for h in (d or {}).get("headings", [])}
        while (d or {}).get("next_commands"):
            nxt = d["next_commands"][0]["command"].split("--start")[-1].strip().strip('"')
            d = run_cli("word.py", "outline", "--target", docx, "--start", nxt)[0]
            heads.update({h["text"]: h for h in (d or {}).get("headings", [])})
        h1, h2 = heads.get(f"K장 {tag}"), heads.get(f"K절 {tag}")
        d2, c, r, e = run_cli("word.py", "read", "--target", docx, "--section", f"K장 {tag}")
        got = [p["text"] for p in (d2 or {}).get("paragraphs", [])]
        want = [f"K장 {tag}", f"K 본문 {tag}", f"K절 {tag}", f"K 세부 {tag}"]
        ok = ok_(d2) and h1 and h2 and h1["level"] == 1 and h2["level"] == 2 and isinstance(h1.get("page"), int) and got == want
        record(ctx, "K1", "워드 outline(수준·쪽) + read --section(하위 절 포함, 다음 장 전까지)", ok,
               f"제목={h1}, {h2}, 장절 단락={got}", d2, c, r, e)

    def k2():
        run_cli("word.py", "write", "--target", docx, "--page-break")
        run_cli("word.py", "write", "--target", docx, "--text", f"K 쪽 확인 {tag}")
        doc = wd_doc("A.docx")

        def page_of():
            rng = doc.Content
            f = rng.Find
            f.ClearFormatting()
            return int(rng.Information(3)) if f.Execute(FindText=f"K 쪽 확인 {tag}", MatchCase=True) else None
        page = com_retry(page_of)
        d, c, r, e = run_cli("word.py", "read", "--target", docx, "--pages", str(page))
        hit = next((p for p in (d or {}).get("paragraphs", []) if p["text"] == f"K 쪽 확인 {tag}"), None)
        d2 = run_cli("word.py", "read", "--target", docx, "--paragraphs", str(hit["no"]))[0] if hit else None
        same = [p["text"] for p in (d2 or {}).get("paragraphs", [])] == [f"K 쪽 확인 {tag}"]
        ok = ok_(d) and page and hit and hit.get("page") == page and same
        record(ctx, "K2", "워드 read --pages(쪽 번호) + --paragraphs(같은 단락 번호)", ok,
               f"Word 쪽={page}, read 결과={hit}, --paragraphs 일치={same}", d, c, r, e)

    def k3():
        doc = wd_doc("A.docx")
        n0 = com_retry(lambda: doc.Comments.Count)
        d, c, r, e = run_cli("word.py", "comment", "--target", docx, "--action", "add",
                             "--anchor", f"K 본문 {tag}", "--text", f"K 검토 메모 {tag}")
        doc = wd_doc("A.docx")
        n1 = com_retry(lambda: doc.Comments.Count)
        scopes = com_retry(lambda: [(doc.Comments(i).Scope.Text, doc.Comments(i).Range.Text) for i in range(1, n1 + 1)])
        d2 = run_cli("word.py", "comment", "--target", docx, "--start", "1")[0]
        listed = any(it.get("text") == f"K 검토 메모 {tag}" for it in (d2 or {}).get("items", []))
        ok = ok_(d) and (d or {}).get("verified") is True and n1 == n0 + 1 \
            and (f"K 본문 {tag}", f"K 검토 메모 {tag}") in scopes and (listed or (d2 or {}).get("truncated"))
        record(ctx, "K3", "워드 메모 달기(기준 글에) + 메모 목록", ok,
               f"메모 {n0}->{n1}, 목록에 있음={listed}", d, c, r, e)

    def k4():
        doc = wd_doc("A.docx")

        def tracked_insert():
            doc.TrackRevisions = True
            try:
                doc.Content.InsertAfter(f"K 변경 문장 {tag}")
            finally:
                doc.TrackRevisions = False
        com_retry(tracked_insert)
        d, c, r, e = run_cli("word.py", "revisions", "--target", docx)
        items = (d or {}).get("items", [])
        hit = next((it for it in items if f"K 변경 문장 {tag}" in it.get("text", "")), None)
        ok = ok_(d) and (d or {}).get("count", 0) >= 1 and hit and hit.get("type") == "삽입"
        record(ctx, "K4", "워드 변경 추적 목록 (삽입 1건 이상, 종류·작성자·쪽)", ok, f"count={(d or {}).get('count')}, 항목={hit}",
               d, c, r, e)
        com_retry(lambda: wd_doc("A.docx").Revisions.AcceptAll())

    def k5():
        d, c, r, e = run_cli("ppt.py", "notes", "--target", "A.pptx", "--slide", "1", "--text", f"K 노트 {tag}\\n둘째 줄",
                             "--overwrite")       # 다시 실행하면 앞 실행의 노트가 있음
        d2 = run_cli("ppt.py", "notes", "--target", "A.pptx", "--slide", "1", "--text", "덧붙인 줄", "--append")[0]

        def notes_text():
            holders = pp_prs("A.pptx").Slides(1).NotesPage.Shapes.Placeholders
            return next(str(holders(i).TextFrame.TextRange.Text) for i in range(1, holders.Count + 1)
                        if holders(i).PlaceholderFormat.Type == 2)
        text = com_retry(notes_text)
        d3 = run_cli("ppt.py", "notes", "--target", "A.pptx", "--slides", "1")[0]
        listed = ((d3 or {}).get("notes") or [{}])[0].get("notes", "")
        want = f"K 노트 {tag}\r둘째 줄\r덧붙인 줄"
        ok = ok_(d) and ok_(d2) and (d2 or {}).get("verified") is True and text == want and listed == want.replace("\r", "\n")
        record(ctx, "K5", "PPT 노트 쓰기 + --append + 읽기", ok, f"노트={text!r}, 읽기={listed!r}", d2, c, r, e)

    for tid, title, fn in [("K1", "목차/장절", k1), ("K2", "쪽/단락", k2), ("K3", "메모", k3), ("K4", "변경 추적", k4),
                           ("K5", "PPT 노트", k5)]:
        run_test(ctx, tid, title, fn)


# ---------------------------------------------------------------------------
# 마무리: 확인용 그림 저장, 테스트 문서 닫기
# ---------------------------------------------------------------------------
def save_shots(ctx):
    """눈으로 볼 결과를 그림으로 남깁니다: A.pptx 모든 슬라이드, A.xlsx 모든 차트."""
    folder = ctx.path("shots")
    shutil.rmtree(folder, ignore_errors=True)
    os.makedirs(folder, exist_ok=True)
    saved = 0
    try:
        prs = pp_prs("A.pptx")
        for i in range(1, prs.Slides.Count + 1):
            com_retry(lambda: prs.Slides(i).Export(os.path.join(folder, f"slide_{i:02d}.png"), "PNG", 960, 720))
            saved += 1
    except Exception as e:
        log(f"[WARN] 슬라이드 그림 저장 실패: {describe_com_error(e)}")
    try:
        wb = xl_wb("A.xlsx")
        for ws in wb.Worksheets:
            for k in range(1, ws.ChartObjects().Count + 1):
                co = ws.ChartObjects(k)
                safe = "".join(ch if ch.isalnum() else "_" for ch in f"{ws.Name}_{co.Name}")
                com_retry(lambda: co.Chart.Export(os.path.join(folder, f"chart_{safe}.png")))
                saved += 1
    except Exception as e:
        log(f"[WARN] 차트 그림 저장 실패: {describe_com_error(e)}")
    return folder, saved


# 별도 프로세스에서 실행합니다: 러너가 COM 참조를 쥐고 있으면 Quit해도 PowerPoint/Excel이 남습니다.
# 끝내도 되는 프로세스의 판단은 tools/win_test_rules.py(테스트 있음)에 있습니다.
#  - 문서가 남아 있는 프로세스(COM 문서 수 > 0, 제목 있는 창, 창 제목에 A.xlsx 등)는 Quit도 WM_CLOSE도 하지 않음
#  - allowed_pids를 주면(A5) 그 PID만, 아니면 이 실행이 시작된 뒤 생긴 프로세스만
CLEANUP_SCRIPT = r"""
import datetime, json, os, sys, time
root = sys.argv[1]
opts = json.loads(sys.argv[2])
sys.path.insert(0, root)
sys.path.insert(0, sys.argv[3] if len(sys.argv) > 3 else os.path.join(root, "tools"))   # win_test_rules.py 위치
import _office
import win_test_rules as rules
try:
    import win32api, win32con, win32gui, win32process
except ImportError:
    win32gui = None
test_dir = os.path.normcase(opts["test_dir"])
keep_untitled = set(opts.get("keep_untitled") or [])   # 실행 전부터 있던 제목 없는 문서는 사용자 것
only_name = opts.get("only_name") or ""                 # 지정하면 그 이름의 문서만 (빈 문자열: 테스트 문서 전체)
allowed = set(opts["allowed_pids"]) if opts.get("allowed_pids") is not None else None
run_start = (datetime.datetime.fromtimestamp(float(opts["run_start"]), datetime.timezone.utc)
             if opts.get("run_start") else None)
protect_names = opts.get("protect_names") or []
closed, actions, doc_pids = [], [], set()
_starts = {}

def started(pid):
    if pid is None or win32gui is None:
        return None
    if pid not in _starts:
        try:
            h = win32api.OpenProcess(0x1000, False, pid)        # PROCESS_QUERY_LIMITED_INFORMATION
            try:
                _starts[pid] = win32process.GetProcessTimes(h)["CreationTime"]
            finally:
                win32api.CloseHandle(h)
        except Exception:
            _starts[pid] = None
    return _starts[pid]

def app_pid(app, kind):
    if win32gui is None:
        return None
    try:
        if kind == "excel":
            hwnd = app.Hwnd
        elif kind == "ppt":
            hwnd = app.HWND
        else:
            hwnd = app.ActiveWindow.Hwnd                         # Word 2013+, 문서가 있을 때만
        return win32process.GetWindowThreadProcessId(int(hwnd))[1]
    except Exception:
        return None

def frame_windows():
    found = []
    if win32gui is None:
        return found
    classes = _office.FRAME_CLASS["excel"] + _office.FRAME_CLASS["ppt"]    # PowerPoint 2007은 PP12FrameClass
    def cb(hwnd, _):
        try:
            cls = win32gui.GetClassName(hwnd)
            if cls in classes:
                found.append({"hwnd": hwnd, "pid": win32process.GetWindowThreadProcessId(hwnd)[1], "cls": cls,
                              "title": win32gui.GetWindowText(hwnd), "visible": bool(win32gui.IsWindowVisible(hwnd))})
        except Exception:
            pass
        return True
    win32gui.EnumWindows(cb, None)
    return found

seen = set()
quit_later = []
for kind, coll in (("excel", "Workbooks"), ("word", "Documents"), ("ppt", "Presentations")):
    for app in _office.find_apps(kind):
        pid = app_pid(app, kind)
        if pid is not None and (kind, pid) in seen:
            actions.append(f"{kind} pid {pid}: 같은 인스턴스가 두 번 나와 건너뜀")
            continue
        seen.add((kind, pid))
        try:
            docs = getattr(app, coll)
            for i in range(docs.Count, 0, -1):
                d = docs(i)
                name, path = str(d.Name), os.path.normcase(str(d.Path or ""))
                ours = name == only_name if only_name else (path == test_dir or (not path and name not in keep_untitled))
                if ours:
                    closed.append(name)
                    d.Close(False) if kind == "excel" else (d.Close(0) if kind == "word" else d.Close())
                    actions.append(f"{kind} pid {pid}: 문서 닫음 {name}")
            left = docs.Count
            if left:
                if pid is not None:
                    doc_pids.add(pid)
                actions.append(f"{kind} pid {pid}: 문서 {left}개 남음 -> 보호 (종료하지 않음)")
            else:
                quit_later.append((kind, pid, app))
        except Exception as e:
            sys.stderr.write(f"[WARN] {kind}: {e}\n")
            actions.append(f"{kind} pid {pid}: 확인 실패 {type(e).__name__}: {e}")
windows = frame_windows()
protected = rules.protected_pids(windows, doc_pids, protect_names)
for kind, pid, app in quit_later:
    ok, why = rules.may_quit_app(pid, 0, protected, allowed, run_start, started(pid))
    if ok:
        try:
            app.Quit()
            actions.append(f"{kind} pid {pid}: Quit ({why})")
        except Exception as e:
            actions.append(f"{kind} pid {pid}: Quit 실패 {type(e).__name__}: {e}")
    else:
        actions.append(f"{kind} pid {pid}: 문서 0개지만 Quit 안 함 ({why})")
app = docs = d = quit_later = None                      # COM 참조를 놓아야 Quit이 실제로 반영됨
# Quit을 받고도 남는 인스턴스(예: excel.exe /x로 띄운 두 번째 Excel)는 창 닫기(WM_CLOSE)로 끝냅니다.
# 대상: 숨은 창 또는 제목이 빈 창이면서 rules.may_end_process를 통과한 프로세스의 창만.
if win32gui is not None and (run_start is not None or allowed is not None):
    time.sleep(2)
    windows = frame_windows()
    protected = rules.protected_pids(windows, doc_pids, protect_names)
    targets, skipped = rules.select_close_targets(windows, protected, allowed, run_start,
                                                  {w["pid"]: started(w["pid"]) for w in windows})
    for w in targets:
        try:
            win32gui.PostMessage(w["hwnd"], win32con.WM_CLOSE, 0, 0)
            closed.append(f"(빈 {w['cls']} pid {w['pid']})")
            actions.append(f"WM_CLOSE {w['cls']} pid {w['pid']} hwnd {w['hwnd']} 제목='{w['title']}' 표시={w['visible']}")
        except Exception as e:
            sys.stderr.write(f"[WARN] window close: {e}\n")
            actions.append(f"WM_CLOSE 실패 pid {w['pid']}: {e}")
    for w, why in skipped:
        actions.append(f"WM_CLOSE 안 함 {w['cls']} pid {w['pid']} 제목='{w['title']}' 표시={w['visible']}: {why}")
    if protected:
        actions.append(f"보호한 PID: {sorted(protected)}")
print(json.dumps({"closed": closed, "actions": actions}, ensure_ascii=False))
"""


def untitled_documents():
    """지금 열려 있는 제목 없는(저장 안 한) 문서 이름들."""
    if ROOT not in sys.path:
        sys.path.insert(0, ROOT)
    import _office
    names = []
    for kind, coll in (("excel", "Workbooks"), ("word", "Documents"), ("ppt", "Presentations")):
        for app in _office.find_apps(kind):
            try:
                docs = getattr(app, coll)
                names += [str(docs(i).Name) for i in range(1, docs.Count + 1) if not docs(i).Path]
            except Exception:
                pass
    return names


def close_test_documents(ctx, only_name="", allowed_pids=None, label="종료 정리"):
    """테스트 폴더 문서와 실행 중에 생긴 제목 없는 문서를 저장하지 않고 닫고, 문서가 남지 않은 앱은 끝냅니다.
    allowed_pids를 주면 그 프로세스만 끝냅니다. 문서가 남아 있는 프로세스는 어떤 경우에도 끝내지 않습니다.
    한 일은 ctx.cleanup_log에 남겨 errors_*.txt에 기록합니다."""
    opts = {"test_dir": ctx.dir, "keep_untitled": ctx.untitled_before, "only_name": only_name,
            "run_start": ctx.start_time, "allowed_pids": sorted(allowed_pids) if allowed_pids is not None else None,
            "protect_names": [] if not only_name else [n for n in rules.TEST_DOCS.values() if n != only_name]}
    args = [PY, "-c", CLEANUP_SCRIPT, ROOT, json.dumps(opts, ensure_ascii=False), TOOLS_DIR]
    stamp = datetime.now().strftime("%H:%M:%S")
    try:
        proc = subprocess.run(args, capture_output=True, timeout=120, env=dict(os.environ, PYTHONIOENCODING="utf-8"))
    except Exception as e:
        ctx.cleanup_log.append(f"{stamp} [{label}] 정리 스크립트 실행 실패: {type(e).__name__}: {e}")
        return []
    err = proc.stderr.decode("utf-8", errors="replace").strip()
    if err:
        log(err[-500:])
    try:
        result = json.loads(proc.stdout.decode("utf-8").strip().splitlines()[-1])
    except Exception:
        ctx.cleanup_log.append(f"{stamp} [{label}] 정리 결과를 읽지 못함 (exit {proc.returncode}): {err[-300:]}")
        return []
    ctx.cleanup_log.append(f"{stamp} [{label}] 대상 문서={only_name or '테스트 폴더 전체'}, "
                           f"끝내도 되는 PID={opts['allowed_pids'] if allowed_pids is not None else '실행 중 생긴 프로세스'}")
    ctx.cleanup_log += [f"{stamp} [{label}]   {a}" for a in result.get("actions", [])]
    return result.get("closed", [])


# ---------------------------------------------------------------------------
def write_report(ctx):
    stamp = ctx.stamp
    path = ctx.path(f"report_{stamp}.txt")
    counts = {}
    for r in ctx.results:
        counts[r["status"]] = counts.get(r["status"], 0) + 1
    lines = [
        f"office-cli 1.0 Windows 테스트 보고서 ({stamp})",
        f"Python {sys.version.split()[0]} / 테스트 폴더 {ctx.dir}",
        "요약: " + ", ".join(f"{k}={v}" for k, v in sorted(counts.items())),
        "",
    ]
    for r in ctx.results:
        lines.append(f"[{r['status']}] {r['id']} {r['title']}")
        if r.get("detail"):
            lines.append(f"    내용: {r['detail']}")
        if r.get("exit_code") is not None:
            lines.append(f"    종료코드: {r['exit_code']}")
        if r.get("json"):
            lines.append(f"    JSON: {r['json']}")
        if r.get("stderr_tail"):
            lines.append("    stderr(끝부분): " + r["stderr_tail"].replace("\n", "\n        "))
        if r.get("trace"):
            lines.append("    trace: " + r["trace"].replace("\n", "\n        "))
        lines.append("")
    with open(path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))
    return path, counts


# ---------------------------------------------------------------------------
# 오류 분석용 상세 로그 (errors_*.txt)
# ---------------------------------------------------------------------------
PROBLEM = ("FAIL", "ERROR")
LIBS = ("pywin32", "openpyxl", "python-docx", "python-pptx", "pymupdf", "olefile", "matplotlib", "numpy")


def tool_version():
    try:
        with open(os.path.join(ROOT, "SKILL.md"), encoding="utf-8") as f:
            for line in f:
                if line.strip().startswith("version:"):
                    return line.split(":", 1)[1].strip().strip('"')
    except Exception:
        pass
    return "?"


def environment_lines(ctx):
    """PC, Python, 라이브러리, Office 버전, 도구 설정. 실패해도 다른 줄은 남깁니다."""
    from importlib import metadata
    lines = [f"office-cli 버전: {tool_version()} ({ROOT})",
             f"Windows: {platform.platform()} / {platform.version()}",
             f"Python: {sys.version.split()[0]} ({sys.executable})"]
    vers = []
    for lib in LIBS:
        try:
            vers.append(f"{lib} {metadata.version(lib)}")
        except Exception:
            vers.append(f"{lib} (없음)")
    lines.append("라이브러리: " + ", ".join(vers))
    try:
        proc = subprocess.run([PY, os.path.join(ROOT, "office.py"), "windows"], cwd=ROOT, capture_output=True,
                              timeout=30, env=dict(os.environ, PYTHONIOENCODING="utf-8", **RUN_ENV))
        out = proc.stdout.decode("utf-8", errors="replace").strip().splitlines()
        data = json.loads(out[-1]) if out else {}
        lines.append("Office (office.py windows의 office): " + json.dumps(data.get("office"), ensure_ascii=False))
    except Exception as e:
        lines.append(f"Office: 확인 실패 ({type(e).__name__}: {e})")
    cfg = os.path.join(ROOT, "config.json")
    try:
        with open(cfg, encoding="utf-8") as f:
            lines.append("config.json: " + json.dumps(json.load(f), ensure_ascii=False))
    except FileNotFoundError:
        lines.append("config.json: 없음 (transfer --probe 전)")
    except Exception as e:
        lines.append(f"config.json: 읽기 실패 ({e})")
    a = ctx.args
    lines.append(f"러너 옵션: only={a.only} auto={a.auto} agent={a.agent} skip_outlook={a.skip_outlook} "
                 f"self_email={'있음' if a.self_email else '없음'} keep_open={a.keep_open}")
    lines.append(f"테스트 폴더: {ctx.dir} / 도구 환경변수: {json.dumps(RUN_ENV, ensure_ascii=False)}")
    lines.append(f"소요 시간: {round(time.time() - ctx.start_time)}초")
    return lines


def _indent(text, pad="        "):
    return pad + str(text).rstrip().replace("\n", "\n" + pad)


def _call_lines(call):
    lines = [f"    $ {call['command']}",
             f"      시작 {call.get('started')} / 종료코드 {call.get('exit_code')} / {call.get('elapsed')}초"]
    out = (call.get("stdout") or "").strip()
    if out:
        try:
            pretty = json.dumps(json.loads(out.splitlines()[-1]), ensure_ascii=False, indent=2)
            rest = "\n".join(out.splitlines()[:-1]).strip()
            lines.append("      stdout(JSON):")
            lines.append(_indent(pretty))
            if rest:
                lines.append("      stdout(JSON 앞의 다른 출력, 계약 위반):")
                lines.append(_indent(rest))
        except Exception:
            lines.append("      stdout(JSON 아님):")
            lines.append(_indent(out))
    else:
        lines.append("      stdout: (없음)")
    err = (call.get("stderr") or "").strip()
    if err:
        lines.append("      stderr:")
        lines.append(_indent(err))
    return lines


def write_error_log(ctx, fatal=None):
    """Claude가 읽고 원인을 판단할 수 있도록 실패 항목의 근거를 빠짐없이 모읍니다.
    실패가 없어도 환경 정보와 요약을 남깁니다(통과 확인용)."""
    path = ctx.path(f"errors_{ctx.stamp}.txt")
    counts = {}
    for r in ctx.results:
        counts[r["status"]] = counts.get(r["status"], 0) + 1
    problems = [r for r in ctx.results if r["status"] in PROBLEM]
    lines = [
        "# office-cli 실기 테스트 오류 로그 (분석용)",
        f"# 생성 {ctx.stamp}. 이 파일 전체를 그대로 Claude에게 전달하세요. 편집하지 마세요.",
        "",
        "## 1. 환경",
    ]
    try:
        lines += environment_lines(ctx)
    except Exception as e:
        lines.append(f"환경 정보 수집 실패: {e}")
    lines += ["", "## 2. 요약",
              "결과: " + (", ".join(f"{k}={v}" for k, v in sorted(counts.items())) or "실행된 항목 없음"),
              "문제 항목: " + (", ".join(f"{r['id']}({r['status']})" for r in problems) or "없음")]
    skipped = [r for r in ctx.results if r["status"] in ("ENV-SKIP", "MANUAL-SKIP")]
    if skipped:
        lines.append("건너뜀: " + ", ".join(f"{r['id']}({r['status']}: {r.get('detail', '')})" for r in skipped))
    if fatal is not None:
        lines += ["", "## 3. 테스트 중단 (준비 단계 또는 러너 오류)",
                  f"단계: {getattr(fatal, 'label', type(fatal).__name__)}",
                  f"내용: {getattr(fatal, 'detail', '') or fatal}"]
        trace = getattr(fatal, "trace", "") or "".join(traceback.format_exception(fatal))
        lines += ["traceback:", _indent(trace)]
        lines += ["중단 직전 도구 실행:"] + [ln for c in CALLS[-3:] for ln in _call_lines(c)]
        lines += ["중단 직전 로그 (최근 200줄):", _indent("\n".join(LOG_LINES[-200:]))]
    sec = 4 if fatal is not None else 3
    lines += ["", f"## {sec}. 러너 자동 복구와 창/프로세스 정리"]
    lines.append("자동 복구: " + ("없음" if not ctx.recoveries else ""))
    lines += [f"  {r}" for r in ctx.recoveries]
    lines.append("정리 기록: " + ("없음 (아직 정리 전)" if not ctx.cleanup_log else ""))
    lines += [f"  {c}" for c in ctx.cleanup_log]
    sec += 1
    lines += ["", f"## {sec}. 문제 항목 상세"]
    if not problems:
        lines.append("없음")
    for r in problems:
        lines += ["", "-" * 70, f"[{r['status']}] {r['id']} {r['title']}"]
        if r.get("detail"):
            lines.append(f"  판정 근거: {r['detail']}")
        if r.get("full_trace") or r.get("trace"):
            lines += ["  테스트 스크립트 traceback:", _indent(r.get("full_trace") or r.get("trace"))]
        calls = r.get("calls") or []
        lines.append(f"  이 항목에서 실행한 도구 명령 {len(calls)}개:")
        for c in calls:
            lines += _call_lines(c)
        if r.get("log"):
            lines += ["  이 항목의 진행 로그:", _indent("\n".join(r["log"]))]
    lines += ["", f"## {sec + 1}. 도구 오류 로그 ({TOOL_ERROR_LOG_NAME})"] + tool_error_log_lines()
    with open(path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines) + "\n")
    return path


TOOL_ERROR_LOG_NAME = r"_data\logs\errors.txt"
TOOL_LOG_LIMIT = 200000   # 도구 오류 로그를 옮겨 담는 상한(글자, 뒤쪽 우선)
TOOL_LOG_OFFSET = [0]     # 실행 시작 때 도구 오류 로그 크기 (이번 실행분만 담기 위해)


def tool_error_log_lines():
    """도구가 실패할 때마다 남긴 <OFFICE_CLI_DATA_DIR>/logs/errors.txt를 그대로 옮겨 담습니다 (한 파일에 모두 모으려고)."""
    data = RUN_ENV.get("OFFICE_CLI_DATA_DIR")
    if not data:
        return ["(도구 데이터 폴더를 정하기 전에 끝남)"]
    path = os.path.join(data, "logs", "errors.txt")
    head = [f"원본: {path} (이번 실행에서 추가된 부분)"]
    try:
        with open(path, "rb") as f:
            raw = f.read()
        start = TOOL_LOG_OFFSET[0]
        if start > len(raw):          # 실행 중 2MB를 넘어 errors.old.txt로 넘어감: 새 파일 전체
            head.append("(실행 중 로그가 회전되어 이전 부분은 errors.old.txt에 있음)")
            start = 0
        text = raw[start:].decode("utf-8", errors="replace")
    except FileNotFoundError:
        return [f"없음 ({path} 없음: 도구 실패가 기록되지 않음)"]
    except Exception as e:
        return [f"읽기 실패 ({path}): {type(e).__name__}: {e}"]
    if not text.strip():
        return [f"없음 (이번 실행에서 {path}에 추가된 기록 없음)"]
    if len(text) > TOOL_LOG_LIMIT:
        head.append(f"(앞부분 {len(text) - TOOL_LOG_LIMIT}자 생략)")
        text = text[-TOOL_LOG_LIMIT:]
    return head + text.rstrip().splitlines()


GROUPS = [("A", test_A), ("B", test_B), ("C", test_C), ("D", test_D), ("E", test_E), ("F", test_F), ("G", test_G),
          ("H", test_H), ("I", test_I), ("J", test_J), ("K", test_K)]


def main():
    parser = argparse.ArgumentParser(description="office-cli 1.0 Windows 실기 테스트 러너")
    parser.add_argument("--dir", default=r"C:\office_test", help="테스트 파일을 만들 폴더 (기본 C:\\office_test)")
    parser.add_argument("--only", help="실행할 그룹 (예: A,C,E)")
    parser.add_argument("--self-email", help="G3 실제 발송 테스트에 쓸 본인 메일 주소 (생략 시 G3 건너뜀)")
    parser.add_argument("--skip-outlook", action="store_true", help="G 그룹(Outlook) 건너뜀")
    parser.add_argument("--auto", action="store_true", help="y/n 질문 없이 자동 판정 항목만 진행")
    parser.add_argument("--agent", action="store_true",
                        help="에이전트(OpenCode 스킬)용: --auto와 같고, 로그는 stderr로, stdout에는 결과 JSON 1줄만 출력")
    parser.add_argument("--keep-open", action="store_true",
                        help="끝난 뒤 테스트 문서를 열어 둠 (기본: 테스트 폴더 문서만 저장하지 않고 닫음)")
    args = parser.parse_args()
    if args.agent:
        args.auto = True

    pythoncom.CoInitialize()
    global CTX
    ctx = CTX = Ctx(args)
    selected = {g.strip().upper() for g in args.only.split(",")} if args.only else {g for g, _ in GROUPS}
    if args.skip_outlook:
        selected.discard("G")

    log("=" * 70)
    log(" office-cli 1.0 Windows 테스트")
    log(f" 대상 그룹: {','.join(sorted(selected))} / 테스트 폴더: {ctx.dir}")
    log(" 실행 전 다른 Excel/Word/PowerPoint 문서는 저장 후 모두 닫아 주세요.")
    log(" 테스트 중에는 Office 창을 직접 조작하지 마세요 (질문 응답 제외).")
    log("=" * 70)
    running = {img: process_count(img) for img in ("EXCEL.EXE", "WINWORD.EXE", "POWERPNT.EXE")}
    if any(v > 0 for v in running.values()):
        log(f"[주의] 실행 중인 Office 프로세스: {running}")
    if not args.auto:
        if input("시작할까요? (y/n): ").strip().lower() != "y":
            return

    try:
        ctx.untitled_before = untitled_documents()   # 실행 전부터 있던 제목 없는 문서는 닫지 않음
        setup_files(ctx)
        open_user_windows(ctx)
        for g, fn in GROUPS:
            if g in selected:
                ctx.mark = (len(CALLS), len(LOG_LINES))
                log(f"\n{'-' * 25} 그룹 {g} {'-' * 25}")
                try:
                    ensure_test_docs(ctx)          # 앞 그룹에서 테스트 문서가 닫혔으면 다시 엶 (JIT 복구)
                    fn(ctx)
                except Exception as e:
                    trace = traceback.format_exc()
                    log(f"[ERROR] 그룹 {g}: {e}")
                    calls, lines = _evidence(ctx)
                    ctx.results.append({"id": g, "title": f"그룹 {g} 준비", "status": "ERROR", "detail": str(e),
                                        "trace": trace[-800:], "full_trace": trace, "calls": calls, "log": lines})
    except KeyboardInterrupt as e:
        log("[ERROR] 사용자가 중단했습니다 (Ctrl+C).")
        err_path = _safe_error_log(ctx, e)
        log(f" 오류 로그(중단 시점까지): {err_path}")
        if args.agent:
            emit_agent_json({"status": "error", "error_type": "interrupted", "message": "사용자가 중단함",
                             "error_log": err_path, "agent_directive": "STOP_RETRY: 사용자가 중단했습니다. 재시도하지 마세요."}, 130)
        sys.exit(130)
    except Exception as e:
        # 준비 단계 실패나 러너 자체 오류: 테스트를 계속할 수 없으므로 오류 로그만 남기고 끝냅니다.
        if not isinstance(e, SetupError):
            log(f"[ERROR] 러너 오류: {describe_com_error(e)}")
            log(traceback.format_exc())
        err_path = _safe_error_log(ctx, e)
        log(f" 오류 로그: {err_path}  <- 이 파일을 그대로 공유해 주세요.")
        if args.agent:
            emit_agent_json({
                "status": "error", "error_type": "setup", "message": str(e), "error_log": err_path,
                "agent_directive": "STOP_RETRY: 테스트를 시작하지 못했습니다. 재시도하지 말고 message와 "
                                   "error_log 경로를 사용자에게 그대로 알리세요."}, 1)
        sys.exit(1)

    shots, n_shots = save_shots(ctx)
    path, counts = write_report(ctx)
    err_path = _safe_error_log(ctx)
    problems = [r["id"] for r in ctx.results if r["status"] in PROBLEM]
    log("\n" + "=" * 70)
    log(" 결과 요약: " + ", ".join(f"{k}={v}" for k, v in sorted(counts.items())))
    for r in ctx.results:
        if r["status"] not in ("PASS",):
            log(f"   [{r['status']}] {r['id']} {r['title']} {r.get('detail', '')}")
    log(f" 보고서: {path}")
    log(f" 오류 로그: {err_path}" + ("  <- 문제 항목이 있습니다. 이 파일을 그대로 공유해 주세요." if problems else ""))
    log(f" 확인용 그림 {n_shots}개: {shots} (슬라이드, 차트)")
    closed = []
    if args.keep_open:
        log(" 테스트용 A 파일들은 열린 채로 남아 있습니다. 저장하지 않고 닫아도 됩니다.")
    else:
        import gc
        gc.collect()
        pythoncom.CoUninitialize()            # 러너가 쥔 COM 참조를 놓아야 Quit한 PowerPoint가 실제로 끝남
        closed = close_test_documents(ctx)
        log(f" 테스트 문서 {len(closed)}개를 저장하지 않고 닫았습니다: {', '.join(closed)}")
        err_path = _safe_error_log(ctx)       # 종료 정리 기록까지 넣어 다시 씀 (같은 파일)
    log("=" * 70)
    if args.agent:
        if problems:
            directive = (f"VERIFIED: 테스트 완료, 문제 항목 {len(problems)}개. 재실행하지 말고 summary와 failed를 보고하고, "
                         "error_log 파일을 Claude에게 전달하라고 사용자에게 안내하세요.")
        else:
            directive = "VERIFIED: 테스트 완료, 문제 항목 없음. summary를 사용자에게 보고하세요."
        emit_agent_json({"status": "ok", "summary": counts, "failed": problems, "report_file": path,
                         "error_log": err_path, "shots_dir": shots, "closed_documents": len(closed),
                         "agent_directive": directive}, 0)


def _safe_error_log(ctx, fatal=None):
    try:
        return write_error_log(ctx, fatal)
    except Exception as e:
        # 오류 로그를 못 쓰는 경우(테스트 폴더 쓰기 불가 등): 현재 폴더에 최소 내용이라도 남깁니다.
        fallback = os.path.join(os.getcwd(), f"errors_{ctx.stamp}.txt")
        try:
            with open(fallback, "w", encoding="utf-8") as f:
                f.write(f"오류 로그 작성 실패: {e}\n중단 원인: {fatal}\n"
                        + "".join(traceback.format_exception(fatal) if fatal else []) + "\n".join(LOG_LINES[-200:]))
            return fallback
        except Exception:
            return f"(오류 로그를 쓰지 못함: {e})"


if __name__ == "__main__":
    main()
