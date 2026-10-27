# -*- coding: utf-8 -*-
"""
Generates the PC test sample files for ml-kit (deterministic, seed fixed).
Usage: python make_samples.py   (writes next to this script)

A: process data (regression, multi-target), Excel report-style layout:
   title rows, header row 4, unit row 5, Excel dates, text category, thousands
   separators, '-' / '미측정' placeholders, a free-text column that must not be used.
B: quality inspection (classification), CP949 CSV like Korean Excel exports:
   string dates, numeric equipment codes (to be forced categorical), multiclass
   grade and imbalanced binary defect flag.
*_new files: unseen rows for prediction, deliberately containing out-of-range values,
   an unseen category, missing values and dates after the training period.
"""
import json
import os
from datetime import datetime, timedelta

import numpy as np
import openpyxl
import pandas as pd

HERE = os.path.dirname(os.path.abspath(__file__))


def _process_rows(n, seed, start):
    rng = np.random.default_rng(seed)
    rows = []
    for i in range(n):
        ts = start + timedelta(hours=8 * i + int(rng.integers(0, 3)))
        line = rng.choice(["A라인", "B라인", "C라인"], p=[0.4, 0.35, 0.25])
        temp = rng.uniform(210, 270)
        power = rng.uniform(1200, 2600)
        press = rng.uniform(2.0, 4.5)
        cat = rng.uniform(0.3, 1.4)
        water = rng.uniform(12, 28)
        flow = rng.uniform(4, 9)
        night = 1.0 if ts.hour >= 20 or ts.hour < 6 else 0.0
        line_eff = {"A라인": 0.0, "B라인": 3.5, "C라인": -2.0}[line]
        eff = (0.22 * temp + 9 * cat - 0.45 * water + 0.002 * power + line_eff - 2.5 * night
               + rng.normal(0, 1.2))
        tensile = 180 + 35 * cat + 4.0 * press - 0.8 * water + {"A라인": 0, "B라인": 6, "C라인": -4}[line] \
            + rng.normal(0, 3.0)
        rows.append([ts, line, round(temp, 1), f"{power:,.0f}", round(press, 2), round(cat, 3),
                     round(water, 1), round(flow, 2), "", round(eff, 2), round(tensile, 1)])
    return rows, rng


def _write_process_xlsx(path, rows, title):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "공정데이터"
    ws.append([title])
    ws.append(["작성: 공정기술팀", None, None, "승인: ____"])
    ws.append([])
    ws.append(["측정일시", "생산라인", "반응온도", "소비전력", "압축압력", "촉매비율", "수분함량", "유입유속",
               "비고", "생산효율", "인장강도"])
    ws.append(["", "", "°C", "kW", "bar", "wt%", "%", "m/s", "", "%", "MPa"])
    for r in rows:
        ws.append(r)
    for row in ws.iter_rows(min_row=6, min_col=1, max_col=1):
        row[0].number_format = "yyyy-mm-dd hh:mm"
    ws.column_dimensions["A"].width = 18
    wb.save(path)


def make_process():
    rows, rng = _process_rows(400, 11, datetime(2025, 1, 1, 6))
    for i in rng.choice(len(rows), 12, replace=False):
        rows[i][int(rng.choice([2, 4, 6]))] = rng.choice(["-", "미측정", None])
    for i in rng.choice(len(rows), 6, replace=False):
        rows[i][8] = rng.choice(["설비 점검 후 재가동", "원료 lot 변경", "야간 교대"])
    for i in rng.choice(len(rows), 3, replace=False):
        rows[i][9] = None  # target missing -> row excluded for that target only
    _write_process_xlsx(os.path.join(HERE, "sample_A_process.xlsx"), rows, "반응기 공정 운전 데이터 (2025년)")

    new_rows, rng2 = _process_rows(30, 99, datetime(2025, 6, 10, 6))  # after training period
    for r in new_rows:
        r[9] = None
        r[10] = None
    new_rows[0][2] = 305.0                 # 반응온도 above training max
    new_rows[1][2] = 180.0                 # below training min
    new_rows[2][1] = "D라인"                # unseen category
    new_rows[3][5] = None                  # missing 촉매비율
    new_rows[4][6] = "미측정"               # placeholder
    _write_process_xlsx(os.path.join(HERE, "sample_A_process_new.xlsx"), new_rows, "신규 배치 운전 조건 (예측용)")


def _quality_frame(n, seed, start):
    rng = np.random.default_rng(seed)
    days = [start + timedelta(days=int(d)) for d in rng.integers(0, 150, n)]
    equip = rng.choice([101, 102, 103, 104, 105], n)
    thick = rng.normal(2.0, 0.08, n)
    hard = rng.normal(320, 18, n)
    rough = rng.gamma(4, 0.08, n)
    temp = rng.uniform(18, 32, n)
    equip_eff = pd.Series(equip).map({101: 0.0, 102: 0.3, 103: -0.4, 104: 0.1, 105: -0.9}).values
    score = (-(np.abs(thick - 2.0) / 0.08) * 0.6 + (hard - 320) / 18 * 0.8 - (rough - 0.32) / 0.16
             - 0.05 * (temp - 25) + equip_eff + rng.normal(0, 0.5, n))
    grade = np.where(score > 0.5, "A", np.where(score > -0.9, "B", "C"))
    defect = np.where(score < np.quantile(score, 0.10), "불량", "양품")
    return pd.DataFrame({
        "검사일자": [d.strftime("%Y-%m-%d") for d in days],
        "설비코드": equip,
        "두께(mm)": thick.round(3),
        "경도(HV)": hard.round(1),
        "표면조도(Ra)": rough.round(3),
        "작업장온도(°C)": temp.round(1),
        "품질등급": grade,
        "불량여부": defect,
    }).sort_values("검사일자").reset_index(drop=True)


def make_quality():
    df = _quality_frame(500, 21, datetime(2025, 1, 2))
    df.to_csv(os.path.join(HERE, "sample_B_quality.csv"), index=False, encoding="cp949")

    new = _quality_frame(20, 77, datetime(2025, 6, 1)).drop(columns=["품질등급", "불량여부"])
    new.loc[0, "설비코드"] = 106          # unseen equipment code (categorical)
    new.loc[1, "경도(HV)"] = 420.0        # out of range
    new.loc[2, "두께(mm)"] = np.nan       # missing
    new.to_csv(os.path.join(HERE, "sample_B_quality_new.csv"), index=False, encoding="cp949")

    single = {"검사일자": "2025-06-15", "설비코드": 103, "두께": 2.12, "경도": 300.5, "표면조도": 0.55, "작업장온도": 29.0}
    # UTF-16 with BOM = what Windows PowerShell 5 'Out-File' writes by default
    with open(os.path.join(HERE, "sample_B_single_input.json"), "w", encoding="utf-16") as f:
        json.dump(single, f, ensure_ascii=False, indent=2)


def make_slots():
    slots = {
        "slot1_residual": "교차검증 잔차는 전 운전 구간에서 대칭에 가깝게 분포하며, 고효율 구간(70% 이상)에서만 소폭 과소 예측 경향이 확인됨.",
        "slot2_feature": "반응온도와 촉매비율이 예측의 대부분을 설명하며, 야간 운전(측정일시_hour) 효과가 독립적으로 관측됨.",
        "slot3_lineage": "베이스라인 대비 저기여 피처 제외와 교차항 추가로 검증 성능이 단계적으로 개선됨.",
        "slot4_correlation": "소비전력과 반응온도 간 상관이 중간 수준으로, 선형 모델 해석 시 주의가 필요함.",
        "slot5_serving": "신규 라인(D라인) 투입 시 범주 미학습 경고가 발생하므로 재학습 후 배포를 권고함.",
    }
    with open(os.path.join(HERE, "sample_slots.json"), "w", encoding="utf-8") as f:
        json.dump(slots, f, ensure_ascii=False, indent=2)


if __name__ == "__main__":
    make_process()
    make_quality()
    make_slots()
    print("[STATUS] samples written to", HERE)
