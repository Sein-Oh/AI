# -*- coding: utf-8 -*-
"""
2.3.1 regression tests: DRM-safe Word template for ml_report.py.

Company PC (corporate DRM): python-docx's templates/default.docx was a substitute document without the
standard styles, so every report failed with "no style with name 'Caption'". The report now opens the
unpacked template folder (templates/default-docx-template, plain XML) first, and repairs a damaged
template (missing styles, placeholder body text) instead of failing.
"""
import json
import os
import subprocess
import sys

import pytest

HERE = os.path.dirname(os.path.abspath(__file__))
KIT = os.path.dirname(HERE)
sys.path.insert(0, HERE)
sys.path.insert(0, KIT)
docx = pytest.importorskip("docx")
pytest.importorskip("matplotlib")
import docx.api  # noqa: E402
from make_sample import make_sample  # noqa: E402
from test_cli_smoke import FEATURES, TARGET, run_tool  # noqa: E402
import ml_predict as mc  # noqa: E402
import ml_report  # noqa: E402

PLACEHOLDER = "보안 정책으로 보호된 문서입니다"
REMOVED_STYLES = ("Caption", "List Bullet", "Heading 2")


def make_damaged_default(path):
    """A docx like the DRM substitute: valid Word file, standard styles missing, placeholder body text."""
    d = docx.Document(ml_report.docx_template_dir())
    styles_el = d.styles.element
    for st in list(d.styles):
        if st.name in REMOVED_STYLES:
            styles_el.remove(st.element)
    d.add_paragraph(PLACEHOLDER)
    d.save(str(path))
    check = docx.Document(str(path))
    assert "Caption" not in {s.name for s in check.styles}
    return str(path)


@pytest.fixture
def recovered(monkeypatch):
    lst = []
    monkeypatch.setattr(mc, "_RECOVERED", lst)
    return lst


def test_directory_template_used_normally(tmp_path, monkeypatch, recovered):
    """With the template folder present, default.docx is never opened (even when it is damaged)."""
    assert os.path.isdir(ml_report.docx_template_dir())
    bad = make_damaged_default(tmp_path / "default.docx")
    monkeypatch.setattr(docx.api, "_default_docx_path", lambda: bad)
    d = ml_report.new_report_document()
    names = {s.name for s in d.styles}
    assert set(ml_report.REPORT_STYLES) <= names
    assert len(d.paragraphs) == 0 and len(d.tables) == 0
    assert recovered == []


def test_damaged_default_without_template_folder(tmp_path, monkeypatch, recovered):
    bad = make_damaged_default(tmp_path / "default.docx")
    monkeypatch.setattr(docx.api, "_default_docx_path", lambda: bad)
    monkeypatch.setattr(ml_report, "docx_template_dir", lambda: str(tmp_path / "없는폴더"))
    d = ml_report.new_report_document()
    assert set(ml_report.REPORT_STYLES) <= {s.name for s in d.styles}
    assert len(d.paragraphs) == 0
    assert any("default.docx 사용" in r for r in recovered)
    assert any("본문 1개" in r for r in recovered)
    assert any("스타일 없음" in r and "Caption" in r for r in recovered)
    d.add_paragraph("x", style="Caption")  # the call that failed on the company PC
    d.add_heading("h", level=2)


def test_both_templates_broken_is_stop_retry(tmp_path, monkeypatch, recovered):
    not_zip = tmp_path / "default.docx"
    not_zip.write_bytes(b"DRM-ENCRYPTED\x00\x01")  # like pptx/templates/default.pptx on the company PC
    monkeypatch.setattr(docx.api, "_default_docx_path", lambda: str(not_zip))
    monkeypatch.setattr(ml_report, "docx_template_dir", lambda: str(tmp_path / "없는폴더"))
    with pytest.raises(mc.InputError) as ei:
        ml_report.new_report_document()
    assert "python -m pip install --force-reinstall python-docx" in str(ei.value)


RUNNER = r"""
import sys
import docx.api
sys.path.insert(0, sys.argv[1])
import ml_report
bad, missing = sys.argv[2], sys.argv[3]
docx.api._default_docx_path = lambda: bad
ml_report.docx_template_dir = lambda: missing
sys.argv = ["ml_report.py"] + sys.argv[4:]
ml_report.main()
"""


def test_report_succeeds_with_damaged_templates(tmp_path):
    """End to end: damaged default.docx + no template folder -> report ok, captions present, no placeholder."""
    data = make_sample(str(tmp_path / "drm.csv"), n=80)
    tr = run_tool("ml_train.py", "-f", data, "-x", FEATURES, "-y", TARGET, "--desc", "drm")
    bad = make_damaged_default(tmp_path / "default.docx")
    out = tmp_path / "report.docx"
    proc = subprocess.run(
        [sys.executable, "-c", RUNNER, KIT, bad, str(tmp_path / "없는폴더"), "-m", tr["version_dir"], "-o", str(out)],
        capture_output=True, text=True, encoding="utf-8", timeout=600,
    )
    lines = [ln for ln in proc.stdout.splitlines() if ln.strip()]
    assert len(lines) == 1, proc.stdout + proc.stderr
    payload = json.loads(lines[0])
    assert proc.returncode == 0 and payload["status"] == "ok", proc.stdout + proc.stderr
    assert any("스타일 없음" in r for r in payload["recovered"])
    assert "[WARN] Word" in proc.stderr
    d = docx.Document(str(out))
    captions = [p.text for p in d.paragraphs if p.style.name == "Caption"]
    assert len(captions) >= payload["figure_count"] + 1
    assert d.paragraphs[0].text == "ML 결과 보고서"
    assert PLACEHOLDER not in "\n".join(p.text for p in d.paragraphs)
