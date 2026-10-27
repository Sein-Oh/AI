"""word.py 테스트: 가짜 Word 문서(문자열 모델)로 삽입 위치, 마크다운 변환, 찾아 바꾸기 반복을 검증합니다."""
import json

import pytest

import _office as off
import word
from fake_office import make_win32com
from fake_word import FakeWordDocument


@pytest.fixture
def wd(monkeypatch):
    holder = {"doc": FakeWordDocument("추진 배경\r2026년 계획입니다.\r일정")}
    fake, _ = make_win32com()
    monkeypatch.setattr(off, "win32com", fake)
    monkeypatch.setattr(off, "pythoncom", None)
    monkeypatch.setattr(off, "list_open", lambda kind: [])
    monkeypatch.setattr(off, "get_active", lambda kind: {"name": holder["doc"].Name, "path": holder["doc"].FullName,
                                                          "app": object(), "doc": holder["doc"]})
    monkeypatch.setattr(off, "save_snapshot", lambda *a, **k: None)
    return holder


def run(capsys, *argv):
    with pytest.raises(SystemExit) as ei:
        word.main(list(argv))
    out = [ln for ln in capsys.readouterr().out.splitlines() if ln.strip()]
    assert len(out) == 1, out
    return ei.value.code, json.loads(out[0])


def paras(doc):
    return doc.text.split("\r")


def test_write_plain_text_at_end(wd, capsys):
    code, out = run(capsys, "write", "--text", "새 문단입니다")
    assert code == 0, out
    assert paras(wd["doc"])[:4] == ["추진 배경", "2026년 계획입니다.", "일정", "새 문단입니다"]


def test_write_after_anchor_goes_after_whole_paragraph(wd, capsys):
    code, out = run(capsys, "write", "--text", "추가 문단", "--anchor", "계획", "--pos", "after")
    assert code == 0
    assert paras(wd["doc"])[:4] == ["추진 배경", "2026년 계획입니다.", "추가 문단", "일정"]   # 단락 중간을 쪼개지 않음


def test_write_before_and_replace_anchor(wd, capsys):
    run(capsys, "write", "--text", "머리말", "--anchor", "추진", "--pos", "before")
    assert paras(wd["doc"])[0] == "머리말"
    run(capsys, "write", "--text", "새 계획", "--anchor", "2026년", "--pos", "replace")
    assert "2026년 계획입니다." not in wd["doc"].text and "새 계획" in paras(wd["doc"])


def test_write_markdown_blocks_styles_and_table(wd, capsys):
    md = "## 결론\n본문 **강조** 끝\n\n- 가\n- 나\n\n| 공정 | 수율 |\n|---|---|\n| 가열 | 98 |"
    code, out = run(capsys, "write", "--text", md)
    assert code == 0
    doc = wd["doc"]
    assert ("style", "결론", word.heading_style(2)) in doc.log
    assert ("style", "가", word.WD_STYLE_LIST_BULLET) in doc.log
    assert "본문 강조 끝" in paras(doc)
    assert any(e[0] == ("font", doc.text.index("강조"), doc.text.index("강조") + 2) and e[1] == "Bold" and e[2] is True
               for e in doc.log)
    assert doc.tables == [["공정\t수율", "가열\t98"]]
    assert out["inserted"] == ["제목2", "본문", "글머리 목록 2개", "표 2x2"]


def test_write_escaped_newlines_and_anchor_not_found(wd, capsys):
    code, out = run(capsys, "write", "--text", "첫줄\\n둘째줄")
    assert "첫줄" in paras(wd["doc"]) and "둘째줄" in paras(wd["doc"])
    code, out = run(capsys, "write", "--text", "x", "--anchor", "없는문구")
    assert code == 1 and "read" in out["message"]


def test_write_needs_content(wd, capsys):
    code, out = run(capsys, "write")
    assert code == 1 and out["error_type"] == "input"


def test_replace_no_infinite_loop_and_counts(wd, capsys):
    wd["doc"] = FakeWordDocument("가 가 가")
    code, out = run(capsys, "replace", "--find", "가", "--replace", "가가")
    assert code == 0 and out["replaced"] == 3 and wd["doc"].text.startswith("가가 가가 가가")


def test_replace_empty_deletes_and_formatting(wd, capsys):
    wd["doc"] = FakeWordDocument("초안 보고서 초안")
    code, out = run(capsys, "replace", "--find", "초안 ", "--replace", "")
    assert out["replaced"] == 1 and wd["doc"].text.startswith("보고서 초안")
    code, out = run(capsys, "replace", "--find", "보고서", "--bold", "--color", "blue")
    assert out["replaced"] == 1
    assert any(len(e) == 3 and e[1] == "Bold" and e[2] is True for e in wd["doc"].log)


def test_replace_map_and_limits(wd, capsys):
    wd["doc"] = FakeWordDocument("Draft 진행중")
    code, out = run(capsys, "replace", "--map", '{"Draft": "Final", "진행중": "완료"}')
    assert out["replaced"] == 2 and wd["doc"].text.startswith("Final 완료")
    code, out = run(capsys, "replace", "--find", "x" * 300, "--replace", "y")
    assert code == 1 and "255" in out["message"]
    code, out = run(capsys, "replace", "--find", "없음", "--replace", "y")
    assert code == 0 and out["replaced"] == 0


def test_read_closed_docx_with_python_docx(tmp_path, capsys, monkeypatch):
    import docx
    p = tmp_path / "계획.docx"
    d = docx.Document()
    d.add_heading("배경", level=1)
    d.add_paragraph("본문입니다")
    d.save(p)
    fake, calls = make_win32com()
    monkeypatch.setattr(off, "win32com", fake)
    monkeypatch.setattr(off, "list_open", lambda kind: [])
    code, out = run(capsys, "read", "--target", str(p))
    assert code == 0 and [x["text"] for x in out["paragraphs"]] == ["배경", "본문입니다"]
    assert out["paragraphs"][1]["heading"] == "배경"
    assert not any(c[0] == "DispatchEx" for c in calls)


def test_highlight_color_suggestion(wd, capsys):
    code, out = run(capsys, "replace", "--find", "계획", "--highlight", "yelow")
    assert code == 1 and "[추천] yellow" in out["message"]
