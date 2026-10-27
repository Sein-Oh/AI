"""ppt.py 테스트: 가짜 PowerPoint로 슬라이드 배치, 마크다운 -> 슬라이드, 치환, 슬라이드 관리, 화면 이동 규칙을 검증합니다."""
import json

import pytest

import _office as off
import ppt
from fake_office import make_win32com
from fake_ppt import FakePresentation, FakeShape


class FakeView:
    def __init__(self):
        self.goto = []

    def GotoSlide(self, n):
        self.goto.append(n)


class FakeWindow:
    def __init__(self, prs):
        self.View = FakeView()
        self.Presentation = prs


class FakePptApp:
    def __init__(self, prs):
        self.ActiveWindow = FakeWindow(prs)


@pytest.fixture
def pp(monkeypatch):
    holder = {"prs": FakePresentation(3)}
    holder["app"] = FakePptApp(holder["prs"])
    fake, _ = make_win32com()
    monkeypatch.setattr(off, "win32com", fake)
    monkeypatch.setattr(off, "pythoncom", None)
    monkeypatch.setattr(off, "list_open", lambda kind: [])
    monkeypatch.setattr(off, "save_snapshot", lambda *a, **k: None)

    def active(kind):
        holder["app"].ActiveWindow.Presentation = holder["prs"]
        return {"name": holder["prs"].Name, "path": holder["prs"].FullName, "app": holder["app"], "doc": holder["prs"]}
    monkeypatch.setattr(off, "get_active", active)
    return holder


def run(capsys, *argv):
    with pytest.raises(SystemExit) as ei:
        ppt.main(list(argv))
    out = [ln for ln in capsys.readouterr().out.splitlines() if ln.strip()]
    assert len(out) == 1, out
    return ei.value.code, json.loads(out[0])


def live(slide):
    return [s for s in slide.Shapes.items if not s.deleted]


def test_write_new_slide_title_placeholder_and_bullets(pp, capsys):
    code, out = run(capsys, "write", "--slide", "new", "--title", "추진 전략", "--text", "- 안정성\\n- 확장성\\n설명")
    assert code == 0, out
    prs = pp["prs"]
    assert prs.Slides.Count == 4 and out["slides_written"] == [4]
    s = prs.Slides(4)
    assert s.Shapes.Title.TextFrame.TextRange.Text == "추진 전략"
    body = live(s)[1]
    assert body.TextFrame.TextRange.Text == "안정성\r확장성\r설명"
    assert body.Top >= 20 + 60                                     # 제목 아래에 쌓임
    assert pp["app"].ActiveWindow.View.goto == [4]                 # 고친 슬라이드로만 이동


def bullet(shape, i):
    return shape.TextFrame.TextRange.paras[i].ParagraphFormat.Bullet


def test_write_bullets_set_type_explicitly(pp, capsys):
    # Visible만 켜면 PowerPoint 2007 글상자는 번호 목록(Type=2)이 되므로 종류를 명시해야 함 (실기 F3)
    code, out = run(capsys, "write", "--slide", "new", "--title", "목록", "--text", "- 안정성\\n2. 확장성\\n설명")
    assert code == 0, out
    body = live(pp["prs"].Slides(4))[1]
    assert body.TextFrame.TextRange.Text == "안정성\r확장성\r설명"
    assert bullet(body, 1).Visible == -1 and bullet(body, 1).Type == ppt.PP_BULLET_UNNUMBERED
    assert bullet(body, 2).Visible == -1 and bullet(body, 2).Type == ppt.PP_BULLET_NUMBERED
    assert bullet(body, 3).Visible == 0 and "Type" not in vars(bullet(body, 3))


def test_info_detail_reports_lists_outside_overlap_and_overflow(pp, capsys):
    run(capsys, "write", "--slide", "new", "--title", "배치", "--text", "- 가\\n1. 나\\n다")
    s = pp["prs"].Slides(4)
    body = live(s)[1]
    tr = body.TextFrame.TextRange
    tr.Paragraphs().Count = 3
    tr.BoundHeight = body.Height + 50                                       # 글자가 상자보다 큼
    s.Shapes.add(FakeShape("그림 밖", left=900, top=100, width=200, height=100, kind=13))   # 슬라이드 너비 960
    s.Shapes.add(FakeShape("겹친 상자", left=body.Left, top=body.Top, width=body.Width, height=body.Height, text="겹침"))
    code, out = run(capsys, "info", "--detail", "--slides", "4")
    assert code == 0 and [x["no"] for x in out["slides"]] == [4] and out["slides_total"] == 4
    det = {d["name"]: d for d in out["slides"][0]["shape_detail"]}
    assert det[body.Name]["lists"] == ["bullet", "numbered", "none"] and det[body.Name]["text_overflow"] is True
    assert det["그림 밖"]["kind"] == "picture" and det["그림 밖"]["outside"] is True
    warnings = " | ".join(out["slides"][0]["warnings"])
    assert "슬라이드 밖" in warnings and "넘칩니다" in warnings and "겹칩니다" in warnings
    assert out["layout_warnings"] == len(out["slides"][0]["warnings"])


def test_markdown_bullets_and_numbered_lists_keep_their_kind():
    sections = ppt._md_sections("## 제목\n\n- 가\n- 나\n\n1. 첫째\n2. 둘째\n")
    assert sections[0]["lines"] == [("가", "bullet"), ("나", "bullet"), ("첫째", "numbered"), ("둘째", "numbered")]


def test_write_markdown_file_makes_one_slide_per_heading(pp, capsys, tmp_path):
    md = tmp_path / "결과.md"
    md.write_text("# 검색 결과\n\n## 보고서.pdf > p.3\n\n**마일스톤** 일정 정리\n\n## 계획.docx > 단락 2\n\n- 3분기 적용\n\n"
                  "| 항목 | 값 |\n|---|---|\n| A | 1 |\n", encoding="utf-8")
    code, out = run(capsys, "write", "--file", str(md))
    assert code == 0
    prs = pp["prs"]
    assert out["slides_written"] == [4, 5, 6] and prs.Slides.Count == 6
    assert prs.Slides(5).Shapes.Title.TextFrame.TextRange.Text == "보고서.pdf > p.3"
    assert live(prs.Slides(5))[1].TextFrame.TextRange.Text == "마일스톤 일정 정리"      # ** 제거
    table = [s for s in live(prs.Slides(6)) if s.HasTable][0].Table
    assert table.Cell(2, 1).Shape.TextFrame.TextRange.Text == "A"


def test_write_long_csv_table_splits_slides_with_header(pp, capsys, tmp_path):
    f = tmp_path / "t.csv"
    f.write_text("이름,값\n" + "".join(f"항목{i},{i}\n" for i in range(1, 21)), encoding="utf-8")
    code, out = run(capsys, "write", "--file", str(f), "--title", "결과")
    assert out["slides_written"] == [4, 5]
    t2 = [s for s in live(pp["prs"].Slides(5)) if s.HasTable][0].Table
    assert t2.Cell(1, 1).Shape.TextFrame.TextRange.Text == "이름" and t2.Rows.Count == 6


def test_write_image_fits_remaining_space_keeping_ratio(pp, capsys, tmp_path):
    img = tmp_path / "chart.png"
    img.write_bytes(b"\x89PNG")
    code, out = run(capsys, "write", "--slide", "2", "--image", str(img))
    assert code == 0
    pic = [s for s in live(pp["prs"].Slides(2)) if s.Type == 13][0]
    assert pic.Width / pic.Height == pytest.approx(1600 / 900)
    assert pic.Top + pic.Height <= 540 - ppt.MARGIN + 0.01 and pic.Left + pic.Width <= 960


def test_write_bad_slide_number_and_no_content(pp, capsys):
    code, out = run(capsys, "write", "--slide", "9", "--text", "x")
    assert code == 1 and "new" in out["message"]
    code, out = run(capsys, "write")
    assert code == 1


def test_cards(pp, capsys):
    code, out = run(capsys, "write", "--slide", "new", "--title", "핵심", "--cards", "안정성:무중단|확장성:모듈")
    cards = [s for s in live(pp["prs"].Slides(4)) if s.Name.startswith("Shape")]
    assert len(cards) == 2 and cards[0].TextFrame.TextRange.Text == "안정성\r무중단"


def test_replace_in_shapes_tables_no_loop(pp, capsys):
    prs = pp["prs"]
    s = prs.Slides(1)
    s.Shapes.add(FakeShape("tb", text="초안 초안 진행중"))
    tbl = s.Shapes.AddTable(1, 1, 0, 0, 10, 10)
    tbl.Table.Cell(1, 1).Shape.TextFrame.TextRange.Text = "초안"
    code, out = run(capsys, "replace", "--find", "초안", "--replace", "초안본", "--bold")
    assert code == 0 and out["replaced"] == 3 and out["slides"] == [1]
    assert live(s)[1].TextFrame.TextRange.Text == "초안본 초안본 진행중"
    code, out = run(capsys, "replace", "--find", "진행중", "--replace", "")
    assert out["replaced"] == 1 and live(s)[1].TextFrame.TextRange.Text == "초안본 초안본 "


def test_read_does_not_move_view(pp, capsys):
    s = pp["prs"].Slides(2)
    s.Shapes.Title.TextFrame.TextRange.Text = "둘째"
    s.Shapes.add(FakeShape("tb", text="본문 내용"))
    code, out = run(capsys, "read", "--slides", "2")
    assert code == 0 and out["slides"][0]["title"] == "둘째" and out["slides"][0]["text"] == "본문 내용"
    code, out = run(capsys, "info")
    assert out["slides_total"] == 3
    assert pp["app"].ActiveWindow.View.goto == []


def test_slide_actions(pp, capsys):
    prs = pp["prs"]
    first = prs.Slides(1)
    code, out = run(capsys, "slide", "--action", "move", "--slide", "1", "--to", "3")
    assert code == 0 and prs.Slides(3) is first
    code, out = run(capsys, "slide", "--action", "duplicate", "--slide", "3")
    assert out["slides_total"] == 4
    code, out = run(capsys, "slide", "--action", "clear")
    assert code == 1 and "--all-slides" in out["message"]
    prs.Slides(2).Shapes.add(FakeShape("x", text="a"))
    code, out = run(capsys, "slide", "--action", "clear", "--slide", "2")
    assert code == 0 and prs.Slides(2).Shapes.Count == 0
    code, out = run(capsys, "slide", "--action", "delete", "--slide", "9")
    assert code == 1


def test_read_closed_pptx_uses_python_pptx(tmp_path, capsys, monkeypatch):
    import pptx
    p = tmp_path / "d.pptx"
    prs = pptx.Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "제목"
    prs.save(p)
    fake, calls = make_win32com()
    monkeypatch.setattr(off, "win32com", fake)
    monkeypatch.setattr(off, "list_open", lambda kind: [])
    code, out = run(capsys, "read", "--target", str(p))
    assert code == 0 and out["slides"][0]["title"] == "제목"
    assert not calls


class com_error(Exception):
    """pywintypes.com_error 흉내 (args: hresult, 설명, excepinfo, argerr)."""


E_FAIL = (-2147352567, "예외가 발생했습니다.", (0, None, None, None, 0, -2147467259), None)


def _pdf_export(pp, behavior, version="16.0"):
    calls = []

    def export(path, fmt, **kw):
        calls.append((path, fmt, kw))
        if "PrintRange" not in kw:      # 실제 pywin32: 생략하면 인자 변환 단계에서 TypeError
            raise TypeError("The Python instance can not be converted to a COM object")
        if behavior == "ok":
            open(path, "wb").close()
        else:
            raise behavior
    pp["prs"].ExportAsFixedFormat = export
    pp["app"].Version = version
    return calls


def test_export_pdf_passes_print_range_none(pp, capsys):
    calls = _pdf_export(pp, "ok")
    code, out = run(capsys, "export", "--save", "deck.pdf")
    assert code == 0 and out["output_file"].endswith("deck.pdf")
    assert calls[0][1] == 2 and calls[0][2] == {"PrintRange": None}


def test_export_pdf_ppt2007_without_addin_is_input_error(pp, capsys, monkeypatch):
    # 추가 기능이 없는 Office 2007은 내보내기를 시도하지 않고 바로 안내
    monkeypatch.setattr(off, "_office12_pdf_addin", lambda: False)
    calls = _pdf_export(pp, com_error(*E_FAIL), version="12.0")
    code, out = run(capsys, "export", "--save", "deck.pdf")
    assert code == 1 and out["error_type"] == "input" and calls == []
    assert out["agent_directive"].startswith("STOP_RETRY:") and ".pptx" in out["message"]


def test_export_pdf_ppt2007_e_fail_with_addin_still_input_error(pp, capsys, monkeypatch):
    # 추가 기능 파일이 있는데도 E_FAIL이면 시도 후 같은 안내 (설치가 깨진 경우)
    monkeypatch.setattr(off, "_office12_pdf_addin", lambda: True)
    calls = _pdf_export(pp, com_error(*E_FAIL), version="12.0")
    code, out = run(capsys, "export", "--save", "deck.pdf")
    assert code == 1 and out["error_type"] == "input" and len(calls) == 1


def test_export_pdf_e_fail_on_newer_office_is_not_relabelled(pp, capsys):
    _pdf_export(pp, com_error(*E_FAIL), version="16.0")
    code, out = run(capsys, "export", "--save", "deck.pdf")
    assert code == 1 and out["error_type"] == "office" and "PDF 내보내기 기능이 없습니다" not in out["message"]
