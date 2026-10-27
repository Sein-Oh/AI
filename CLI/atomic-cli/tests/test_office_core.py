"""_office.py 테스트: 세션(사용자 창 미간섭, 읽기 전용, 화면 이동 규칙), 스냅샷, 대량 쓰기 설정 복원."""
import os

import pytest

import _office as off
import _runtime as rt
from fake_office import FakeExcelApp, FakePptApp, FakeWordDoc, FakeWorkbook, make_win32com


@pytest.fixture
def fake_com(monkeypatch):
    """win32com을 가짜로 바꾸고, 열린 문서/활성 문서는 테스트가 지정합니다."""
    state = {"open": {"excel": [], "word": [], "ppt": []}, "active": {}}
    monkeypatch.setattr(off, "pythoncom", None)
    monkeypatch.setattr(off, "list_open", lambda kind: state["open"][kind])

    def get_active(kind):
        if kind not in state["active"]:
            raise rt.InputError("열려 있는 문서가 없습니다.")
        return state["active"][kind]
    monkeypatch.setattr(off, "get_active", get_active)

    def install(**kw):
        fake, calls = make_win32com(**kw)
        monkeypatch.setattr(off, "win32com", fake)
        state["calls"] = calls
        return calls
    state["install"] = install
    install()
    return state


def test_excel_file_read_uses_dedicated_instance_readonly(tmp_path, fake_com):
    path = tmp_path / "a.xlsx"
    path.write_bytes(b"x")
    app = FakeExcelApp()
    calls = fake_com["install"](dispatch_ex=app)
    with off.open_target("excel", str(path)) as t:
        assert t.source == "file" and t.readonly
        t.doc.Saved = False
    assert ("DispatchEx", "Excel.Application") in calls
    assert app.Visible is False and app.Workbooks.open_kwargs["ReadOnly"] is True
    wb = app.Workbooks.items[0]
    assert ("Save",) not in wb.calls and ("Close", False) in wb.calls
    assert app.quit_called


class _Doc2007:
    """Word 2007 문서: SaveAs2가 없음 (pywin32 동적 디스패치는 AttributeError)."""
    def __init__(self):
        self.calls = []

    def SaveAs(self, path, **kw):
        self.calls.append(("SaveAs", path, kw))


class _Doc2010(_Doc2007):
    def SaveAs2(self, path, **kw):
        self.calls.append(("SaveAs2", path, kw))


def test_word_save_as_falls_back_to_saveas_on_word2007():
    old, new = _Doc2007(), _Doc2010()
    off.word_save_as(old, "C:/x/a.docx", 16)
    off.word_save_as(new, "C:/x/a.docx", 16)
    off.word_save_as(new, "C:/x/b.docx")
    assert old.calls == [("SaveAs", "C:/x/a.docx", {"FileFormat": 16})]
    assert new.calls == [("SaveAs2", "C:/x/a.docx", {"FileFormat": 16}), ("SaveAs2", "C:/x/b.docx", {})]


@pytest.mark.parametrize("version,addin,expected", [
    ("16.0", False, True), ("14.0", False, True), ("12.0", False, False), ("12.0", True, True), ("11.0", True, False)])
def test_office_caps_pdf_export(monkeypatch, version, addin, expected):
    monkeypatch.setattr(off, "_office12_pdf_addin", lambda: addin)
    app = type("App", (), {"Version": version})()
    caps = off.office_caps(app)
    assert caps["version"] == version and caps["major"] == int(float(version)) and caps["pdf_export"] is expected


def test_require_pdf_export(monkeypatch):
    monkeypatch.setattr(off, "_office12_pdf_addin", lambda: False)
    with pytest.raises(rt.InputError) as ei:
        off.require_pdf_export(type("App", (), {"Version": "12.0"})(), ".docx")
    assert "PDF 내보내기 기능이 없습니다" in str(ei.value) and ".docx" in str(ei.value) and "2007" in str(ei.value)
    off.require_pdf_export(type("App", (), {"Version": "16.0"})(), ".docx")      # 가능: 예외 없음
    off.require_pdf_export(object(), ".docx")                                    # 버전 모름: 시도하도록 통과


def test_missing_file_readonly_suggests_open_documents(fake_com):
    fake_com["open"]["excel"] = [{"name": "10월보고.xlsx", "path": "C:/x/10월보고.xlsx", "app": None, "doc": None}]
    with pytest.raises(rt.InputError) as ei:
        with off.open_target("excel", "10월 보고.xlsx"):
            pass
    assert "[추천] 10월보고.xlsx" in str(ei.value)


def test_write_create_new_file_with_matching_format(tmp_path, fake_com):
    app = FakeExcelApp()
    fake_com["install"](dispatch_ex=app)
    target = tmp_path / "new.xlsm"
    with off.open_target("excel", str(target), readonly=False, create=True) as t:
        assert t.source == "new"
    wb = app.Workbooks.items[0]
    assert ("SaveAs", str(target), 52) in wb.calls


def test_write_without_create_does_not_make_files(tmp_path, fake_com):
    with pytest.raises(rt.InputError):
        with off.open_target("excel", str(tmp_path / "none.xlsx"), readonly=False):
            pass


def test_open_document_matched_by_name_is_not_closed(fake_com):
    wb = FakeWorkbook("C:/x/보고.xlsx")
    fake_com["open"]["excel"] = [{"name": "보고.xlsx", "path": "C:/x/보고.xlsx", "app": FakeExcelApp(), "doc": wb}]
    with off.open_target("excel", "보고.XLSX", readonly=False, snapshot=False) as t:
        assert t.source == "open" and t.doc is wb
        t.doc.Saved = False
    assert wb.calls == []          # 저장도 닫기도 하지 않음 (사용자 문서)


def test_sheet_lookup_rules(fake_com):
    wb = FakeWorkbook("C:/x/a.xlsx", sheets=("매출", "요약"))
    fake_com["active"]["excel"] = {"name": "a.xlsx", "path": "C:/x/a.xlsx", "app": FakeExcelApp(), "doc": wb}
    with off.open_target("excel") as t:
        assert t.sheet("요약 ").Name == "요약"
        assert t.sheet().Name == "매출"
        with pytest.raises(rt.InputError) as ei:
            t.sheet("매출2", create=True)            # 읽기 세션은 만들지 않음
        assert "[추천] 매출" in str(ei.value)
    with off.open_target("excel", readonly=False, snapshot=False) as t:
        assert t.sheet("신규", create=True).Name == "신규"


def test_read_session_never_moves_view_write_moves_to_changed_sheet(fake_com):
    wb = FakeWorkbook("C:/x/a.xlsx", sheets=("A", "B"))
    fake_com["active"]["excel"] = {"name": "a.xlsx", "path": "C:/x/a.xlsx", "app": FakeExcelApp(), "doc": wb}
    with off.open_target("excel") as t:
        t.set_focus(sheet=t.sheet("B"))
    assert wb.Worksheets("B").activated == 0
    with off.open_target("excel", readonly=False, snapshot=False, focus=False) as t:
        t.set_focus(sheet=t.sheet("B"))
    assert wb.Worksheets("B").activated == 0
    with off.open_target("excel", readonly=False, snapshot=False) as t:
        t.set_focus(sheet=t.sheet("B"))
        assert t.resolved()["sheet"] == "B"
    assert wb.Worksheets("B").activated == 1


def test_exception_in_session_does_not_save_file(tmp_path, fake_com):
    path = tmp_path / "a.xlsx"
    path.write_bytes(b"x")
    app = FakeExcelApp()
    fake_com["install"](dispatch_ex=app)
    with pytest.raises(RuntimeError):
        with off.open_target("excel", str(path), readonly=False, snapshot=False) as t:
            t.doc.Saved = False
            raise RuntimeError("도중 실패")
    wb = app.Workbooks.items[0]
    assert ("Save",) not in wb.calls and ("Close", False) in wb.calls and app.quit_called


def test_ppt_file_mode_leaves_users_instance_alone(tmp_path, fake_com):
    path = tmp_path / "deck.pptx"
    path.write_bytes(b"x")
    app = FakePptApp()
    fake_com["install"](dispatch=app, active=app)
    with off.open_target("ppt", str(path)) as t:
        pass
    assert app.touched == [] and app.DisplayAlerts == 2 and not app.quit_called
    assert app.Presentations.open_args[1:] == (-1, 0, 0)
    assert t is not None


def test_write_session_takes_snapshot_first(tmp_path, fake_com):
    wb = FakeWorkbook(str(tmp_path / "보고.xlsm"))
    (tmp_path / "보고.xlsm").write_bytes(b"x")
    fake_com["active"]["excel"] = {"name": wb.Name, "path": wb.FullName, "app": FakeExcelApp(), "doc": wb}
    with off.open_target("excel", readonly=False, tool="excel", command="write"):
        pass
    snaps = off.list_snapshots("excel")
    assert len(snaps) == 1 and snaps[0]["snapshot_file"].endswith(".xlsm")
    assert snaps[0]["by"] == "excel write" and snaps[0]["original_path"] == wb.FullName


def test_word_snapshot_unsaved_uses_flat_xml_without_touching_original(tmp_path):
    orig = tmp_path / "a.docx"
    orig.write_bytes(b"disk")
    off.save_snapshot("word", FakeWordDoc(str(orig), saved=False))
    snap = off.list_snapshots("word")[0]
    assert snap["snapshot_format"] == "flat_xml" and snap["snapshot_file"].endswith(".xml")
    off.save_snapshot("word", FakeWordDoc(str(orig), saved=True))
    assert off.list_snapshots("word")[0]["snapshot_file"].endswith(".docx")


def _xlsx(path, created):
    import datetime as dt
    import openpyxl
    wb = openpyxl.Workbook()
    wb.properties.created = dt.datetime(*created)
    wb.save(path)


def test_snapshots_of_a_replaced_document_are_dropped(tmp_path):
    # 실기 C7: 같은 경로에 파일을 새로 만들어도 이전 파일의 스냅샷이 undo 대상에 남았음
    path = str(tmp_path / "A.xlsx")
    _xlsx(path, (2026, 1, 1, 9, 0, 0))
    off.save_snapshot("excel", FakeWorkbook(path))
    snaps = off.list_snapshots("excel")
    assert len(snaps) == 1 and snaps[0]["doc_created"] == off.doc_identity(path) and "2026-01-01" in snaps[0]["doc_created"]

    import openpyxl                                         # 같은 문서를 다시 저장: 만든 시각 유지 -> 스냅샷 유지
    wb = openpyxl.load_workbook(path)
    wb.active["A1"] = "수정"
    wb.save(path)
    assert len(off.list_snapshots("excel")) == 1

    _xlsx(path, (2026, 5, 5, 10, 0, 0))                     # 같은 경로에 새 문서 -> 옛 스냅샷 정리
    assert off.list_snapshots("excel") == []
    assert not [f for f in os.listdir(off._snapshot_dir()) if f.startswith("excel_")]


def test_snapshot_identity_unknown_keeps_snapshot(tmp_path):
    path = str(tmp_path / "old.xls")                        # OOXML이 아니면 비교하지 않음
    with open(path, "wb") as f:
        f.write(b"x")
    off.save_snapshot("excel", FakeWorkbook(path))
    snaps = off.list_snapshots("excel")
    assert len(snaps) == 1 and snaps[0]["doc_created"] is None


def test_snapshot_rotation_and_size_limit(tmp_path, monkeypatch):
    path = tmp_path / "a.xlsx"
    path.write_bytes(b"x")
    for _ in range(off.SNAPSHOT_KEEP + 3):
        off.save_snapshot("excel", FakeWorkbook(str(path)))
    assert len(off.list_snapshots("excel")) == off.SNAPSHOT_KEEP
    monkeypatch.setattr(off, "SNAPSHOT_MAX_BYTES", 0)
    assert off.save_snapshot("excel", FakeWorkbook(str(path))) is None


def test_excel_batch_restores_settings():
    app = FakeExcelApp()
    with off.excel_batch(app):
        assert app.ScreenUpdating is False and app.Calculation == -4135 and app.EnableEvents is False
    assert app.ScreenUpdating is True and app.Calculation == -4105 and app.EnableEvents is True
    with pytest.raises(ValueError):
        with off.excel_batch(app):
            raise ValueError()
    assert app.ScreenUpdating is True


def test_range_to_matrix():
    assert off.range_to_matrix(3) == [[3]]
    assert off.range_to_matrix((("a", "b"),)) == [["a", "b"]]
    assert off.range_to_matrix((("a",), ("b",))) == [["a"], ["b"]]


def test_require_com_message(monkeypatch):
    monkeypatch.setattr(off, "win32com", None)
    with pytest.raises(rt.InputError) as ei:
        off.require_com()
    assert "reader.py" in str(ei.value)
