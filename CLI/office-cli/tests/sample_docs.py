"""테스트용 새 docx/pptx 만들기.

회사 PC(DRM)에서는 site-packages 안의 라이브러리 기본 템플릿(pptx/templates/default.pptx,
docx/templates/default.docx)이 암호화되어 zip이 아니게 될 수 있습니다. 그러면 인자 없는
pptx.Presentation() / docx.Document()가 PackageNotFoundError로 실패합니다.
- docx: 같은 내용의 폴더 템플릿(docx/templates/default-docx-template)으로 대신 만듭니다.
  회사 PC에서는 폴더 템플릿의 docProps/thumbnail.jpeg도 지워져 있었으므로, 빠진 파일을 가리키는 관계는 빼고 엽니다.
- pptx: 폴더 템플릿이 없으므로 그 테스트만 건너뜁니다 (나머지 테스트는 그대로 실행).
"""
import os
import zipfile

import pytest

PPTX_SKIP_REASON = "python-pptx 기본 템플릿이 손상됨 (DRM 추정)"
DOCX_SKIP_REASON = "python-docx 기본 템플릿이 손상됨 (DRM 추정)"


def docx_template_paths():
    """(기본 .docx 파일, 폴더 템플릿) 경로."""
    import docx
    base = os.path.join(os.path.dirname(docx.__file__), "templates")
    return os.path.join(base, "default.docx"), os.path.join(base, "default-docx-template")


def pptx_template_path():
    import pptx
    return os.path.join(os.path.dirname(pptx.__file__), "templates", "default.pptx")


def pptx_template_ok():
    return zipfile.is_zipfile(pptx_template_path())


def docx_from_folder(folder):
    """폴더 템플릿을 메모리 속 .docx로 묶어 엽니다. 관계(.rels)가 가리키는데 없는 파일은 관계째 뺍니다."""
    import io
    import posixpath
    import xml.etree.ElementTree as ET
    import docx
    files = {}
    for root, _dirs, names in os.walk(folder):
        for n in names:
            full = os.path.join(root, n)
            with open(full, "rb") as f:
                files[os.path.relpath(full, folder).replace(os.sep, "/")] = f.read()
    ET.register_namespace("", "http://schemas.openxmlformats.org/package/2006/relationships")
    for rels in [k for k in files if k.endswith(".rels")]:
        base = posixpath.dirname(posixpath.dirname(rels))
        tree = ET.fromstring(files[rels])
        for rel in list(tree):
            target = rel.get("Target", "")
            part = target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join(base, target))
            if rel.get("TargetMode") != "External" and part not in files:
                tree.remove(rel)
        files[rels] = ET.tostring(tree, xml_declaration=True, encoding="UTF-8")
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for name in sorted(files, key=lambda k: (k != "[Content_Types].xml", k)):
            z.writestr(name, files[name])
    buf.seek(0)
    return docx.Document(buf)


def new_docx():
    """빈 python-docx 문서. 기본 .docx가 zip이 아니면 폴더 템플릿을 씁니다."""
    import docx
    default, folder = docx_template_paths()
    if zipfile.is_zipfile(default):
        return docx.Document()
    if os.path.isdir(folder):
        return docx_from_folder(folder)
    pytest.skip(DOCX_SKIP_REASON)


def new_pptx():
    """빈 python-pptx 프레젠테이션. 기본 템플릿이 zip이 아니면 이 테스트를 건너뜁니다."""
    import pptx
    if not pptx_template_ok():
        pytest.skip(PPTX_SKIP_REASON)
    return pptx.Presentation()
