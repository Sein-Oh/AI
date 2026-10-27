"""reader.py 테스트: 실제 샘플 파일(docx/pptx/xlsx/pdf/eml/hwpx/csv/txt)로 파싱, 캐시, 검색, CLI 계약을 검증합니다."""
import json
import os
import subprocess
import sys
import zipfile

import pytest

import _runtime as rt
import reader

TOOL_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


# ---------------------------------------------------------------------------
# 샘플 파일 생성
# ---------------------------------------------------------------------------

def make_docx(path):
    import docx
    d = docx.Document()
    d.add_heading("추진 배경", level=1)
    d.add_paragraph("2026년 공정 최적화 마일스톤을 정리합니다.")
    t = d.add_table(rows=2, cols=3)
    a = t.cell(0, 0).merge(t.cell(0, 1))
    a.text = "병합 헤더"
    t.cell(0, 2).text = "값"
    t.cell(1, 0).text = "가"
    t.cell(1, 1).text = "나"
    t.cell(1, 2).text = "100"
    d.add_heading("일정", level=1)
    d.add_paragraph("3분기 양산 적용")
    d.save(path)


def make_pptx(path):
    import pptx
    from pptx.util import Inches
    prs = pptx.Presentation()
    for i, title in enumerate(["첫 슬라이드", "두 번째 슬라이드", "세 번째 슬라이드"], 1):
        s = prs.slides.add_slide(prs.slide_layouts[5])  # 제목만
        s.shapes.title.text = title
        box = s.shapes.add_textbox(Inches(1), Inches(2), Inches(5), Inches(1))
        tf = box.text_frame
        p = tf.paragraphs[0]
        r1 = p.add_run()
        r1.text = "Hel"
        r2 = p.add_run()
        r2.text = f"lo 본문{i}"
        if i == 2:
            rows = s.shapes.add_table(2, 2, Inches(1), Inches(3), Inches(4), Inches(1)).table
            rows.cell(0, 0).text = "항목"
            rows.cell(0, 1).text = "케이블카"
            s.notes_slide.notes_text_frame.text = "발표 노트"
    # 슬라이드 순서 바꾸기: 3번째를 맨 앞으로 (파일 번호 순서와 실제 순서가 달라짐)
    sld_ids = prs.slides._sldIdLst
    last = sld_ids[-1]
    sld_ids.remove(last)
    sld_ids.insert(0, last)
    prs.save(path)


def make_xlsx(path):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "매출"
    ws.append(["지역", "금액"])
    ws.append(["서울", 1200])
    ws.append([None, None])
    ws.append(["부산", 800.0])
    ws2 = wb.create_sheet("요약")
    ws2.append(["합계", "=SUM(매출!B2:B4)"])
    wb.save(path)


def make_pdf(path):
    import pymupdf
    doc = pymupdf.open()
    for text in ("Phu Quoc cable car 7.9km", "Safari tour schedule"):
        page = doc.new_page()
        page.insert_text((72, 72), text)
    doc.save(path)
    doc.close()


def make_eml(path):
    body = "안녕하세요.\n\n견적 단가를 회신드립니다.".encode("cp949")
    raw = (
        b"Subject: =?ks_c_5601-1987?B?" + __import__("base64").b64encode("견적 회신".encode("cp949")) + b"?=\r\n"
        b"From: " + "김팀장".encode("cp949") + b" <kim@example.com>\r\n"
        b"To: lee@example.com\r\n"
        b"Date: Tue, 27 Oct 2026 09:00:00 +0900\r\n"
        b"MIME-Version: 1.0\r\n"
        b"Content-Type: multipart/mixed; boundary=XX\r\n\r\n"
        b"--XX\r\nContent-Type: text/plain\r\nContent-Transfer-Encoding: 8bit\r\n\r\n" + body + b"\r\n"
        b"--XX\r\nContent-Type: message/rfc822\r\n\r\nSubject: old\r\n\r\nforwarded body should not appear\r\n"
        b"--XX\r\nContent-Type: application/octet-stream\r\nContent-Disposition: attachment; filename=\"a.xlsx\"\r\n\r\nAAAA\r\n"
        b"--XX--\r\n"
    )
    with open(path, "wb") as f:
        f.write(raw)


def make_hwpx(path):
    ns = 'xmlns:hp="http://www.hancom.co.kr/hwpml/2011/paragraph"'
    sec0 = (f'<?xml version="1.0" encoding="UTF-8"?><hs:sec xmlns:hs="x" {ns}>'
            '<hp:p><hp:run><hp:t>결재 </hp:t><hp:t>일자</hp:t></hp:run>'
            '<hp:run><hp:tbl><hp:tr><hp:tc><hp:subList><hp:p><hp:run><hp:t>표 안 문단</hp:t></hp:run></hp:p>'
            '</hp:subList></hp:tc></hp:tr></hp:tbl></hp:run></hp:p>'
            '<hp:p><hp:run><hp:t>둘째<hp:tab/>문단</hp:t></hp:run></hp:p></hs:sec>')
    sec10 = f'<?xml version="1.0" encoding="UTF-8"?><hs:sec xmlns:hs="x" {ns}><hp:p><hp:run><hp:t>열번째 구역</hp:t></hp:run></hp:p></hs:sec>'
    sec2 = f'<?xml version="1.0" encoding="UTF-8"?><hs:sec xmlns:hs="x" {ns}><hp:p><hp:run><hp:t>두번째 구역</hp:t></hp:run></hp:p></hs:sec>'
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("mimetype", "application/hwp+zip")
        z.writestr("Contents/section0.xml", sec0)
        z.writestr("Contents/section10.xml", sec10)
        z.writestr("Contents/section2.xml", sec2)


@pytest.fixture
def samples(tmp_path):
    files = {
        "docx": tmp_path / "계획서.docx", "pptx": tmp_path / "발표.pptx", "xlsx": tmp_path / "매출.xlsx",
        "pdf": tmp_path / "여행.pdf", "eml": tmp_path / "메일.eml", "hwpx": tmp_path / "양식.hwpx",
    }
    make_docx(str(files["docx"]))
    make_pptx(str(files["pptx"]))
    make_xlsx(str(files["xlsx"]))
    make_pdf(str(files["pdf"]))
    make_eml(str(files["eml"]))
    make_hwpx(str(files["hwpx"]))
    return {k: str(v) for k, v in files.items()}


# ---------------------------------------------------------------------------
# 파서
# ---------------------------------------------------------------------------

def texts(units):
    return [u["text"] for u in units]


def test_docx_order_heading_and_merged_cells(samples):
    units, meta = reader.parse_docx(samples["docx"])
    assert texts(units)[0] == "추진 배경"
    table = [u for u in units if u["loc"] == "표 1"][0]
    assert table["text"].splitlines()[0] == "병합 헤더 | 값"      # 병합 셀 중복 없음
    assert table["heading"] == "추진 배경"
    assert [u for u in units if u["text"] == "3분기 양산 적용"][0]["heading"] == "일정"
    assert meta["tables"] == 1


def test_pptx_real_slide_order_runs_table_notes(samples):
    units, meta = reader.parse_pptx(samples["pptx"])
    assert [u["heading"] for u in units] == ["세 번째 슬라이드", "첫 슬라이드", "두 번째 슬라이드"]
    assert "Hello 본문1" in units[1]["text"]          # 서식 경계에서 단어가 쪼개지지 않음
    assert "항목 | 케이블카" in units[2]["text"]
    assert "[노트] 발표 노트" in units[2]["text"]


def test_xlsx_rows_skip_empty_and_numbers(samples):
    units, meta = reader.parse_xlsx(samples["xlsx"])
    assert meta["sheets"] == ["매출", "요약"]
    assert [u["loc"] for u in units[:3]] == ["매출!1행", "매출!2행", "매출!4행"]
    assert units[2]["text"] == "부산\t800"


def test_pdf_pages(samples):
    units, meta = reader.parse_pdf(samples["pdf"])
    assert meta["pages"] == 2 and units[0]["loc"] == "p.1"
    assert "7.9km" in units[0]["text"]


def test_eml_charset_headers_and_forwarded(samples):
    units, meta = reader.parse_eml(samples["eml"])
    assert meta["subject"] == "견적 회신"
    assert meta["from"].startswith("김팀장")
    joined = "\n".join(texts(units))
    assert "견적 단가를 회신드립니다." in joined
    assert "forwarded body" not in joined
    assert "a.xlsx" in meta["attachments"]


def test_hwpx_paragraphs_nested_tables_and_section_order(samples):
    units, _ = reader.parse_hwpx(samples["hwpx"])
    assert texts(units) == ["결재 일자", "표 안 문단", "둘째\t문단", "두번째 구역", "열번째 구역"]


def test_hwp_para_text_control_characters():
    def wchar(c):
        return c.to_bytes(2, "little")
    ext_ctrl = wchar(11) + b"\x00" * 14        # 표/개체 (확장 컨트롤, 8 WCHAR)
    sec_ctrl = wchar(2) + b"dces" + b"\x00" * 10  # 구역 정의 (확장 컨트롤)
    rec = (sec_ctrl + "가나".encode("utf-16-le") + ext_ctrl + "다".encode("utf-16-le")
           + wchar(30) + "라".encode("utf-16-le") + wchar(9) + b"\x00" * 14 + "마".encode("utf-16-le") + wchar(13))
    assert reader._hwp_para_text(rec) == "가나다 라\t마"


def test_encrypted_office_file_gives_drm_hint(tmp_path):
    fake = tmp_path / "보안.docx"
    fake.write_bytes(b"DRMONE\x00\x01" + os.urandom(64))
    with pytest.raises(rt.InputError) as ei:
        reader.parse_docx(str(fake))
    assert "DRM" in str(ei.value)


def test_old_office_format_is_redirected():
    with pytest.raises(rt.InputError) as ei:
        reader.file_type("a.xls")
    assert "excel.py" in str(ei.value)


def test_matcher_and_or_and_snippet():
    m = reader.build_matcher("케이블카 AND 사파리 OR 단가", False, False)
    assert m("케이블카와 사파리") == ["케이블카", "사파리"]
    assert m("케이블카만") is None
    assert m("견적 단가") == ["단가"]
    snip = reader.make_snippet("가" * 200 + "단가표" + "나" * 200, ["단가"], 10, False)
    assert snip.startswith("...") and "**단가**" in snip


def test_cache_reused_and_invalidated(samples, monkeypatch):
    calls = []
    orig = reader.PARSERS["docx"]
    monkeypatch.setitem(reader.PARSERS, "docx", lambda p: (calls.append(p), orig(p))[1])
    reader.load_document(samples["docx"])
    reader.load_document(samples["docx"])
    assert len(calls) == 1
    os.utime(samples["docx"], (1, 1))
    reader.load_document(samples["docx"])
    assert len(calls) == 2
    assert os.path.abspath(samples["docx"]) in reader.load_index()


# ---------------------------------------------------------------------------
# CLI 계약
# ---------------------------------------------------------------------------

def run(tmp_env, *args):
    proc = subprocess.run([sys.executable, os.path.join(TOOL_DIR, "reader.py"), *args],
                          capture_output=True, timeout=120, env=tmp_env)
    lines = [ln for ln in proc.stdout.decode("utf-8").splitlines() if ln.strip()]
    assert len(lines) == 1, proc.stdout.decode("utf-8") + proc.stderr.decode("utf-8", "replace")
    return proc.returncode, json.loads(lines[0])


@pytest.fixture
def env(tmp_path):
    return dict(os.environ, OFFICE_CLI_DATA_DIR=str(tmp_path / "data"), OFFICE_CLI_CACHE_DIR=str(tmp_path / "cache"))


def test_cli_read_pages_and_summary(samples, env):
    code, out = run(env, "read", "--file", samples["pptx"], "--pages", "2")
    assert code == 0 and out["status"] == "ok"
    assert out["units_total"] == 3 and [u["no"] for u in out["units"]] == [2]
    code, out = run(env, "read", "--file", samples["pdf"], "--format", "summary")
    assert code == 0 and "units" not in out and out["meta"]["pages"] == 2


def test_cli_read_save_gives_write_commands(samples, env):
    code, out = run(env, "read", "--file", samples["docx"], "--save")
    assert code == 0 and os.path.exists(out["output_file"]) and out["output_file"].endswith(".md")
    assert "word.py" in out["next_commands"][0]["command"]
    text = open(out["output_file"], encoding="utf-8").read()
    assert "## 표 1 - 추진 배경" in text


def test_cli_read_truncation_offers_next_page(tmp_path, env):
    big = tmp_path / "긴문서.txt"
    big.write_text("\n\n".join(f"{i}번 단락 " + "가" * 500 for i in range(1, 60)), encoding="utf-8")
    code, out = run(env, "read", "--file", str(big))
    assert out["truncated"] is True
    assert "--pages" in out["next_commands"][0]["command"]


def test_cli_search_index_folder_save_and_no_match(samples, env, tmp_path):
    for k in ("docx", "pptx", "pdf"):
        run(env, "read", "--file", samples[k], "--format", "summary")
    code, out = run(env, "search", "--query", "케이블카")
    assert code == 0 and out["files_searched"] == 3 and out["total_matches"] == 1
    assert out["matches"][0]["file"] == "발표.pptx" and out["matches"][0]["loc"] == "슬라이드 3"
    code, out = run(env, "search", "--query", "cable OR 양산")
    assert out["total_matches"] == 2

    code, out = run(env, "search", "--query", "마일스톤", "--save")
    assert out["matched"] == 1 and os.path.exists(out["output_file"])
    assert "word.py" in out["next_commands"][0]["command"]

    code, out = run(env, "search", "--query", "단가", "--file", str(tmp_path), "--save", "결과.csv")
    assert out["total_matches"] >= 1 and out["output_file"].endswith("결과.csv")
    assert "excel.py" in out["next_commands"][0]["command"]

    code, out = run(env, "search", "--query", "존재하지않는단어")
    assert code == 0 and out["total_matches"] == 0 and out["agent_directive"].startswith("VERIFIED:")


def test_cli_search_without_index_is_stop_retry(env):
    code, out = run(env, "search", "--query", "무엇")
    assert code == 1 and out["agent_directive"].startswith("STOP_RETRY:")


def test_cli_list(samples, env):
    run(env, "read", "--file", samples["eml"], "--format", "summary")
    code, out = run(env, "list")
    assert code == 0 and out["total"] == 1 and out["documents"][0]["file"] == "메일.eml"


def test_cli_missing_file_suggests(samples, env, tmp_path):
    data = tmp_path / "data"
    data.mkdir(exist_ok=True)
    (data / "보고서_최종.pdf").write_bytes(b"x")
    code, out = run(env, "read", "--file", "보고서_최종.docx")
    assert code == 1 and "[추천] 보고서_최종.pdf" in out["message"]


def test_cli_argument_error(env):
    code, out = run(env, "read")
    assert code == 2 and out["status"] == "error"
