"""reader.py 1.3.0 테스트: 공통 캐시 형식, 장절/번호 인식, 쪽 번호, 표 행 단위, 캐시 기준 검색과 근거 출력.

샘플 파일은 테스트 안에서 만듭니다 (바이너리 고정 파일 없음).
"""
import json
import os
import subprocess
import sys
import zipfile

import pytest

import reader
from sample_docs import new_docx, new_pptx

TOOL_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
W_NS = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
UNIT_KEYS = {"id", "kind", "loc", "page", "page_source", "section", "number", "heading", "text"}


def run(*args):
    """reader.py를 하위 프로세스로 실행 (캐시/데이터 폴더는 conftest가 정한 환경변수)."""
    proc = subprocess.run([sys.executable, os.path.join(TOOL_DIR, "reader.py"), *args],
                          capture_output=True, timeout=120, env=dict(os.environ))
    lines = [ln for ln in proc.stdout.decode("utf-8").splitlines() if ln.strip()]
    assert len(lines) == 1, proc.stdout.decode("utf-8") + proc.stderr.decode("utf-8", "replace")
    return proc.returncode, json.loads(lines[0])


def by_text(units, start):
    return [u for u in units if u["text"].startswith(start)][0]


# ---------------------------------------------------------------------------
# 샘플 만들기
# ---------------------------------------------------------------------------

def make_pdf(path, pages, font="korea"):
    """pages: [[(글, 크기, 굵게)], ...]. 표는 ("TABLE", rows) 항목으로 선을 그려 넣습니다."""
    import pymupdf
    doc = pymupdf.open()
    for items in pages:
        page = doc.new_page()
        y = 72
        for item in items:
            if item[0] == "TABLE":
                rows = item[1]
                x0, y0, cw, rh = 72, y, 150, 22
                for i in range(len(rows) + 1):
                    page.draw_line((x0, y0 + i * rh), (x0 + len(rows[0]) * cw, y0 + i * rh))
                for j in range(len(rows[0]) + 1):
                    page.draw_line((x0 + j * cw, y0), (x0 + j * cw, y0 + len(rows) * rh))
                for i, r in enumerate(rows):
                    for j, c in enumerate(r):
                        page.insert_text((x0 + j * cw + 4, y0 + i * rh + 15), c, fontname=font, fontsize=10)
                y = y0 + len(rows) * rh + 20
                continue
            text, size, bold = item
            fontname = ("hebo" if bold else "helv") if font == "helv" else font
            page.insert_text((72, y), text, fontname=fontname, fontsize=size)
            y += size + 8
    doc.save(str(path))
    doc.close()


def korean_pdf(path):
    make_pdf(path, [
        [("III. 개발 내용", 16, False), ("1. 개발 내용 종합", 14, False), ("가. 검증기술", 12, False),
         ("O 프로그램 전체 연구 목표", 10, False), ("검증 절차를 수행한다.", 10, False),
         ("TABLE", [["항목", "기준", "시험방법"], ["내전압", "1.5 kV 1분", "KS C 1234"], ["절연저항", "100 MΩ 이상", "메거"]])],
        [("나. 시험평가", 12, False), ("시험평가 본문입니다.", 10, False), ("2. 기대 효과", 14, False),
         ("기대 효과 본문입니다.", 10, False), ("IV. 활용 계획", 16, False), ("1. 사업화", 14, False)],
    ])


def english_pdf(path):
    make_pdf(path, [
        [("1. SCOPE", 16, True), ("This spec covers data.", 10, False), ("1.1. Scope", 13, True),
         ("1.1.1. Data", 11, True), ("Data body text.", 10, False)],
        [("1.2. Purpose", 13, True), ("Purpose body.", 10, False), ("2 REQUIREMENTS", 16, True),
         ("2.1 General", 13, True), ("Requirement body.", 10, False)],
    ], font="helv")


def hash_pdf(path):
    make_pdf(path, [
        [("1.0# Introduction", 14, False), ("intro text here", 10, False), ("1.1# Purpose", 12, False),
         ("1.1.1# Information", 11, False), ("info text", 10, False)],
        [("2.0# Design", 14, False), ("design text", 10, False)],
    ], font="helv")


def _numbering_part(d):
    """문서의 번호 매기기 파트. 없으면 만듭니다: 회사 PC처럼 python-docx 기본 템플릿이 보안 프로그램으로 바뀌어
    numbering.xml이 없으면 python-docx가 NotImplementedError를 냄 (2026-10-29 실기)."""
    try:
        return d.part.numbering_part
    except (KeyError, NotImplementedError):
        from docx.opc.constants import CONTENT_TYPE as CT, RELATIONSHIP_TYPE as RT
        from docx.opc.packuri import PackURI
        from docx.oxml import parse_xml
        from docx.parts.numbering import NumberingPart
        part = NumberingPart(PackURI("/word/numbering.xml"), CT.WML_NUMBERING, parse_xml(f"<w:numbering {W_NS}/>"),
                             d.part.package)
        d.part.relate_to(part, RT.NUMBERING)
        return d.part.numbering_part


def add_numbering(d):
    """numbering.xml에 3.로 시작하는 다단계 번호(3. / 3.1. / 3.1.1 / 가))를 넣고 numId 90을 돌려줍니다."""
    from docx.oxml import parse_xml
    numbering = _numbering_part(d).element
    numbering.insert(0, parse_xml(
        f'<w:abstractNum {W_NS} w:abstractNumId="90">'
        '<w:lvl w:ilvl="0"><w:start w:val="3"/><w:numFmt w:val="decimal"/><w:lvlText w:val="%1."/></w:lvl>'
        '<w:lvl w:ilvl="1"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:lvlText w:val="%1.%2."/></w:lvl>'
        '<w:lvl w:ilvl="2"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:lvlText w:val="%1.%2.%3"/></w:lvl>'
        '<w:lvl w:ilvl="3"><w:start w:val="1"/><w:numFmt w:val="ganada"/><w:lvlText w:val="%4)"/></w:lvl>'
        '</w:abstractNum>'))
    numbering.append(parse_xml(f'<w:num {W_NS} w:numId="90"><w:abstractNumId w:val="90"/></w:num>'))


def numbered(d, text, ilvl):
    from docx.oxml import parse_xml
    p = d.add_paragraph(text)
    p._p.get_or_add_pPr().append(parse_xml(f'<w:numPr {W_NS}><w:ilvl w:val="{ilvl}"/><w:numId w:val="90"/></w:numPr>'))
    return p


def spec_docx(path, saved_layout=False):
    """제목 스타일 + 자동 번호 + 일반 번호 단락 + 강제 쪽 나눔 + 요구사항 표."""
    from docx.enum.text import WD_BREAK
    from docx.oxml import parse_xml
    d = new_docx()
    add_numbering(d)
    d.add_heading("시험 규격", level=1)
    d.add_paragraph("이 규격은 시험 방법을 정한다.")
    numbered(d, "일반 사항", 0)
    numbered(d, "적용 범위", 1)
    numbered(d, "세부 조건", 2)
    numbered(d, "온도 조건", 3)
    d.add_paragraph("시험 온도는 25도로 한다.")
    d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    d.add_heading("요구사항", level=1)
    d.add_paragraph("가. 검증기술")
    d.add_paragraph("O 프로그램 전체 연구 목표")
    t = d.add_table(rows=3, cols=3)
    for i, row in enumerate([["항목", "기준", "시험방법"], ["내전압", "1.5 kV 1분", "KS C\n1234"],
                             ["절연저항", "100 MΩ 이상", "메거"]]):
        for j, v in enumerate(row):
            t.cell(i, j).text = v
    d.add_paragraph("마지막 쪽 문단")
    if saved_layout:
        # Word가 저장한 쪽 경계 흉내: 3쪽 시작 (2쪽은 쪽 나눔 뒤에 Word가 기록)
        for p in d.paragraphs:
            if p.text in ("요구사항", "마지막 쪽 문단"):
                p.runs[0]._r.insert(0, parse_xml(f'<w:lastRenderedPageBreak {W_NS}/>'))
    d.save(str(path))


# ---------------------------------------------------------------------------
# 번호 분류와 제목 판정
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("line,family,number,depth", [
    ("III. 개발 내용", "roman", "III", 1),
    ("\u2162. 개발 내용", "roman", "III", 1),                 # 전각 로마 숫자
    ("1. SCOPE", "arabic", "1", 1),
    ("1.1. Scope", "arabic", "1.1", 2),
    ("1.1 Scope", "arabic", "1.1", 2),
    ("1.1.1. Data", "arabic", "1.1.1", 3),
    ("1.0# Introduction", "arabic", "1.0", 1),
    ("1.1# Purpose", "arabic", "1.1", 2),
    ("가. 검증기술", "ganada", "가", 1),
    ("(1) 시험 조건", "paren_arabic", "(1)", 1),
    ("1) 시험 조건", "rparen_arabic", "1)", 1),
    ("(가) 세부", "paren_ganada", "(가)", 1),
    ("가) 세부", "rparen_ganada", "가)", 1),
    ("\u2460 첫째 항", "circled", "(1)", 1),                   # 원문자 입력 -> 출력은 (1)
    ("제1장 총칙", "ko_jang", "제1장", 1),
    ("제2절 적용", "ko_jeol", "제2절", 1),
    ("제12조(목적) 이 규정은 시험 절차를 정한다.", "ko_jo", "제12조", 1),
    ("Chapter 1 Overview", "chapter", "Chapter 1", 1),
    ("Annex A Test data", "annex", "Annex A", 1),
    ("부록 1 시험 양식", "annex_ko", "부록 1", 1),
    ("Q3. 만족도는?", "question", "Q3", 1),
    ("A. Purpose", "alpha_upper", "A", 1),
])
def test_classify_number_families(line, family, number, depth):
    info = reader.classify_number(line)
    assert info and info["family"] == family and info["number"] == number and info["depth"] == depth


@pytest.mark.parametrize("line", ["O 프로그램 전체 연구 목표", "- 하이픈 항목", "\u25cb 동그라미 항목", "\u25a1 네모 항목",
                                  "\u00b7 가운뎃점 항목", "* 별표 항목"])
def test_bullets_have_no_number(line):
    info = reader.classify_number(line)
    assert info["family"] == "bullet" and info["number"] is None
    assert reader.heading_title(info, line) is None


@pytest.mark.parametrize("line", ["2026. 10. 28. 회의", "3.5kg 이하", "본문 1. 안의 번호", "2026년 계획"])
def test_numbers_inside_text_or_dates_are_not_numbers(line):
    info = reader.classify_number(line)
    assert info is None or info["plain"] or info["number"] is None


@pytest.mark.parametrize("line,is_heading", [
    ("1. 개발 내용 종합", True),
    ("1. 시스템은 입력을 검증한다.", False),            # 문장 끝
    ("가. 성능을 확인하였음.", False),
    ("가. 시험 결과를 별도로 정리함", False),           # 개조식 문장 끝
    ("가. 소음", True),                                # 짧은 명사는 제목
    ("1. 개요", True),
    ("1. " + "가" * 90, False),                         # 너무 김
    ("2 REQUIREMENTS", False),                          # 구분 기호 없는 숫자는 굵게/큰 글씨일 때만
])
def test_heading_guard(line, is_heading):
    assert bool(reader.heading_title(reader.classify_number(line), line)) is is_heading


def test_plain_number_heading_needs_strong_signal():
    line = "2 REQUIREMENTS"
    assert reader.heading_title(reader.classify_number(line), line, strong=True) == line


def test_article_heading_keeps_number_and_title_only():
    line = "제12조(목적) 이 규정은 시험 절차를 정한다."
    assert reader.heading_title(reader.classify_number(line), line) == "제12조(목적)"


def test_outline_first_appearance_levels():
    b = reader.UnitBuilder()
    for t in ["III. 개발 내용", "1. 개발 내용 종합", "가. 검증기술", "O 프로그램 전체 연구 목표", "나. 시험", "2. 기대 효과",
              "IV. 활용", "1. 사업화"]:
        b.add_text("x", t)
    sec = {u["text"]: u["section"] for u in b.units}
    assert sec["O 프로그램 전체 연구 목표"] == ["III. 개발 내용", "1. 개발 내용 종합", "가. 검증기술"]
    assert sec["나. 시험"] == ["III. 개발 내용", "1. 개발 내용 종합", "나. 시험"]
    assert sec["2. 기대 효과"] == ["III. 개발 내용", "2. 기대 효과"]
    assert sec["1. 사업화"] == ["IV. 활용", "1. 사업화"]
    bullet = [u for u in b.units if u["text"].startswith("O ")][0]
    assert bullet["kind"] == "paragraph" and bullet["number"] is None


# ---------------------------------------------------------------------------
# 표 행
# ---------------------------------------------------------------------------

def test_table_rows_header_labels_merged_and_multiline():
    rows = reader.table_rows([["항목", "기준", "", "시험방법"],
                              ["내전압", "1.5 kV", "1분", "KS C\n1234"],
                              ["", "", "", ""],
                              ["절연저항", "100 MΩ 이상", None, "메거"]])
    assert [r["row"] for r in rows] == [2, 3]
    assert rows[0]["text"] == "항목: 내전압 | 기준: 1.5 kV | 기준: 1분 | 시험방법: KS C 1234"
    assert rows[1]["text"] == "항목: 절연저항 | 기준: 100 MΩ 이상 | 시험방법: 메거"


def test_table_rows_without_header_and_row_number():
    rows = reader.table_rows([["1", "내전압 1.5 kV 이상이어야 하며 1분간 견딜 것"], ["REQ-002", "절연저항 100 MΩ 이상"]])
    assert rows[0]["text"] == "1 | 내전압 1.5 kV 이상이어야 하며 1분간 견딜 것" and rows[0]["number"] == "1"
    assert rows[1]["number"] == "REQ-002"


# ---------------------------------------------------------------------------
# 형식별 파서: 공통 키, 장절, 번호, 쪽
# ---------------------------------------------------------------------------

def test_pdf_korean_sections_pages_table_rows(tmp_path):
    f = tmp_path / "계획서.pdf"
    korean_pdf(f)
    units, meta = reader.parse_pdf(str(f))
    assert all(set(u) == UNIT_KEYS for u in units)
    assert [u["id"] for u in units] == list(range(1, len(units) + 1))
    bullet = by_text(units, "O 프로그램")
    assert bullet["section"] == ["III. 개발 내용", "1. 개발 내용 종합", "가. 검증기술"] and bullet["number"] is None
    assert bullet["page"] == 1 and bullet["page_source"] == "pdf" and bullet["loc"] == "p.1"
    row = [u for u in units if u["kind"] == "row" and "내전압" in u["text"]][0]
    assert row["text"] == "항목: 내전압 | 기준: 1.5 kV 1분 | 시험방법: KS C 1234"
    assert row["section"][-1] == "가. 검증기술" and row["page"] == 1
    assert not [u for u in units if u["kind"] == "paragraph" and "KS C 1234" in u["text"]]   # 표 글은 단락에 중복 없음
    h = by_text(units, "나. 시험평가")
    assert h["kind"] == "heading" and h["page"] == 2 and h["number"] == "나"
    assert by_text(units, "기대 효과 본문")["section"] == ["III. 개발 내용", "2. 기대 효과"]
    assert by_text(units, "1. 사업화")["section"] == ["IV. 활용 계획", "1. 사업화"]
    assert meta["pages"] == 2


def test_pdf_english_dotted_with_font_signal(tmp_path):
    f = tmp_path / "spec.pdf"
    english_pdf(f)
    units, _ = reader.parse_pdf(str(f))
    assert by_text(units, "Data body")["section"] == ["1. SCOPE", "1.1. Scope", "1.1.1. Data"]
    purpose = by_text(units, "1.2. Purpose")
    assert purpose["section"] == ["1. SCOPE", "1.2. Purpose"] and purpose["number"] == "1.2" and purpose["page"] == 2
    assert by_text(units, "Requirement body")["section"] == ["2 REQUIREMENTS", "2.1 General"]   # '2 ' 굵게 큰 글씨


def test_pdf_hash_numbering(tmp_path):
    f = tmp_path / "hash.pdf"
    hash_pdf(f)
    units, _ = reader.parse_pdf(str(f))
    info = by_text(units, "info text")
    assert info["section"] == ["1.0# Introduction", "1.1# Purpose", "1.1.1# Information"]
    assert by_text(units, "1.1.1# Information")["number"] == "1.1.1"
    design = by_text(units, "design text")
    assert design["section"] == ["2.0# Design"] and design["page"] == 2


def test_docx_styles_auto_numbering_page_breaks_and_table_rows(tmp_path):
    f = tmp_path / "규격.docx"
    spec_docx(f)
    units, meta = reader.parse_docx(str(f))
    assert all(set(u) == UNIT_KEYS for u in units)
    assert meta["page_source"] == "page_break"
    sub = by_text(units, "3.1.1 세부 조건")                    # Word 자동 번호를 글 앞에 붙임
    assert sub["number"] == "3.1.1"
    assert sub["section"] == ["시험 규격", "3. 일반 사항", "3.1. 적용 범위", "3.1.1 세부 조건"]
    assert by_text(units, "가) 온도 조건")["number"] == "가)"
    body = by_text(units, "시험 온도는")
    assert body["page"] == 1 and body["section"][-1] == "가) 온도 조건"
    h = by_text(units, "요구사항")
    assert h["kind"] == "heading" and h["page"] == 2 and h["section"] == ["요구사항"]
    assert by_text(units, "O 프로그램")["section"] == ["요구사항", "가. 검증기술"]
    row = [u for u in units if u["kind"] == "row" and "내전압" in u["text"]][0]
    assert row["text"] == "항목: 내전압 | 기준: 1.5 kV 1분 | 시험방법: KS C 1234"
    assert row["loc"] == "표 1 2행" and row["page"] == 2 and row["section"] == ["요구사항", "가. 검증기술"]


def test_docx_saved_layout_pages(tmp_path):
    f = tmp_path / "저장됨.docx"
    spec_docx(f, saved_layout=True)
    units, meta = reader.parse_docx(str(f))
    assert meta["page_source"] == "word_saved_layout"
    last = by_text(units, "마지막 쪽 문단")
    assert last["page"] == 3 and last["page_source"] == "word_saved_layout"
    assert by_text(units, "시험 온도는")["page"] == 1


def test_docx_without_page_info_has_null_pages(tmp_path):
    d = new_docx()
    d.add_paragraph("1. 개요")
    d.add_paragraph("본문입니다.")
    d.save(str(tmp_path / "a.docx"))
    units, meta = reader.parse_docx(str(tmp_path / "a.docx"))
    assert all(u["page"] is None and u["page_source"] is None for u in units)
    assert units[1]["section"] == ["1. 개요"]


def test_pptx_slide_section_and_page(tmp_path):
    prs = new_pptx()
    for title in ["1. 개요", "2. 일정"]:
        s = prs.slides.add_slide(prs.slide_layouts[5])
        s.shapes.title.text = title
    prs.save(str(tmp_path / "발표.pptx"))
    units, _ = reader.parse_pptx(str(tmp_path / "발표.pptx"))
    assert [(u["kind"], u["page"], u["page_source"], u["section"], u["number"]) for u in units] == [
        ("slide", 1, "slide", ["1. 개요"], "1"), ("slide", 2, "slide", ["2. 일정"], "2")]


def test_xlsx_rows_section_is_sheet(tmp_path):
    import openpyxl
    wb = openpyxl.Workbook()
    wb.active.title = "매출"
    wb.active.append(["1.1", "서울"])
    wb.save(str(tmp_path / "a.xlsx"))
    units, _ = reader.parse_xlsx(str(tmp_path / "a.xlsx"))
    u = units[0]
    assert (u["kind"], u["section"], u["page"], u["number"], u["loc"]) == ("row", ["매출"], None, None, "매출!1행")


def test_hwpx_outline_levels_and_numbers(tmp_path):
    hp = 'xmlns:hp="http://www.hancom.co.kr/hwpml/2011/paragraph"'
    hh = 'xmlns:hh="http://www.hancom.co.kr/hwpml/2011/head"'
    header = (f'<?xml version="1.0" encoding="UTF-8"?><hh:head {hh}><hh:refList>'
              '<hh:numberings><hh:numbering id="1" start="0">'
              '<hh:paraHead start="1" level="1" numFormat="DIGIT">^1.</hh:paraHead>'
              '<hh:paraHead start="1" level="2" numFormat="HANGUL_SYLLABLE">^2.</hh:paraHead>'
              '</hh:numbering></hh:numberings><hh:paraProperties>'
              '<hh:paraPr id="1"><hh:heading type="OUTLINE" idRef="1" level="0"/></hh:paraPr>'
              '<hh:paraPr id="2"><hh:heading type="OUTLINE" idRef="1" level="1"/></hh:paraPr>'
              '<hh:paraPr id="0"><hh:heading type="NONE" idRef="0" level="0"/></hh:paraPr>'
              '</hh:paraProperties></hh:refList></hh:head>')
    sec = (f'<?xml version="1.0" encoding="UTF-8"?><hs:sec xmlns:hs="x" {hp}>'
           '<hp:p paraPrIDRef="1"><hp:run><hp:t>개요</hp:t></hp:run></hp:p>'
           '<hp:p paraPrIDRef="2"><hp:run><hp:t>목적</hp:t></hp:run></hp:p>'
           '<hp:p paraPrIDRef="0"><hp:run><hp:t>본문 내용</hp:t></hp:run></hp:p>'
           '<hp:p paraPrIDRef="1" pageBreak="1"><hp:run><hp:t>범위</hp:t></hp:run></hp:p></hs:sec>')
    path = tmp_path / "양식.hwpx"
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("Contents/header.xml", header)
        z.writestr("Contents/section0.xml", sec)
    units, _ = reader.parse_hwpx(str(path))
    assert [u["text"] for u in units] == ["1. 개요", "가. 목적", "본문 내용", "2. 범위"]
    assert units[2]["section"] == ["1. 개요", "가. 목적"]
    assert units[3]["section"] == ["2. 범위"] and units[3]["number"] == "2" and units[3]["page"] == 2


# ---------------------------------------------------------------------------
# CLI: index / search / list
# ---------------------------------------------------------------------------

@pytest.fixture
def docs(tmp_path):
    folder = tmp_path / "규격"
    (folder / "sub").mkdir(parents=True)
    korean_pdf(folder / "계획서.pdf")
    english_pdf(folder / "spec.pdf")
    spec_docx(folder / "sub" / "규격.docx")
    (folder / "메일.eml").write_bytes(b"Subject: x\r\n\r\nbody")
    (folder / "그림.png").write_bytes(b"\x89PNG")
    return folder


def test_index_folder_counts_and_unsupported_skipped(docs):
    code, out = run("index", "--folder", str(docs))
    assert code == 0 and out["status"] == "ok"
    assert (out["count"], out["indexed"], out["updated"], out["unchanged"], out["failed"]) == (3, 3, 0, 0, 0)
    assert out["skipped_types"] == {".eml": 1, ".png": 1} and out["truncated"] is False
    code, out = run("index", "--folder", str(docs))
    assert (out["indexed"], out["unchanged"]) == (0, 3)
    os.utime(docs / "spec.pdf", (1, 1))
    code, out = run("index", "--folder", str(docs), "--depth", "0")
    assert (out["count"], out["updated"], out["unchanged"]) == (2, 1, 1)


def test_index_max_files_continues(docs):
    code, out = run("index", "--folder", str(docs), "--max-files", "1")
    assert out["indexed"] == 1 and out["remaining"] == 2 and out["truncated"] is True
    assert "--max-files" in out["next_commands"][0]["command"] and "읽기 전용" in out["agent_directive"]
    code, out = run("index", "--folder", str(docs), "--max-files", "1")
    assert out["indexed"] == 1 and out["unchanged"] == 1 and out["remaining"] == 1


def test_search_without_file_uses_cache_only(docs):
    code, out = run("search", "--query", "내전압")
    assert code == 1 and out["error_type"] == "not_indexed" and out["agent_directive"].startswith("STOP_RETRY:")
    run("index", "--folder", str(docs))
    code, out = run("search", "--query", "내전압")
    assert code == 0 and out["total_matches"] == 2 and "indexed_now" not in out
    hit = [m for m in out["matches"] if m["file"] == "규격.docx"][0]
    assert hit["section"] == ["요구사항", "가. 검증기술"] and hit["section_text"] == "요구사항 > 가. 검증기술"
    assert hit["page"] == 2 and hit["page_source"] == "page_break" and hit["loc"] == "표 1 2행"
    assert "기준: 1.5 kV 1분" in hit["snippet"]
    assert hit["evidence"] == "규격.docx p.2(추정), 요구사항 > 가. 검증기술"
    pdf_hit = [m for m in out["matches"] if m["file"] == "계획서.pdf"][0]
    assert pdf_hit["evidence"] == "계획서.pdf p.1, III. 개발 내용 > 1. 개발 내용 종합 > 가. 검증기술"


def test_search_file_auto_caches_then_uses_cache(docs):
    code, out = run("search", "--query", "Purpose", "--file", str(docs / "spec.pdf"))
    assert code == 0 and out["indexed_now"] == ["spec.pdf"] and out["from_cache"] == 0
    m = out["matches"][0]
    assert (m["page"], m["number"], m["section_text"]) == (2, "1.2", "1. SCOPE > 1.2. Purpose")
    assert m["evidence"] == "spec.pdf p.2, 1. SCOPE > 1.2. Purpose (1.2)"
    code, out = run("search", "--query", "Purpose", "--file", str(docs / "spec.pdf"))
    assert out["indexed_now"] == [] and out["from_cache"] == 1
    code, out = run("search", "--query", "내전압", "--file", str(docs))           # 폴더: 남은 2개만 새로 등록
    assert out["indexed_now_count"] == 2 and out["from_cache"] == 1 and out["skipped_types"][".eml"] == 1
    code, out = run("list", "--name", "SPEC")
    assert out["count"] == 1 and out["documents"][0]["file"] == "spec.pdf"


def test_search_folder_respects_max_files(docs):
    code, out = run("search", "--query", "본문", "--file", str(docs), "--max-files", "1")
    assert out["indexed_now_count"] == 1 and out["not_indexed_remaining"] == 2
    assert any("index" in c["command"] and "--folder" in c["command"] for c in out["next_commands"])


def test_search_pages_and_section_filters(docs):
    run("index", "--folder", str(docs))
    code, out = run("search", "--query", "본문", "--file", str(docs / "계획서.pdf"), "--pages", "2")
    assert out["total_matches"] == 2 and all(m["page"] == 2 for m in out["matches"])
    code, out = run("search", "--query", "본문", "--section", "기대 효과")
    assert [m["section_text"] for m in out["matches"]] == ["III. 개발 내용 > 2. 기대 효과"]
    code, out = run("search", "--query", "온도", "--pages", "1-3")
    assert out["total_matches"] >= 1 and all(m["page"] in (1, 2, 3) for m in out["matches"])
    code, out = run("search", "--query", "본문", "--section", "없는 장절")
    assert code == 0 and out["count"] == 0 and out["agent_directive"].startswith("VERIFIED: 0건.")


def test_search_save_xlsx_columns_and_md_evidence(docs, tmp_path):
    import openpyxl
    run("index", "--folder", str(docs))
    out_path = tmp_path / "근거.xlsx"
    code, out = run("search", "--query", "내전압 OR Purpose", "--save", str(out_path), "--top", "1")
    assert code == 0 and out["output_file"] == str(out_path) and out["saved"] == out["total_matches"] == 4
    wb = openpyxl.load_workbook(out["output_file"])
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    assert list(rows[0]) == ["파일", "쪽", "장절", "번호", "위치", "내용(문장/단락)", "문맥"]
    assert ws["A1"].font.bold and ws["F2"].alignment.wrap_text
    docx_row = [r for r in rows if r[0] == "규격.docx"][0]
    assert docx_row[1] == 2 and docx_row[2] == "요구사항 > 가. 검증기술" and docx_row[4] == "표 1 2행"
    purpose = [r for r in rows if r[0] == "spec.pdf" and r[3] == "1.2"][0]
    assert purpose[1] == 2 and purpose[2] == "1. SCOPE > 1.2. Purpose"
    code, out = run("search", "--query", "Purpose", "--save", str(tmp_path / "근거.md"))
    text = open(out["output_file"], encoding="utf-8").read()
    assert "근거: spec.pdf p.2, 1. SCOPE > 1.2. Purpose (1.2)" in text
    code, out = run("search", "--query", "Purpose", "--save", str(tmp_path / "근거.csv"))
    assert out["output_file"].endswith(".csv") and "excel.py" in out["next_commands"][0]["command"]


def test_search_truncated_offers_full_save(docs):
    run("index", "--folder", str(docs))
    code, out = run("search", "--query", "본문 OR body", "--top", "1")
    assert out["truncated"] is True and out["shown"] == 1 and out["count"] > 1
    assert "--save" in out["next_commands"][-1]["command"]


def _old_cache(src, text="옛날 단락 내용"):
    src.write_text(text, encoding="utf-8")
    st = os.stat(src)
    path = os.path.abspath(str(src))
    old = {"file": path, "type": "pdf", "mtime": st.st_mtime, "size": st.st_size, "parser_version": 1,
           "parsed_at": "20261001_000000", "meta": {"pages": 3},
           "units": [{"loc": "p.3", "heading": "", "text": "옛날 캐시의 검증 문장"},
                     {"loc": "슬라이드 2", "heading": "제목", "text": "검증 슬라이드"}]}
    with open(reader._doc_cache_path(path), "w", encoding="utf-8") as f:
        json.dump(old, f, ensure_ascii=False)
    with open(reader._index_path(), "w", encoding="utf-8") as f:
        json.dump({path: {"type": "txt", "units": 2, "parsed_at": "20261001_000000"}}, f, ensure_ascii=False)
    return path


def test_old_cache_is_normalized_and_marked_stale_when_original_is_gone(tmp_path, isolated_dirs):
    """parser_version 1 캐시(쪽/장절 키 없음)도 검색되고, 원본이 없어 다시 읽을 수 없으면 stale로 표시하고 UNVERIFIED."""
    src = tmp_path / "old.txt"
    _old_cache(src)
    os.remove(src)
    code, out = run("search", "--query", "검증")
    assert code == 0 and out["total_matches"] == 2 and out["stale_cache_count"] == 1
    assert out["matches"][0]["page"] == 3 and out["matches"][0]["page_source"] == "pdf"
    assert out["matches"][1]["section"] == ["제목"] and out["matches"][1]["page"] == 2
    assert out["matches"][0]["stale"] is True and out["agent_directive"].startswith("UNVERIFIED:")
    code, out = run("list")
    assert out["documents"][0]["stale"] is True and out["documents"][0]["missing"] is True


def test_stale_cache_is_reread_during_search(tmp_path, isolated_dirs):
    """원본이 있는 낡은 캐시는 검색할 때 그 자리에서 다시 읽어 새 내용으로 찾습니다 (recovered로 알림)."""
    src = tmp_path / "old.txt"
    _old_cache(src, "보증 기간 3년. 새 내용의 검증 문장")
    code, out = run("search", "--query", "검증")
    assert code == 0 and out["total_matches"] == 1 and "stale_cache_count" not in out
    assert out["matches"][0]["snippet"].startswith("보증 기간 3년") and "stale" not in out["matches"][0]
    assert out["recovered"] == ["원본이 바뀐 캐시 1개 -> 다시 읽고 검색"]
    assert out["agent_directive"].startswith("VERIFIED:")
    code, out = run("list")
    assert out["documents"][0]["stale"] is False and out["stale_count"] == 0


def test_list_filters_missing_and_zero(docs):
    run("index", "--folder", str(docs))
    code, out = run("list", "--folder", str(docs / "sub"))
    assert out["count"] == 1 and out["documents"][0]["file"] == "규격.docx"
    d = out["documents"][0]
    assert d["pages"] == 2 and d["sections"] >= 3 and d["missing"] is False and d["cached_at"]
    os.remove(docs / "spec.pdf")
    code, out = run("list", "--name", "spec")
    assert out["documents"][0]["missing"] is True and out["missing_count"] == 1
    code, out = run("list", "--name", "없는이름")
    assert code == 0 and out["count"] == 0 and out["truncated"] is False and out["agent_directive"].startswith("VERIFIED: 0건.")


def test_read_shows_page_section_number_and_page_filter(docs):
    code, out = run("read", "--file", str(docs / "계획서.pdf"), "--pages", "2")
    assert code == 0 and {u["page"] for u in out["units"]} == {2}
    u = [x for x in out["units"] if x["text"].startswith("2. 기대 효과")][0]
    assert u["number"] == "2" and u["section"] == ["III. 개발 내용", "2. 기대 효과"] and u["kind"] == "heading"
    code, out = run("read", "--file", str(docs / "계획서.pdf"), "--save", str(docs / "out.md"))
    assert "근거: 계획서.pdf p.1, III. 개발 내용" in open(out["output_file"], encoding="utf-8").read()
    code, out = run("list")
    assert out["documents"][0]["pages"] == 2 and out["documents"][0]["parser_version"] == reader.PARSER_VERSION


def test_eml_is_unsupported_like_any_unknown_type(tmp_path):
    f = tmp_path / "메일.eml"
    f.write_bytes(b"Subject: x\r\n\r\nbody")
    for args in (("read", "--file", str(f)), ("search", "--query", "x", "--file", str(f)), ("index", "--file", str(f))):
        code, out = run(*args)
        assert code == 1 and out["agent_directive"].startswith("STOP_RETRY:")
        assert "지원하지 않는 형식입니다: .eml" in out["message"] and ".pdf" in out["message"]
    assert not hasattr(reader, "parse_eml") and ".eml" not in reader.SUPPORTED


# ---------------------------------------------------------------------------
# 점검 반영: PDF 머리글/바닥글, 2단 편집
# ---------------------------------------------------------------------------

def make_pdf_at(path, pages):
    """pages: [[(x, y, 글, 크기)], ...] 위치를 정해 글을 넣습니다 (A4 595x842)."""
    import pymupdf
    doc = pymupdf.open()
    for items in pages:
        page = doc.new_page()
        for x, y, text, size in items:
            page.insert_text((x, y), text, fontname="korea", fontsize=size)
    doc.save(str(path))
    doc.close()


def test_pdf_repeated_header_and_page_number_do_not_break_sections(tmp_path):
    p = tmp_path / "hdr.pdf"
    head = (72, 40, "제2장 시스템 요구사항", 9)
    make_pdf_at(p, [
        [head, (72, 120, "2.1 전기 요구사항", 14), (72, 150, "2.1.1 정격", 12), (72, 180, "정격 전압은 220 V로 한다.", 10),
         (290, 815, "- 1 -", 9)],
        [head, (72, 120, "정격 전류는 10 A 이하로 한다.", 10), (72, 150, "10.5 kV 이하 기기는 별도로 시험한다.", 10),
         (290, 815, "- 2 -", 9)],
        [head, (72, 120, "2.2 기계 요구사항", 14), (72, 150, "기계 본문입니다.", 10), (290, 815, "- 3 -", 9)],
    ])
    units, meta = reader.parse_pdf(str(p))
    assert meta["header_footer_removed"] == 6
    assert not [u for u in units if "시스템 요구사항" in u["text"] or u["text"].startswith("- ")]
    cur = by_text(units, "정격 전류는")
    assert cur["section"] == ["2.1 전기 요구사항", "2.1.1 정격"] and cur["page"] == 2      # 머리글이 장절을 끊지 않음
    kv = by_text(units, "10.5 kV")
    assert kv["kind"] != "heading" and kv["section"] == ["2.1 전기 요구사항", "2.1.1 정격"]   # 수치 줄은 제목 아님
    assert by_text(units, "기계 본문")["section"] == ["2.2 기계 요구사항"]


def test_pdf_two_columns_read_left_column_first(tmp_path):
    p = tmp_path / "col.pdf"
    make_pdf_at(p, [[
        (72, 80, "보고서 제목 전체 폭으로 길게 쓴 줄입니다 가나다라마바사아자차카타파하", 12),
        (60, 130, "1. 왼쪽 장", 13), (60, 160, "왼쪽 본문 첫째 문장이다.", 10), (60, 190, "왼쪽 본문 둘째 문장이다.", 10),
        (320, 130, "3. 오른쪽 장", 13), (320, 160, "오른쪽 본문 첫째 문장이다.", 10), (320, 190, "오른쪽 본문 둘째 문장이다.", 10),
    ]])
    units, _ = reader.parse_pdf(str(p))
    texts = [u["text"] for u in units]
    assert texts.index("1. 왼쪽 장") < texts.index("왼쪽 본문 둘째 문장이다.") < texts.index("3. 오른쪽 장")
    assert by_text(units, "왼쪽 본문 첫째")["section"] == ["1. 왼쪽 장"]
    assert by_text(units, "오른쪽 본문 둘째")["section"] == ["3. 오른쪽 장"]


@pytest.mark.parametrize("line", ["10.5 kV 이하", "1.5 kV 1분", "3 mm 이상", "2 개월", "1.2 설계 기준 ........ 12",
                                  "(3) Verification Review; and", "가) 시험 장비,", "1) 설계 검토 및",
                                  "1. 사업 개요 ······ 3", "2.1 범위\t5"])
def test_measurements_and_toc_lines_are_not_headings(line):
    assert reader.heading_title(reader.classify_number(line), line, strong=True) is None


@pytest.mark.parametrize("line", ["1. A형 시험", "2.1 개요", "3. 분석 방법", "1 일반 사항"])
def test_real_headings_still_detected(line):
    assert reader.heading_title(reader.classify_number(line), line, strong=True) == line


# ---------------------------------------------------------------------------
# 점검 반영: DOCX 변경 추적·콘텐츠 컨트롤·여러 쪽 표·빈 번호 단락
# ---------------------------------------------------------------------------

def test_docx_tracked_insertions_content_controls_and_table_row_pages(tmp_path):
    from docx.oxml import parse_xml
    d = new_docx()
    p = d.add_paragraph("납기는 ")
    p._p.append(parse_xml(f'<w:ins {W_NS} w:id="1" w:author="a"><w:r><w:t>30일 이내</w:t></w:r></w:ins>'))
    p._p.append(parse_xml(f'<w:del {W_NS} w:id="2" w:author="a"><w:r><w:delText>60일</w:delText></w:r></w:del>'))
    body = d.element.body
    body.insert(len(body) - 1, parse_xml(
        f'<w:sdt {W_NS}><w:sdtContent><w:p><w:r><w:t>콘텐츠 컨트롤 안의 보증 기간 조항</w:t></w:r></w:p></w:sdtContent></w:sdt>'))
    t = d.add_table(rows=4, cols=2)
    for i, row in enumerate([["번호", "요구사항"], ["R-1", "첫째"], ["R-2", "둘째"], ["R-3", "셋째"]]):
        for j, v in enumerate(row):
            t.cell(i, j).text = v
    # R-3 행 앞에서 쪽이 바뀜 (Word가 저장한 쪽 경계)
    t.cell(3, 0).paragraphs[0].runs[0]._r.insert(0, parse_xml(f'<w:lastRenderedPageBreak {W_NS}/>'))
    path = tmp_path / "trk.docx"
    d.save(str(path))
    units, _ = reader.parse_docx(str(path))
    assert by_text(units, "납기는")["text"] == "납기는 30일 이내"                   # 지운 글(60일)은 없음
    assert by_text(units, "콘텐츠 컨트롤 안의")
    rows = {u["number"]: u for u in units if u["kind"] == "row"}
    assert rows["R-1"]["page"] == 1 and rows["R-3"]["page"] == 2


def test_docx_empty_numbered_paragraph_keeps_counting(tmp_path):
    d = new_docx()
    add_numbering(d)
    numbered(d, "일반 사항", 0)          # 3.
    numbered(d, "", 0)                   # 4. (빈 줄이어도 Word는 번호를 매김)
    numbered(d, "시험 방법", 0)          # 5.
    path = tmp_path / "num.docx"
    d.save(str(path))
    units, _ = reader.parse_docx(str(path))
    assert by_text(units, "5. 시험 방법")


# ---------------------------------------------------------------------------
# 점검 반영: 검색 문법, 검색 범위, 색인 복구
# ---------------------------------------------------------------------------

def test_parse_query_quotes_case_insensitive_operators():
    assert reader.parse_query('"설계 기준" and 온도 OR 습도') == [["설계 기준", "온도"], ["습도"]]
    assert reader.parse_query("A AND B or C D") == [["A", "B"], ["C D"]]
    assert reader.parse_query('""') == []


def test_matcher_ignores_spacing_and_empty_regex():
    m = reader.build_matcher("설계기준", False, False)
    assert m("본 설계 기준에 따른다") == ["설계기준"] and m("설계기\n준") == ["설계기준"]
    assert reader.build_matcher('"설계 기준"', False, False)("설계기준을 따른다")
    assert reader.build_matcher("a*", True, False)("xyz") is None
    assert "**설계 기준**" in reader.make_snippet("본 설계 기준에 따른다", ["설계기준"], 20, False)


def test_zero_hits_with_unindexed_files_is_unverified(tmp_path):
    d = tmp_path / "d"
    d.mkdir()
    for i in range(4):
        (d / f"f{i}.txt").write_text("하자 보수" if i == 3 else f"일반 문서 {i}", encoding="utf-8")
    code, out = run("search", "--query", "하자", "--file", str(d), "--max-files", "2")
    assert code == 0 and out["count"] == 0 and out["not_indexed_remaining"] == 2
    assert out["agent_directive"].startswith("UNVERIFIED:") and "등록하지 못한 파일 2개" in out["agent_directive"]
    assert out["verified"] is False


def test_folder_search_reports_excluded_subfolders(tmp_path):
    d = tmp_path / "d"
    (d / "_보관").mkdir(parents=True)
    (d / "a.txt").write_text("일반", encoding="utf-8")
    (d / "_보관" / "b.txt").write_text("하자", encoding="utf-8")
    code, out = run("search", "--query", "하자", "--file", str(d))
    assert code == 0 and out["count"] == 0 and any("하위 폴더 1개" in g for g in out["not_searched"])


def test_index_is_repaired_when_cache_exists_but_index_entry_is_missing(tmp_path):
    f = tmp_path / "a.txt"
    f.write_text("보증 기간 조항", encoding="utf-8")
    run("index", "--file", str(f))
    os.remove(reader._index_path())                      # 캐시는 있는데 색인만 없음 (저장 전 중단)
    code, out = run("index", "--file", str(f))
    assert code == 0 and out["unchanged"] == 1
    code, out = run("search", "--query", "보증")
    assert code == 0 and out["count"] == 1


def test_index_keys_differing_only_in_case_are_merged(tmp_path):
    a, b = str(tmp_path / "A.txt"), str(tmp_path / "A.txt").lower()
    with open(reader._index_path(), "w", encoding="utf-8") as fp:
        json.dump({a: {"parsed_at": "1"}, b: {"parsed_at": "2"}}, fp)
    idx = reader.load_index()
    if os.path.normcase(a) == os.path.normcase(b):          # Windows: 같은 파일 -> 최근 것 하나
        assert list(idx) == [b]
    else:
        assert len(idx) == 2


def test_hwp_text_pairs_surrogates_and_replaces_lone_ones():
    rec = "\U00020BB7田".encode("utf-16-le") + b"\x00\xd8"        # 확장 한자 + 짝 없는 서로게이트
    text = reader._hwp_para_text(rec)
    assert text == "\U00020BB7田\ufffd"
    text.encode("utf-8")


def test_read_pages_on_docx_uses_real_pages(tmp_path):
    from docx.enum.text import WD_BREAK
    d = new_docx()
    d.add_paragraph("1쪽 문단")
    d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    d.add_paragraph("2쪽 문단")
    path = tmp_path / "p.docx"
    d.save(str(path))
    code, out = run("read", "--file", str(path), "--pages", "2")
    assert code == 0 and out["pages_mode"] == "page" and [u["text"] for u in out["units"]] == ["2쪽 문단"]


def test_numbering_helper_creates_part_when_template_has_none(tmp_path):
    """회사 PC 기본 템플릿처럼 numbering.xml이 없는 문서에서도 번호 단락을 만들 수 있어야 함 (테스트 도우미 자체 확인)."""
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    d = new_docx()
    for rid, rel in list(d.part.rels.items()):
        if rel.reltype == RT.NUMBERING:
            del d.part.rels[rid]
    d.part.__dict__.pop("numbering_part", None)
    add_numbering(d)
    numbered(d, "일반 사항", 0)
    path = tmp_path / "n.docx"
    d.save(str(path))
    units, _ = reader.parse_docx(str(path))
    assert by_text(units, "3. 일반 사항")


@pytest.mark.parametrize("query,text,hit", [
    ("SOI 1", "Activities for SOI #1 are listed", True),        # 2026-10-29: 'SOI 1'로 'SOI #1'을 못 찾던 문제
    ("SOI #1", "SOI-1 검토", True),
    ("SOI1", "SOI #1", True),
    ("SOI 1", "SOI #10 활동", False),                            # 숫자 경계는 지킴
    ("SOI 1", "SOI 21", False),
    ("1.2", "제 1.2절", True),
    ("1.2", "11.2절", False),
    ("설계기준", "설계 기준에 따른다", True),
    ("#4", "SOI #4 활동", True),
    ("#4", "4개 항목", False),                                     # 기호를 빼면 숫자만 남는 검색어는 기호 그대로
    ("C++", "C++ 코딩 규칙", True),
])
def test_search_ignores_symbols_and_spacing_but_keeps_number_boundaries(query, text, hit):
    assert (reader.build_matcher(query, False, False)(text) is not None) == hit


def test_snippet_highlights_symbol_insensitive_match():
    assert "**SOI #1**" in reader.make_snippet("Activities for SOI #1 are listed", ["SOI 1"], 30, False)


# ---------------------------------------------------------------------------
# clear: 캐시 지우기
# ---------------------------------------------------------------------------

def _cached(tmp_path, *names):
    paths = []
    for n in names:
        f = tmp_path / n
        f.write_text(f"{n} 내용", encoding="utf-8")
        run("index", "--file", str(f))
        paths.append(f)
    return paths


def test_clear_requires_target_and_dry_run_changes_nothing(tmp_path):
    a, b = _cached(tmp_path, "a규격.txt", "b보고.txt")
    code, out = run("clear")
    assert code == 1 and "--dry-run" in out["message"]
    code, out = run("clear", "--name", "규격", "--dry-run")
    assert code == 0 and out["dry_run"] is True and out["count"] == 1 and out["planned"] == ["캐시 삭제: a규격.txt"]
    assert "clear --name 규격" in out["next_commands"][0]["command"] and "--dry-run" not in out["next_commands"][0]["command"]
    assert len(reader.load_index()) == 2


def test_clear_by_name_removes_index_and_cache_file_keeps_original(tmp_path):
    a, b = _cached(tmp_path, "a규격.txt", "b보고.txt")
    cache_a = reader._doc_cache_path(str(a))
    code, out = run("clear", "--name", "규격")
    assert code == 0 and out["verified"] is True and out["remaining_in_cache"] == 1
    assert not os.path.exists(cache_a) and a.exists()                     # 원본은 그대로
    assert list(reader.load_index()) == [str(b)]
    code, out = run("search", "--query", "내용")
    assert out["count"] == 1                                             # 지운 문서는 검색 대상에서 빠짐


def test_clear_missing_and_all(tmp_path):
    a, b = _cached(tmp_path, "a.txt", "b.txt")
    os.remove(a)
    code, out = run("clear", "--missing")
    assert code == 0 and [d["file"] for d in out["documents"]] == ["a.txt"]
    code, out = run("clear", "--missing")
    assert code == 0 and out["count"] == 0 and "0건" in out["agent_directive"]
    code, out = run("clear", "--all")
    assert code == 0 and out["remaining_in_cache"] == 0 and reader.load_index() == {}
