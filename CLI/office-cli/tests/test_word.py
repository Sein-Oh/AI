"""word.py 테스트: 가짜 Word 문서(문자열 모델)로 삽입 위치, 마크다운 변환, 찾아 바꾸기 반복을 검증합니다."""
import json

import pytest

import _office as off
import word
from fake_office import make_win32com
from fake_word import FakeWordDocument


@pytest.fixture
def wd(monkeypatch):
    holder = {"doc": FakeWordDocument("추진 배경\r2026년 계획입니다.\r일정")}
    fake, _ = make_win32com()
    monkeypatch.setattr(off, "win32com", fake)
    monkeypatch.setattr(off, "pythoncom", None)
    monkeypatch.setattr(off, "list_open", lambda kind: [])
    monkeypatch.setattr(off, "get_active", lambda kind: {"name": holder["doc"].Name, "path": holder["doc"].FullName,
                                                          "app": object(), "doc": holder["doc"]})
    monkeypatch.setattr(off, "save_snapshot", lambda *a, **k: None)
    return holder


def run(capsys, *argv):
    with pytest.raises(SystemExit) as ei:
        word.main(list(argv))
    out = [ln for ln in capsys.readouterr().out.splitlines() if ln.strip()]
    assert len(out) == 1, out
    return ei.value.code, json.loads(out[0])


def paras(doc):
    return doc.text.split("\r")


class _Shapes:
    def __init__(self, starts):
        self.starts = starts
        self.Count = len(starts)

    def Item(self, i):
        return type("S", (), {"Range": type("R", (), {"Start": self.starts[i - 1]})()})()


class _Para:
    """COM 단락: Range.Text, Range.Start, Range.InlineShapes(그림 위치), OutlineLevel."""
    def __init__(self, start, text, pictures=(), level=10):
        self.Range = type("R", (), {"Text": text, "Start": start, "InlineShapes": _Shapes([start + p for p in pictures])})()
        self.OutlineLevel = level


def test_com_units_inline_picture_becomes_marker_not_slash():
    # Word 2007은 본문 그림을 '/', 새 버전은 '\x01'로 돌려줌 (실기: 'Word 내용' 시트에 '/'가 본문으로 들어감)
    doc = type("D", (), {"Paragraphs": [
        _Para(0, "첫 문단\r"),
        _Para(5, "/\r", pictures=[0]),                     # Word 2007 그림만 있는 단락
        _Para(7, "그래프 \x01 참고\r", pictures=[4]),        # 새 버전, 글 사이 그림
        _Para(16, "A/B 비율\r"),                            # 그림 없는 진짜 '/'
        _Para(23, "A/B /\r", pictures=[4]),                 # 진짜 '/'와 그림이 같은 단락
    ]})()
    texts = [u["text"] for u in word._com_units(doc)]
    assert texts == ["첫 문단", "[그림]", "그래프 [그림] 참고", "A/B 비율", "A/B [그림]"]


def test_write_plain_text_at_end(wd, capsys):
    code, out = run(capsys, "write", "--text", "새 문단입니다")
    assert code == 0, out
    assert paras(wd["doc"])[:4] == ["추진 배경", "2026년 계획입니다.", "일정", "새 문단입니다"]


def test_write_after_anchor_goes_after_whole_paragraph(wd, capsys):
    code, out = run(capsys, "write", "--text", "추가 문단", "--anchor", "계획", "--pos", "after")
    assert code == 0
    assert paras(wd["doc"])[:4] == ["추진 배경", "2026년 계획입니다.", "추가 문단", "일정"]   # 단락 중간을 쪼개지 않음


def test_anchor_in_several_paragraphs_is_not_guessed(wd, capsys):
    wd["doc"] = FakeWordDocument("계획 개요\r본문\r세부 계획\r끝")
    before = wd["doc"].text
    code, out = run(capsys, "write", "--text", "새 문단", "--anchor", "계획", "--pos", "replace")
    assert code == 1 and out["error_type"] == "ambiguous" and out["agent_directive"].startswith("STOP_RETRY:")
    assert out["candidates"] == ["1: 계획 개요", "2: 세부 계획"] and "--occurrence" in out["message"]
    assert wd["doc"].text == before                                       # 어느 단락도 지우지 않음
    code, out = run(capsys, "write", "--text", "새 문단", "--anchor", "계획", "--pos", "replace", "--occurrence", "2")
    assert code == 0 and paras(wd["doc"])[:4] == ["계획 개요", "본문", "새 문단", "끝"]
    code, out = run(capsys, "write", "--text", "x", "--anchor", "개요", "--occurrence", "5")
    assert code == 1 and "1~1" in out["message"]


def test_anchor_same_paragraph_twice_is_one_candidate(wd, capsys):
    wd["doc"] = FakeWordDocument("계획과 계획\r끝")
    code, out = run(capsys, "write", "--text", "추가", "--anchor", "계획")
    assert code == 0 and paras(wd["doc"])[:3] == ["계획과 계획", "추가", "끝"]


def test_anchor_not_found_recommends_similar_text(wd, capsys):
    code, out = run(capsys, "write", "--text", "x", "--anchor", "2026 계획")
    assert code == 1 and out["error_type"] == "input" and out["message"].endswith("[추천] 2026년 계획")
    code, out = run(capsys, "write", "--text", "x", "--anchor", "2026년 계획")    # 추천대로 다시 하면 찾음
    assert code == 0


def test_anchor_not_found_does_not_recommend_different_text(wd, capsys):
    """글자가 바뀐 문구('2027' -> '2026')는 뜻이 다를 수 있어 [추천]하지 않음. replace는 끼어든 글자도 [추천] 안 함."""
    before = wd["doc"].text
    code, out = run(capsys, "write", "--text", "x", "--anchor", "2027년 계획")
    assert code == 1 and "[추천]" not in out["message"] and "비슷한 문구: '2026년 계획'" in out["message"]
    code, out = run(capsys, "write", "--text", "x", "--anchor", "2026 계획", "--pos", "replace")
    assert code == 1 and "[추천]" not in out["message"] and "비슷한 문구" in out["message"]
    assert wd["doc"].text == before


def test_safe_anchor_fix_rules():
    assert word._safe_anchor_fix("2026년계획", "2026년 계획", "replace")       # 공백만 다름
    assert word._safe_anchor_fix("2026 계획", "2026년 계획", "after")          # 끼어든 글자 1자
    assert not word._safe_anchor_fix("2026 계획", "2026년 계획", "replace")
    for anchor, similar in [("Q2 매출 보고", "Q1 매출 보고"), ("결론 분석", "결과 분석"),
                            ("Sales 2025", "Sales 202"), ("3. 결과요약", "2. 결과")]:
        assert not word._safe_anchor_fix(anchor, similar, "after"), (anchor, similar)

def test_write_before_and_replace_anchor(wd, capsys):
    run(capsys, "write", "--text", "머리말", "--anchor", "추진", "--pos", "before")
    assert paras(wd["doc"])[0] == "머리말"
    run(capsys, "write", "--text", "새 계획", "--anchor", "2026년", "--pos", "replace")
    assert "2026년 계획입니다." not in wd["doc"].text and "새 계획" in paras(wd["doc"])


def test_write_markdown_blocks_styles_and_table(wd, capsys):
    md = "## 결론\n본문 **강조** 끝\n\n- 가\n- 나\n\n| 공정 | 수율 |\n|---|---|\n| 가열 | 98 |"
    code, out = run(capsys, "write", "--text", md)
    assert code == 0
    doc = wd["doc"]
    assert ("style", "결론", word.heading_style(2)) in doc.log
    assert ("style", "가", word.WD_STYLE_LIST_BULLET) in doc.log
    assert "본문 강조 끝" in paras(doc)
    assert any(e[0] == ("font", doc.text.index("강조"), doc.text.index("강조") + 2) and e[1] == "Bold" and e[2] is True
               for e in doc.log)
    assert doc.tables == [["공정\t수율", "가열\t98"]]
    assert out["inserted"] == ["제목2", "본문", "글머리 목록 2개", "표 2x2"]


def test_write_escaped_newlines_and_anchor_not_found(wd, capsys):
    code, out = run(capsys, "write", "--text", "첫줄\\n둘째줄")
    assert "첫줄" in paras(wd["doc"]) and "둘째줄" in paras(wd["doc"])
    code, out = run(capsys, "write", "--text", "x", "--anchor", "없는문구")
    assert code == 1 and "read" in out["message"]


def test_write_needs_content(wd, capsys):
    code, out = run(capsys, "write")
    assert code == 1 and out["error_type"] == "input"


def test_replace_no_infinite_loop_and_counts(wd, capsys):
    wd["doc"] = FakeWordDocument("가 가 가")
    code, out = run(capsys, "replace", "--find", "가", "--replace", "가가")
    assert code == 0 and out["replaced"] == 3 and wd["doc"].text.startswith("가가 가가 가가")


def test_replace_empty_deletes_and_formatting(wd, capsys):
    wd["doc"] = FakeWordDocument("초안 보고서 초안")
    code, out = run(capsys, "replace", "--find", "초안 ", "--replace", "")
    assert out["replaced"] == 1 and wd["doc"].text.startswith("보고서 초안")
    code, out = run(capsys, "replace", "--find", "보고서", "--bold", "--color", "blue")
    assert out["replaced"] == 1
    assert any(len(e) == 3 and e[1] == "Bold" and e[2] is True for e in wd["doc"].log)


def test_replace_map_and_limits(wd, capsys):
    wd["doc"] = FakeWordDocument("Draft 진행중")
    code, out = run(capsys, "replace", "--map", '{"Draft": "Final", "진행중": "완료"}')
    assert out["replaced"] == 2 and wd["doc"].text.startswith("Final 완료")
    code, out = run(capsys, "replace", "--find", "x" * 300, "--replace", "y")
    assert code == 1 and "255" in out["message"]
    code, out = run(capsys, "replace", "--find", "없음", "--replace", "y")
    assert code == 0 and out["replaced"] == 0


def test_read_closed_docx_with_python_docx(tmp_path, capsys, monkeypatch):
    import docx
    p = tmp_path / "계획.docx"
    d = docx.Document()
    d.add_heading("배경", level=1)
    d.add_paragraph("본문입니다")
    d.save(p)
    fake, calls = make_win32com()
    monkeypatch.setattr(off, "win32com", fake)
    monkeypatch.setattr(off, "list_open", lambda kind: [])
    code, out = run(capsys, "read", "--target", str(p))
    assert code == 0 and [x["text"] for x in out["paragraphs"]] == ["배경", "본문입니다"]
    assert out["paragraphs"][1]["heading"] == "배경"
    assert not any(c[0] == "DispatchEx" for c in calls)


def test_highlight_color_suggestion(wd, capsys):
    code, out = run(capsys, "replace", "--find", "계획", "--highlight", "yelow")
    assert code == 1 and "[추천] yellow" in out["message"]
