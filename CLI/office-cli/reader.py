"""reader.py - 문서 파일 읽기, 등록(캐시), 근거 검색 (Office 프로그램 불필요)

명령:
    read    파일 하나를 읽고 캐시에 등록합니다. PDF, HWP, HWPX, DOCX, PPTX, XLSX, CSV, TXT, MD
    index   폴더(또는 파일)의 문서를 내용 출력 없이 한꺼번에 캐시에 등록합니다.
    search  캐시에 등록된 문서에서 검색합니다. 결과마다 근거(쪽, 장절, 번호)를 줍니다.
            --file을 주면 그 파일/폴더에서 아직 등록되지 않았거나 바뀐 문서를 먼저 등록한 뒤 검색합니다.
    list    캐시에 등록된 문서 목록 (--name, --folder로 거르기)

캐시 형식(모든 형식 공통): 문서마다 units 목록, 단위마다
    id, kind, loc, page, page_source, section(장절 경로), number(항목 번호), heading, text
결과를 다른 문서에 넣으려면 --save로 파일을 만든 뒤 word.py/ppt.py write --file 또는 excel.py write --file로 넘깁니다.

예:
    python reader.py index --folder "C:\\docs\\규격"
    python reader.py read --file "C:\\docs\\보고서.pdf" --pages 1-3
    python reader.py search --query "마일스톤 AND 일정" --save "근거.xlsx"
    python reader.py search --query "단가" --file "C:\\docs\\견적" --pages 3-7 --section "SCOPE"
    python reader.py list --name 규격
"""
import hashlib
import json
import os
import re
import sys
import unicodedata
import zipfile
import zlib
from collections import Counter
from typing import Any, Dict, List, Optional, Tuple
from xml.etree import ElementTree as ET

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import _runtime as rt  # noqa: E402

TOOL = "reader"
SUPPORTED = {
    ".pdf": "pdf", ".hwp": "hwp", ".hwpx": "hwpx",
    ".docx": "docx", ".docm": "docx", ".pptx": "pptx", ".pptm": "pptx",
    ".xlsx": "xlsx", ".xlsm": "xlsx", ".csv": "csv", ".txt": "txt", ".md": "txt",
}
UNIT_NAMES = {"pdf": "단락", "pptx": "슬라이드", "xlsx": "행", "csv": "행"}
PAGE_TYPES = ("pdf", "pptx")   # read --pages가 실제 쪽/슬라이드 번호인 형식 (그 밖은 항목 번호)
PARSER_VERSION = 4          # 파서가 바뀌면 올려서 기존 캐시를 무효화 (2: 공통 단위 형식, 장절/번호/쪽, 3: PDF 머리글·2단, DOCX 변경 추적·표 행 쪽, 4: 나열 항목은 제목 아님)
MAX_SHEET_ROWS = 20000      # xlsx 시트당 최대 행 수
MAX_FOLDER_FILES = 300      # 한 번에 새로 등록(파싱)할 최대 파일 수
MAX_FOLDER_DEPTH = 3        # 폴더 기본 탐색 깊이
MAX_FOLDER_SCAN = 20000     # 폴더에서 살펴볼 최대 파일 수
MAX_LIST_ITEMS = 20         # JSON에 담는 실패/새 등록 목록 상한
MAX_SAVE_HITS = 5000        # --save 파일에 담는 최대 검색 결과 수
STALE_REFRESH_MAX = 20      # 검색 때 원본이 바뀐 캐시를 그 자리에서 다시 읽는 최대 문서 수
HEADING_MAX_CHARS = 80      # 번호로 시작하는 짧은 줄만 제목 후보 (굵게/큰 글씨는 120)
LEGACY_OFFICE = (".xls", ".doc", ".ppt")


# ==============================================================================
# 1. 번호와 장절 인식 (모든 형식 공통)
# ==============================================================================
# 번호 토큰을 "계열"로 나누고, 계열이 문서에서 처음 나온 순서로 수준을 정합니다.
#   - 처음 보는 계열은 현재 가장 깊은 제목의 한 단계 아래
#   - 이미 본 계열은 그 계열의 수준으로 돌아감 (그보다 깊은 제목은 닫힘)
#   - 1.1.1 같은 점 번호는 숫자 개수가 깊이 (끝의 .0은 빼고 셈: 1.0# -> 깊이 1)
# 글머리 기호(O, -, 동그라미/네모 기호 등)는 번호가 아니고 제목도 아닙니다.

GANADA = "가나다라마바사아자차카타파하"
CHOSUNG = "ㄱㄴㄷㄹㅁㅂㅅㅇㅈㅊㅋㅌㅍㅎ"
_ROMAN_FULLWIDTH = {chr(0x2160 + i): r for i, r in enumerate(
    ["I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X", "XI", "XII"])}
# 글머리 기호: 영문 O/o, -, *, 가운뎃점, 그리고 유니코드 기호들(코드로만 적음)
_BULLET_SPACE = "Oo-*"
_BULLET_SYMBOLS = ("\u00b7\u2022\u2023\u2043\u2219\u25cb\u25cf\u25ce\u25e6\u25a1\u25a0\u25aa\u25ab\u25c7\u25c6"
                   "\u25b6\u25b7\u25ba\u25b8\u25b9\u203b\u3147\u2218\u27a2\u2713\u2714\u221a\u2192\u21d2\u2794"
                   "\uf0a7\uf0b7\uf06c\uf0d8\uf0fc\uf076")

_GA = f"[{GANADA}]"
_NUM_PATTERNS = [
    # (계열, 정규식, 번호 표기 함수)
    ("paren_arabic", re.compile(r"\(\s*(\d{1,3})\s*\)"), lambda m: f"({m.group(1)})"),
    ("paren_ganada", re.compile(rf"\(\s*({_GA})\s*\)"), lambda m: f"({m.group(1)})"),
    ("paren_alpha", re.compile(r"\(\s*([a-zA-Z])\s*\)"), lambda m: f"({m.group(1)})"),
    ("rparen_arabic", re.compile(r"(\d{1,3})\)"), lambda m: f"{m.group(1)})"),
    ("rparen_ganada", re.compile(rf"({_GA})\)"), lambda m: f"{m.group(1)})"),
    ("rparen_alpha", re.compile(r"([a-z])\)"), lambda m: f"{m.group(1)})"),
]
_KO_UNIT = {"편": "ko_pyeon", "장": "ko_jang", "절": "ko_jeol", "관": "ko_gwan", "조": "ko_jo", "항": "ko_hang",
            "호": "ko_ho"}
_KO_RE = re.compile(r"제\s*(\d{1,4})\s*(편|장|절|관|조|항|호)(?:\s*의\s*(\d{1,3}))?(?=[\s(.:]|$)")
_WORD_RE = re.compile(r"(?i)(chapter|part|section|article)\s+(\d{1,3}|[IVX]{1,6})(?=[\s.:]|$)[.:]?")
_ANNEX_RE = re.compile(r"(?i)(annex|appendix|attachment)\s+([A-Z]|\d{1,3})(?=[\s.:)]|$)[.:]?")
_ANNEX_KO_RE = re.compile(r"(부록|별표|별지|첨부)\s*(\d{1,3}|[A-Z])(?=[\s.:)]|$)[.:]?|(부록)(?=[\s.:]|$)")
_Q_RE = re.compile(r"(Q|문항|질문|문)\s*\.?\s*(\d{1,3})(?=[\s.):]|$)[.)]?")
_ROMAN_RE = re.compile(r"([IVX]{1,6})[.)](?=\s|$)")
_ROMAN_LOWER_RE = re.compile(r"([ivx]{1,6})[.)](?=\s|$)")
_ARABIC_RE = re.compile(r"(\d{1,2}(?:\.\d{1,3}){0,6})(\.|#)?(?=\s|$)")
_GANADA_RE = re.compile(rf"({_GA})\.(?=\s|$)")
_ALPHA_UPPER_RE = re.compile(r"([A-Z])\.(?=\s)")
_ALPHA_LOWER_RE = re.compile(r"([a-z])\.(?=\s)")
_SENTENCE_END = re.compile(r"(?:[.\u3002!;]|다|니다|어요|아요|해요|세요)\s*$")
_NOUN_LIKE_END = re.compile(r"(?:음|함|임|됨)\s*$")
_LETTER = re.compile(r"[A-Za-z\uac00-\ud7a3]")
# 번호 뒤가 단위면 번호가 아니라 수치 (10.5 kV 이하, 1.5 kV 1분, 3 mm, 2 개월). '1. A형'처럼 뒤에 글자가 붙으면 단위 아님
_UNIT_AFTER = re.compile(r"\s*(?:%|\u2103|\u00b0|\u2126|\u03a9|"
                         r"[kKMGmun\u03bc]?(?:V|A|W|Hz|Pa|bar|g|m|L|s|dB|lx|ppm|ohm)(?=[\s/,.)\]]|$)|"
                         r"(?:초|분|시간|일|주|개월|년|회|개|배|도|kg|mm|cm|km|mA|kA)(?=[\s/,.)\]]|$))")
# 목차 줄: 점선 리더나 탭 뒤에 쪽 번호 (1.2 설계 기준 ........ 12)
# 문장 중간에서 끝나는 나열 항목: "(3) Verification Review; and", "가) 시험 장비,", "1) 설계 검토 및"
_LIST_CONTINUES = re.compile(r"(?:[;,]|[;,]?\s+(?:and|or|및|또는|그리고))\s*$", re.IGNORECASE)
_TOC_LINE = re.compile(r"(?:[.\u00b7\u2026\u2025\u30fb\u2219_-]{3,}|\t)\s*\d{1,4}\s*$")


def _info(family: str, number: Optional[str], rest: str, depth: int = 1, plain: bool = False) -> Dict[str, Any]:
    return {"family": family, "number": number, "rest": rest.strip(), "depth": depth, "plain": plain}


def classify_number(line: str) -> Optional[Dict[str, Any]]:
    """줄 머리의 번호 토큰을 찾습니다. 반환: {family, number, rest, depth, plain} 또는 None.
    family 'bullet'은 글머리 기호(번호 아님). plain=True는 구분 기호 없는 숫자 하나('1 Scope')로, 굵게/큰 글씨일 때만 제목."""
    s = (line or "").strip()
    if not s:
        return None
    c, o = s[0], ord(s[0])
    if 0x2460 <= o <= 0x2473:                        # 동그라미 숫자 1~20
        return _info("circled", f"({o - 0x245F})", s[1:])
    if 0x2474 <= o <= 0x2487:                        # 괄호 숫자 1~20
        return _info("paren_arabic", f"({o - 0x2473})", s[1:])
    if 0x326E <= o <= 0x327B:                        # 동그라미 가나다
        return _info("circled_ganada", f"({GANADA[o - 0x326E]})", s[1:])
    if 0x3260 <= o <= 0x326D:                        # 동그라미 자음
        return _info("circled_jamo", f"({CHOSUNG[o - 0x3260]})", s[1:])
    if c in _ROMAN_FULLWIDTH and (len(s) == 1 or s[1] in " .)\t"):
        rest = s[1:].lstrip(".)")
        return _info("roman", _ROMAN_FULLWIDTH[c], rest)
    if c in _BULLET_SYMBOLS or (c in _BULLET_SPACE and (len(s) == 1 or s[1] in " \t")):
        return _info("bullet", None, s[1:])
    for family, pat, fmt in _NUM_PATTERNS:
        m = pat.match(s)
        if m:
            return _info(family, fmt(m), s[m.end():])
    m = _KO_RE.match(s)
    if m:
        num = f"제{m.group(1)}{m.group(2)}" + (f"의{m.group(3)}" if m.group(3) else "")
        return _info(_KO_UNIT[m.group(2)], num, s[m.end():])
    m = _WORD_RE.match(s)
    if m:
        return _info(m.group(1).lower(), f"{m.group(1)} {m.group(2)}", s[m.end():])
    m = _ANNEX_RE.match(s)
    if m:
        return _info("annex", f"{m.group(1)} {m.group(2)}", s[m.end():])
    m = _ANNEX_KO_RE.match(s)
    if m:
        num = f"{m.group(1)} {m.group(2)}" if m.group(1) else m.group(3)
        return _info("annex_ko", num, s[m.end():])
    m = _Q_RE.match(s)
    if m:
        return _info("question", f"{m.group(1)}{m.group(2)}", s[m.end():])
    m = _ROMAN_RE.match(s)
    if m:
        return _info("roman", m.group(1), s[m.end():])
    m = _ROMAN_LOWER_RE.match(s)
    if m:
        return _info("roman_lower", m.group(1), s[m.end():])
    m = _ARABIC_RE.match(s)
    if m:
        rest = s[m.end():]
        if rest.strip()[:1].isdigit() or _UNIT_AFTER.match(rest):   # 날짜/수치 (예: 10. 28., 10.5 kV)
            return None
        parts = m.group(1).split(".")
        depth_parts = list(parts)
        while len(depth_parts) > 1 and depth_parts[-1] == "0":   # 1.0# -> 깊이 1
            depth_parts.pop()
        plain = len(parts) == 1 and not m.group(2)
        return _info("arabic", m.group(1), rest, depth=len(depth_parts), plain=plain)
    m = _GANADA_RE.match(s)
    if m:
        return _info("ganada", m.group(1), s[m.end():])
    m = _ALPHA_UPPER_RE.match(s)
    if m:
        return _info("alpha_upper", m.group(1), s[m.end():])
    m = _ALPHA_LOWER_RE.match(s)
    if m:
        return _info("alpha_lower", m.group(1), s[m.end():])
    return None


def heading_title(info: Optional[Dict[str, Any]], line: str, strong: bool = False) -> Optional[str]:
    """번호로 시작하는 줄이 제목이면 장절 경로에 넣을 제목 글을, 아니면 None을 돌려줍니다.
    제목 조건: 짧은 한 줄, 문장 끝(. / 다 / 요 등)으로 끝나지 않음, 번호 뒤에 글자가 있음."""
    if info is None or info["family"] == "bullet":
        return None
    s = line.strip()
    if "\n" in s or _TOC_LINE.search(s):            # 목차 줄은 제목이 아님 (목차 쪽이 장절로 남지 않게)
        return None
    rest = info["rest"]
    if info["family"].startswith("ko_") and info["family"] in ("ko_pyeon", "ko_jang", "ko_jeol", "ko_gwan", "ko_jo"):
        # 제N조(목적) 본문... 처럼 조문은 문장이어도 구조이므로 번호와 괄호 제목만 장절로 씁니다.
        m = re.match(r"\s*(\([^)]{1,40}\))", rest)
        if len(s) > HEADING_MAX_CHARS or _SENTENCE_END.search(s):
            return info["number"] + (m.group(1) if m else "")
        return s if _LETTER.search(rest) or m else info["number"]
    if info["plain"] and not strong:
        return None
    if not _LETTER.search(rest):
        return None
    if len(s) > (120 if strong else HEADING_MAX_CHARS):
        return None
    if _SENTENCE_END.search(s) or _LIST_CONTINUES.search(s):
        return None
    if _NOUN_LIKE_END.search(s) and len(rest) >= 8:   # 개조식 문장 끝 (~함, ~임). 짧은 명사(소음)는 제외
        return None
    return s


class Outline:
    """장절 경로 스택. 계열별 수준은 문서에서 처음 나온 위치로 정합니다."""

    def __init__(self) -> None:
        self.stack: List[Dict[str, Any]] = []
        self.levels: Dict[str, int] = {}

    def path(self) -> List[str]:
        return [e["title"] for e in self.stack]

    def _place(self, level: int, title: str, family: str, size: Optional[float]) -> None:
        level = max(1, level)
        while self.stack and self.stack[-1]["level"] >= level:
            self.stack.pop()
        self.stack.append({"level": level, "title": title, "family": family, "size": size})

    def register(self, family: str, depth: int, level: int) -> None:
        """스타일(제목 N, 개요 N)로 수준이 정해진 제목의 번호 계열을 기록 (같은 계열의 일반 단락이 맞는 수준을 쓰게)."""
        if family not in self.levels:
            self.levels[family] = max(1, level - (depth - 1))

    def push_level(self, level: int, title: str, size: Optional[float] = None) -> None:
        self._place(level, title, "style", size)

    def push_family(self, family: str, depth: int, title: str, size: Optional[float] = None) -> None:
        if family in self.levels:
            level = self.levels[family] + depth - 1
        else:
            base = (self.stack[-1]["level"] if self.stack else 0) + 1
            self.levels[family] = base
            level = base + depth - 1
        # 글자 크기 확인: 상위 제목보다 분명히 큰 제목은 그 상위의 하위가 아님 -> 더 큰 제목 아래까지 올라감
        if size:
            parents = [e for e in self.stack if e["level"] < level]
            if parents and parents[-1]["size"] and size > parents[-1]["size"] + 0.5:
                keep = [e for e in parents if e["size"] and e["size"] > size + 0.5]
                new_level = (keep[-1]["level"] if keep else 0) + 1
                self.levels[family] = max(1, new_level - (depth - 1))
                level = new_level
        self._place(level, title, family, size)


class UnitBuilder:
    """모든 형식이 같은 키의 단위를 만들게 합니다."""

    def __init__(self) -> None:
        self.units: List[Dict[str, Any]] = []
        self.outline = Outline()

    def add(self, kind: str, loc: str, text: str, page: Optional[int] = None, page_source: Optional[str] = None,
            number: Optional[str] = None, section: Optional[List[str]] = None) -> Dict[str, Any]:
        sec = list(section) if section is not None else self.outline.path()
        unit = {"id": len(self.units) + 1, "kind": kind, "loc": loc, "page": page,
                "page_source": page_source if page is not None else None,
                "section": sec, "number": number, "heading": sec[-1] if sec else "", "text": text}
        self.units.append(unit)
        return unit

    def add_text(self, loc: str, text: str, page: Optional[int] = None, page_source: Optional[str] = None, *,
                 kind: str = "paragraph", strong: bool = False, size: Optional[float] = None,
                 style_level: Optional[int] = None, list_family: Optional[str] = None,
                 list_depth: int = 1, allow_heading: bool = True) -> Dict[str, Any]:
        """글 단위 하나를 넣으며 제목이면 장절 경로를 갱신합니다."""
        s = text.strip()
        first = s.split("\n", 1)[0]
        info = classify_number(first)
        number = info["number"] if info and info["family"] != "bullet" and not info["plain"] else None
        if style_level:
            if info and info["family"] != "bullet":
                number = info["number"]
                self.outline.register(list_family or info["family"], list_depth if list_family else info["depth"],
                                      style_level)
            self.outline.push_level(style_level, first[:120], size)
            return self.add("heading", loc, s, page, page_source, number)
        title = heading_title(info, s, strong) if allow_heading else None
        if title:
            assert info is not None
            fam, depth = (list_family, list_depth) if list_family else (info["family"], info["depth"])
            self.outline.push_family(fam, depth, title[:120], size)
            return self.add("heading", loc, s, page, page_source, info["number"])
        return self.add(kind, loc, s, page, page_source, number)


_ROW_ID_RE = re.compile(r"^(?:\d{1,3}(?:\.\d{1,3})*\.?|[A-Za-z]{1,8}[-_]?\d[\w.-]{0,12})$")


def _clean_cell(v: Any) -> str:
    return " ".join(str(v or "").split())


def _looks_like_header(row: List[str]) -> bool:
    cells = [c for c in row if c]
    if not cells or len(cells) < max(1, len([1 for c in row]) // 2):
        return False
    return all(len(c) <= 30 and not re.fullmatch(r"[\d.,%\s-]+", c) for c in cells)


def table_rows(rows: List[List[Any]], keys: Optional[List[Any]] = None) -> List[Dict[str, Any]]:
    """표를 행 단위 글로 바꿉니다. 첫 행이 머리글로 보이면 '머리글: 값 | 머리글: 값' 형식.
    병합/빈 머리글 칸은 앞 칸의 머리글을 이어 쓰고, 여러 줄 칸은 공백으로 잇습니다.
    keys를 주면 행마다 같은 순서의 값을 결과의 key로 돌려줍니다 (예: 행의 쪽 번호).
    반환: [{"row": 행 번호(1부터), "text", "number", "key"}] (머리글 행은 단독 표가 아니면 넣지 않음)"""
    cleaned = [([_clean_cell(c) for c in r], (keys[i] if keys is not None and i < len(keys) else None))
               for i, r in enumerate(rows)]
    cleaned = [(r, k) for r, k in cleaned if any(r)]
    grid = [r for r, _ in cleaned]
    row_keys = [k for _, k in cleaned]
    if not grid:
        return []
    header: Optional[List[str]] = None
    if len(grid) >= 2 and _looks_like_header(grid[0]):
        header, prev = [], ""
        for c in grid[0]:
            prev = c or prev
            header.append(prev)
    out: List[Dict[str, Any]] = []
    for i, r in enumerate(grid, 1):
        if header is not None and i == 1:
            continue
        parts: List[str] = []
        for j, v in enumerate(r):
            if not v or (j > 0 and v == r[j - 1] and header is not None and header[j] == header[j - 1]):
                continue                                     # 빈 칸, 가로 병합으로 반복된 칸
            h = header[j] if header is not None and j < len(header) else ""
            parts.append(f"{h}: {v}" if h and h != v else v)
        if not parts:
            continue
        first = next((c for c in r if c), "")
        number = None
        if len(first) <= 20 and _ROW_ID_RE.match(first):
            number = first.rstrip(".")
        else:
            info = classify_number(first)
            if info and info["family"] != "bullet" and not info["plain"]:
                number = info["number"]
        out.append({"row": i, "text": " | ".join(parts), "number": number, "key": row_keys[i - 1]})
    return out


# ==============================================================================
# 2. 형식별 파서: 결과는 (units, meta). units는 UnitBuilder 형식
# ==============================================================================

def _encrypted_hint(path: str) -> rt.InputError:
    return rt.InputError(
        f"파일을 열 수 없습니다. 문서 보안(DRM)으로 암호화되었거나 손상된 파일일 수 있습니다: {os.path.basename(path)}. "
        "Office 문서라면 excel.py / word.py / ppt.py read --target \"<경로>\"로 Office에서 직접 읽어 보세요."
    )


def _pdf_tables(page) -> List[Dict[str, Any]]:
    """선으로 그린 표를 찾습니다 (pymupdf find_tables). 실패하거나 없으면 빈 목록 (단락으로만 처리)."""
    try:
        found = page.find_tables()
        tables = []
        for t in getattr(found, "tables", []) or []:
            rows = t.extract() or []
            if len(rows) >= 2 and max((len(r) for r in rows), default=0) >= 2:
                tables.append({"bbox": tuple(t.bbox), "rows": rows})
        return tables
    except Exception as e:
        rt.debug(f"PDF 표 인식 실패 (단락으로 처리): {e}")
        return []


def _inside(bbox: Tuple[float, ...], box: Tuple[float, ...]) -> bool:
    cx, cy = (bbox[0] + bbox[2]) / 2, (bbox[1] + bbox[3]) / 2
    return box[0] - 1 <= cx <= box[2] + 1 and box[1] - 1 <= cy <= box[3] + 1


def _pdf_lines(page, tables: Optional[List[Dict[str, Any]]] = None) -> List[Any]:
    """페이지의 내용을 읽는 순서로: 글 블록은 (줄 글, 글자 크기, 굵게) 목록, 표는 {"table": 행 목록}."""
    return _order_items(_pdf_items(page, tables), float(page.rect.width))


def _pdf_items(page, tables: Optional[List[Dict[str, Any]]] = None) -> List[Tuple[Tuple[float, ...], Any]]:
    """페이지의 (bbox, 내용) 목록. 표 영역 안의 글 블록은 표 행과 겹치므로 뺍니다."""
    tables = tables or []
    items: List[Tuple[Tuple[float, ...], Any]] = [(tuple(t["bbox"]), {"table": t["rows"]}) for t in tables]
    for b in page.get_text("dict").get("blocks", []):
        if b.get("type", 0) != 0:
            continue
        bbox = tuple(b.get("bbox", (0, 0, 0, 0)))
        if any(_inside(bbox, t["bbox"]) for t in tables):
            continue
        lines: List[Tuple[str, float, bool]] = []
        for ln in b.get("lines", []):
            spans = [sp for sp in ln.get("spans", []) if sp.get("text", "").strip()]
            if not spans:
                continue
            text = "".join(sp.get("text", "") for sp in ln.get("spans", [])).strip()
            size = max(float(sp.get("size", 0) or 0) for sp in spans)
            bold = all((int(sp.get("flags", 0)) & 16) or "bold" in str(sp.get("font", "")).lower() for sp in spans)
            lines.append((text, size, bool(bold)))
        if lines:
            items.append((bbox, lines))
    return items


def _order_items(items: List[Tuple[Tuple[float, ...], Any]], width: float) -> List[Any]:
    """읽는 순서. 기본은 위에서 아래. 2단 편집(왼쪽 단과 오른쪽 단이 세로로 겹침)이면 가로로 걸친 블록(제목, 표) 사이마다
    왼쪽 단을 먼저, 오른쪽 단을 다음에 읽습니다 (y로만 정렬하면 두 단의 줄이 섞여 장절이 틀어짐)."""
    def by_y(xs):
        return sorted(xs, key=lambda it: (round(it[0][1], 1), it[0][0]))
    if not items or width <= 0:
        return [it[1] for it in by_y(items)]
    mid, gap = width / 2, width * 0.04
    full = [it for it in items if it[0][0] < mid - gap and it[0][2] > mid + gap]
    left = [it for it in items if it[0][2] <= mid + gap and it not in full]
    right = [it for it in items if it[0][0] >= mid - gap and it not in full and it not in left]
    overlap = any(lb[0][1] < rb[0][3] and rb[0][1] < lb[0][3] for lb in left for rb in right)
    if len(left) < 2 or len(right) < 2 or not overlap:
        return [it[1] for it in by_y(items)]
    out: List[Any] = []
    rest = left + right
    for f in by_y(full):
        band = [it for it in rest if it[0][1] < f[0][1]]
        rest = [it for it in rest if it not in band]
        out += [it[1] for it in by_y([b for b in band if b in left]) + by_y([b for b in band if b in right])]
        out.append(f[1])
    out += [it[1] for it in by_y([b for b in rest if b in left]) + by_y([b for b in rest if b in right])]
    return out


_MARGIN = 0.08          # 쪽 위/아래 8% 안에 있는 줄만 머리글/바닥글 후보


def _margin_key(item: Tuple[Tuple[float, ...], Any], height: float) -> Optional[str]:
    """머리글/바닥글 후보 블록의 비교용 글 (숫자는 #으로: '- 3 -'과 '- 4 -'는 같은 쪽 번호 줄)."""
    bbox, payload = item
    if isinstance(payload, dict) or height <= 0:
        return None
    if not (bbox[3] <= height * _MARGIN or bbox[1] >= height * (1 - _MARGIN)):
        return None
    text = " ".join(t for t, _, _ in payload)
    return re.sub(r"\d+", "#", re.sub(r"\s+", " ", text)).strip() or None


def _drop_headers_footers(pages: List[Tuple[List[Any], float]]) -> int:
    """여러 쪽의 위/아래 여백에 반복되는 줄(문서 제목 머리글, 쪽 번호)을 뺍니다. 뺀 블록 수를 돌려줍니다.
    남겨 두면 매 쪽 장절 경로가 머리글로 바뀌고, 머리글 문구가 쪽마다 검색 결과에 걸립니다."""
    if len(pages) < 2:
        return 0
    seen: Counter = Counter()
    for items, height in pages:
        seen.update({k for k in (_margin_key(it, height) for it in items) if k})
    repeated = {k for k, n in seen.items() if n >= max(2, (len(pages) + 1) // 2)}
    if not repeated:
        return 0
    removed = 0
    for items, height in pages:
        keep = [it for it in items if _margin_key(it, height) not in repeated]
        removed += len(items) - len(keep)
        items[:] = keep
    return removed


def _import_pymupdf():
    """pymupdf를 불러오고, 라이브러리 안내/경고가 stdout(1줄 JSON 자리)에 나오지 않게 stderr로 돌립니다."""
    try:
        import pymupdf
    except ImportError:
        raise rt.InputError("PDF를 읽으려면 pymupdf가 필요합니다 (python -m pip install pymupdf).")
    for fn, kwargs in (("set_messages", {"stream": sys.stderr}), ("set_log", {"stream": sys.stderr}),
                       ("no_recommend_layout", {})):
        try:
            getattr(pymupdf, fn)(**kwargs)
        except Exception:
            pass
    return pymupdf


def parse_pdf(path: str) -> Tuple[List[Dict[str, Any]], Dict[str, Any]]:
    import contextlib
    with contextlib.redirect_stdout(sys.stderr):    # pymupdf의 print 안내도 stderr로
        return _parse_pdf(path)


def _parse_pdf(path: str) -> Tuple[List[Dict[str, Any]], Dict[str, Any]]:
    pymupdf = _import_pymupdf()
    try:
        doc = pymupdf.open(path)
    except Exception:
        raise _encrypted_hint(path)
    try:
        if doc.needs_pass:
            raise rt.InputError(f"암호가 걸린 PDF입니다: {os.path.basename(path)}")
        raw = [(_pdf_items(page, _pdf_tables(page)), float(page.rect.height), float(page.rect.width)) for page in doc]
        meta = {"pages": doc.page_count, "title": (doc.metadata or {}).get("title") or ""}
    finally:
        doc.close()
    removed = _drop_headers_footers([(items, h) for items, h, _ in raw])
    if removed:
        meta["header_footer_removed"] = removed
    pages = [_order_items(items, w) for items, _, w in raw]

    sizes: Counter = Counter()
    bold_chars = total_chars = 0
    for blocks in pages:
        for lines in blocks:
            if isinstance(lines, dict):
                continue
            for text, size, bold in lines:
                sizes[round(size, 1)] += len(text)
                total_chars += len(text)
                bold_chars += len(text) if bold else 0
    body = sizes.most_common(1)[0][0] if sizes else 0.0
    bold_is_signal = bold_chars < total_chars * 0.5
    meta["body_font_size"] = body

    b = UnitBuilder()
    for pno, blocks in enumerate(pages, 1):
        loc = f"p.{pno}"
        table_no = 0
        for lines in blocks:
            if isinstance(lines, dict):           # 표: 행 단위 (머리글: 값)
                table_no += 1
                for r in table_rows(lines["table"]):
                    b.add("row", f"p.{pno} 표 {table_no} {r['row']}행", r["text"], pno, "pdf", r["number"])
                continue
            buf: List[str] = []

            def flush() -> None:
                if buf:
                    b.add_text(loc, " ".join(buf), pno, "pdf", allow_heading=False)
                    buf.clear()

            for text, size, bold in lines:
                strong = size >= body + 1.0 or (bold_is_signal and bold)
                info = classify_number(text)
                if heading_title(info, text, strong):
                    flush()
                    b.add_text(loc, text, pno, "pdf", strong=strong, size=size)
                elif info is not None:          # 번호/글머리로 시작하는 줄은 새 단락
                    flush()
                    buf.append(text)
                else:
                    buf.append(text)
            flush()
    if not b.units:
        meta["ocr_needed"] = True
        rt.log("WARN", "텍스트가 없는 PDF입니다 (스캔 이미지일 가능성). OCR이 필요합니다.")
    return b.units, meta


# HWP 5.0 제어 문자: 1 WCHAR 짜리 문자 컨트롤. 나머지 1~31은 인라인/확장 컨트롤로 8 WCHAR(16바이트)를 차지합니다.
_HWP_CHAR_CONTROLS = {0, 10, 13, 24, 25, 26, 27, 28, 29, 30, 31}
_HWP_CHAR_MAP = {10: "\n", 24: "-", 30: " ", 31: " "}
_HWPTAG_PARA_TEXT = 67


def _hwp_para_text(rec: bytes) -> str:
    out: List[str] = []
    i, n = 0, len(rec)
    while i + 1 < n:
        code = rec[i] | (rec[i + 1] << 8)
        if 0xD800 <= code <= 0xDBFF and i + 3 < n:          # UTF-16 서로게이트 쌍 (확장 한자 등 BMP 밖 글자)
            low = rec[i + 2] | (rec[i + 3] << 8)
            if 0xDC00 <= low <= 0xDFFF:
                out.append(chr(0x10000 + ((code - 0xD800) << 10) + (low - 0xDC00)))
                i += 4
                continue
        if 0xD800 <= code <= 0xDFFF:                         # 짝이 없는 서로게이트: 캐시 저장(UTF-8)이 실패하지 않게
            out.append("\ufffd")
            i += 2
        elif code >= 32:
            out.append(chr(code))
            i += 2
        elif code in _HWP_CHAR_CONTROLS:
            out.append(_HWP_CHAR_MAP.get(code, ""))
            i += 2
        else:
            if code == 9:
                out.append("\t")
            i += 16
    return "".join(out)


def parse_hwp(path: str) -> Tuple[List[Dict[str, Any]], Dict[str, Any]]:
    try:
        import olefile
    except ImportError:
        raise rt.InputError("HWP를 읽으려면 olefile이 필요합니다 (python -m pip install olefile).")
    if not olefile.isOleFile(path):
        raise _encrypted_hint(path)
    ole = olefile.OleFileIO(path)
    try:
        header = ole.openstream("FileHeader").read()
        props = int.from_bytes(header[36:40], "little") if len(header) >= 40 else 0
        compressed = bool(props & 0x1)
        if props & 0x2:
            raise rt.InputError(f"암호가 걸린 HWP 문서입니다: {os.path.basename(path)}")
        if props & 0x4:
            raise rt.InputError(f"배포용 HWP 문서라 본문을 추출할 수 없습니다: {os.path.basename(path)}")

        sections = [d for d in ole.listdir() if len(d) == 2 and d[0] == "BodyText" and d[1].startswith("Section")]
        sections.sort(key=lambda d: int(re.sub(r"\D", "", d[1]) or 0))
        paragraphs: List[str] = []
        for sec in sections:
            data = ole.openstream(sec).read()
            if compressed:
                data = zlib.decompress(data, -15)
            pos, size_total = 0, len(data)
            while pos + 4 <= size_total:
                h = int.from_bytes(data[pos:pos + 4], "little")
                pos += 4
                tag, size = h & 0x3FF, (h >> 20) & 0xFFF
                if size == 0xFFF:
                    size = int.from_bytes(data[pos:pos + 4], "little")
                    pos += 4
                if tag == _HWPTAG_PARA_TEXT:
                    text = _hwp_para_text(data[pos:pos + size]).strip()
                    if text:
                        paragraphs.append(text)
                pos += size
    finally:
        ole.close()
    b = UnitBuilder()
    for i, t in enumerate(paragraphs, 1):
        b.add_text(f"단락 {i}", t)
    return b.units, {"format": "HWP 5.0", "sections": len(sections)}


def _local(tag: Any) -> str:
    return str(tag).rsplit("}", 1)[-1]


# 개요 번호 형식 (HWPX numFormat / Word numFmt 공통으로 씀)
def _roman(n: int) -> str:
    vals = [(1000, "M"), (900, "CM"), (500, "D"), (400, "CD"), (100, "C"), (90, "XC"), (50, "L"), (40, "XL"),
            (10, "X"), (9, "IX"), (5, "V"), (4, "IV"), (1, "I")]
    out = ""
    for v, r in vals:
        while n >= v:
            out += r
            n -= v
    return out or str(n)


def _letters(n: int) -> str:
    out = ""
    while n > 0:
        n, r = divmod(n - 1, 26)
        out = chr(65 + r) + out
    return out or "A"


def format_counter(n: int, fmt: str) -> str:
    """번호 형식 이름(Word numFmt, HWPX numFormat)으로 숫자를 글자로 바꿉니다. 원문자 형식은 '(n)'으로 씁니다."""
    f = (fmt or "decimal").lower().replace("_", "")
    if n <= 0:
        n = 1
    if f in ("upperroman", "romancapital"):
        return _roman(n)
    if f in ("lowerroman", "romansmall"):
        return _roman(n).lower()
    if f in ("upperletter", "latincapital"):
        return _letters(n)
    if f in ("lowerletter", "latinsmall"):
        return _letters(n).lower()
    if f in ("ganada", "hangulsyllable"):
        return GANADA[(n - 1) % len(GANADA)]
    if f in ("chosung", "hanguljamo"):
        return CHOSUNG[(n - 1) % len(CHOSUNG)]
    if f in ("circledhangulsyllable",):
        return f"({GANADA[(n - 1) % len(GANADA)]})"
    if "circle" in f or "enclosed" in f:
        return f"({n})"
    if f == "decimalzero":
        return f"{n:02d}"
    return str(n)


def _hwpx_header_info(z: zipfile.ZipFile) -> Tuple[Dict[str, Tuple[str, str, int]], Dict[str, Dict[int, Tuple[int, str, str]]]]:
    """header.xml에서 (문단모양 id -> (머리 종류, 번호 id, 수준)), (번호 id -> {수준: (시작, 형식, 글자틀)})."""
    para: Dict[str, Tuple[str, str, int]] = {}
    numbering: Dict[str, Dict[int, Tuple[int, str, str]]] = {}
    names = [n for n in z.namelist() if n.lower().endswith("header.xml")]
    if not names:
        return para, numbering
    root = ET.fromstring(z.read(names[0]))
    for el in root.iter():
        tag = _local(el.tag)
        if tag == "paraPr" and el.get("id") is not None:
            for sub in el.iter():
                if _local(sub.tag) == "heading":
                    try:
                        para[el.get("id", "")] = (sub.get("type", "NONE").upper(), sub.get("idRef", ""),
                                                  int(sub.get("level", "0")))
                    except ValueError:
                        pass
                    break
        elif tag == "numbering" and el.get("id") is not None:
            levels: Dict[int, Tuple[int, str, str]] = {}
            for sub in el:
                if _local(sub.tag) == "paraHead":
                    try:
                        levels[int(sub.get("level", "1"))] = (int(sub.get("start", "1")), sub.get("numFormat", "DIGIT"),
                                                              sub.text or "")
                    except ValueError:
                        pass
            numbering[el.get("id", "")] = levels
    return para, numbering


def _hwpx_walk(elem: ET.Element, buf: Optional[List[str]], out: List[Dict[str, Any]]) -> None:
    """문단(<p>)마다 자기 텍스트(<t>)만 모읍니다. 표 안의 문단은 별도 문단으로, 문서 순서를 유지합니다."""
    tag = _local(elem.tag)
    if tag == "p":
        slot = len(out)
        out.append({"text": "", "para": elem.get("paraPrIDRef", ""), "page_break": elem.get("pageBreak") == "1"})
        own: List[str] = []
        for child in elem:
            _hwpx_walk(child, own, out)
        out[slot]["text"] = "".join(own)
        return
    if tag == "t" and buf is not None:
        if elem.text:
            buf.append(elem.text)
        for child in elem:
            if _local(child.tag) == "tab":
                buf.append("\t")
            if child.tail:
                buf.append(child.tail)
        return
    for child in elem:
        _hwpx_walk(child, buf, out)


def parse_hwpx(path: str) -> Tuple[List[Dict[str, Any]], Dict[str, Any]]:
    if not zipfile.is_zipfile(path):
        raise _encrypted_hint(path)
    with zipfile.ZipFile(path) as z:
        names = [n for n in z.namelist() if re.match(r"Contents/section\d+\.xml$", n)]
        if not names:
            raise rt.InputError(f"HWPX 본문(Contents/section*.xml)이 없습니다: {os.path.basename(path)}")
        names.sort(key=lambda n: int(re.sub(r"\D", "", n)))
        try:
            para_heads, numbering = _hwpx_header_info(z)
        except Exception as e:                      # 개요 정보가 없으면 번호 패턴으로만 인식
            rt.debug(f"HWPX header.xml 해석 실패: {e}")
            para_heads, numbering = {}, {}
        paragraphs: List[Dict[str, Any]] = []
        for name in names:
            _hwpx_walk(ET.fromstring(z.read(name)), None, paragraphs)
    use_pages = any(p["page_break"] for p in paragraphs)
    page = 1
    counters: Dict[str, List[int]] = {}
    b = UnitBuilder()
    n = 0
    for p in paragraphs:
        if p["page_break"] and b.units:
            page += 1
        text = (p["text"] or "").strip()
        if not text:
            continue
        n += 1
        kind_level = para_heads.get(p["para"])
        style_level = None
        if kind_level and kind_level[0] == "OUTLINE":
            style_level = kind_level[2] + 1
            levels = numbering.get(kind_level[1]) or {}
            fmt = levels.get(style_level)
            if fmt:                                  # 개요 번호(자동 번호)를 글 앞에 붙임: ^1.^2 -> 1.2
                cnt = counters.setdefault(kind_level[1], [0] * 11)
                cnt[style_level] += 1
                for k in range(style_level + 1, 11):
                    cnt[k] = 0

                def repl(m: "re.Match[str]") -> str:
                    lv = int(m.group(1))
                    start, numfmt, _ = levels.get(lv, (1, "DIGIT", ""))
                    return format_counter(max(cnt[lv], 1) + start - 1, numfmt)
                label = re.sub(r"\^(\d+)", repl, fmt[2]).strip()
                if label and not text.startswith(label):
                    text = f"{label} {text}"
        b.add_text(f"단락 {n}", text, page if use_pages else None, "page_break" if use_pages else None,
                   style_level=style_level)
    return b.units, {"format": "HWPX", "sections": len(names)}


# ---------------------------------------------------------------- DOCX

_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _wval(el: Any, path: str) -> Optional[str]:
    """요소 아래 path(w: 태그를 '/'로 이은 것)의 w:val 값."""
    cur = el
    for part in path.split("/"):
        if cur is None:
            return None
        cur = cur.find(_W + part)
    return None if cur is None else cur.get(_W + "val")


class _DocxNumbering:
    """numbering.xml로 Word 자동 번호 글자(예: 3.2.1.)를 계산합니다."""

    def __init__(self, document: Any) -> None:
        self.abstract: Dict[str, Dict[int, Dict[str, Any]]] = {}
        self.nums: Dict[str, Tuple[str, Dict[int, int]]] = {}
        self.counters: Dict[str, List[Optional[int]]] = {}
        self.seen: set = set()
        try:
            root = document.part.numbering_part.element
        except Exception:
            return
        for an in root.findall(_W + "abstractNum"):
            aid = an.get(_W + "abstractNumId", "")
            lvls: Dict[int, Dict[str, Any]] = {}
            for lvl in an.findall(_W + "lvl"):
                try:
                    ilvl = int(lvl.get(_W + "ilvl", "0"))
                    start = int(_wval(lvl, "start") or 1)
                except ValueError:
                    continue
                lvls[ilvl] = {"start": start, "fmt": _wval(lvl, "numFmt") or "decimal",
                              "text": _wval(lvl, "lvlText") or "", "lgl": lvl.find(_W + "isLgl") is not None,
                              "style": _wval(lvl, "pStyle")}
            self.abstract[aid] = lvls
        for num in root.findall(_W + "num"):
            nid = num.get(_W + "numId", "")
            aid = _wval(num, "abstractNumId") or ""
            overrides: Dict[int, int] = {}
            for ov in num.findall(_W + "lvlOverride"):
                so = _wval(ov, "startOverride")
                if so is not None:
                    try:
                        overrides[int(ov.get(_W + "ilvl", "0"))] = int(so)
                    except ValueError:
                        pass
            self.nums[nid] = (aid, overrides)

    def level_for_style(self, num_id: str, style_id: str) -> Optional[int]:
        aid = self.nums.get(num_id, ("", {}))[0]
        for ilvl, lv in self.abstract.get(aid, {}).items():
            if lv.get("style") == style_id:
                return ilvl
        return None

    def label(self, num_id: str, ilvl: int) -> Optional[Tuple[str, str]]:
        """(번호 글자, 계열 이름). 글머리 기호나 번호 없음은 ('', 계열)."""
        if not num_id or num_id == "0" or num_id not in self.nums:
            return None
        aid, overrides = self.nums[num_id]
        lvls = self.abstract.get(aid)
        if not lvls or ilvl not in lvls:
            return None
        cnt = self.counters.setdefault(aid, [None] * 10)
        if num_id not in self.seen:
            self.seen.add(num_id)
            for lv, start in overrides.items():
                if 0 <= lv < 10:
                    cnt[lv] = start - 1
                    for k in range(lv + 1, 10):
                        cnt[k] = None
        cur = cnt[ilvl]
        cnt[ilvl] = (cur + 1) if cur is not None else lvls[ilvl]["start"]
        for k in range(ilvl + 1, 10):
            cnt[k] = None
        info = lvls[ilvl]
        fam = f"wlist:{aid}"
        if info["fmt"] in ("bullet", "none"):
            return "", fam

        def repl(m: "re.Match[str]") -> str:
            k = int(m.group(1)) - 1
            lv = lvls.get(k, {"start": 1, "fmt": "decimal"})
            value = cnt[k] if cnt[k] is not None else lv["start"]
            return format_counter(value, "decimal" if info["lgl"] else lv["fmt"])
        return re.sub(r"%(\d)", repl, info["text"]).strip(), fam


def _style_chain(style: Any) -> List[Any]:
    chain = []
    while style is not None and len(chain) < 10:
        chain.append(style)
        style = getattr(style, "base_style", None)
    return chain


def _docx_style_info(para: Any, numbering: _DocxNumbering) -> Tuple[Optional[int], bool, Optional[Tuple[str, int]]]:
    """(제목 수준, 문서 제목 여부, (numId, ilvl)). 제목 수준: Heading N / 제목 N / 개요 수준(outlineLvl)."""
    style = para.style
    name = (style.name if style is not None else "") or ""
    level: Optional[int] = None
    is_title = name.strip().lower() in ("title", "subtitle") or name.strip() in ("제목", "부제")
    m = re.match(r"(?i)(heading|제목)\s*(\d)$", name.strip())
    if m:
        level = int(m.group(2))
    ppr = para._p.pPr
    outline = _wval(ppr, "outlineLvl") if ppr is not None else None
    chain = _style_chain(style)
    if outline is None:
        for st in chain:
            spr = st.element.find(_W + "pPr")
            outline = _wval(spr, "outlineLvl") if spr is not None else None
            if outline is not None:
                break
    if level is None and outline is not None and outline.isdigit() and int(outline) < 9:
        level = int(outline) + 1
    num: Optional[Tuple[str, int]] = None
    if ppr is not None and ppr.find(_W + "numPr") is not None:
        nid = _wval(ppr, "numPr/numId")
        ilvl = _wval(ppr, "numPr/ilvl")
        if nid:
            num = (nid, int(ilvl) if ilvl and ilvl.isdigit() else 0)
    if num is None:
        for st in chain:
            spr = st.element.find(_W + "pPr")
            nid = _wval(spr, "numPr/numId") if spr is not None else None
            if nid:
                ilvl = _wval(spr, "numPr/ilvl")
                lv = numbering.level_for_style(nid, st.style_id)
                num = (nid, lv if lv is not None else (int(ilvl) if ilvl and ilvl.isdigit() else 0))
                break
    return level, is_title, num


class _DocxPages:
    """저장된 쪽 정보로 단락의 쪽 번호를 셉니다.
    w:lastRenderedPageBreak(Word가 저장할 때 기록한 쪽 경계)가 있으면 그것만 쓰고(word_saved_layout),
    없으면 강제 쪽 나눔(w:br type=page, pageBreakBefore, 다음 쪽 구역 나누기)만 셉니다(page_break, 실제보다 적을 수 있음)."""

    def __init__(self, body: Any) -> None:
        self.mode: Optional[str] = None
        if any(True for _ in body.iter(_W + "lastRenderedPageBreak")):
            self.mode = "word_saved_layout"
        elif (any(br.get(_W + "type") == "page" for br in body.iter(_W + "br"))
              or any(True for _ in body.iter(_W + "pageBreakBefore"))
              or any(self._section_break(s) for s in body.iter(_W + "sectPr") if s.getparent().tag == _W + "pPr")):
            self.mode = "page_break"
        self.page = 1
        self.pending = 0
        self.started = False

    @staticmethod
    def _section_break(sect: Any) -> bool:
        t = _wval(sect, "type")
        return t not in ("continuous",)

    def element_page(self, el: Any) -> Optional[int]:
        """요소(단락/표)의 첫 글자가 있는 쪽. 요소 안의 쪽 경계도 반영해 다음 요소로 넘깁니다."""
        if self.mode is None:
            return None
        self.page += self.pending
        self.pending = 0
        if self.mode == "page_break" and self.started:
            ppr = el.find(_W + "pPr") if el.tag == _W + "p" else None
            if ppr is not None and ppr.find(_W + "pageBreakBefore") is not None:
                v = ppr.find(_W + "pageBreakBefore").get(_W + "val")
                if v not in ("0", "false"):
                    self.page += 1
        first: Optional[int] = None
        for node in el.iter():
            tag = node.tag
            if tag == _W + "t" and node.text and node.text.strip():
                if first is None:
                    first = self.page
                self.started = True
            elif self.mode == "word_saved_layout" and tag == _W + "lastRenderedPageBreak":
                if self.started or first is not None:
                    self.page += 1
            elif self.mode == "page_break" and tag == _W + "br" and node.get(_W + "type") == "page":
                if self.started or first is not None:
                    self.page += 1
            elif (self.mode == "page_break" and tag == _W + "sectPr" and node.getparent().tag == _W + "pPr"
                  and self._section_break(node)):
                self.pending += 1
        return first if first is not None else self.page


_SKIP_TEXT_IN = {_W + "del", _W + "moveFrom", _W + "txbxContent"}   # 지운 글, 옮기기 전 자리, 글상자(따로 단락)


def _p_text(el: Any) -> str:
    """단락(w:p)의 글. python-docx의 Paragraph.text와 달리 변경 추적으로 삽입된 글(w:ins)과 옮긴 글(w:moveTo)도 넣고,
    지운 글(w:del)은 뺍니다. 탭과 줄바꿈은 글자로."""
    out: List[str] = []

    def walk(node: Any) -> None:
        for ch in node:
            tag = ch.tag
            if tag in _SKIP_TEXT_IN:
                continue
            if tag == _W + "t":
                out.append(ch.text or "")
            elif tag == _W + "tab":
                out.append("\t")
            elif tag in (_W + "br", _W + "cr"):
                out.append("\n")
            else:
                walk(ch)
    walk(el)
    return "".join(out)


def _body_blocks(parent: Any):
    """본문의 단락(w:p)과 표(w:tbl)를 문서 순서로. 콘텐츠 컨트롤(w:sdt)과 customXml 안에 든 것도 꺼냅니다."""
    for child in parent.iterchildren():
        tag = child.tag
        if tag in (_W + "p", _W + "tbl"):
            yield child
        elif tag == _W + "sdt":
            content = child.find(_W + "sdtContent")
            if content is not None:
                yield from _body_blocks(content)
        elif tag == _W + "customXml":
            yield from _body_blocks(child)


def parse_docx(path: str) -> Tuple[List[Dict[str, Any]], Dict[str, Any]]:
    if not zipfile.is_zipfile(path):
        raise _encrypted_hint(path)
    try:
        import docx
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError:
        raise rt.InputError("DOCX를 읽으려면 python-docx가 필요합니다 (python -m pip install python-docx).")
    try:
        document = docx.Document(path)
    except Exception:
        raise _encrypted_hint(path)
    numbering = _DocxNumbering(document)
    pages = _DocxPages(document.element.body)
    b = UnitBuilder()
    para_no = table_no = 0
    numbering_failed = False
    for child in _body_blocks(document.element.body):
        tag = _local(child.tag)
        if tag == "p":
            para = Paragraph(child, document)
            page = pages.element_page(child)
            text = _p_text(child).strip()
            if not text:
                # 빈 번호 단락도 Word는 번호를 매김 ("3." 빈 줄): 번호 카운터를 진행해야 뒤 조항 번호가 밀리지 않음
                try:
                    _, _, num = _docx_style_info(para, numbering)
                    if num:
                        numbering.label(*num)
                except Exception:
                    pass
                continue
            para_no += 1
            loc = f"단락 {para_no}"
            try:
                level, is_title, num = _docx_style_info(para, numbering)
                lab = numbering.label(*num) if num else None
            except Exception as e:                  # 번호 해석 실패 -> 글의 번호 패턴만 사용
                if not numbering_failed:
                    rt.debug(f"Word 번호 해석 실패, 글 패턴으로 대신함: {e}")
                numbering_failed = True
                level, is_title, lab = None, False, None
            list_family, list_depth = None, 1
            if lab:
                label, list_family = lab
                list_depth = (num[1] + 1) if num else 1
                if label and not text.startswith(label):
                    text = f"{label} {text}"
            if is_title and not level:
                b.add("heading", loc, text, page, pages.mode, None, section=[text[:120]])
                continue
            b.add_text(loc, text, page, pages.mode, style_level=level, list_family=list_family, list_depth=list_depth)
        elif tag == "tbl":
            table_no += 1
            grid: List[List[str]] = []
            row_pages: List[Optional[int]] = []
            for row in Table(child, document).rows:
                row_pages.append(pages.element_page(row._tr))      # 여러 쪽에 걸친 표: 행마다 그 행이 있는 쪽
                cells: List[str] = []
                prev_tc = None
                for cell in row.cells:
                    # 가로 병합 칸은 같은 칸이 반복되므로 빈 칸으로 (머리글은 앞 칸을 이어 씀)
                    cells.append("" if cell._tc is prev_tc else
                                 "\n".join(_p_text(p) for p in cell._tc.iter(_W + "p")))
                    prev_tc = cell._tc
                grid.append(cells)
            for r in table_rows(grid, row_pages):
                b.add("row", f"표 {table_no} {r['row']}행", r["text"], r["key"], pages.mode, r["number"])
    core = document.core_properties
    meta: Dict[str, Any] = {"title": core.title or "", "paragraphs": para_no, "tables": table_no,
                            "page_source": pages.mode}
    if pages.mode:
        meta["pages"] = max([u["page"] for u in b.units if u["page"]] or [1])
    return b.units, meta


# ---------------------------------------------------------------- PPTX / XLSX / CSV / TXT

def _pptx_shape_lines(shape, out: List[str]) -> None:
    """도형 하나의 텍스트를 줄 단위로 모읍니다 (그룹은 재귀, 표는 행 단위)."""
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        for sub in shape.shapes:
            _pptx_shape_lines(sub, out)
    elif getattr(shape, "has_table", False):
        for row in shape.table.rows:
            out.append(" | ".join(c.text.strip().replace("\n", " ") for c in row.cells))
    elif getattr(shape, "has_text_frame", False):
        for para in shape.text_frame.paragraphs:
            text = para.text.replace("\v", " ").strip()
            if text:
                out.append(text)


def parse_pptx(path: str) -> Tuple[List[Dict[str, Any]], Dict[str, Any]]:
    """슬라이드 하나 = 단위 하나 (ppt.py/office.py가 슬라이드 번호로 씀). 장절 = [슬라이드 제목]."""
    if not zipfile.is_zipfile(path):
        raise _encrypted_hint(path)
    try:
        import pptx
    except ImportError:
        raise rt.InputError("PPTX를 읽으려면 python-pptx가 필요합니다 (python -m pip install python-pptx).")
    try:
        prs = pptx.Presentation(path)
    except Exception:
        raise _encrypted_hint(path)
    b = UnitBuilder()
    for idx, slide in enumerate(prs.slides, 1):   # presentation.xml의 실제 슬라이드 순서
        title_shape = slide.shapes.title
        title = title_shape.text.strip() if title_shape is not None and title_shape.has_text_frame else ""
        lines: List[str] = []
        for shape in slide.shapes:
            if title_shape is not None and shape.shape_id == title_shape.shape_id:
                continue
            _pptx_shape_lines(shape, lines)
        if not title and lines:
            title = lines[0]
        notes = ""
        if slide.has_notes_slide:
            notes = (slide.notes_slide.notes_text_frame.text or "").strip()
        body = "\n".join(([title] if title and (not lines or lines[0] != title) else []) + lines)
        if notes:
            body += f"\n[노트] {notes}"
        title_line = title.replace("\n", " ").strip()
        info = classify_number(title_line)
        number = info["number"] if info and info["family"] != "bullet" and not info["plain"] else None
        b.add("slide", f"슬라이드 {idx}", body.strip(), idx, "slide", number,
              section=[title_line] if title_line else [])
    return b.units, {"slides": len(prs.slides), "pages": len(prs.slides)}


def _cell_str(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if hasattr(v, "isoformat"):
        return v.isoformat(sep=" ") if hasattr(v, "hour") else v.isoformat()
    return str(v)


def parse_xlsx(path: str) -> Tuple[List[Dict[str, Any]], Dict[str, Any]]:
    if not zipfile.is_zipfile(path):
        raise _encrypted_hint(path)
    try:
        import openpyxl
    except ImportError:
        raise rt.InputError("XLSX를 읽으려면 openpyxl이 필요합니다 (python -m pip install openpyxl).")
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception:
        raise _encrypted_hint(path)
    b = UnitBuilder()
    sheets: List[str] = []
    truncated = False
    try:
        for ws in wb.worksheets:
            sheets.append(ws.title)
            count = 0
            for r_idx, row in enumerate(ws.iter_rows(values_only=True), 1):
                cells = [_cell_str(v) for v in row]
                while cells and cells[-1] == "":
                    cells.pop()
                if not any(cells):
                    continue
                count += 1
                if count > MAX_SHEET_ROWS:
                    truncated = True
                    break
                b.add("row", f"{ws.title}!{r_idx}행", "\t".join(cells), section=[ws.title])
    finally:
        wb.close()
    meta: Dict[str, Any] = {"sheets": sheets}
    if truncated:
        meta["rows_truncated"] = True
        rt.log("WARN", f"시트당 {MAX_SHEET_ROWS}행까지만 읽었습니다.")
    return b.units, meta


def parse_csv(path: str) -> Tuple[List[Dict[str, Any]], Dict[str, Any]]:
    rows = rt.read_csv_rows(path)
    b = UnitBuilder()
    for i, r in enumerate(rows, 1):
        if any(c.strip() for c in r):
            b.add("row", f"{i}행", "\t".join(r), section=[])
    return b.units, {"rows": len(rows)}


def parse_txt(path: str) -> Tuple[List[Dict[str, Any]], Dict[str, Any]]:
    """빈 줄로 나눈 단락. 단락 첫 줄이 제목(마크다운 #, 또는 번호 제목)이면 제목 단위로 따로 둡니다."""
    text = rt.read_text_file(path)
    b = UnitBuilder()
    n = 0
    for block in [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]:
        lines = block.splitlines()
        first = lines[0].strip()
        md = re.match(r"^(#{1,6})\s+(.*)$", first)
        if md:
            n += 1
            b.add_text(f"단락 {n}", md.group(2).strip().rstrip("#").strip(), style_level=len(md.group(1)))
            rest = "\n".join(lines[1:]).strip()
        elif len(lines) > 1 and heading_title(classify_number(first), first):
            n += 1
            b.add_text(f"단락 {n}", first)
            rest = "\n".join(lines[1:]).strip()
        else:
            rest = block
        if rest:
            n += 1
            b.add_text(f"단락 {n}", rest, allow_heading=len(rest.splitlines()) == 1)
    return b.units, {}


PARSERS = {"pdf": parse_pdf, "hwp": parse_hwp, "hwpx": parse_hwpx, "docx": parse_docx,
           "pptx": parse_pptx, "xlsx": parse_xlsx, "csv": parse_csv, "txt": parse_txt}


def file_type(path: str) -> str:
    ext = os.path.splitext(path)[1].lower()
    if ext in LEGACY_OFFICE:
        raise rt.InputError(f"구형 Office 형식({ext})은 라이브러리로 읽을 수 없습니다. "
                            f"excel.py / word.py / ppt.py read --target \"{path}\"로 Office에서 읽으세요.")
    if ext not in SUPPORTED:
        raise rt.InputError(f"지원하지 않는 형식입니다: {ext or '(확장자 없음)'} (지원: {', '.join(sorted(SUPPORTED))})")
    return SUPPORTED[ext]


# ==============================================================================
# 3. 캐시와 색인
# ==============================================================================

def _doc_cache_path(path: str) -> str:
    key = hashlib.sha1(os.path.normcase(os.path.abspath(path)).encode("utf-8")).hexdigest()[:16]
    folder = os.path.join(rt.get_cache_dir(), "docs")
    os.makedirs(folder, exist_ok=True)
    return os.path.join(folder, f"{key}.json")


def _index_path() -> str:
    return os.path.join(rt.get_cache_dir(), "index.json")


def load_index() -> Dict[str, Dict[str, Any]]:
    """색인 {원본 경로: 정보}. Windows에서 대소문자만 다른 같은 파일(C:\\Docs, c:\\docs)은 최근 것 하나로 합칩니다."""
    p = _index_path()
    if not os.path.exists(p):
        return {}
    try:
        data = json.loads(rt.read_text_file(p))
    except Exception:
        rt.log("WARN", "색인 파일이 손상되어 새로 만듭니다.")
        return {}
    if not isinstance(data, dict):
        return {}
    merged: Dict[str, Tuple[str, Dict[str, Any]]] = {}
    for path, entry in data.items():
        key = os.path.normcase(os.path.abspath(path))
        if key not in merged or str(entry.get("parsed_at", "")) > str(merged[key][1].get("parsed_at", "")):
            merged[key] = (path, entry)
    return {path: entry for path, entry in merged.values()}


def _save_index(index: Dict[str, Dict[str, Any]]) -> None:
    rt.safe_write_text(_index_path(), json.dumps(index, ensure_ascii=False))


def _index_put(index: Dict[str, Dict[str, Any]], path: str, entry: Dict[str, Any]) -> None:
    key = os.path.normcase(path)
    for old in [k for k in index if os.path.normcase(k) == key]:
        del index[old]
    index[path] = entry


def _index_entry(doc: Dict[str, Any], cache_file: str) -> Dict[str, Any]:
    return {"type": doc["type"], "units": len(doc["units"]), "parsed_at": doc.get("parsed_at", ""),
            "chars": sum(len(u["text"]) for u in doc["units"]), "cache_file": os.path.basename(cache_file),
            "mtime": doc.get("mtime"), "size": doc.get("size"), "parser_version": doc.get("parser_version"),
            "pages": _doc_pages(doc), "sections": _doc_sections(doc)}


def _ensure_indexed(path: str, doc: Dict[str, Any], cache_file: str) -> None:
    """캐시는 최신인데 색인에 없거나 다르면(색인 저장 전 중단, 동시 실행) 색인을 고칩니다 (JIT 복구, 실패는 무시)."""
    try:
        index = load_index()
        entry = next((e for k, e in index.items() if os.path.normcase(k) == os.path.normcase(path)), None)
        if entry and all(entry.get(k) == doc.get(k) for k in ("mtime", "size", "parser_version")):
            return
        _index_put(index, path, _index_entry(doc, cache_file))
        _save_index(index)
        rt.log("WARN", f"캐시는 있는데 색인에서 빠진 문서를 색인에 다시 넣었습니다: {os.path.basename(path)}")
    except Exception:
        pass


_KIND_FROM_LOC = (("표", "table"), ("슬라이드", "slide"), ("헤더", "header"))  # 이전 캐시의 loc로 종류 추정


def normalize_doc(doc: Dict[str, Any]) -> Dict[str, Any]:
    """이전 파서 버전의 캐시도 같은 키로 검색되게 빈 키를 채웁니다 (loc에서 쪽/슬라이드 번호 추정)."""
    for i, u in enumerate(doc.get("units") or [], 1):
        loc = str(u.get("loc", ""))
        u.setdefault("id", i)
        if "kind" not in u:
            kind = "row" if loc.endswith("행") else "paragraph"
            for prefix, k in _KIND_FROM_LOC:
                if loc.startswith(prefix):
                    kind = k
            u["kind"] = kind
        if "page" not in u:
            m = re.fullmatch(r"p\.(\d+)", loc)
            s = re.fullmatch(r"슬라이드 (\d+)", loc)
            u["page"] = int(m.group(1)) if m else (int(s.group(1)) if s else None)
            u["page_source"] = "pdf" if m else ("slide" if s else None)
        u.setdefault("page_source", None)
        if "section" not in u:
            u["section"] = [u["heading"]] if u.get("heading") else []
        u.setdefault("number", None)
        u.setdefault("heading", u["section"][-1] if u["section"] else "")
        u.setdefault("text", "")
    return doc


def is_stale_doc(doc: Dict[str, Any], st: Optional[os.stat_result] = None) -> bool:
    """캐시가 이전 파서 버전이거나, 원본 파일이 캐시한 뒤 바뀌었으면 True (원본이 없으면 파서 버전만 봄)."""
    if doc.get("parser_version") != PARSER_VERSION:
        return True
    if st is None:
        try:
            st = os.stat(doc.get("file", ""))
        except OSError:
            return False
    return doc.get("mtime") != st.st_mtime or doc.get("size") != st.st_size


def _doc_pages(doc: Dict[str, Any]) -> Optional[int]:
    meta = doc.get("meta") or {}
    if isinstance(meta.get("pages"), int):
        return meta["pages"]
    pages = [u.get("page") for u in doc.get("units", []) if isinstance(u.get("page"), int)]
    return max(pages) if pages else None


def _doc_sections(doc: Dict[str, Any]) -> int:
    units = doc.get("units", [])
    heads = sum(1 for u in units if u.get("kind") == "heading")
    if heads:
        return heads
    return len({tuple(u.get("section") or []) for u in units if u.get("section")})


def load_cached_only(path: str) -> Optional[Dict[str, Any]]:
    cache_file = _doc_cache_path(path)
    if not os.path.exists(cache_file):
        return None
    try:
        doc = json.loads(rt.read_text_file(cache_file))
        return normalize_doc(doc) if isinstance(doc, dict) and isinstance(doc.get("units"), list) else None
    except Exception:
        return None


def load_document_status(path: str, use_cache: bool = True, check_index: bool = True) -> Tuple[Dict[str, Any], str]:
    """파일을 캐시에서 불러오거나(수정 시각/크기/파서 버전이 같으면) 파싱해 캐시와 색인에 등록합니다.
    반환: (문서, 상태) 상태는 'cached'(캐시 사용), 'new'(처음 등록), 'updated'(바뀌어서 다시 등록)."""
    path = os.path.abspath(path)
    ftype = file_type(path)
    st = os.stat(path)
    cache_file = _doc_cache_path(path)
    had_cache = False
    if use_cache and os.path.exists(cache_file):
        had_cache = True
        try:
            cached = json.loads(rt.read_text_file(cache_file))
            if not is_stale_doc(cached, st):
                rt.debug(f"캐시 사용: {os.path.basename(path)}")
                if check_index:
                    _ensure_indexed(path, cached, cache_file)
                return normalize_doc(cached), "cached"
        except Exception:
            rt.log("WARN", f"캐시가 손상되어 다시 읽습니다: {os.path.basename(path)}")
    rt.log("PROGRESS", f"문서 읽는 중: {os.path.basename(path)}")
    units, meta = PARSERS[ftype](path)
    doc = {"file": path, "type": ftype, "mtime": st.st_mtime, "size": st.st_size,
           "parser_version": PARSER_VERSION, "parsed_at": rt.timestamp(), "meta": meta, "units": units}
    try:
        rt.safe_write_text(cache_file, json.dumps(doc, ensure_ascii=False))
        index = load_index()
        _index_put(index, path, _index_entry(doc, cache_file))
        _save_index(index)
    except OSError as e:
        rt.log("WARN", f"캐시 저장 실패 (결과에는 영향 없음): {e}")
    return doc, ("updated" if had_cache else "new")


def load_document(path: str, use_cache: bool = True) -> Dict[str, Any]:
    return load_document_status(path, use_cache)[0]


def _under(path: str, folder: str) -> bool:
    p = os.path.normcase(os.path.abspath(path))
    f = os.path.normcase(os.path.abspath(folder)).rstrip("\\/") + os.sep
    return p.startswith(f)


def _collect_folder(folder: str, depth: int = MAX_FOLDER_DEPTH,
                    excluded: Optional[List[str]] = None) -> Tuple[List[str], Dict[str, int]]:
    """폴더의 지원 형식 파일 목록과 건너뛴 형식별 개수 (예: {'.png': 2}). depth 0 = 그 폴더만.
    excluded를 주면 들어가지 않은 하위 폴더(깊이 제한, '_'/'.'로 시작)를 거기에 담습니다."""
    files: List[str] = []
    skipped: Counter = Counter()
    base_depth = folder.rstrip("\\/").count(os.sep)
    for root, dirs, names in os.walk(folder):
        if root.count(os.sep) - base_depth >= depth:
            if excluded is not None:
                excluded += [os.path.join(root, d) for d in dirs]
            dirs[:] = []
        if excluded is not None:
            excluded += [os.path.join(root, d) for d in dirs if d.startswith(("_", "~$"))]
        dirs[:] = sorted(d for d in dirs if not d.startswith((".", "~$", "_")))
        for n in sorted(names):
            if n.startswith(("~$", ".")):
                continue
            ext = os.path.splitext(n)[1].lower()
            if ext not in SUPPORTED:
                skipped[ext or "(확장자 없음)"] += 1
                continue
            files.append(os.path.join(root, n))
            if len(files) >= MAX_FOLDER_SCAN:
                rt.log("WARN", f"폴더에서 최대 {MAX_FOLDER_SCAN}개 파일까지만 살펴봅니다.")
                return files, dict(skipped)
    return files, dict(skipped)


def _folder_arg(raw: str) -> str:
    folder = os.path.abspath(os.path.expandvars(os.path.expanduser(str(raw).strip().strip('"'))))
    if not os.path.isdir(folder):
        raise rt.InputError(f"폴더를 찾을 수 없습니다: {folder}")
    return folder


def _index_files(files: List[str], max_files: int) -> Dict[str, Any]:
    """파일들을 캐시에 등록합니다. 새로 파싱하는 파일 수만 max_files로 제한합니다 (이미 최신인 파일은 세지 않음).
    반환: docs(불러온 문서), new/updated/cached(파일 이름 목록), failures, remaining."""
    out: Dict[str, Any] = {"docs": [], "new": [], "updated": [], "cached": [], "failures": [], "remaining": 0}
    parsed = 0
    index = load_index()
    total = len(files)
    for i, path in enumerate(files, 1):
        try:
            st = os.stat(path)
        except OSError as e:
            out["failures"].append({"file": path, "reason": str(e)})
            continue
        key = os.path.normcase(os.path.abspath(path))
        entry = next((e for k, e in index.items() if os.path.normcase(k) == key), None)
        fresh = bool(entry and entry.get("parser_version") == PARSER_VERSION
                     and entry.get("mtime") == st.st_mtime and entry.get("size") == st.st_size)
        if not fresh and parsed >= max_files:
            out["remaining"] += 1
            continue
        try:
            if not fresh:
                rt.log("PROGRESS", f"[{i}/{total}] 등록 중: {os.path.basename(path)}")
            doc, status = load_document_status(path, check_index=not fresh)
        except Exception as e:
            out["failures"].append({"file": path, "reason": str(e)[:300]})
            parsed += 0 if fresh else 1
            continue
        if status != "cached":
            parsed += 1
        out["docs"].append(doc)
        out[status].append(path)
    return out


# ==============================================================================
# 4. 공통 출력 도우미
# ==============================================================================

def parse_pages(spec: Optional[str], total: int, option: str = "--pages") -> List[int]:
    """'3', '1-3', '1,4,6-8' -> 1부터 센 번호 목록. option은 오류 문구에 쓸 인자 이름 (ppt.py는 --slides)."""
    if not spec:
        return list(range(1, total + 1))
    result: List[int] = []
    for part in str(spec).replace(" ", "").split(","):
        if not part:
            continue
        m = re.fullmatch(r"(\d+)(?:-(\d+))?", part)
        if not m:
            raise rt.InputError(f"{option} 형식이 잘못되었습니다: '{spec}' (예: 3, 1-3, 1,4,6-8)")
        a = int(m.group(1))
        b = int(m.group(2) or a)
        result += [n for n in range(min(a, b), max(a, b) + 1) if 1 <= n <= total]
    if not result:
        raise rt.InputError(f"{option} '{spec}'에 해당하는 항목이 없습니다 (전체 {total}개).")
    return sorted(set(result))


def page_set(spec: Optional[str]) -> Optional[set]:
    """검색 --pages: '3-7', '2,5' -> 쪽 번호 집합 (상한 없음)."""
    if not spec:
        return None
    pages: set = set()
    for part in str(spec).replace(" ", "").split(","):
        if not part:
            continue
        m = re.fullmatch(r"(\d+)(?:-(\d+))?", part)
        if not m:
            raise rt.InputError(f"--pages 형식이 잘못되었습니다: '{spec}' (예: 3, 3-7, 1,4,6-8)")
        a, b = int(m.group(1)), int(m.group(2) or m.group(1))
        if abs(b - a) > 100000:
            raise rt.InputError(f"--pages 범위가 너무 넓습니다: '{spec}'")
        pages.update(range(min(a, b), max(a, b) + 1))
    return pages


def section_text(unit: Dict[str, Any]) -> str:
    return " > ".join(unit.get("section") or [])


def where_text(unit: Dict[str, Any]) -> str:
    """근거의 위치 부분: PDF/저장된 Word 쪽은 'p.N', 쪽 나눔만 센 Word 쪽은 'p.N(추정)', 그 밖은 loc."""
    page, src = unit.get("page"), unit.get("page_source")
    if page is not None and src in ("pdf", "word_saved_layout"):
        return f"p.{page}"
    if page is not None and src == "page_break":
        return f"p.{page}(추정)"
    return str(unit.get("loc", ""))


def evidence_text(file: str, unit: Dict[str, Any]) -> str:
    """'규격서.pdf p.12, 1. SCOPE > 1.1. Scope (1.1)' 형식의 근거 문구."""
    s = f"{os.path.basename(file)} {where_text(unit)}".strip()
    sec = section_text(unit)
    if sec:
        s += f", {sec}"
    if unit.get("number"):
        s += f" ({unit['number']})"
    return s


def units_to_markdown(title: str, units: List[Dict[str, Any]], file: Optional[str] = None) -> str:
    """단위 목록을 마크다운으로. reader 단위(section 키 있음)는 '근거:' 줄을 붙입니다 (word.py/ppt.py 단위는 그대로)."""
    lines = [f"# {title}", ""]
    for u in units:
        head = u["loc"] + (f" - {u['heading']}" if u.get("heading") and u["heading"] != u["text"] else "")
        lines += [f"## {head}", ""]
        if file and "section" in u:
            lines += [f"근거: {evidence_text(file, u)}", ""]
        lines += [u["text"], ""]
    return "\n".join(lines)


def _write_next_commands(output_file: str) -> List[Dict[str, str]]:
    low = output_file.lower()
    if low.endswith(".xlsx"):
        return []
    if low.endswith(".csv"):
        return [rt.next_command("excel.py", ["write", "--target", "active", "--file", output_file],
                                "저장한 표를 열린 엑셀에 넣기")]
    return [rt.next_command("word.py", ["write", "--target", "active", "--file", output_file],
                            "저장한 내용을 열린 워드 문서 끝에 넣기"),
            rt.next_command("ppt.py", ["write", "--target", "active", "--file", output_file],
                            "저장한 내용을 열린 PPT에 새 슬라이드로 넣기")]


def _save_ext(save: Any) -> str:
    if save is True or not save:
        return ".md"
    ext = os.path.splitext(str(save).strip().strip('"'))[1].lower()
    return ext if ext in (".xlsx", ".csv", ".md") else ".md"


def _xlsx_safe(value: Any) -> Any:
    if not isinstance(value, str):
        return value
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    return ILLEGAL_CHARACTERS_RE.sub("", value)[:32000]


def write_xlsx_table(path: str, header: List[str], rows: List[List[Any]], widths: List[int],
                     sheet: str = "결과") -> str:
    """표를 xlsx로 저장합니다 (머리글 굵게, 줄바꿈, 열 너비, 첫 행 고정). 저장 후 다시 열어 행 수를 확인합니다."""
    try:
        import openpyxl
        from openpyxl.styles import Alignment, Font
    except ImportError:
        raise rt.InputError("xlsx로 저장하려면 openpyxl이 필요합니다 (python -m pip install openpyxl). .csv나 .md로 저장할 수도 있습니다.")

    def _write(p: str) -> None:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = sheet
        ws.append(header)
        for row in rows:
            ws.append([_xlsx_safe(v) for v in row])
        wrap = Alignment(wrap_text=True, vertical="top")
        for row in ws.iter_rows():
            for cell in row:
                if cell.data_type == "f":           # 문서 글이 '='로 시작해도 수식이 아닌 글로
                    cell.data_type = "s"
                cell.alignment = wrap
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for i, w in enumerate(widths):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i + 1)].width = w
        ws.freeze_panes = "A2"
        wb.save(p)

    out = rt.save_with_lock_fallback(path, _write, "엑셀 파일")
    wb = openpyxl.load_workbook(out, read_only=True)
    try:
        n = sum(1 for _ in wb.active.iter_rows(values_only=True))
    finally:
        wb.close()
    if n != len(rows) + 1:
        raise RuntimeError(f"저장한 xlsx의 행 수가 다릅니다 (기대 {len(rows) + 1}, 실제 {n}): {out}")
    return out


def _unit_view(n: int, u: Dict[str, Any], text: str) -> Dict[str, Any]:
    return {"no": n, "loc": u["loc"], "kind": u.get("kind"), "page": u.get("page"), "section": u.get("section") or [],
            "number": u.get("number"), "heading": u.get("heading", ""), "text": text}


# ==============================================================================
# 5. 명령: read
# ==============================================================================

def cmd_read(args) -> Dict[str, Any]:
    path = rt.resolve_input_path(args.file, "문서 파일")
    doc = load_document(path)
    units = doc["units"]
    total = len(units)
    name = os.path.basename(path)
    unit_name = UNIT_NAMES.get(doc["type"], "항목")
    # 쪽 정보가 있는 형식(docx/hwpx의 저장된 쪽 경계 포함)은 --pages가 실제 쪽: search의 근거 'p.12'를 그대로 읽을 수 있게
    by_page = doc["type"] in PAGE_TYPES or any(u.get("page") for u in units)
    if args.units:
        numbers = parse_pages(args.units, total) if total else []
    elif by_page and args.pages:
        pages = set(parse_pages(args.pages, _doc_pages(doc) or 0)) if total else set()
        numbers = [i for i, u in enumerate(units, 1) if u.get("page") in pages]
    else:
        numbers = parse_pages(args.pages, total) if total else []
    selected = [units[n - 1] for n in numbers]
    result: Dict[str, Any] = {"file": path, "type": doc["type"], "units_total": total, "count": len(selected),
                              "unit_name": unit_name, "pages_total": _doc_pages(doc),
                              "pages_mode": "page" if by_page else "item",
                              "sections_total": _doc_sections(doc), "meta": doc.get("meta", {})}

    if args.save:
        ext = _save_ext(args.save)
        save_arg = None if args.save is True else args.save
        out = rt.resolve_output_path(save_arg, os.path.splitext(name)[0], ext)
        header = ["위치", "쪽", "장절", "번호", "제목", "내용"]
        rows = [[u["loc"], u.get("page") if u.get("page") is not None else "", section_text(u), u.get("number") or "",
                 u.get("heading", ""), u["text"]] for u in selected]
        if ext == ".csv":
            out = rt.write_csv_rows(out, [header] + rows)
        elif ext == ".xlsx":
            out = write_xlsx_table(out, header, rows, [14, 6, 40, 10, 30, 80], "내용")
        else:
            out = rt.write_text(out, units_to_markdown(name, selected, path))
        result.update(output_file=out, next_commands=_write_next_commands(out))
        preview = f"{name}: {len(selected)}개 {unit_name}을(를) 파일로 저장했습니다 ({out})"
        return rt.ok(TOOL, "read", preview, "VERIFIED: 저장 완료. 다른 문서에 넣으려면 next_commands를 그대로 실행하세요.", **result)

    if args.format == "summary":
        preview = (f"{name}: {doc['type']} 문서, {unit_name} {total}개, 쪽 {_doc_pages(doc) or '-'}, "
                   f"장절 {_doc_sections(doc)}개, 글자 수 {sum(len(u['text']) for u in units)}")
        return rt.ok(TOOL, "read", preview, **result)

    if not selected:
        result.update(units=[], truncated=False, shown=0)
        return rt.ok(TOOL, "read", f"{name}: 글이 없습니다 (0건)",
                     "VERIFIED: 0건. 이 문서에는 읽을 글이 없습니다. 사용자에게 그대로 알리고 다시 실행하지 마세요.", **result)

    # 본문: 상한까지 담고, 넘으면 다음 구간 명령 제공
    budget = rt.DEFAULT_MAX_CHARS
    shown: List[Dict[str, Any]] = []
    used = 0
    for n, u in zip(numbers, selected):
        size = len(u["text"]) + len(u["loc"]) + len(section_text(u)) + 24
        if shown and used + size > budget:
            break
        text, _ = rt.clamp_text(u["text"], budget)
        shown.append(_unit_view(n, u, text))
        used += size
    truncated = len(shown) < len(selected)
    result["units"] = shown
    result["truncated"] = truncated
    result["shown"] = len(shown)
    if truncated:
        nxt_i = len(shown)
        nxt, last = numbers[nxt_i], numbers[-1]
        cut = selected[nxt_i]
        if not by_page:
            cont = ["--pages", f"{nxt}-{last}"]
        elif cut.get("page") != selected[nxt_i - 1].get("page") and not args.units:
            cont = ["--pages", f"{cut.get('page')}-{selected[-1].get('page')}"]
        else:                                    # 쪽 중간에서 잘림 -> 단위 번호로 이어 읽기
            cont = ["--units", f"{nxt}-{last}"]
        result["next_commands"] = [
            rt.next_command("reader.py", ["read", "--file", path] + cont, f"나머지 {unit_name} 이어서 읽기"),
            rt.next_command("reader.py", ["read", "--file", path, "--save"], "전체 내용을 파일로 저장 (다른 문서에 넣을 때)"),
        ]
    preview = f"{name}: {unit_name} {len(shown)}개 표시 (전체 {total}개)" + (" - 길이 제한으로 일부만 표시" if truncated else "")
    directive = ("VERIFIED: 읽기 완료. 일부만 표시되었으니 필요하면 next_commands로 이어서 읽으세요."
                 if truncated else "VERIFIED: 읽기 완료. 내용을 인용할 때는 page, section, number로 근거(파일 p.N, 장절 (번호))를 함께 알리세요.")
    return rt.ok(TOOL, "read", preview, directive, **result)


# ==============================================================================
# 6. 명령: index
# ==============================================================================

def cmd_index(args) -> Dict[str, Any]:
    if sum(bool(x) for x in (args.folder, args.file, args.all)) != 1:
        raise rt.InputError("--folder, --file, --all 중 하나만 지정하세요. 예: reader.py index --folder \"C:\\docs\\규격\"")
    skipped_types: Dict[str, int] = {}
    missing = 0
    target = ""
    if args.folder:
        target = _folder_arg(args.folder)
        depth = 50 if args.recursive else max(0, args.depth)
        files, skipped_types = _collect_folder(target, depth)
    elif args.file:
        files = [rt.resolve_input_path(args.file, "등록할 파일")]
        file_type(files[0])
        target = files[0]
    else:
        index = load_index()
        files = [p for p in index if os.path.exists(p)]
        missing = len(index) - len(files)
        target = "(등록된 문서 전체)"
    res = _index_files(files, max(1, args.max_files))
    failures = [{"file": os.path.basename(f["file"]), "path": f["file"], "reason": f["reason"]} for f in res["failures"]]
    result: Dict[str, Any] = {
        "target": target, "count": len(files), "indexed": len(res["new"]), "updated": len(res["updated"]),
        "unchanged": len(res["cached"]), "failed": len(failures), "failures": failures[:MAX_LIST_ITEMS],
        "remaining": res["remaining"], "truncated": res["remaining"] > 0 or len(failures) > MAX_LIST_ITEMS,
        "total_units": sum(len(d["units"]) for d in res["docs"]),
    }
    if skipped_types:
        result["skipped_types"] = skipped_types
    if missing:
        result["missing"] = missing
    nexts: List[Dict[str, str]] = []
    if res["remaining"]:
        same = ["index"] + (["--folder", target] if args.folder else ["--all"]) + ["--max-files", str(args.max_files)]
        if args.folder and args.recursive:
            same.append("--recursive")
        elif args.folder:
            same += ["--depth", str(args.depth)]
        nexts.append(rt.next_command("reader.py", same, f"남은 {res['remaining']}개 파일 이어서 등록 (읽기 전용)"))
    search_target = ["--file", target] if (args.folder or args.file) else []
    nexts.append(rt.next_command("reader.py", ["search", "--query", "<검색어>"] + search_target,
                                 "등록한 문서에서 검색 (<검색어>를 사용자 검색어로 바꿈)"))
    result["next_commands"] = nexts
    if not files:
        what = "폴더에 지원 형식 문서가 없습니다" if args.folder else "등록된 문서가 없습니다"
        return rt.ok(TOOL, "index", f"0건: {what}",
                     f"VERIFIED: 0건. {what}. 사용자에게 그대로 알리고 다시 실행하지 마세요.",
                     **{k: v for k, v in result.items() if k != "next_commands"})
    preview = (f"문서 {len(files)}개: 새로 등록 {result['indexed']}, 다시 등록 {result['updated']}, "
               f"변경 없음 {result['unchanged']}, 실패 {result['failed']}"
               + (f", 남음 {res['remaining']} (--max-files {args.max_files})" if res["remaining"] else ""))
    if res["remaining"]:
        directive = ("VERIFIED: 일부만 등록했습니다(캐시에 저장 확인). 남은 파일은 next_commands[0]을 그대로 실행해 이어서 등록하세요 "
                     "(문서를 바꾸지 않는 읽기 전용 작업). 실패한 파일은 failures를 사용자에게 알리세요.")
    else:
        directive = ("VERIFIED: 등록 완료(캐시에 저장 확인). 이제 reader.py search로 검색할 수 있습니다. "
                     "실패한 파일이 있으면 failures를 사용자에게 알리세요.")
    return rt.ok(TOOL, "index", preview, directive, **result)


# ==============================================================================
# 7. 명령: search (캐시 기준)
# ==============================================================================

def build_matcher(query: str, regex: bool, case_sensitive: bool):
    """검색 조건을 함수로 만듭니다. 'A AND B', 'A OR B' 조합 지원 (AND가 OR보다 먼저 묶임).
    반환: match_fn(text) -> 강조할 단어 목록 또는 None"""
    flags = 0 if case_sensitive else re.IGNORECASE
    if regex:
        try:
            pat = re.compile(query, flags)
        except re.error as e:
            raise rt.InputError(f"정규식 오류: {e}")

        def match_regex(text: str):
            m = next((m for m in pat.finditer(text) if m.group(0)), None)   # 빈 일치('a*')는 일치로 보지 않음
            return [m.group(0)] if m else None
        return match_regex

    groups = parse_query(query)
    if not groups:
        raise rt.InputError("검색어가 비어 있습니다.")

    pats = [[_term_pattern(unicodedata.normalize("NFC", t), flags) for t in g] for g in groups]

    def match_terms(text: str):
        t = unicodedata.normalize("NFC", text)
        for g, ps in zip(groups, pats):
            if all(p.search(t) for p in ps):
                return g
        return None
    return match_terms


def parse_query(query: str) -> List[List[str]]:
    """'A AND B OR "C D"' -> [['A', 'B'], ['C D']]. AND/OR는 대소문자 무관(AND가 OR보다 먼저 묶임),
    따옴표로 묶은 글은 구문 그대로(따옴표는 검색어가 아님), 연산자 없이 이어진 단어는 한 구문."""
    groups: List[List[str]] = [[]]
    words: List[str] = []

    def push() -> None:
        if words:
            groups[-1].append(" ".join(words))
            words.clear()
    for tok in re.findall(r'"[^"]*"|\u201c[^\u201d]*\u201d|\S+', unicodedata.normalize("NFC", query or "")):
        if tok.upper() == "AND":
            push()
        elif tok.upper() == "OR":
            push()
            groups.append([])
        elif tok[:1] in ('"', "\u201c") and len(tok) >= 2:
            push()
            if tok[1:-1].strip():
                groups[-1].append(tok[1:-1].strip())
        else:
            words.append(tok)
    push()
    return [g for g in groups if g]


_SEP = r"[\W_]*"          # 글자 사이에 끼어도 되는 것: 띄어쓰기, 줄바꿈, 기호(# - . : / ( ) 등)


def _term_pattern(term: str, flags: int) -> "re.Pattern[str]":
    """검색어를 글 안에서 찾는 정규식. 띄어쓰기·줄바꿈·기호가 달라도 찾습니다:
    'SOI 1' = 'SOI #1' = 'SOI-1' = 'SOI1', '설계기준' = '설계 기준', PDF 줄바꿈으로 갈라진 글.
    숫자 경계는 지킵니다: 'SOI 1'은 'SOI #10'이나 'SOI 21'과 맞지 않습니다.
    기호를 빼면 숫자만 남는 검색어('#4')는 기호를 그대로 찾습니다 (모든 '4'와 맞지 않게)."""
    compact = re.sub(r"\s+", "", term)
    core = re.sub(r"[\W_]+", "", compact)
    if not core or (core.isdigit() and core != compact):
        chars, sep = compact or term, r"\s*"      # 기호까지 그대로 ('1.2'는 '12'와 다름)
    else:
        chars, sep = core, _SEP
    body = sep.join(re.escape(c) for c in chars)
    if chars[0].isdigit():
        body = r"(?<![\d.])" + body              # '1.2'가 '11.2', '3.1.2'의 끝과 맞지 않게
    if chars[-1].isdigit():
        body += r"(?!\d)"
    return re.compile(body, flags)


def make_snippet(text: str, terms: List[str], context: int, case_sensitive: bool) -> str:
    flags = 0 if case_sensitive else re.IGNORECASE
    first = None
    for term in terms:
        m = _term_pattern(term, flags).search(text)
        if m and (first is None or m.start() < first.start()):
            first = m
    if first is None:
        return text[: context * 2]
    start = max(0, first.start() - context)
    end = min(len(text), first.end() + context)
    snippet = ("..." if start > 0 else "") + text[start:end] + ("..." if end < len(text) else "")
    pattern = re.compile("|".join(_term_pattern(t, 0).pattern for t in sorted(terms, key=len, reverse=True)), flags)
    return pattern.sub(lambda m: f"**{m.group(0)}**", snippet).replace("\n", " ")


_SENT_SPLIT = re.compile(r"(?<=[.!?\u3002])\s+|\n+")


def match_sentence(text: str, terms: List[str], case_sensitive: bool, limit: int = 1000) -> str:
    """일치한 단어가 든 문장 (문장으로 못 나누면 단락 앞부분)."""
    flags = 0 if case_sensitive else re.IGNORECASE
    for sent in _SENT_SPLIT.split(text):
        if any(_term_pattern(t, flags).search(sent) for t in terms):
            return sent.strip()[:limit]
    return text.strip()[:limit]


def _not_indexed_error() -> rt.InputError:
    return rt.InputError(
        "캐시에 등록된 문서가 없어 검색할 대상이 없습니다. 문서를 먼저 등록하세요: "
        "reader.py index --folder \"<폴더>\" 또는 reader.py read --file \"<파일>\".",
        directive=("검색할 등록 문서가 없습니다. 다시 실행하지 말고 사용자에게 검색할 폴더나 파일을 물어본 뒤 "
                   "search에 --file \"<폴더 또는 파일>\"을 붙여 실행하세요 (처음 한 번 자동 등록)."),
        error_type="not_indexed",
        fields={"count": 0})


def _search_targets(args) -> Tuple[List[Dict[str, Any]], Dict[str, Any]]:
    """검색할 문서를 정합니다. --file이 있으면 그 파일/폴더의 미등록/변경 문서를 먼저 등록, 없으면 캐시 전체만.
    반환: (문서 목록, 정보: indexed_now, from_cache, skipped, remaining, stale, missing, skipped_types)"""
    info: Dict[str, Any] = {"indexed_now": [], "from_cache": 0, "skipped": [], "remaining": 0, "stale": [],
                            "missing": 0, "skipped_types": {}, "folder": None, "excluded_dirs": [], "refreshed": 0}
    docs: List[Dict[str, Any]] = []
    if args.file:
        raw = os.path.expandvars(os.path.expanduser(str(args.file).strip().strip('"')))
        if os.path.isdir(raw):
            folder = os.path.abspath(raw)
            info["folder"] = folder
            files, info["skipped_types"] = _collect_folder(folder, excluded=info["excluded_dirs"])
        else:
            path = rt.resolve_input_path(raw, "검색할 파일")
            file_type(path)
            files = [path]
        res = _index_files(files, max(1, args.max_files))
        if len(files) == 1 and not info["folder"] and res["failures"]:
            raise rt.InputError(res["failures"][0]["reason"])
        docs = res["docs"]
        info["indexed_now"] = res["new"] + res["updated"]
        info["from_cache"] = len(res["cached"])
        info["remaining"] = res["remaining"]
        info["skipped"] = [{"file": f["file"], "reason": f["reason"]} for f in res["failures"]]
        return docs, info

    index = load_index()
    if not index:
        raise _not_indexed_error()
    for path in sorted(index, key=lambda p: index[p].get("parsed_at", ""), reverse=True):
        doc = load_cached_only(path)
        if doc is None:
            info["skipped"].append({"file": path, "reason": "캐시 파일이 없거나 손상됨 -> 다시 등록 필요"})
            info["stale"].append(path)
            continue
        if not os.path.exists(path):
            doc["file_missing"] = True
            info["missing"] += 1
        if is_stale_doc(doc) and not doc.get("file_missing") and info["refreshed"] < STALE_REFRESH_MAX:
            try:
                doc, _ = load_document_status(path)     # 원본이 바뀐 캐시는 그 자리에서 다시 읽음 (JIT 복구)
                info["refreshed"] += 1
            except Exception as e:
                rt.debug(f"다시 읽기 실패, 이전 캐시로 검색: {e}")
        if is_stale_doc(doc):
            doc["_stale"] = True
            info["stale"].append(path)
        docs.append(doc)
    info["from_cache"] = len(docs)
    if info["refreshed"]:
        rt.note_recovered(f"원본이 바뀐 캐시 {info['refreshed']}개 -> 다시 읽고 검색")
    return docs, info


def cmd_search(args) -> Dict[str, Any]:
    matcher = build_matcher(args.query, args.regex, args.case_sensitive)
    pages = page_set(args.pages)
    sec_filter = re.sub(r"\s+", " ", args.section).strip().lower() if args.section else ""
    docs, info = _search_targets(args)
    hits: List[Dict[str, Any]] = []
    page_unknown = 0
    for doc in docs:
        for u in doc["units"]:
            if sec_filter and sec_filter not in re.sub(r"\s+", " ", section_text(u)).lower():
                continue
            if pages is not None:
                if u.get("page") is None:
                    page_unknown += 1
                    continue
                if u["page"] not in pages:
                    continue
            terms = matcher(u["text"])
            if terms is None:
                continue
            hits.append({"doc": doc, "unit": u, "terms": terms,
                         "snippet": make_snippet(u["text"], terms, args.context, args.case_sensitive)})

    total = len(hits)
    shown = hits[: max(0, args.top)]
    result: Dict[str, Any] = {"query": args.query, "files_searched": len(docs), "total_matches": total,
                              "count": total, "matched": len(shown), "shown": len(shown),
                              "truncated": total > len(shown)}
    filters = {k: v for k, v in (("file", args.file), ("pages", args.pages), ("section", args.section)) if v}
    if filters:
        result["filters"] = filters
    if pages is not None and page_unknown:
        result["page_unknown_skipped"] = page_unknown
    if args.file:
        result["indexed_now_count"] = len(info["indexed_now"])
        result["indexed_now"] = [os.path.basename(p) for p in info["indexed_now"][:MAX_LIST_ITEMS]]
        result["from_cache"] = info["from_cache"]
        if info["remaining"]:
            result["not_indexed_remaining"] = info["remaining"]
    if info["skipped_types"]:
        result["skipped_types"] = info["skipped_types"]
    if info["skipped"]:
        result["skipped"] = [{"file": os.path.basename(s["file"]), "path": s["file"], "reason": s["reason"]}
                             for s in info["skipped"][:MAX_LIST_ITEMS]]
    if info["stale"]:
        result["stale_cache_count"] = len(info["stale"])
    if info["missing"]:
        result["missing_file_count"] = info["missing"]

    nexts: List[Dict[str, str]] = []
    if args.save and hits:
        stem = "search_" + re.sub(r"[^\w가-힣]+", "_", args.query)[:30].strip("_")
        ext = _save_ext(args.save)
        save_arg = None if args.save is True else args.save
        out = rt.resolve_output_path(save_arg, stem, ext)
        saved = hits[:MAX_SAVE_HITS]
        header = ["파일", "쪽", "장절", "번호", "위치", "내용(문장/단락)", "문맥"]
        rows = []
        for h in saved:
            u = h["unit"]
            rows.append([os.path.basename(h["doc"]["file"]), u.get("page") if u.get("page") is not None else "",
                         section_text(u), u.get("number") or "", u["loc"],
                         match_sentence(u["text"], h["terms"], args.case_sensitive),
                         rt.clamp_text(u["text"], 3000)[0]])
        if ext == ".csv":
            out = rt.write_csv_rows(out, [header] + rows)
        elif ext == ".xlsx":
            out = write_xlsx_table(out, header, rows, [22, 6, 45, 10, 14, 60, 80], "검색결과")
        else:
            lines = [f"# 검색 결과: {args.query}", ""]
            for h in saved:
                u = h["unit"]
                sec = section_text(u)
                lines += [f"## {os.path.basename(h['doc']['file'])} > {u['loc']}" + (f" ({sec})" if sec else ""), "",
                          f"근거: {evidence_text(h['doc']['file'], u)}", "",
                          u["text"] if args.full else h["snippet"].replace("**", ""), ""]
            out = rt.write_text(out, "\n".join(lines))
        result.update(output_file=out, saved=len(saved))
        nexts += _write_next_commands(out)

    view = []
    for h in shown:
        u, doc = h["unit"], h["doc"]
        item = {"file": os.path.basename(doc["file"]), "path": doc["file"], "page": u.get("page"),
                "page_source": u.get("page_source"), "section": u.get("section") or [],
                "section_text": section_text(u), "number": u.get("number"), "loc": u["loc"],
                "heading": u.get("heading", ""), "evidence": evidence_text(doc["file"], u), "snippet": h["snippet"]}
        if args.full:
            item["text"], _ = rt.clamp_text(u["text"], 2000)
        if doc.get("file_missing"):
            item["file_missing"] = True
        if doc.get("_stale"):
            item["stale"] = True             # 원본이 바뀌기 전 내용 (근거로 쓰기 전에 다시 등록 필요)
        view.append(item)
    result["matches"] = view

    if info["stale"] and not args.file:
        nexts.append(rt.next_command("reader.py", ["index", "--all"],
                                     f"바뀌었거나 이전 형식인 캐시 {len(info['stale'])}개 다시 등록 (읽기 전용)"))
    if info["remaining"] and info["folder"]:
        nexts.append(rt.next_command("reader.py", ["index", "--folder", info["folder"], "--max-files", str(args.max_files)],
                                     f"아직 등록하지 못한 {info['remaining']}개 파일 등록 (읽기 전용)"))
    if result["truncated"] and not args.save:
        cmd = ["search", "--query", args.query]
        for flag, val in (("--file", args.file), ("--pages", args.pages), ("--section", args.section)):
            if val:
                cmd += [flag, val]
        if args.regex:
            cmd.append("--regex")
        if args.case_sensitive:
            cmd.append("--case-sensitive")
        nexts.append(rt.next_command("reader.py", cmd + ["--save", "근거.xlsx"], f"전체 {total}건을 엑셀 파일로 저장"))
    if nexts:
        result["next_commands"] = nexts

    scope = f"문서 {len(docs)}개"
    if args.file and info["indexed_now"]:
        scope += f"(새로 등록 {len(info['indexed_now'])}개)"
    notes = ""
    if info["stale"] and not args.file:
        notes += (f" 캐시 {len(info['stale'])}개는 원본이 바뀌었거나 이전 형식입니다. 최신 내용이 필요하면 "
                  "next_commands의 index --all(읽기 전용)을 실행한 뒤 같은 검색을 1번 다시 하세요.")
    if info["indexed_now"]:
        notes += " 새로 등록한 파일(indexed_now)을 사용자에게 알리세요."
    gaps = _search_gaps(docs, info)
    if gaps:
        result["not_searched"] = gaps
    if total == 0:
        preview = f"'{args.query}': {scope}에서 일치 항목 없음 (0건)"
        if gaps:
            # 일부를 검색하지 못했으면 "근거 없음"으로 단정하지 않음 (규격 대조에서 잘못된 판정이 됨)
            return rt.ok_unverified(TOOL, "search", preview + ", 검색 범위 불완전",
                                    ["0건이지만 검색하지 못한 부분이 있음: " + "; ".join(gaps)], **result)
        return rt.ok(TOOL, "search", preview,
                     "VERIFIED: 0건. 조건에 맞는 항목이 없습니다. 사용자에게 그대로 알리고 조건을 바꿔 다시 실행하지 마세요." + notes,
                     **result)
    preview = f"'{args.query}': {scope}에서 {total}건 일치" + (f" (상위 {len(shown)}건 표시)" if total > len(shown) else "")
    evid = " 각 결과의 근거는 evidence(파일 p.N, 장절 (번호))를 그대로 알리세요. page가 null이면 장절과 번호로 알립니다."
    if "output_file" in result:
        preview += f", 저장: {result['output_file']}"
        directive = "VERIFIED: 검색과 저장 완료(저장 파일 확인)." + evid
        if result.get("next_commands") and _write_next_commands(result["output_file"]):
            directive += " 다른 문서에 넣으려면 next_commands를 그대로 실행하세요."
    else:
        directive = "VERIFIED: 검색 완료." + evid + " 전체 결과가 필요하면 --save \"근거.xlsx\"로 저장하세요."
    if gaps:
        payload = rt.ok_unverified(TOOL, "search", preview + ", 검색 범위 불완전",
                                   ["검색하지 못한 부분이 있어 결과가 전부가 아닐 수 있음: " + "; ".join(gaps)], **result)
        payload["agent_directive"] += evid
        return payload
    return rt.ok(TOOL, "search", preview, directive + notes, **result)


def _search_gaps(docs: List[Dict[str, Any]], info: Dict[str, Any]) -> List[str]:
    """검색 범위에서 빠진 것: 결과가 없거나 적어도 '없다'고 단정하면 안 되는 이유들."""
    gaps: List[str] = []
    if info["remaining"]:
        gaps.append(f"아직 등록하지 못한 파일 {info['remaining']}개 (next_commands의 index로 등록)")
    if info["skipped"]:
        gaps.append(f"읽지 못한 파일 {len(info['skipped'])}개 (skipped)")
    if info["stale"]:
        gaps.append(f"원본이 바뀐 뒤 다시 읽지 못한 캐시 {len(info['stale'])}개 (옛 내용으로 검색, 결과의 stale 표시)")
    ocr = [os.path.basename(d["file"]) for d in docs if (d.get("meta") or {}).get("ocr_needed")]
    if ocr:
        gaps.append(f"글자가 없는 PDF {len(ocr)}개 (스캔 이미지, OCR 필요): {', '.join(ocr[:5])}")
    cut = [os.path.basename(d["file"]) for d in docs if (d.get("meta") or {}).get("rows_truncated")]
    if cut:
        gaps.append(f"행이 많아 {MAX_SHEET_ROWS}행까지만 읽은 엑셀 {len(cut)}개: {', '.join(cut[:5])}")
    if info["excluded_dirs"]:
        gaps.append(f"들어가지 않은 하위 폴더 {len(info['excluded_dirs'])}개 (깊이 제한 또는 '_'로 시작): "
                    + ", ".join(os.path.basename(d) for d in info["excluded_dirs"][:5]))
    return gaps


# ==============================================================================
# 8. 명령: list (캐시된 문서 조회)
# ==============================================================================

def cmd_list(args) -> Dict[str, Any]:
    index = load_index()
    name_filter = str(args.name or "").strip().lower()
    folder = _folder_arg(args.folder) if args.folder else None
    items = []
    for path, info in sorted(index.items(), key=lambda kv: kv[1].get("parsed_at", ""), reverse=True):
        if name_filter and name_filter not in os.path.basename(path).lower() and name_filter not in path.lower():
            continue
        if folder and not _under(path, folder):
            continue
        exists = os.path.exists(path)
        stale = info.get("parser_version") != PARSER_VERSION
        if exists and not stale:
            try:
                st = os.stat(path)
                stale = info.get("mtime") != st.st_mtime or info.get("size") != st.st_size
            except OSError:
                pass
        items.append({"file": os.path.basename(path), "path": path, "type": info.get("type"),
                      "units": info.get("units"), "pages": info.get("pages"), "sections": info.get("sections"),
                      "chars": info.get("chars"), "cached_at": info.get("parsed_at"), "read_at": info.get("parsed_at"),
                      "parser_version": info.get("parser_version", 1),
                      "cache": "stale" if stale else "current", "stale": stale, "missing": not exists,
                      "exists": exists})
    limit = 100
    shown = items[:limit]
    stale_n = sum(1 for i in items if i["stale"])
    missing_n = sum(1 for i in items if i["missing"])
    result: Dict[str, Any] = {"documents": shown, "total": len(items), "count": len(items), "shown": len(shown),
                              "truncated": len(items) > limit, "stale_count": stale_n, "missing_count": missing_n}
    if result["truncated"]:
        out = os.path.join(rt.get_data_dir(), "results", f"cached_documents_{rt.timestamp()}.csv")
        os.makedirs(os.path.dirname(out), exist_ok=True)
        cols = ["file", "path", "type", "units", "pages", "sections", "cached_at", "stale", "missing"]
        result["full_result_file"] = rt.write_csv_rows(out, [cols] + [[i[c] for c in cols] for i in items])
    if stale_n:
        result["next_commands"] = [rt.next_command("reader.py", ["index", "--all"],
                                                   f"바뀌었거나 이전 형식인 캐시 {stale_n}개 다시 등록 (읽기 전용)")]
    cond = " ".join(x for x in (f"이름 '{args.name}'" if args.name else "", f"폴더 {folder}" if folder else "") if x)
    if not items:
        what = f"{cond}에 맞는 캐시 문서가 없습니다" if cond else "캐시에 등록된 문서가 없습니다"
        return rt.ok(TOOL, "list", f"0건: {what}",
                     f"VERIFIED: 0건. {what}. 사용자에게 그대로 알리고 조건을 바꿔 다시 실행하지 마세요.", **result)
    preview = (f"캐시 문서 {len(items)}개" + (f" ({cond})" if cond else "") + (f", 최근 {limit}개 표시" if result["truncated"] else "")
               + (f", 원본 변경/이전 형식 {stale_n}개" if stale_n else "") + (f", 원본 없음 {missing_n}개" if missing_n else ""))
    directive = "VERIFIED: 캐시 목록 조회 완료. documents의 file, pages, cached_at, stale, missing을 사용자에게 알리세요."
    if result["truncated"]:
        directive += " 전체 목록은 full_result_file에 있습니다."
    return rt.ok(TOOL, "list", preview, directive, **result)


# ==============================================================================
# 9. 진입점
# ==============================================================================

def build_cli():
    parser, sub = rt.build_parser(TOOL, "문서 파일 읽기, 등록(캐시), 근거 검색 (Office 불필요)")

    p = sub.add_parser("read", help="파일 하나 읽기 (캐시에 등록)")
    p.add_argument("--file", "-f", required=True, help="문서 경로 (파일명만 쓰면 data 폴더에서 찾음)")
    p.add_argument("--pages", "-p", help="PDF 쪽/PPT 슬라이드 번호 (그 외 형식은 항목 번호). 예: 3, 1-3, 1,4")
    p.add_argument("--units", help="항목(단위) 번호로 고르기. 예: 57-120 (이어 읽기 명령이 씀)")
    p.add_argument("--format", choices=["text", "summary"], default="text", help="text: 본문, summary: 요약 정보만")
    p.add_argument("--save", nargs="?", const=True, help="결과를 파일로 저장 (.md 기본, .csv/.xlsx 가능). 경로 생략 시 작업 폴더")

    x = sub.add_parser("index", help="폴더/파일의 문서를 내용 출력 없이 캐시에 등록")
    x.add_argument("--folder", "-d", help="등록할 폴더")
    x.add_argument("--file", "-f", help="등록할 파일 하나")
    x.add_argument("--all", action="store_true", help="등록된 문서 전체를 다시 확인 (바뀐 파일, 이전 형식 캐시만 다시 읽음)")
    x.add_argument("--depth", type=int, default=MAX_FOLDER_DEPTH, help=f"하위 폴더 깊이 (기본 {MAX_FOLDER_DEPTH}, 0=그 폴더만)")
    x.add_argument("--recursive", action="store_true", help="하위 폴더 전체")
    x.add_argument("--max-files", type=int, default=MAX_FOLDER_FILES,
                   help=f"이번에 새로 읽을 최대 파일 수 (기본 {MAX_FOLDER_FILES}, 이미 최신인 파일은 세지 않음)")

    s = sub.add_parser("search", help="캐시에 등록된 문서에서 검색 (--file이면 그 파일/폴더를 먼저 자동 등록)")
    s.add_argument("--query", "-q", required=True, help="검색어. 'A AND B', 'A OR B' 사용 가능")
    s.add_argument("--file", "-f", help="검색할 파일 또는 폴더 (미등록/바뀐 문서는 자동 등록). 생략 시 캐시 전체")
    s.add_argument("--pages", "-p", help="쪽 번호로 거르기. 예: 3-7 (쪽 정보가 없는 단위는 제외)")
    s.add_argument("--section", help="장절 경로 일부로 거르기 (대소문자 무시)")
    s.add_argument("--regex", "-r", action="store_true", help="정규식으로 검색")
    s.add_argument("--case-sensitive", "-c", action="store_true", help="대소문자 구분")
    s.add_argument("--top", "-n", type=int, default=10, help="표시할 최대 건수 (기본 10, --save 파일에는 전체)")
    s.add_argument("--context", "-C", type=int, default=80, help="스니펫 앞뒤 글자 수 (기본 80)")
    s.add_argument("--full", action="store_true", help="스니펫 대신 일치한 단락/페이지 전체")
    s.add_argument("--max-files", type=int, default=MAX_FOLDER_FILES,
                   help=f"--file 폴더에서 이번에 새로 등록할 최대 파일 수 (기본 {MAX_FOLDER_FILES})")
    s.add_argument("--save", nargs="?", const=True, help="결과를 파일로 저장 (.md 기본, .xlsx/.csv 가능)")

    lst = sub.add_parser("list", help="캐시에 등록된 문서 목록")
    lst.add_argument("--name", help="파일 이름/경로 일부로 거르기 (대소문자 무시)")
    lst.add_argument("--folder", "-d", help="이 폴더 아래 문서만")
    return parser


def main(argv: Optional[List[str]] = None) -> None:
    args = build_cli().parse_args(argv)
    handlers = {"read": cmd_read, "index": cmd_index, "search": cmd_search, "list": cmd_list}
    rt.run_cli(TOOL, args.command, lambda: handlers[args.command](args))


if __name__ == "__main__":
    main()
