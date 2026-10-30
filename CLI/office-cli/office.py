"""office.py - Office 공통 작업

명령:
    windows   열린 Excel/Word/PowerPoint 문서 목록
    focus     특정 앱/문서 창을 화면 맨 앞으로
    transfer  앱 간 전달
              - 엑셀 차트 -> 워드/PPT (그림)          --from excel --chart <이름> --to word,ppt
              - 엑셀 범위 -> 워드/PPT (표)             --from excel --range A1:D10 --to word
              - 워드/PPT 내용 -> 엑셀 (새 시트)        --from word --what text --to excel
              - 워드/PPT 표 -> 엑셀                    --from word --what table --table 1 --to excel
    undo      직전 작업 되돌리기 (--list로 목록, --pick으로 선택)

차트 그림 전달 방식 (회사 DRM 대응, 자동으로 순서대로 시도):
    export  엑셀이 PNG로 내보냄 (파일이 암호화되면 건너뜀)
    redraw  차트 데이터를 읽어 Python(matplotlib)으로 다시 그림 (클립보드 사용 안 함)
    relay   엑셀 -> 워드 -> PPT 클립보드 경유
    처음 한 번 `office.py transfer --probe`를 실행하면 이 PC에서 되는 방식을 시험해 저장합니다.
"""
import os
import sys
import time
from datetime import datetime
from typing import Any, Dict, List, Optional, Sequence, Tuple

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import _office as off  # noqa: E402
import _runtime as rt  # noqa: E402

TOOL = "office"
PNG_SIGNATURE = b"\x89PNG\r\n\x1a\n"
METHODS = ("export", "redraw", "relay")
APP_ALIASES = {"excel": "excel", "엑셀": "excel", "word": "word", "워드": "word", "ppt": "ppt", "powerpoint": "ppt",
               "파워포인트": "ppt"}


def _app(name: str) -> str:
    key = str(name).strip().lower()
    if key not in APP_ALIASES:
        raise rt.not_found("앱", name, ["excel", "word", "ppt"], vocab=True)
    return APP_ALIASES[key]


# ==============================================================================
# 1. windows / focus
# ==============================================================================

def cmd_windows(args) -> Dict[str, Any]:
    docs = off.open_documents_summary()
    by_app: Dict[str, int] = {}
    for d in docs:
        by_app[d["app"]] = by_app.get(d["app"], 0) + 1
    preview = ", ".join(f"{k} {v}개" for k, v in by_app.items()) or "열린 Office 문서가 없습니다"
    caps = off.running_office_caps()
    versions = sorted({c["version"] for c in caps.values() if c["version"]})
    if versions:
        preview += f" / Office {', '.join(versions)}"
        no_pdf = [k for k, c in caps.items() if c["pdf_export"] is False]
        if no_pdf:
            preview += f" (PDF 내보내기 불가: {', '.join(no_pdf)})"
    return rt.ok(TOOL, "windows", preview, documents=docs, office=caps)


def bring_to_front(hwnd: int) -> bool:
    """최소화된 창만 복원하고(최대화 상태 유지) 맨 앞으로 가져옵니다."""
    try:
        import win32api
        import win32con
        import win32gui
        if win32gui.IsIconic(hwnd):
            win32gui.ShowWindow(hwnd, win32con.SW_RESTORE)
        try:
            win32gui.SetForegroundWindow(hwnd)
        except Exception:
            win32api.keybd_event(0x12, 0, 0, 0)                 # Alt를 눌렀다 떼면 포그라운드 제한이 풀림
            try:
                win32gui.SetForegroundWindow(hwnd)
            finally:
                win32api.keybd_event(0x12, 0, 2, 0)             # 실패해도 Alt를 눌린 채로 두지 않음
        return True
    except Exception as e:
        rt.debug(f"창 앞으로 가져오기 실패: {e}")
        return False


def cmd_focus(args) -> Dict[str, Any]:
    kind = _app(args.app)
    docs = off.list_open(kind)
    if args.name:
        hit = off.match_open(kind, args.name, docs)
        if hit is None:
            raise rt.not_found(f"열린 {off.APP_LABEL[kind]} 문서", args.name, [d["name"] for d in docs])
    else:
        hit = off.get_active(kind)
    doc, app = hit["doc"], hit["app"]
    hwnd = None
    if kind == "excel":
        doc.Activate()
        hwnd = int(app.Hwnd)
    elif kind == "word":
        doc.Activate()
        try:
            hwnd = int(doc.ActiveWindow.Hwnd)
        except Exception:
            hwnd = None
    else:
        doc.Windows(1).Activate()
        hwnd = int(app.HWND)
    front = bring_to_front(hwnd) if hwnd else False
    msg = f"{off.APP_LABEL[kind]} '{hit['name']}' 창을 " + ("앞으로 가져왔습니다" if front else "활성화했습니다 (맨 앞 전환은 Windows가 막았을 수 있음)")
    return rt.ok(TOOL, "focus", msg, app=off.APP_LABEL[kind], name=hit["name"], brought_to_front=front)


# ==============================================================================
# 2. 차트 그림 만들기 (export / redraw / relay)
# ==============================================================================

def is_plain_png(path: str) -> bool:
    try:
        with open(path, "rb") as f:
            return f.read(8) == PNG_SIGNATURE
    except OSError:
        return False


def image_by_export(chart_obj: Any, stem: str) -> str:
    path = rt.resolve_output_path(None, stem + "_export", ".png")
    chart_obj.Chart.Export(os.path.normpath(path), "PNG")
    if not is_plain_png(path):
        raise RuntimeError("엑셀이 저장한 PNG가 보안 프로그램으로 암호화되어 다른 앱이 읽을 수 없습니다")
    return path


_LINE = {4, 63, 64, 65, 66, 67}
_COLUMN = {51, 52, 53}
_BAR = {57, 58, 59}
_PIE = {5, 69, 70, -4102, -4120}
_AREA = {1, 76, 77}
_SCATTER = {-4169, 72, 73, 74, 75}


def chart_data(chart_obj: Any) -> Dict[str, Any]:
    """COM으로 차트의 숫자 데이터만 읽습니다 (파일/클립보드 사용 안 함)."""
    chart = chart_obj.Chart
    title = ""
    try:
        if chart.HasTitle:
            title = str(chart.ChartTitle.Text)
    except Exception:
        pass
    series = []
    coll = chart.SeriesCollection()
    for i in range(1, coll.Count + 1):
        s = coll(i)
        raw = s.Values                                   # 계열 값은 1차원 튜플로 옴
        values = list(raw) if isinstance(raw, (tuple, list)) else [raw]
        try:
            xs = s.XValues
            xs = list(xs) if isinstance(xs, (tuple, list)) else [xs]
        except Exception:
            xs = list(range(1, len(values) + 1))
        series.append({"name": str(s.Name), "values": list(values), "x": xs})
    return {"type": int(chart.ChartType), "title": title, "series": series,
            "width": float(chart_obj.Width), "height": float(chart_obj.Height)}


def render_chart(data: Dict[str, Any], path: str) -> str:
    """matplotlib으로 차트를 그려 PNG로 저장합니다 (Python이 직접 쓴 파일)."""
    try:
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
    except ImportError:
        raise RuntimeError("matplotlib이 없어 차트를 다시 그릴 수 없습니다")
    import logging
    logging.getLogger("matplotlib.font_manager").setLevel(logging.ERROR)
    plt.rcParams["font.family"] = ["Malgun Gothic", "AppleGothic", "NanumGothic", "DejaVu Sans"]
    plt.rcParams["axes.unicode_minus"] = False
    fig, ax = plt.subplots(figsize=(max(data["width"], 200) / 72, max(data["height"], 120) / 72), dpi=150)
    kind = data["type"]
    series = data["series"]
    colors = ["#1B365D", "#D9381E", "#4A777A", "#E6A700", "#6A4C93"]

    def nums(vals):
        return [float(v) if isinstance(v, (int, float)) else float("nan") for v in vals]

    if kind in _PIE and series:
        s = series[0]
        pairs = [(str(x), abs(v)) for x, v in zip(s["x"], nums(s["values"])) if v == v and v != 0]   # 엑셀처럼 음수는 절댓값
        if not pairs:
            raise RuntimeError("원형 차트에 그릴 값이 없습니다")
        ax.pie([v for _, v in pairs], labels=[x for x, _ in pairs], autopct="%1.0f%%", colors=colors)
        ax.axis("equal")
    else:
        n = len(series)
        for i, s in enumerate(series):
            y = nums(s["values"])
            xs = s["x"]
            numeric_x = all(isinstance(v, (int, float)) for v in xs)
            color = colors[i % len(colors)]
            if kind in _COLUMN or kind in _BAR:
                pos = [k + (i - (n - 1) / 2) * (0.8 / n) for k in range(len(y))]
                (ax.barh if kind in _BAR else ax.bar)(pos, y, 0.8 / n, label=s["name"], color=color)
                ticks = range(len(y))
                (ax.set_yticks if kind in _BAR else ax.set_xticks)(list(ticks))
                (ax.set_yticklabels if kind in _BAR else ax.set_xticklabels)([str(v) for v in xs])
            elif kind in _AREA:
                x = xs if numeric_x else range(len(y))
                ax.fill_between(list(x), y, alpha=0.5, label=s["name"], color=color)
            elif kind == -4169:
                ax.scatter(xs if numeric_x else range(len(y)), y, label=s["name"], color=color, s=12)
            else:
                ax.plot(xs if numeric_x else range(len(y)), y, label=s["name"], color=color, linewidth=1.8)
                if not numeric_x:
                    ax.set_xticks(list(range(len(y))))
                    ax.set_xticklabels([str(v) for v in xs], rotation=0)
        if n > 1:
            ax.legend(loc="best", fontsize=8)
        ax.grid(alpha=0.3)
    if data["title"]:
        ax.set_title(data["title"])
    import warnings
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")              # 한글 글꼴이 없는 환경의 경고는 출력하지 않음
        fig.tight_layout()
        fig.savefig(path, format="png")
    plt.close(fig)
    return path


def image_by_redraw(chart_obj: Any, stem: str) -> str:
    path = rt.resolve_output_path(None, stem + "_redraw", ".png")
    return render_chart(chart_data(chart_obj), path)


# ==============================================================================
# 3. 삽입 도우미
# ==============================================================================

def insert_picture_word(target: str, image: str, anchor: Optional[str], occurrence: Optional[int] = None) -> Dict[str, Any]:
    import word
    with off.open_target("word", target, readonly=False, create=True, tool=TOOL, command="transfer") as t:
        start, _ = word._insert_position(t.doc, anchor, "after" if anchor else "end", occurrence)
        w = word.Writer(t.doc, start, None)
        w.image(image, None)
        t.set_focus()
        return t.resolved()


def _ppt_slide(prs: Any, slide: Optional[str], title: str) -> Tuple[Any, int]:
    import ppt
    no, new = ppt._resolve_slide(prs, slide, default_new=slide is None)
    if new:
        no = int(prs.Slides.Count) + 1
        s = prs.Slides.Add(no, ppt.PP_LAYOUT["title_only"] if title else ppt.PP_LAYOUT["blank"])
        if title:
            ppt.SlideWriter(prs, s, off.get_theme("navy")).title(title)
        return s, no
    return prs.Slides(no), no


def insert_picture_ppt(target: str, image: str, slide: Optional[str], title: str) -> Dict[str, Any]:
    import ppt
    with off.open_target("ppt", target, readonly=False, create=True, tool=TOOL, command="transfer") as t:
        s, no = _ppt_slide(t.doc, slide, title)
        ppt.SlideWriter(t.doc, s, off.get_theme("navy")).image(image, None)
        t.set_focus(slide=no)
        return t.resolved()


def relay_chart(chart_obj: Any, to_word: Optional[str], to_ppt: Optional[str], anchor: Optional[str],
                slide: Optional[str], title: str, occurrence: Optional[int] = None) -> Dict[str, Dict[str, Any]]:
    """클립보드 경유: 엑셀 차트 -> 워드(대상 또는 임시 문서) -> PPT."""
    import word
    done: Dict[str, Dict[str, Any]] = {}
    chart_obj.CopyPicture(1, -4147)                      # xlScreen, xlPicture(메타파일)
    time.sleep(0.3)
    temp_app = temp_doc = None
    try:
        if to_word:
            with off.open_target("word", to_word, readonly=False, create=True, tool=TOOL, command="transfer") as t:
                start, _ = word._insert_position(t.doc, anchor, "after" if anchor else "end", occurrence)
                rng = t.doc.Range(start, start)
                rng.InsertAfter("\r")
                rng = t.doc.Range(start, start)
                rng.PasteSpecial(DataType=3)              # wdPasteMetafilePicture
                pasted = t.doc.Range(start, start + 1).InlineShapes
                if pasted.Count == 0:
                    raise RuntimeError("워드에 붙여넣기 실패 (클립보드)")
                if to_ppt:
                    pasted.Item(1).Range.CopyAsPicture()
                t.set_focus()
                done["word"] = t.resolved()
        elif to_ppt:
            temp_app = off.win32com.client.DispatchEx("Word.Application")
            temp_app.Visible = False
            temp_doc = temp_app.Documents.Add(Visible=False)
            temp_doc.Range(0, 0).PasteSpecial(DataType=3)
            if temp_doc.InlineShapes.Count == 0:
                raise RuntimeError("임시 워드 문서에 붙여넣기 실패 (클립보드)")
            temp_doc.InlineShapes.Item(1).Range.CopyAsPicture()
        if to_ppt:
            import ppt
            time.sleep(0.3)
            with off.open_target("ppt", to_ppt, readonly=False, create=True, tool=TOOL, command="transfer") as t:
                s, no = _ppt_slide(t.doc, slide, title)
                w = ppt.SlideWriter(t.doc, s, off.get_theme("navy"))
                top = w.next_top()
                shape_range = s.Shapes.Paste()
                if shape_range.Count == 0:
                    raise RuntimeError("PPT에 붙여넣기 실패 (클립보드)")
                pic = shape_range.Item(1)
                max_w, max_h = w.W - 2 * ppt.MARGIN, w.H - ppt.MARGIN - top
                ratio = min(max_w / float(pic.Width), max_h / float(pic.Height), 1.0)
                pic.LockAspectRatio = -1
                pic.Width = float(pic.Width) * ratio
                pic.Left = (w.W - float(pic.Width)) / 2
                pic.Top = top
                t.set_focus(slide=no)
                done["ppt"] = t.resolved()
    finally:
        if temp_doc is not None:
            try:
                temp_doc.Close(SaveChanges=0)
            except Exception:
                pass
        if temp_app is not None:
            try:
                temp_app.Quit(SaveChanges=0)
            except Exception:
                pass
    return done


def method_order(kind: str, forced: Optional[str]) -> List[str]:
    if forced and forced != "auto":
        return [forced]
    saved = rt.load_config().get("transfer_method", {}).get(kind)
    if isinstance(saved, list) and saved:
        return [m for m in saved if m in METHODS] + [m for m in METHODS if m not in saved]
    return list(METHODS)


# ==============================================================================
# 4. transfer
# ==============================================================================

def _excel_chart_transfer(args, targets: List[str]) -> Dict[str, Any]:
    import excel
    unknown = [t for t in targets if t not in ("word", "ppt")]
    if unknown:
        raise rt.InputError("엑셀 차트는 word, ppt로만 보낼 수 있습니다.")
    results: Dict[str, Any] = {}
    errors: Dict[str, List[str]] = {}
    with off.open_target("excel", args.from_target, readonly=True) as src:
        ws, chart_obj = excel.find_chart(src, args.chart, args.sheet)
        title = excel._chart_title(chart_obj) or str(chart_obj.Name)
        stem = f"chart_{str(chart_obj.Name)}"[:40]
        images: Dict[str, str] = {}
        pending = list(targets)
        attempt = 0
        while pending and attempt < len(METHODS) + 1:
            attempt += 1
            # 대상마다 아직 시도하지 않은 다음 방식으로
            for kind in list(pending):
                order = method_order(kind, args.method)
                tried = errors.setdefault(kind, [])
                remaining = [m for m in order if not any(e.startswith(m + ":") for e in tried)]
                if not remaining:
                    pending.remove(kind)
                    continue
                method = remaining[0]
                try:
                    if method == "relay":
                        done = relay_chart(chart_obj, args.to_target if kind == "word" else None,
                                           args.to_target if kind == "ppt" else None, args.anchor, args.slide, title,
                                           args.occurrence)
                        results[kind] = {"method": "relay", "target_resolved": done[kind]}
                    else:
                        if method not in images:
                            images[method] = (image_by_export if method == "export" else image_by_redraw)(chart_obj, stem)
                        img = images[method]
                        resolved = (insert_picture_word(args.to_target, img, args.anchor, args.occurrence) if kind == "word"
                                    else insert_picture_ppt(args.to_target, img, args.slide, title))
                        results[kind] = {"method": method, "target_resolved": resolved, "image_file": img}
                    pending.remove(kind)
                except Exception as e:
                    if isinstance(e, rt.InputError) and e.no_fallback:
                        raise                           # 앵커/대상 문제는 다른 방식으로 바꿔도 같음
                    msg = rt.error_payload(e)["message"]
                    rt.log("WARN", f"{off.APP_LABEL[kind]}로 전달 실패 ({method}): {msg} -> 다음 방식 시도")
                    tried.append(f"{method}: {msg}")
    failed = [k for k in targets if k not in results]
    payload = {"chart": title, "results": results}
    if failed:
        payload["failures"] = {k: errors.get(k, []) for k in failed}
        if not results:
            raise rt.InputError(f"차트를 {', '.join(off.APP_LABEL[k] for k in failed)}에 넣지 못했습니다. "
                                f"시도한 방식: {'; '.join(sum((errors.get(k, []) for k in failed), []))}. "
                                "office.py transfer --probe로 이 PC에서 되는 방식을 확인하세요.")
    used = ", ".join(f"{off.APP_LABEL[k]}({v['method']})" for k, v in results.items())
    preview = f"차트 '{title}'를 넣었습니다: {used}" + (f" / 실패: {', '.join(failed)}" if failed else "")
    directive = "VERIFIED: 전달 완료." if not failed else \
        "VERIFIED: 일부만 전달했습니다. failures 내용을 사용자에게 알리세요."
    return rt.ok(TOOL, "transfer", preview, directive, **payload)


def _excel_range_rows(args) -> Tuple[List[List[Any]], str]:
    import excel
    closed = excel.is_closed_xlsx(args.from_target)
    if closed:
        sheet, _, _, matrix, _ = excel._read_matrix_closed(closed, args.sheet, args.range)
    else:
        with off.open_target("excel", args.from_target, readonly=True) as t:
            ws = t.sheet(args.sheet)
            rng = excel.target_range(t, ws, args.range)
            matrix = off.range_to_matrix(rng.Value)
            sheet = str(rng.Worksheet.Name)
    rows = [[excel.out_value(v) for v in r] for r in matrix]
    if not rows or not any(any(str(v).strip() for v in r) for r in rows):
        raise rt.InputError(f"범위 {args.range}에 데이터가 없습니다.")
    return rows, sheet


def _excel_range_transfer(args, targets: List[str]) -> Dict[str, Any]:
    import ppt
    import word
    rows, sheet = _excel_range_rows(args)
    results: Dict[str, Any] = {}
    for kind in targets:
        if kind == "word":
            with off.open_target("word", args.to_target, readonly=False, create=True, tool=TOOL, command="transfer") as t:
                start, _ = word._insert_position(t.doc, args.anchor, "after" if args.anchor else "end", args.occurrence)
                word.Writer(t.doc, start, off.get_theme(args.theme)).table(rows)
                t.set_focus()
                results["word"] = {"method": "table", "target_resolved": t.resolved()}
        elif kind == "ppt":
            with off.open_target("ppt", args.to_target, readonly=False, create=True, tool=TOOL, command="transfer") as t:
                header, body = rows[0], rows[1:] or [[""] * len(rows[0])]
                slides = []
                for k in range(0, len(body), ppt.TABLE_ROWS_PER_SLIDE):
                    s, no = _ppt_slide(t.doc, args.slide if not slides else "new", "")
                    ppt.SlideWriter(t.doc, s, off.get_theme(args.theme or "navy")).table([header] + body[k:k + ppt.TABLE_ROWS_PER_SLIDE])
                    slides.append(no)
                t.set_focus(slide=slides[-1])
                results["ppt"] = {"method": "table", "slides": slides, "target_resolved": t.resolved()}
        else:
            raise rt.InputError("엑셀 범위는 word, ppt로 보낼 수 있습니다. 엑셀 안에서 옮기려면 excel.py range --action copy를 쓰세요.")
    return rt.ok(TOOL, "transfer", f"{sheet}!{args.range} ({len(rows)}행)을 표로 넣었습니다: {', '.join(results)}",
                 rows=len(rows), results=results)


def _word_rows(args) -> Tuple[List[List[Any]], str]:
    import word
    if args.what == "text":
        closed = word._closed_docx(args.from_target)
        if closed:
            import reader
            units, _ = reader.parse_docx(closed)
            rows = [["번호", "구분", "상위 제목", "내용"]] + [[i, "표" if u["loc"].startswith("표") else "본문", u["heading"], u["text"]]
                                                         for i, u in enumerate(units, 1)]
            return rows, os.path.basename(closed)
        with off.open_target("word", args.from_target, readonly=True) as t:
            units = word._com_units(t.doc)
            name = t.name
        rows = [["번호", "구분", "상위 제목", "내용"]] + [[i, f"제목{u['level']}" if u["level"] else "본문", u["heading"], u["text"]]
                                                     for i, u in enumerate(units, 1)]
        return rows, name
    n = args.table or 1
    closed = word._closed_docx(args.from_target)
    if closed:
        import docx
        d = docx.Document(closed)
        if not 1 <= n <= len(d.tables):
            raise rt.InputError(f"표 {n}번이 없습니다 (전체 {len(d.tables)}개).")
        rows = []
        for row in d.tables[n - 1].rows:
            seen, cells = set(), []
            for c in row.cells:
                if id(c._tc) not in seen:
                    seen.add(id(c._tc))
                    cells.append(c.text.strip())
            rows.append(cells)
        return rows, os.path.basename(closed)
    with off.open_target("word", args.from_target, readonly=True) as t:
        if not 1 <= n <= t.doc.Tables.Count:
            raise rt.InputError(f"표 {n}번이 없습니다 (전체 {t.doc.Tables.Count}개).")
        grid: Dict[Tuple[int, int], str] = {}
        for cell in t.doc.Tables(n).Range.Cells:            # 병합 셀도 안전하게 순회
            grid[(int(cell.RowIndex), int(cell.ColumnIndex))] = str(cell.Range.Text).rstrip("\r\x07").replace("\r", " ").strip()
        name = t.name
    nrows = max(r for r, _ in grid)
    ncols = max(c for _, c in grid)
    return [[grid.get((r, c), "") for c in range(1, ncols + 1)] for r in range(1, nrows + 1)], name


def _ppt_rows(args) -> Tuple[List[List[Any]], str]:
    import ppt
    if args.what == "text":
        closed = ppt._closed_pptx(args.from_target)
        if closed:
            import reader
            units, _ = reader.parse_pptx(closed)
            rows = [["슬라이드", "제목", "내용"]] + [[i, u["heading"], u["text"]] for i, u in enumerate(units, 1)]
            return rows, os.path.basename(closed)
        with off.open_target("ppt", args.from_target, readonly=True) as t:
            us = [ppt.slide_units(t.doc.Slides(i), i) for i in range(1, t.doc.Slides.Count + 1)]
            name = t.name
        return [["슬라이드", "제목", "내용", "노트"]] + [[u["no"], u["title"], u["text"], u["notes"]] for u in us], name
    with off.open_target("ppt", args.from_target, readonly=True) as t:
        prs = t.doc
        numbers = [int(args.slide)] if args.slide and str(args.slide).isdigit() else list(range(1, prs.Slides.Count + 1))
        tables = []
        for n in numbers:
            s = prs.Slides(n)
            for k in range(1, s.Shapes.Count + 1):
                sh = s.Shapes(k)
                if ppt._has(sh, "HasTable"):
                    tables.append(sh.Table)
        idx = args.table or 1
        if not 1 <= idx <= len(tables):
            raise rt.InputError(f"표 {idx}번이 없습니다 (찾은 표 {len(tables)}개).")
        tb = tables[idx - 1]
        rows = [[str(tb.Cell(r, c).Shape.TextFrame.TextRange.Text).strip() for c in range(1, tb.Columns.Count + 1)]
                for r in range(1, tb.Rows.Count + 1)]
        name = t.name
    return rows, name


def _to_excel(args, rows: List[List[Any]], sheet_base: str) -> Dict[str, Any]:
    with off.open_target("excel", args.to_target, readonly=False, create=True, tool=TOOL, command="transfer") as t:
        wb = t.doc
        names = [str(wb.Worksheets(i).Name) for i in range(1, wb.Worksheets.Count + 1)]
        name = (args.sheet or sheet_base)[:31]
        k = 2
        while rt.find_name(name, names) and not args.sheet:
            name = f"{sheet_base[:27]} {k}"
            k += 1
        existing = rt.find_name(name, names)
        if existing is not None and not args.overwrite:
            # --sheet는 "새 시트 이름": 이미 있는 시트에 A1부터 덮어쓰지 않습니다
            free, k = name, 2
            while rt.find_name(free, names):
                free = f"{name[:27]} {k}"
                k += 1
            raise rt.InputError(f"시트 '{existing}'이(가) 이미 있어 덮어쓰지 않았습니다. 새 이름으로 받거나, 기존 시트에 덮어쓰려면 "
                                f"사용자 확인 후 --overwrite를 붙이세요. [추천] {free}")
        ws = t.sheet(existing or name, create=existing is None)
        width = max(len(r) for r in rows)
        import excel
        # Word/PPT text stays text: "001", "1-2", "(100)" would become numbers/dates, "=..." a formula
        cell = lambda v: "'" + v if isinstance(v, str) and v.startswith("=") else excel.xl_text(v)  # noqa: E731
        data = tuple(tuple(cell(v) for v in list(r) + [""] * (width - len(r))) for r in rows)
        with off.excel_batch(t.app):
            ws.Range(ws.Cells(1, 1), ws.Cells(len(data), width)).Value = data
            header = ws.Range(ws.Cells(1, 1), ws.Cells(1, width))
            header.Font.Bold = True
            try:
                ws.Range(ws.Cells(1, 1), ws.Cells(len(data), width)).Columns.AutoFit()
            except Exception:
                pass
        t.set_focus(sheet=ws)
        return {"sheet": str(ws.Name), "rows": len(data), "cols": width, "target_resolved": t.resolved()}


def cmd_transfer(args) -> Dict[str, Any]:
    if args.probe:
        return probe()
    if not args.source or not args.to:
        raise rt.InputError("--from과 --to를 지정하세요 (예: --from excel --chart 차트이름 --to word,ppt).")
    src = _app(args.source)
    targets = list(dict.fromkeys(_app(x) for x in args.to.split(",") if x.strip()))
    if src in targets:
        raise rt.InputError("보내는 앱과 받는 앱이 같습니다.")
    if src == "excel":
        if args.chart:
            return _excel_chart_transfer(args, targets)
        if args.range:
            return _excel_range_transfer(args, targets)
        raise rt.InputError("엑셀에서 보낼 때는 --chart <차트 이름> 또는 --range <범위>를 지정하세요. 차트 이름은 excel.py chart --action list로 확인.")
    if targets != ["excel"]:
        raise rt.InputError("워드/PPT의 내용은 엑셀로 보낼 수 있습니다 (--to excel). 워드<->PPT는 read --save 후 write --file을 쓰세요.")
    rows, name = _word_rows(args) if src == "word" else _ppt_rows(args)
    if len(rows) <= 1 and args.what == "text":
        raise rt.InputError(f"'{name}'에서 옮길 내용이 없습니다.")
    base = f"{'Word' if src == 'word' else 'PPT'} {'내용' if args.what == 'text' else '표'}"
    info = _to_excel(args, rows, base)
    return rt.ok(TOOL, "transfer", f"'{name}'의 {'내용' if args.what == 'text' else '표'}을 엑셀 시트 '{info['sheet']}'에 넣었습니다 "
                 f"({info['rows']}행 x {info['cols']}열)", "VERIFIED: 전달 완료. 다시 읽어서 확인하지 마세요.", **info)


# ==============================================================================
# 5. probe: 이 PC에서 되는 차트 전달 방식 시험
# ==============================================================================

def probe() -> Dict[str, Any]:
    off.require_com()
    off.init_com()
    results: Dict[str, Dict[str, str]] = {"word": {}, "ppt": {}}
    xl = off.win32com.client.DispatchEx("Excel.Application")
    wd = off.win32com.client.DispatchEx("Word.Application")
    pp = None
    pp_started = False
    wb = wdoc = prs = None
    try:
        xl.Visible = False
        xl.DisplayAlerts = False
        wd.Visible = False
        wd.DisplayAlerts = 0
        wb = xl.Workbooks.Add()
        ws = wb.Worksheets(1)
        ws.Range("A1:B4").Value = (("x", "y"), (1, 2), (2, 4), (3, 3))
        chart_obj = ws.ChartObjects().Add(10, 10, 300, 180)
        chart_obj.Chart.ChartType = 4
        chart_obj.Chart.SetSourceData(ws.Range("A1:B4"))
        wdoc = wd.Documents.Add(Visible=False)
        try:
            off.win32com.client.GetActiveObject("PowerPoint.Application")
        except Exception:
            pp_started = True
        pp = off.win32com.client.Dispatch("PowerPoint.Application")
        prs = pp.Presentations.Add(0)
        slide = prs.Slides.Add(1, 12)

        for method in ("export", "redraw"):
            try:
                img = (image_by_export if method == "export" else image_by_redraw)(chart_obj, "probe")
                try:
                    wdoc.InlineShapes.AddPicture(FileName=img, LinkToFile=False, SaveWithDocument=True)
                    results["word"][method] = "ok"
                except Exception as e:
                    results["word"][method] = f"실패: {rt.error_payload(e)['message']}"
                try:
                    slide.Shapes.AddPicture(FileName=img, LinkToFile=0, SaveWithDocument=-1, Left=0, Top=0)
                    results["ppt"][method] = "ok"
                except Exception as e:
                    results["ppt"][method] = f"실패: {rt.error_payload(e)['message']}"
            except Exception as e:
                msg = f"실패: {rt.error_payload(e)['message']}"
                results["word"][method] = results["ppt"][method] = msg
        try:
            chart_obj.CopyPicture(1, -4147)
            time.sleep(0.3)
            before = wdoc.InlineShapes.Count
            wdoc.Range(0, 0).PasteSpecial(DataType=3)
            if wdoc.InlineShapes.Count <= before:
                raise RuntimeError("워드에 붙여넣기 안 됨")
            results["word"]["relay"] = "ok"
            wdoc.InlineShapes.Item(1).Range.CopyAsPicture()
            time.sleep(0.3)
            n0 = slide.Shapes.Count
            slide.Shapes.Paste()
            results["ppt"]["relay"] = "ok" if slide.Shapes.Count > n0 else "실패: PPT에 붙여넣기 안 됨"
        except Exception as e:
            msg = f"실패: {rt.error_payload(e)['message']}"
            results["word"].setdefault("relay", msg)
            results["ppt"].setdefault("relay", msg)
    finally:
        for closer in (lambda: prs.Close() if prs is not None else None,
                       lambda: pp.Quit() if pp is not None and pp_started and pp.Presentations.Count == 0 else None,
                       lambda: wdoc.Close(SaveChanges=0) if wdoc is not None else None,
                       lambda: wd.Quit(SaveChanges=0),
                       lambda: wb.Close(SaveChanges=False) if wb is not None else None,
                       lambda: xl.Quit()):
            try:
                closer()
            except Exception:
                pass
    order = {k: [m for m in METHODS if v.get(m) == "ok"] for k, v in results.items()}
    rt.set_config("transfer_method", order)
    preview = "; ".join(f"{off.APP_LABEL[k]}: {', '.join(v) or '되는 방식 없음'}" for k, v in order.items())
    return rt.ok(TOOL, "transfer", f"차트 전달 방식 시험 결과 저장 - {preview}",
                 "VERIFIED: 시험 결과를 저장했습니다. 이후 transfer는 이 순서로 자동 시도합니다.",
                 probe=results, saved_order=order)


# ==============================================================================
# 6. undo
# ==============================================================================

def _backup_path(kind: str, ext: str) -> str:
    folder = os.path.join(rt.get_cache_dir(), "snapshots", "undo_backups")
    os.makedirs(folder, exist_ok=True)
    return os.path.normpath(os.path.join(folder, f"{kind}_{rt.timestamp()}_{os.getpid()}{ext}"))


def _word_from_flat_xml(app: Any, xml_path: str, dest: str) -> Any:
    doc = app.Documents.Open(os.path.normpath(xml_path), ConfirmConversions=False, AddToRecentFiles=False)
    fmt = off.EXTENSIONS["word"].get(os.path.splitext(dest)[1].lower())
    off.word_save_as(doc, os.path.normpath(dest), fmt)
    return doc


def _open_snapshot_untitled(kind: str, app: Any, data: str) -> Any:
    """스냅샷 내용을 '제목 없는 새 문서'로 엽니다 (캐시 파일 자체를 열지 않음)."""
    if kind == "excel":
        return app.Workbooks.Add(os.path.normpath(data))
    if kind == "ppt":
        return app.Presentations.Open(os.path.normpath(data), 0, -1, -1)
    return app.Documents.Add(Template=os.path.normpath(data))


def _doc_name(doc: Any, default: str = "") -> str:
    try:
        return str(doc.Name)
    except Exception:
        return default


def _word_replace_content(doc: Any, data: str, fmt: str) -> None:
    """열린 워드 문서의 내용을 스냅샷 내용으로 바꿉니다 (같은 창 유지).
    flat_xml(WordOpenXML) 스냅샷은 InsertXML로 넣어 스타일을 살리고, 그 밖에는 InsertFile."""
    if fmt == "flat_xml":
        doc.Range().InsertXML(rt.read_text_file(data))
    else:
        doc.Content.Delete()
        doc.Content.InsertFile(os.path.normpath(data))


def _restore_untitled(snap: Dict[str, Any]) -> Dict[str, Any]:
    """저장한 적 없는 문서 되돌리기 (시나리오 12: 새 워드 문서에 차트를 넣은 뒤 undo).
    - 바뀐 문서가 하나 열려 있으면: 워드는 같은 창에서 내용을 바꾸고(restored_in_place),
      엑셀/PPT는 되돌린 내용을 새 문서로 연 뒤 바뀐 문서를 저장하지 않고 닫는다(replaced_untitled_document).
      어느 쪽이든 되돌리기 직전 상태를 pre_undo_backup에 먼저 보관한다.
    - 바뀐 문서가 닫혔거나 같은 이름의 저장 안 된 문서가 둘 이상이면: 고르지 않고 새 문서로만 연다(new_untitled_document)."""
    kind, name, data = snap["app"], snap.get("name", ""), snap["snapshot_file"]
    label = off.APP_LABEL[kind]
    hits = off.open_untitled(kind, name) if name else []
    if len(hits) == 1:
        app, doc = hits[0]["app"], hits[0]["doc"]
        if kind == "word":
            backup = _backup_path(kind, ".xml")
            with open(backup, "w", encoding="utf-8") as f:
                f.write(doc.WordOpenXML)
            _word_replace_content(doc, data, snap.get("snapshot_format", ""))
            try:
                doc.Activate()
            except Exception:
                pass
            return {"restored_as": "restored_in_place", "pre_undo_backup": backup, "restored_name": _doc_name(doc, name),
                    "detail": f"'{name}'을(를) 같은 창에서 되돌렸습니다 (저장하지 않은 문서라 저장은 하지 않았습니다)"}
        backup = _backup_path(kind, ".xlsx" if kind == "excel" else ".pptx")
        doc.SaveCopyAs(backup)
        new = _open_snapshot_untitled(kind, app, data)
        if kind == "excel":
            doc.Close(SaveChanges=False)
        else:
            try:
                doc.Saved = -1                     # 닫을 때 저장할지 묻지 않도록 (내용은 pre_undo_backup에 보관)
            except Exception:
                pass
            doc.Close()
        try:
            new.Activate()
        except Exception:
            pass
        new_name = _doc_name(new)
        return {"restored_as": "replaced_untitled_document", "pre_undo_backup": backup, "restored_name": new_name,
                "detail": f"저장한 적 없는 {label} 문서라 되돌린 내용을 새 창 '{new_name}'(으)로 열고, "
                          f"바뀐 '{name}'은(는) 저장하지 않고 닫았습니다"}
    if len(hits) > 1:
        reason = f"이름이 '{name}'인 저장 안 된 {label} 문서가 {len(hits)}개라 어느 것인지 정할 수 없어"
    else:
        reason = f"'{name}'이(가) 열려 있지 않아"
    app = hits[0]["app"] if hits else off.get_active(kind)["app"]
    new_name = _doc_name(_open_snapshot_untitled(kind, app, data))
    return {"restored_as": "new_untitled_document", "pre_undo_backup": "", "restored_name": new_name,
            "detail": f"{reason} 되돌린 내용을 새 창 '{new_name}'(으)로 열었습니다"
                      + (" (바뀐 문서는 그대로 열려 있습니다)" if hits else "")}


def restore_snapshot(snap: Dict[str, Any]) -> Dict[str, Any]:
    import shutil
    kind = snap["app"]
    orig = snap.get("original_path", "")
    data = snap["snapshot_file"]
    if not orig:
        return _restore_untitled(snap)
    if not os.path.isdir(os.path.dirname(orig)):
        raise rt.InputError(f"원본 폴더가 없습니다: {os.path.dirname(orig)}")
    ext = os.path.splitext(orig)[1]
    open_item = None
    try:
        open_item = off.match_open(kind, orig)
    except Exception:
        pass
    if open_item is not None:
        app, doc = open_item["app"], open_item["doc"]
        opener = {"word": lambda: app.Documents, "excel": lambda: app.Workbooks, "ppt": lambda: app.Presentations}[kind]
        if kind == "word":
            backup = _backup_path(kind, ".xml")
            with open(backup, "w", encoding="utf-8") as f:
                f.write(doc.WordOpenXML)
            doc.Close(SaveChanges=0)
            if snap.get("snapshot_format") == "flat_xml":
                _word_from_flat_xml(app, data, orig)
                return {"restored_as": "reopened_in_office", "pre_undo_backup": backup}
        else:
            backup = _backup_path(kind, ext)
            doc.SaveCopyAs(backup)
            doc.Close(SaveChanges=False) if kind == "excel" else doc.Close()
        try:
            # Office가 닫은 직후 잠금을 늦게 풀거나 백신이 잡고 있을 수 있음
            rt.retry_fs(lambda: shutil.copy2(data, orig))
        except OSError as e:
            try:
                opener().Open(os.path.normpath(orig))       # 사용자 문서를 닫힌 채로 두지 않음
            except Exception:
                pass
            raise rt.InputError(f"되돌린 파일을 원본 자리에 복사하지 못해 되돌리지 않았습니다 ({type(e).__name__}). 원본은 그대로 다시 열었고, "
                                f"되돌리기 직전 상태는 pre_undo_backup에 있습니다.",
                                directive="되돌리지 못했습니다. message와 pre_undo_backup 경로를 사용자에게 그대로 알리세요.",
                                error_type="undo_failed", fields={"pre_undo_backup": backup})
        opener().Open(os.path.normpath(orig))
        return {"restored_as": "reopened_in_office", "pre_undo_backup": backup}
    backup = ""
    if os.path.exists(orig):
        backup = _backup_path(kind, ext)
        shutil.copy2(orig, backup)
    if kind == "word" and snap.get("snapshot_format") == "flat_xml":
        wd = off.win32com.client.DispatchEx("Word.Application")
        try:
            wd.Visible = False
            wd.DisplayAlerts = 0
            _word_from_flat_xml(wd, data, orig).Close(SaveChanges=0)
        finally:
            wd.Quit(SaveChanges=0)
    else:
        rt.retry_fs(lambda: shutil.copy2(data, orig))
    return {"restored_as": "file_replaced", "pre_undo_backup": backup}


def cmd_undo(args) -> Dict[str, Any]:
    kind = _app(args.app) if args.app else None
    snaps = off.list_snapshots(kind)
    if args.list:
        # 번호는 --app 필터에 따라 달라지므로 그대로 실행할 명령(같은 --app 포함)을 항목마다 줍니다
        app_arg = ["--app", args.app] if args.app else []
        items = [{"pick": i, "app": off.APP_LABEL[s["app"]], "name": s["name"], "created_at": s.get("created_at", "")[:15],
                  "by": s.get("by", "")}
                 | ({"undo_available": False, "reason": s.get("reason", "")} if s.get("barrier")
                    else {"command": rt.next_command("office.py", ["undo", "--pick", i] + app_arg, "")["command"]})
                 for i, s in enumerate(snaps[:20], 1)]
        return rt.ok(TOOL, "undo", f"되돌릴 수 있는 지점 {sum(1 for s in snaps if not s.get('barrier'))}개",
                     "VERIFIED: 목록을 사용자에게 보여 주세요. 사용자가 고른 항목의 command를 글자 그대로 실행하세요 (번호만 바꿔 직접 만들지 않음).",
                     snapshots=items, count=len(snaps), shown=len(items), truncated=len(snaps) > len(items))
    if not snaps:
        return rt.ok(TOOL, "undo", "되돌릴 스냅샷이 없습니다",
                     "VERIFIED: 되돌릴 수 있는 작업이 없습니다. 사용자에게 알리세요.", restored=False)
    pick = args.pick or 1
    if not 1 <= pick <= len(snaps):
        raise rt.InputError(f"--pick은 1~{len(snaps)} 사이여야 합니다 (office.py undo --list로 확인).")
    snap = snaps[pick - 1]
    if snap.get("barrier"):
        if args.pick:
            raise rt.InputError(f"{pick}번은 되돌릴 지점이 없는 작업입니다 ({snap.get('reason', '')}). office.py undo --list에서 다른 번호를 고르세요.")
        raise rt.UserDecisionError(
            f"직전 작업('{snap['name']}' {snap.get('by', '')})은 되돌릴 지점이 없습니다: {snap.get('reason', '')}. "
            "그보다 앞 지점으로 되돌리면 그 사이의 작업도 함께 사라집니다. 되돌리지 않았습니다.",
            "되돌리지 않았습니다. message를 사용자에게 알리고, 사용자가 더 앞 지점으로 되돌리길 원하면 "
            "office.py undo --list를 보여 준 뒤 사용자가 고른 번호로 --pick을 붙여 실행하세요.")
    orig = snap.get("original_path", "")
    created = float(snap.get("created_ts") or 0)
    if orig and created and not args.overwrite and os.path.exists(orig) and os.path.getmtime(orig) > created + 120:
        changed = datetime.fromtimestamp(os.path.getmtime(orig)).strftime("%Y-%m-%d %H:%M")
        raise rt.UserDecisionError(
            f"'{snap['name']}'은(는) 이 되돌리기 지점({snap.get('created_at', '')[:15]}) 뒤에 다시 저장되었습니다 (마지막 저장 {changed}). "
            "되돌리면 그 사이에 저장한 내용이 사라집니다. 되돌리지 않았습니다.",
            "되돌리지 않았습니다. message를 사용자에게 알리고, 사용자가 그래도 되돌리길 원하면 같은 명령에 --overwrite를 붙여 실행하세요.")
    off.require_com()
    off.init_com()
    info = restore_snapshot(snap)
    off.delete_snapshot(snap["meta_file"])                 # 사용한 스냅샷은 소비 -> 다음 undo는 그 이전으로
    detail = info.pop("detail", "")
    preview = f"'{snap['name']}'을(를) '{snap.get('by') or '이전 작업'}' 이전 상태로 되돌렸습니다 ({info['restored_as']})"
    if detail:
        preview += f". {detail}"
    directive = "VERIFIED: 되돌리기 완료. " + (f"{detail}. " if detail else "") + (
        "되돌리기 직전 상태는 pre_undo_backup에 보관했습니다. preview를 사용자에게 그대로 알리세요."
        if info.get("pre_undo_backup") else "preview를 사용자에게 그대로 알리세요.")
    return rt.ok(TOOL, "undo", preview, directive,
                 restored=True, app=off.APP_LABEL[snap["app"]], name=snap["name"], original_path=snap.get("original_path", ""), **info)


# ==============================================================================
# 7. 진입점
# ==============================================================================

def build_cli():
    parser, sub = rt.build_parser(TOOL, "Office 공통 작업")
    sub.add_parser("windows", help="열린 Office 문서 목록")

    p = sub.add_parser("focus", help="창을 맨 앞으로")
    p.add_argument("--app", required=True, help="excel, word, ppt")
    p.add_argument("--name", help="문서 이름 (생략 시 그 앱의 활성 문서)")

    p = sub.add_parser("transfer", help="앱 간 전달")
    p.add_argument("--from", dest="source", help="보내는 앱: excel, word, ppt")
    p.add_argument("--from-target", default="active", help="보내는 문서 (active, 열린 문서 이름, 경로)")
    p.add_argument("--to", help="받는 앱 (쉼표로 여러 개: word,ppt)")
    p.add_argument("--to-target", default="active", help="받는 문서 (기본: 각 앱의 활성 문서)")
    p.add_argument("--chart", help="엑셀 차트 이름 또는 제목 (excel.py chart --action list로 확인)")
    p.add_argument("--range", help="엑셀 범위 (표로 전달)")
    p.add_argument("--sheet", help="엑셀 시트 (보낼 때: 원본 시트 / 엑셀로 받을 때: 새 시트 이름)")
    p.add_argument("--overwrite", action="store_true",
                   help="엑셀로 받을 때 --sheet가 이미 있으면 그 시트의 A1부터 덮어쓰기 (범위 밖 기존 셀은 남음, 시트가 없으면 새로 만듦)")
    p.add_argument("--what", choices=["text", "table"], default="text", help="워드/PPT에서 보낼 것 (기본 text)")
    p.add_argument("--table", type=int, help="표 번호 (1부터)")
    p.add_argument("--slide", help="PPT 슬라이드: 받을 때 번호 또는 new(기본), 보낼 때 표를 찾을 슬라이드")
    p.add_argument("--anchor", help="워드에서 이 글이 있는 단락 뒤에 넣기 (기본: 문서 끝)")
    p.add_argument("--occurrence", type=int, help="--anchor가 여러 단락에 있을 때 몇 번째인지 (1부터, 오류의 [후보] 번호)")
    p.add_argument("--method", choices=["auto"] + list(METHODS), default="auto", help="차트 그림 방식 (기본 auto)")
    p.add_argument("--theme", help=f"표 색: {', '.join(off.THEMES)}")
    p.add_argument("--probe", action="store_true", help="이 PC에서 되는 차트 전달 방식을 시험해 저장")

    p = sub.add_parser("undo", help="직전 작업 되돌리기")
    p.add_argument("--app", help="이 앱의 작업만 (excel, word, ppt)")
    p.add_argument("--list", action="store_true", help="되돌릴 수 있는 지점 목록")
    p.add_argument("--pick", type=int, help="--list의 번호 (기본 1 = 가장 최근)")
    p.add_argument("--overwrite", action="store_true",
                   help="되돌리기 지점 뒤에 파일이 다시 저장되었어도 되돌리기 (그 사이 저장한 내용이 사라짐, 사용자가 원할 때만)")
    return parser


HANDLERS = {"windows": cmd_windows, "focus": cmd_focus, "transfer": cmd_transfer, "undo": cmd_undo}


def main(argv: Optional[Sequence[str]] = None) -> None:
    args = build_cli().parse_args(argv)
    rt.run_cli(TOOL, args.command, lambda: HANDLERS[args.command](args))


if __name__ == "__main__":
    main()
