"""excel.py - 엑셀 작업

명령:
    read     내용 읽기 (범위, 검색, --save로 CSV 저장)
    info     통합문서 요약 (시트, 사용 범위, 차트, 선택 영역). --all: 열린 통합문서 전체 목록
    write    값/표/수식 쓰기 (--value, --values, --file, --formula, --header, --append)
    replace  찾아 바꾸기 (수식 셀은 보존)
    style    서식 (테마, 배경/글꼴/정렬/숫자서식/테두리, 조건부 강조)
    export   CSV / XLSX / PDF로 저장 (원본 창은 그대로)
    sheet    시트 목록/추가/이름변경/복사/삭제/선택
    range    너비/높이/자동맞춤/병합/틀고정/숨김/지우기/복사/이동/테두리
    chart    차트 만들기/목록/삭제

--target: active(기본, 지금 보고 있는 통합문서) | 열린 통합문서 이름 | 파일 경로
닫힌 xlsx/xlsm 파일 읽기는 Excel을 띄우지 않고 openpyxl로 읽습니다.

예:
    python excel.py write --range A1:B101 --header "x,sin(x)" --formula "=(ROW()-2)*4*PI()/99" "=SIN(A2)"
    python excel.py chart --range A1:B101 --type line --title "사인 2주기"
    python excel.py write --file "결과.csv" --sheet "검색결과" --new-sheet
    python excel.py read --range A1:D20
"""
import json
import os
import re
import sys
import zipfile
from typing import Any, Dict, List, Optional, Sequence, Tuple

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import _office as off  # noqa: E402
import _runtime as rt  # noqa: E402

TOOL = "excel"

XL_CHART = {"column": 51, "bar": 57, "line": 4, "line_markers": 65, "pie": 5, "area": 1,
            "scatter": -4169, "scatter_lines": 75, "combo": 51}   # -4169 표식만, 75 선만(표식 없음)
XL_LEGEND = {"bottom": -4107, "right": -4152, "top": -4160, "left": -4131}
XL_ALIGN = {"left": -4131, "center": -4108, "right": -4152}
XL_WEIGHT = {"hairline": 1, "thin": 2, "medium": -4138, "thick": 4}
XL_OPERATOR = {">": 5, ">=": 7, "<": 6, "<=": 8, "=": 3, "<>": 4, "!=": 4}


# ==============================================================================
# 1. 주소와 값 도우미
# ==============================================================================

def col_letter(n: int) -> str:
    s = ""
    while n > 0:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def col_number(letters: str) -> int:
    n = 0
    for ch in letters.upper():
        n = n * 26 + (ord(ch) - 64)
    return n


_CELL_RE = re.compile(r"^\$?([A-Za-z]{1,3})\$?(\d+)$")


def parse_cell(ref: str) -> Tuple[int, int]:
    """'B12' -> (12, 2)"""
    m = _CELL_RE.match(ref.strip())
    if not m:
        raise rt.InputError(f"셀 주소 형식이 아닙니다: '{ref}' (예: A1, B12)")
    return int(m.group(2)), col_number(m.group(1))


def parse_area(cells: str, max_row: Optional[int], max_col: Optional[int]) -> Tuple[int, int, Optional[int], Optional[int]]:
    """'A1:C5', 'B:B', 'B:D', '5:7', 'A1' -> (r1, c1, r2, c2). 열/행 전체는 데이터 끝까지 (None = 끝까지)."""
    parts = cells.replace("$", "").strip().split(":")
    a, b = parts[0], parts[-1]
    if re.fullmatch(r"[A-Za-z]{1,3}", a) and re.fullmatch(r"[A-Za-z]{1,3}", b):
        return 1, col_number(a), max_row, col_number(b)
    if a.isdigit() and b.isdigit():
        return int(a), 1, int(b), max_col
    r1, c1 = parse_cell(a)
    r2, c2 = parse_cell(b)
    return r1, c1, r2, c2


def split_sheet(address: str) -> Tuple[Optional[str], str]:
    """"'My Sheet'!A1:B2" -> ('My Sheet', 'A1:B2'), 'A1' -> (None, 'A1')"""
    if "!" not in address:
        return None, address.strip()
    sheet, cells = address.rsplit("!", 1)
    sheet = sheet.strip()
    if len(sheet) >= 2 and sheet[0] == "'" and sheet[-1] == "'":
        sheet = sheet[1:-1].replace("''", "'")
    return sheet, cells.strip()


def addr(r1: int, c1: int, r2: Optional[int] = None, c2: Optional[int] = None) -> str:
    a = f"{col_letter(c1)}{r1}"
    if r2 is None or (r2 == r1 and c2 == c1):
        return a
    return f"{a}:{col_letter(c2)}{r2}"


def out_value(v: Any) -> Any:
    """Excel 값 -> JSON/CSV용 값 (정수형 실수는 정수로, 날짜는 ISO 문자열로)."""
    if v is None:
        return ""
    if isinstance(v, bool):
        return v
    if isinstance(v, float):
        return int(v) if v.is_integer() and abs(v) < 1e15 else v
    if hasattr(v, "isoformat") and hasattr(v, "year"):
        try:
            if getattr(v, "hour", 0) or getattr(v, "minute", 0) or getattr(v, "second", 0):
                return v.strftime("%Y-%m-%d %H:%M:%S")
            return v.strftime("%Y-%m-%d")
        except Exception:
            return str(v)
    return v


_NUM_RE = re.compile(r"^[+-]?(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d+)?$")


def coerce(s: Any) -> Any:
    """CSV 문자열 -> Excel 값. 숫자는 숫자로, 앞자리 0이 있는 코드(001)와 15자리를 넘는 숫자는 문자 그대로."""
    if not isinstance(s, str):
        return s
    t = s.strip()
    if t == "":
        return None
    if _NUM_RE.match(t):
        digits = t.lstrip("+-").replace(",", "")
        if len(digits) > 1 and digits[0] == "0" and not digits.startswith("0."):
            return t
        if len(digits.replace(".", "")) > 15:       # Excel 숫자 정밀도(15자리)를 넘으면 값이 바뀜
            return t
        num = float(digits) * (-1 if t.startswith("-") else 1)
        return int(num) if num.is_integer() and "." not in digits else num
    return s


# Excel이 COM으로 받은 글자를 직접 입력한 것처럼 다시 해석해 값이 바뀌는 모양:
# 앞자리 0 코드(01234 -> 1234), 15자리 초과 숫자, 월-일/분수처럼 보이는 글(1-2, 3/4 -> 날짜), 괄호 숫자((100) -> -100), 지수(1E5)
_REINTERPRETED = re.compile(r"^(?:0\d+(?:\.\d+)?|[+-]?\d{16,}(?:\.\d+)?|\d{1,2}\s*[-/]\s*\d{1,2}|\(\s*[\d.,]+\s*\)|\d+(?:\.\d+)?[eE][+-]?\d+)$")


def xl_text(v: Any) -> Any:
    """셀에 쓸 값. 글자인데 Excel이 숫자/날짜로 바꿔 버릴 모양이면 앞에 '를 붙여 글자 그대로 둡니다 ('는 셀 값에 남지 않음)."""
    if isinstance(v, str) and _REINTERPRETED.match(v.strip()):
        return "'" + v
    return v


def rect(rows: List[List[Any]]) -> List[List[Any]]:
    width = max((len(r) for r in rows), default=0)
    return [list(r) + [None] * (width - len(r)) for r in rows]


def load_table_file(path: str) -> List[List[Any]]:
    """CSV/TSV, 마크다운 표(.md의 첫 번째 표), 일반 텍스트(한 줄 = 한 행)를 표로 읽습니다."""
    ext = os.path.splitext(path)[1].lower()
    if ext in (".csv", ".tsv"):
        return [[coerce(c) for c in row] for row in rt.read_csv_rows(path)]
    text = rt.read_text_file(path)
    table = [ln.strip() for ln in text.splitlines() if ln.strip().startswith("|")]
    if table:
        rows = []
        for ln in table:
            cells = [c.strip() for c in ln.strip("|").split("|")]
            if all(re.fullmatch(r":?-{3,}:?", c) for c in cells if c):
                continue
            rows.append([coerce(c) for c in cells])
        return rows
    return [[ln.rstrip()] for ln in text.splitlines() if ln.strip()]


def parse_values_arg(raw: str) -> List[List[Any]]:
    try:
        data = json.loads(raw)
    except json.JSONDecodeError as e:
        raise rt.InputError(f"--values는 JSON 2차원 배열이어야 합니다 (예: [[\"이름\",\"값\"],[\"가\",1]]): {e}")
    if not isinstance(data, list) or not data:
        raise rt.InputError("--values는 비어 있지 않은 JSON 배열이어야 합니다.")
    return [list(r) if isinstance(r, list) else [r] for r in data]


# ==============================================================================
# 2. 대상 열기
# ==============================================================================

def is_closed_xlsx(target: str) -> Optional[str]:
    """--target이 열려 있지 않은 xlsx/xlsm 파일 경로면 그 경로 (openpyxl로 읽기), 아니면 None."""
    if not target or target.strip().lower() == "active":
        return None
    raw = os.path.abspath(os.path.expandvars(target.strip().strip('"')))
    if not os.path.isfile(raw):
        cand = os.path.join(rt.get_data_dir(), target.strip().strip('"'))
        if not os.path.isfile(cand):
            return None
        raw = cand
    if os.path.splitext(raw)[1].lower() not in (".xlsx", ".xlsm") or not zipfile.is_zipfile(raw):
        return None   # .xls, DRM 암호화 파일 등은 Excel(COM)로 읽음
    try:
        if off.win32com is not None and off.match_open("excel", raw) is not None:
            return None   # 열려 있으면 화면의 최신(미저장 포함) 내용을 읽음
    except rt.InputError:
        return None   # 같은 이름의 열린 문서가 여럿: 파일을 읽지 않고 COM 경로에서 [후보] 오류로 안내
    except Exception:
        pass
    return raw


def closed_csv(target: str) -> Optional[str]:
    """--target이 CSV/TSV 파일이면 그 경로 (Excel 없이 글자로 읽음), 아니면 None."""
    if not target or target.strip().lower() == "active":
        return None
    raw = os.path.abspath(os.path.expandvars(target.strip().strip('"')))
    if not os.path.isfile(raw):
        cand = os.path.join(rt.get_data_dir(), target.strip().strip('"'))
        raw = cand if os.path.isfile(cand) else ""
    return raw if raw and os.path.splitext(raw)[1].lower() in (".csv", ".tsv") else None


def _read_matrix_csv(path: str, address: Optional[str]) -> Tuple[str, int, int, List[List[Any]], List[str]]:
    """CSV를 엑셀 시트처럼 읽습니다 (인코딩 자동 판별, 숫자는 숫자로). --range가 있으면 그 범위만."""
    rows = [[coerce(v) for v in r] for r in rt.read_csv_rows(path)]
    while rows and all(v is None for v in rows[-1]):
        rows.pop()
    name = os.path.splitext(os.path.basename(path))[0]
    r1, c1 = 1, 1
    if address:
        width = max((len(r) for r in rows), default=0)
        r1, c1, r2, c2 = parse_area(split_sheet(address)[1], len(rows), width)
        rows = [(r + [None] * (width - len(r)))[c1 - 1:c2] for r in rows[r1 - 1:r2]]
    return name, r1, c1, rows, [name]


def write_session(args, command: str):
    return off.open_target("excel", args.target, readonly=False, focus=not args.no_focus,
                           create=True, tool=TOOL, command=command)


def target_range(t: off.OfficeTarget, ws: Any, address: Optional[str]) -> Any:
    """주소 해석: 'A1:B2', 'Sheet!A1', 생략 시 UsedRange."""
    if not address:
        return ws.UsedRange
    sheet, cells = split_sheet(address)
    if sheet:
        ws = t.sheet(sheet)
    try:
        return ws.Range(cells)
    except Exception:
        raise rt.InputError(f"범위 주소가 잘못되었습니다: '{address}' (예: A1, A1:D10, B:B, 'Sheet1'!A1)")


def range_info(rng: Any) -> Dict[str, Any]:
    return {"sheet": str(rng.Worksheet.Name), "range": str(rng.Address).replace("$", ""),
            "rows": int(rng.Rows.Count), "cols": int(rng.Columns.Count)}


# ==============================================================================
# 3. read / info
# ==============================================================================

def _read_matrix_closed(path: str, sheet: Optional[str], address: Optional[str]) -> Tuple[str, int, int, List[List[Any]], List[str]]:
    """닫힌 파일을 openpyxl로 읽습니다. (시트명, 시작행, 시작열, 값, 전체 시트명)"""
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        names = list(wb.sheetnames)
        if sheet:
            found = rt.find_name(sheet, names)
            if found is None:
                raise rt.not_found("시트", sheet, names)
            ws = wb[found]
        else:
            ws = wb[names[0]] if wb.active is None else wb.active
        if address:
            sh, cells = split_sheet(address)
            if sh:
                found = rt.find_name(sh, names)
                if found is None:
                    raise rt.not_found("시트", sh, names)
                ws = wb[found]
            r1, c1, r2, c2 = parse_area(cells, ws.max_row, ws.max_column)
        else:
            r1, c1 = 1, 1
            r2, c2 = ws.max_row, ws.max_column
        if r2 is None or c2 is None:
            # 크기 정보(<dimension>)가 없는 파일(다른 프로그램이 만든 xlsx): 1x1로 읽으면 데이터를 잃으므로 끝까지 읽음
            rows = [list(r) for r in ws.iter_rows(min_row=r1, min_col=c1, values_only=True)]
            width = max((len(r) for r in rows), default=0)
            rows = [r + [None] * (width - len(r)) for r in rows]
        else:
            rows = [list(r) for r in ws.iter_rows(min_row=r1, max_row=r2, min_col=c1, max_col=c2, values_only=True)]
        if not address:   # 끝쪽 빈 행 제거
            while rows and all(v is None for v in rows[-1]):
                rows.pop()
        return ws.title, r1, c1, rows, names
    finally:
        wb.close()


def _search_matrix(matrix: List[List[Any]], r0: int, c0: int, text: str, limit: int = 50) -> Tuple[List[Dict[str, Any]], int]:
    hits: List[Dict[str, Any]] = []
    total = 0
    key = text.lower()
    for i, row in enumerate(matrix):
        for j, v in enumerate(row):
            if v is None:
                continue
            if key in str(out_value(v)).lower():
                total += 1
                if len(hits) < limit:
                    hits.append({"cell": addr(r0 + i, c0 + j), "value": out_value(v)})
    return hits, total


def cmd_read(args) -> Dict[str, Any]:
    csv_path = closed_csv(args.target)
    closed = None if csv_path else is_closed_xlsx(args.target)
    if csv_path:
        # CSV도 같은 명령으로 (행 수 세기, 머리글 확인): 이전에는 확장자 오류가 나 LLM이 pandas 코드를 직접 짰음
        sheet, r0, c0, matrix, names = _read_matrix_csv(csv_path, args.range)
        source = {"app": "Excel", "name": os.path.basename(csv_path), "path": csv_path, "source": "file(csv)"}
    elif closed:
        sheet, r0, c0, matrix, names = _read_matrix_closed(closed, args.sheet, args.range)
        source = {"app": "Excel", "name": os.path.basename(closed), "path": closed, "source": "file(openpyxl)"}
    else:
        with off.open_target("excel", args.target, readonly=True) as t:
            ws = t.sheet(args.sheet)
            rng = target_range(t, ws, args.range)
            matrix = off.range_to_matrix(rng.Value)
            sheet, r0, c0 = str(rng.Worksheet.Name), int(rng.Row), int(rng.Column)
            source = t.resolved()
    nrows = len(matrix)
    ncols = max((len(r) for r in matrix), default=0)
    full_addr = addr(r0, c0, r0 + max(nrows, 1) - 1, c0 + max(ncols, 1) - 1)
    result: Dict[str, Any] = {"target_resolved": source, "sheet": sheet, "range": full_addr, "rows": nrows, "cols": ncols}

    if args.search:
        hits, total = _search_matrix(matrix, r0, c0, args.search)
        result.update(matches=hits, total_matches=total)
        preview = f"'{args.search}': {sheet}!{full_addr}에서 {total}개 셀 일치"
        return rt.ok(TOOL, "read", preview, "" if total else
                     "VERIFIED: 일치하는 셀이 없습니다. 검색어를 임의로 바꿔 재시도하지 말고 사용자에게 알리세요.", **result)

    values = [[out_value(v) for v in row] for row in matrix]
    if args.save:
        out = rt.resolve_output_path(None if args.save is True else args.save, f"{sheet}_{full_addr.replace(':', '-')}", ".csv")
        out = rt.write_csv_rows(out, values)
        result.update(output_file=out, next_commands=[
            rt.next_command("word.py", ["write", "--file", out], "저장한 표를 열린 워드 문서에 넣기"),
            rt.next_command("ppt.py", ["write", "--file", out], "저장한 표를 열린 PPT에 넣기")])
        return rt.ok(TOOL, "read", f"{sheet}!{full_addr} ({nrows}행 x {ncols}열)를 CSV로 저장: {out}", **result)

    shown, truncated = rt.clamp_rows(values, args.max_rows)
    result.update(values=shown, truncated=truncated)
    if truncated:
        next_r = r0 + len(shown)
        last_r = r0 + nrows - 1
        rest = f"{sheet}!" if " " not in sheet else f"'{sheet}'!"
        rest += addr(next_r, c0, last_r, c0 + ncols - 1)
        result["next_commands"] = [
            rt.next_command("excel.py", ["read", "--target", args.target, "--range", rest], "나머지 행 이어서 읽기"),
            rt.next_command("excel.py", ["read", "--target", args.target, "--range", f"{full_addr}", "--sheet", sheet, "--save"],
                            "전체를 CSV로 저장 (다른 문서에 넣을 때)")]
    preview = f"{sheet}!{full_addr}: {nrows}행 x {ncols}열" + (f" (앞 {len(shown)}행만 표시)" if truncated else "")
    return rt.ok(TOOL, "read", preview, "VERIFIED: 읽기 완료. 일부만 표시되었으면 next_commands로 이어서 읽으세요."
                 if truncated else "", **result)


XL_CHART_NAMES = {51: "column", 57: "bar", 4: "line", 65: "line_markers", 5: "pie", 1: "area",
                  -4169: "scatter", 75: "scatter_lines", 74: "scatter_lines_markers"}


def _chart_detail(obj: Any, used_box: Optional[Dict[str, float]], warnings: List[str]) -> Dict[str, Any]:
    """차트 종류·계열 수·범례·위치와, 데이터 셀을 가리는지."""
    ch = obj.Chart
    code = int(ch.ChartType)
    d: Dict[str, Any] = {"chart_type": XL_CHART_NAMES.get(code, str(code)), "chart_type_code": code,
                         "series": int(ch.SeriesCollection().Count), "legend": bool(ch.HasLegend), **off.shape_box(obj)}
    if used_box is not None and off.overlap_ratio(d, used_box) > 0:
        d["covers_data"] = True
        warnings.append(f"차트 '{obj.Name}'이(가) 데이터 셀을 가립니다")
    return d


def _sheet_summary(ws: Any, detail: bool = False) -> Dict[str, Any]:
    used = ws.UsedRange
    info: Dict[str, Any] = {"name": str(ws.Name), "used_range": str(used.Address).replace("$", ""),
                            "rows": int(used.Rows.Count), "cols": int(used.Columns.Count)}
    charts = []
    warnings: List[str] = []
    used_box = None
    if detail and int(used.Rows.Count) * int(used.Columns.Count) > 1:
        try:
            used_box = off.shape_box(used)
        except Exception:
            pass
    try:
        objs = ws.ChartObjects()
        for i in range(1, objs.Count + 1):
            ch = objs.Item(i)
            title = ""
            try:
                if ch.Chart.HasTitle:
                    title = str(ch.Chart.ChartTitle.Text)
            except Exception:
                pass
            item = {"name": str(ch.Name), "title": title}
            if detail:
                try:
                    item.update(_chart_detail(ch, used_box, warnings))
                except Exception as e:
                    rt.debug(f"차트 상세 생략: {e}")
            charts.append(item)
    except Exception:
        pass
    info["charts"] = charts
    if detail:
        info["warnings"] = warnings
    try:
        info["shapes"] = int(ws.Shapes.Count) - len(charts)
    except Exception:
        pass
    return info


def cmd_info(args) -> Dict[str, Any]:
    if args.all:
        books = [{"name": d["name"], "path": d["path"], "saved": d["saved"], "active": d["active"]}
                 for d in off.list_open("excel")]
        return rt.ok(TOOL, "info", f"열린 통합문서 {len(books)}개", workbooks=books)
    with off.open_target("excel", args.target, readonly=True) as t:
        wb = t.doc
        sheets = [_sheet_summary(wb.Worksheets(i), args.detail) for i in range(1, wb.Worksheets.Count + 1)]
        result: Dict[str, Any] = {"target_resolved": t.resolved(), "sheets": sheets,
                                  "saved": off.is_saved(wb)}
        if args.detail:
            result["layout_warnings"] = sum(len(s["warnings"]) for s in sheets)
        try:
            result["active_sheet"] = str(wb.ActiveSheet.Name)
            if t.source == "active":
                result["selection"] = str(t.app.Selection.Address).replace("$", "")
        except Exception:
            pass
    preview = f"{result['target_resolved']['name']}: 시트 {len(sheets)}개 (" + ", ".join(
        f"{s['name']} {s['used_range']}" + (f" 차트{len(s['charts'])}" if s["charts"] else "") for s in sheets[:8]) + ")"
    if args.detail:
        preview += f", 배치 경고 {result['layout_warnings']}건" if result["layout_warnings"] else ", 배치 경고 없음"
    return rt.ok(TOOL, "info", preview, **result)


# ==============================================================================
# 4. write
# ==============================================================================

def _last_used_row(ws: Any, col: int) -> int:
    """해당 열 기준 마지막 사용 행 (없으면 0)."""
    try:
        cell = ws.Cells(ws.Rows.Count, col).End(-4162)   # xlUp
        if cell.Row == 1 and ws.Cells(1, col).Value is None:
            return 0
        return int(cell.Row)
    except Exception:
        used = ws.UsedRange
        return int(used.Row + used.Rows.Count - 1)


def _write_sheet(t: Any, name: Optional[str], new: bool) -> Any:
    """쓸 시트. 없는 이름은 --new-sheet가 있을 때만 새로 만듭니다 (오타로 새 시트가 생기지 않도록)."""
    if new:
        if not name:
            raise rt.InputError("--new-sheet에는 새 시트 이름이 필요합니다. --sheet 이름 또는 --range '이름'!A1로 지정하세요.")
        names = [str(t.doc.Worksheets(i).Name) for i in range(1, t.doc.Worksheets.Count + 1)]
        existing = rt.find_name(name, names)
        if existing is not None:
            raise rt.InputError(f"시트 '{existing}'이(가) 이미 있어 새로 만들지 않았습니다. 기존 시트에 쓰려면 --new-sheet를 빼고, "
                                "새 시트는 다른 이름을 지정하세요.")
    return t.sheet(name, create=new)


def cmd_write(args) -> Dict[str, Any]:
    sources = [x for x in (args.value is not None, args.values, args.file, args.formula) if x]
    if not sources:
        raise rt.InputError("쓸 내용이 없습니다. --value, --values, --file, --formula 중 하나를 지정하세요.")
    if len(sources) > 1:
        raise rt.InputError("--value, --values, --file, --formula는 하나만 지정하세요 (--header는 함께 사용 가능).")
    table: Optional[List[List[Any]]] = None
    if args.values:
        table = parse_values_arg(args.values)
    elif args.file:
        table = load_table_file(rt.resolve_input_path(args.file, "입력 표 파일"))
        if not table:
            raise rt.InputError(f"입력 파일에 데이터가 없습니다: {args.file}")
    header = [h.strip() for h in args.header.split(",")] if args.header else None

    with write_session(args, "write") as t:
        start = args.range or "A1"
        sh, cells = split_sheet(start)
        ws = _write_sheet(t, sh or args.sheet, args.new_sheet)
        parts = cells.replace("$", "").split(":")
        r1, c1 = parse_cell(parts[0])
        r2c2 = parse_cell(parts[1]) if len(parts) > 1 else None

        if args.append:
            width = max(len(header) if header else 0, len(rect(table)[0]) if table else 1)
            # 쓸 열 전체에서 가장 아래 행 (한 열만 보면 그 열이 비어 있는 행에 다른 열 데이터를 덮어씀)
            last = max(_last_used_row(ws, c) for c in range(c1, c1 + width))
            if last >= r1:
                # 같은 머리글은 반복하지 않음: 표의 머리글 행(시작 행)과 비교 (--header도 같음)
                existing = [str(out_value(v)) for v in off.range_to_matrix(
                    ws.Range(ws.Cells(r1, c1), ws.Cells(r1, c1 + width - 1)).Value)[0]]
                if header and existing[:len(header)] == [str(h) for h in header]:
                    header = None
                if table and existing[:len(table[0])] == [str(out_value(x)) for x in table[0]]:
                    table = table[1:]
                    if not table:
                        raise rt.InputError("추가할 데이터 행이 없습니다 (입력 파일에 헤더만 있음).")
                if r2c2:
                    r2c2 = (r2c2[0] - r1 + last + 1, r2c2[1])
                r1 = last + 1

        # 1) plan every block and check it BEFORE writing anything (a failed check leaves the sheet untouched)
        ops: List[Tuple[int, int, int, int, str, Any]] = []      # (r1, c1, r2, c2, kind, payload)
        top = r1
        if header:
            ops.append((r1, c1, r1, c1 + len(header) - 1, "values", (tuple(xl_text(h) for h in header),)))
            r1 += 1
        if args.formula:
            formulas = [f if f.startswith("=") else "=" + f for f in args.formula]
            if all(_is_number(f[1:]) for f in formulas):
                nums = [float(f[1:].replace(",", "")) for f in formulas]
                one_col = r2c2 is not None and r2c2[1] == c1
                rows = [[x] for x in nums] if one_col else [nums]
                rec = f"--value {formulas[0][1:]}" if len(nums) == 1 else f"--values '{json.dumps(rows)}'"
                raise rt.InputError(f"--formula에 숫자만 있습니다. 값은 {rec.split()[0]}로 쓰세요. [추천] {rec}")
            if r2c2 is None:
                if len(formulas) != 1:
                    rec = addr(top, c1, top + 100 if header else top + 99, c1 + len(formulas) - 1)
                    raise rt.InputError("수식이 여러 개면 --range로 범위를 지정하세요"
                                        + (f" (머리글이 있으면 머리글 행부터, 수식은 {top + 1}행 기준으로)" if header else "")
                                        + f". [추천] --range {rec}")
                r2, c2 = r1, c1
            else:
                r2, c2 = r2c2
            if r2 < r1:
                raise rt.InputError("--range 범위가 헤더 행만 포함합니다. 데이터 행까지 포함하세요.")
            ncols = c2 - c1 + 1
            if len(formulas) == 1:
                ops.append((r1, c1, r2, c2, "formula", formulas[0]))          # 상대참조 자동 조정
            elif len(formulas) == ncols:
                ops += [(r1, c1 + j, r2, c1 + j, "formula", f) for j, f in enumerate(formulas)]
            else:
                raise rt.InputError(f"수식 {len(formulas)}개와 범위의 열 수 {ncols}가 다릅니다. 열마다 수식 1개씩 주세요. "
                                    f"[추천] --range {addr(top, c1, r2, c1 + len(formulas) - 1)}")
        elif args.value is not None:
            v = xl_text(coerce(args.value))
            ops.append((r1, c1, *(r2c2 or (r1, c1)), "value", v))
        else:
            data = rect(table)
            nrows, ncols = len(data), len(data[0])
            ops.append((r1, c1, r1 + nrows - 1, c1 + ncols - 1, "values", tuple(tuple(xl_text(v) for v in r) for r in data)))

        # 2) never overwrite existing values without the user's consent (--append writes below the data)
        if not args.overwrite and not args.append:
            filled = [fc for (a, b, c, d, _, _) in ops for fc in _filled_cells(ws, ws.Range(ws.Cells(a, b), ws.Cells(c, d)))]
            if filled:
                raise rt.UserDecisionError(
                    f"{ws.Name}!{addr(ops[0][0], ops[0][1], ops[-1][2], max(o[3] for o in ops))}에 값이 있는 칸 {len(filled)}개가 덮어써집니다: "
                    + ", ".join(f"{c}={_short(v)}" for c, v in filled[:5]) + (" ..." if len(filled) > 5 else ""),
                    "쓰지 않았습니다. 덮어쓸 칸을 사용자에게 보여 주고, 사용자가 덮어써도 된다고 하면 --overwrite를 붙여 다시 실행하세요. "
                    "아니면 비어 있는 다른 --range를 쓰거나 --append로 아래에 이어 쓰세요.")

        # 3) write
        written: List[str] = []
        with off.excel_batch(t.app):
            for a, b, c, d, kind, payload in ops:
                cells_rng = ws.Range(ws.Cells(a, b), ws.Cells(c, d))
                if kind == "formula":
                    cells_rng.Formula = payload
                else:
                    cells_rng.Value = payload
                written.append(addr(a, b, c, d))

        problems = _verify_written(ws, ops)
        first = parse_cell(written[0].split(":")[0])
        last_ref = written[-1].split(":")[-1]
        last = parse_cell(last_ref)
        full = addr(first[0], first[1], last[0], last[1])
        if not args.no_autofit and (table is not None or header):
            try:
                ws.Range(ws.Cells(first[0], first[1]), ws.Cells(last[0], last[1])).Columns.AutoFit()
            except Exception:
                pass
        t.set_focus(sheet=ws)
        sheet_name = str(ws.Name)
        resolved = t.resolved()
        if args.target.lower() != "active" and t.source in ("file", "new"):
            tgt = t.path
        else:
            tgt = args.target

    sheet_ref = sheet_name if " " not in sheet_name else f"'{sheet_name}'"
    result = {"target_resolved": resolved, "sheet": sheet_name, "range_written": full,
              "rows": last[0] - first[0] + 1, "cols": last[1] - first[1] + 1}
    if table is not None or args.formula:
        result["next_commands"] = [rt.next_command("excel.py", ["chart", "--target", tgt, "--range", f"{sheet_ref}!{full}",
                                                                "--type", "line", "--title", "차트"],
                                                   "방금 쓴 범위로 차트 만들기 (--type, --title은 원하는 값으로)")]
    preview = f"{sheet_name}!{full}에 썼습니다 ({result['rows']}행 x {result['cols']}열)"
    if problems:
        return rt.ok_unverified(TOOL, "write", preview, problems, **result)
    return rt.ok(TOOL, "write", preview, "VERIFIED: 쓴 칸을 다시 읽어 확인했습니다. 같은 명령을 다시 실행하지 마세요.", verified=True, **result)


# ==============================================================================
# 5. replace
# ==============================================================================

def _load_map(args) -> Dict[str, str]:
    if args.map:
        raw = args.map.strip()
        if not raw.startswith("{"):
            raw = rt.read_text_file(rt.resolve_input_path(raw, "치환 목록 파일"))
        try:
            data = json.loads(raw)
        except json.JSONDecodeError as e:
            raise rt.InputError(f"--map은 JSON 객체여야 합니다 (예: {{\"초안\": \"최종\"}}): {e}")
        if not isinstance(data, dict) or not data:
            raise rt.InputError("--map은 비어 있지 않은 JSON 객체여야 합니다.")
        return {str(k): ("" if v is None else str(v)) for k, v in data.items()}
    if args.find is None:
        raise rt.InputError("--find 또는 --map을 지정하세요.")
    return {args.find: args.replace if args.replace is not None else args.find}


def cmd_replace(args) -> Dict[str, Any]:
    mapping = _load_map(args)
    fmt = _format_spec(args)
    replaced: List[str] = []
    skipped_formulas: List[str] = []
    with write_session(args, "replace") as t:
        sheets = [t.doc.Worksheets(i) for i in range(1, t.doc.Worksheets.Count + 1)] if args.all_sheets else [t.sheet(args.sheet)]
        with off.excel_batch(t.app):
            for ws in sheets:
                rng = target_range(t, ws, args.range)
                values = off.range_to_matrix(rng.Value)
                formulas = off.range_to_matrix(rng.Formula)
                r0, c0 = int(rng.Row), int(rng.Column)
                for i, row in enumerate(values):
                    for j, v in enumerate(row):
                        if v is None:
                            continue
                        s = str(out_value(v))
                        new = None
                        for find, rep in mapping.items():
                            if (args.whole and s == find) or (not args.whole and find in s):
                                new = rep if args.whole else (new or s).replace(find, rep)
                        if new is None:
                            continue
                        cell_ref = f"{ws.Name}!{addr(r0 + i, c0 + j)}"
                        f = formulas[i][j]
                        cell = ws.Cells(r0 + i, c0 + j)
                        if isinstance(f, str) and f.startswith("="):
                            skipped_formulas.append(cell_ref)      # 수식은 보존, 서식만
                        elif new != s:
                            cell.Value = xl_text(coerce(new) if not isinstance(v, str) else new)
                        if fmt:
                            _apply_format(cell, fmt)
                        replaced.append(cell_ref)
                if replaced:
                    t.set_focus(sheet=ws)
        resolved = t.resolved()
    result = {"target_resolved": resolved, "changed_cells": len(replaced), "cells": replaced[:30]}
    if skipped_formulas:
        result["formula_cells_kept"] = skipped_formulas[:30]
    if not replaced:
        return rt.ok(TOOL, "replace", "일치하는 셀이 없습니다",
                     "VERIFIED: 일치하는 셀이 없습니다. 검색어를 임의로 바꿔 재시도하지 말고 사용자에게 알리세요.", **result)
    preview = f"{len(replaced)}개 셀 처리" + (f" (수식 셀 {len(skipped_formulas)}개는 내용 유지, 서식만 적용)" if skipped_formulas else "")
    return rt.ok(TOOL, "replace", preview, **result)


# ==============================================================================
# 6. style
# ==============================================================================

def _format_spec(args) -> Dict[str, Any]:
    spec: Dict[str, Any] = {}
    if getattr(args, "bg", None):
        spec["bg"] = off.parse_color(args.bg)
    if getattr(args, "color", None):
        spec["color"] = off.parse_color(args.color)
    for key in ("bold", "italic"):
        if getattr(args, key, None) is not None:
            spec[key] = getattr(args, key)
    if getattr(args, "size", None):
        spec["size"] = float(args.size)
    if getattr(args, "align", None):
        spec["align"] = XL_ALIGN[args.align]
    if getattr(args, "number_format", None):
        spec["number_format"] = args.number_format
    if getattr(args, "border", None):
        spec["border"] = args.border
    if getattr(args, "wrap", None):
        spec["wrap"] = True
    return spec


def _apply_format(rng: Any, spec: Dict[str, Any]) -> None:
    if "bg" in spec:
        rng.Interior.Color = spec["bg"]
    if "color" in spec:
        rng.Font.Color = spec["color"]
    if "bold" in spec:
        rng.Font.Bold = spec["bold"]
    if "italic" in spec:
        rng.Font.Italic = spec["italic"]
    if "size" in spec:
        rng.Font.Size = spec["size"]
    if "align" in spec:
        rng.HorizontalAlignment = spec["align"]
    if "number_format" in spec:
        rng.NumberFormat = spec["number_format"]
    if "wrap" in spec:
        rng.WrapText = True
    if "border" in spec:
        _apply_border(rng, spec["border"])


def _apply_border(rng: Any, weight: str, box: bool = False) -> None:
    if weight == "none":
        for b in (7, 8, 9, 10, 11, 12):
            rng.Borders(b).LineStyle = -4142   # xlNone
        return
    w = XL_WEIGHT[weight]
    for b in ((7, 8, 9, 10) if box else (7, 8, 9, 10, 11, 12)):   # 좌/상/하/우(+내부 세로/가로)
        try:
            border = rng.Borders(b)
            border.LineStyle = 1      # 실선 먼저, 두께는 나중 (순서가 바뀌면 두께가 초기화될 수 있음)
            border.Weight = w
        except Exception:
            pass


def _apply_theme(ws: Any, rng: Any, theme: Dict[str, Any], zebra: bool = True) -> Dict[str, Any]:
    """표 범위에 테마 적용: 헤더, 줄무늬, 숫자 서식, 테두리, 합계 행."""
    r0, c0 = int(rng.Row), int(rng.Column)
    nrows, ncols = int(rng.Rows.Count), int(rng.Columns.Count)
    values = off.range_to_matrix(rng.Value)
    header = ws.Range(ws.Cells(r0, c0), ws.Cells(r0, c0 + ncols - 1))
    header.Interior.Color = off.parse_color(theme["header_bg"])
    header.Font.Color = off.parse_color(theme["header_fg"])
    header.Font.Bold = True
    header.HorizontalAlignment = XL_ALIGN["center"]

    summary = nrows >= 3 and any(k in str(values[-1][0] or "") for k in ("합계", "총계", "소계", "평균", "Total"))
    data_last = r0 + nrows - (2 if summary else 1)
    stripes = 0
    if zebra and data_last > r0 + 1:
        color = off.parse_color(theme["zebra"])
        block: List[str] = []
        for r in range(r0 + 2, data_last + 1, 2):   # 여러 행을 한 번에 칠해 COM 호출 수 축소 (주소는 255자 이내)
            part = addr(r, c0, r, c0 + ncols - 1)
            if block and len(",".join(block + [part])) > 240:
                ws.Range(",".join(block)).Interior.Color = color
                block = []
            block.append(part)
            stripes += 1
        if block:
            ws.Range(",".join(block)).Interior.Color = color

    numeric_cols = []
    for j in range(ncols):
        col_vals = [values[i][j] for i in range(1, len(values)) if values[i][j] is not None]
        if col_vals and all(isinstance(v, (int, float)) and not isinstance(v, bool) for v in col_vals):
            numeric_cols.append(j)
            ws.Range(ws.Cells(r0 + 1, c0 + j), ws.Cells(r0 + nrows - 1, c0 + j)).NumberFormat = theme["number_format"]

    _apply_border(rng, "thin")
    header.Borders(9).Weight = XL_WEIGHT["medium"]
    if summary:
        srow = ws.Range(ws.Cells(r0 + nrows - 1, c0), ws.Cells(r0 + nrows - 1, c0 + ncols - 1))
        srow.Interior.Color = off.parse_color(theme["summary_bg"])
        srow.Font.Bold = True
        srow.Borders(9).LineStyle = -4119      # 이중선
    try:
        rng.Columns.AutoFit()
    except Exception:
        pass
    return {"theme": theme["name"], "zebra_rows": stripes, "number_columns": [col_letter(c0 + j) for j in numeric_cols],
            "summary_row": summary}


def _parse_highlight(rule: str) -> Tuple[str, str, int]:
    """'>=100:yellow' -> ('>=', '100', BGR). '=완료:green', '포함:지연:red'(텍스트 포함)."""
    if ":" not in rule:
        raise rt.InputError(f"--highlight 형식: '조건:색' (예: '>=100:yellow', '=완료:lightgreen', '포함:지연:pink'). 입력: {rule}")
    cond, color = rule.rsplit(":", 1)
    cond = cond.strip()
    if cond.startswith("포함:"):
        return "contains", cond[3:], off.parse_color(color)
    for op in (">=", "<=", "<>", "!=", ">", "<", "="):
        if cond.startswith(op):
            return op, cond[len(op):].strip(), off.parse_color(color)
    return "contains", cond, off.parse_color(color)


def cmd_style(args) -> Dict[str, Any]:
    fmt = _format_spec(args)
    theme = off.get_theme(args.theme)
    if not fmt and not theme and not args.highlight and not args.clear:
        raise rt.InputError("적용할 서식이 없습니다. --theme, --bg, --color, --bold, --number-format, --border, --highlight 등을 지정하세요.")
    with write_session(args, "style") as t:
        ws = t.sheet(args.sheet)
        rng = target_range(t, ws, args.range)
        ws = rng.Worksheet
        applied: Dict[str, Any] = {}
        with off.excel_batch(t.app):
            if args.clear:
                rng.ClearFormats()
                rng.FormatConditions.Delete()
                applied["cleared"] = True
            if theme:
                applied.update(_apply_theme(ws, rng, theme, zebra=not args.no_zebra))
            if fmt:
                _apply_format(rng, fmt)
                applied["format"] = sorted(fmt)
            if args.highlight:
                op, val, color = _parse_highlight(args.highlight)
                if op == "contains":
                    top_left = str(rng.Cells(1, 1).Address).replace("$", "")
                    fc = rng.FormatConditions.Add(Type=2, Formula1=f'=ISNUMBER(SEARCH("{val}",{top_left}))')
                else:
                    is_num = bool(re.fullmatch(r"[+-]?\d+(\.\d+)?", val))
                    formula = val if is_num else f'="{val}"'
                    fc = rng.FormatConditions.Add(Type=1, Operator=XL_OPERATOR[op], Formula1=f"={formula}" if is_num else formula)
                fc.Interior.Color = color
                applied["highlight"] = args.highlight
        t.set_focus(sheet=ws)
        info = range_info(rng)
        resolved = t.resolved()
    return rt.ok(TOOL, "style", f"{info['sheet']}!{info['range']}에 서식 적용: {', '.join(map(str, applied))}",
                 target_resolved=resolved, applied=applied, **info)


# ==============================================================================
# 7. export
# ==============================================================================

def cmd_export(args) -> Dict[str, Any]:
    ext = os.path.splitext(args.save)[1].lower() if args.save else ".csv"
    if ext not in (".csv", ".xlsx", ".pdf"):
        raise rt.InputError(f"--save 확장자는 .csv, .xlsx, .pdf 중 하나여야 합니다: {args.save}")
    closed = is_closed_xlsx(args.target)
    if ext == ".csv" and closed and not args.columns:
        sheet, r0, c0, matrix, _ = _read_matrix_closed(closed, args.sheet, args.range)
        out = rt.resolve_output_path(args.save, sheet, ".csv")
        out = rt.write_csv_rows(out, [[out_value(v) for v in r] for r in matrix], args.encoding)
        return rt.ok(TOOL, "export", f"{sheet} ({len(matrix)}행)를 CSV로 저장: {out}", output_file=out, rows=len(matrix))

    with off.open_target("excel", args.target, readonly=True) as t:
        ws = t.sheet(args.sheet)
        rng = target_range(t, ws, args.range)
        ws = rng.Worksheet
        stem = f"{t.name.rsplit('.', 1)[0]}_{ws.Name}"
        out = rt.resolve_output_path(args.save, stem, ext)
        if ext == ".csv":
            matrix = off.range_to_matrix(rng.Value)
            if args.columns:
                c0 = int(rng.Column)
                idx = [col_number(c.strip()) - c0 for c in args.columns.split(",") if c.strip()]
                bad = [args.columns.split(",")[k] for k, i in enumerate(idx) if i < 0 or i >= len(matrix[0])]
                if bad:
                    raise rt.InputError(f"--columns의 열이 범위 밖입니다: {bad}")
                matrix = [[row[i] for i in idx] for row in matrix]
            out = rt.write_csv_rows(out, [[out_value(v) for v in r] for r in matrix], args.encoding)
            detail = f"{len(matrix)}행"
        elif ext == ".pdf":
            off.require_pdf_export(t.app, ".csv 또는 .xlsx")
            rng.ExportAsFixedFormat(0, os.path.normpath(out))   # xlTypePDF
            detail = "PDF"
        else:
            app = t.app
            prev_alerts = app.DisplayAlerts
            new_wb = None
            try:
                app.DisplayAlerts = False
                new_wb = app.Workbooks.Add()
                try:
                    new_wb.Windows(1).Visible = False
                except Exception:
                    pass
                dest = new_wb.Worksheets(1)
                rng.Copy(Destination=dest.Range("A1"))
                dest.Name = str(ws.Name)[:31]
                dest.Columns.AutoFit()
                try:
                    new_wb.Windows(1).Visible = True
                except Exception:
                    pass
                new_wb.SaveAs(os.path.normpath(out), FileFormat=51)
            finally:
                if new_wb is not None:
                    try:
                        new_wb.Close(SaveChanges=False)
                    except Exception:
                        pass
                app.DisplayAlerts = prev_alerts
            detail = "XLSX"
        info = range_info(rng)
        resolved = t.resolved()
    return rt.ok(TOOL, "export", f"{info['sheet']}!{info['range']}를 {detail}로 저장: {out}",
                 "VERIFIED: 저장 완료. 원본 통합문서는 그대로입니다.", target_resolved=resolved, output_file=out, **info)


# ==============================================================================
# 8. sheet / range
# ==============================================================================

def cmd_sheet(args) -> Dict[str, Any]:
    if args.action == "list":
        with off.open_target("excel", args.target, readonly=True) as t:
            names = [str(t.doc.Worksheets(i).Name) for i in range(1, t.doc.Worksheets.Count + 1)]
            active = str(t.doc.ActiveSheet.Name)
            resolved = t.resolved()
        return rt.ok(TOOL, "sheet", f"시트 {len(names)}개: {', '.join(names)}", target_resolved=resolved,
                     sheets=names, active_sheet=active)
    if not args.name:
        raise rt.InputError(f"sheet {args.action}에는 --name이 필요합니다.")
    with write_session(args, "sheet") as t:
        sheets = t.doc.Worksheets
        names = [str(sheets(i).Name) for i in range(1, sheets.Count + 1)]
        existing = rt.find_name(args.name, names)
        if args.action == "add":
            if existing:
                ws = sheets(existing)
                msg = f"시트 '{existing}'이(가) 이미 있어 그대로 사용합니다"
            else:
                ws = t.sheet(args.name, create=True)
                msg = f"시트 '{args.name}'을(를) 추가했습니다"
            t.set_focus(sheet=ws)
        else:
            if existing is None:
                raise rt.not_found("시트", args.name, names)
            ws = sheets(existing)
            if args.action == "select":
                t.set_focus(sheet=ws)
                msg = f"시트 '{existing}'을(를) 선택했습니다"
            elif args.action == "rename":
                if not args.new_name:
                    raise rt.InputError("rename에는 --new-name이 필요합니다.")
                if rt.find_name(args.new_name, names):
                    raise rt.InputError(f"'{args.new_name}' 이름의 시트가 이미 있습니다.")
                ws.Name = args.new_name
                t.set_focus(sheet=ws)
                msg = f"시트 '{existing}' -> '{args.new_name}'"
            elif args.action == "copy":
                new_name = args.new_name or f"{existing} 복사"
                if rt.find_name(new_name, names):
                    raise rt.InputError(f"'{new_name}' 이름의 시트가 이미 있습니다.")
                ws.Copy(After=sheets(sheets.Count))
                copied = sheets(sheets.Count)
                copied.Name = new_name
                t.set_focus(sheet=copied)
                msg = f"시트 '{existing}'을(를) '{new_name}'(으)로 복사했습니다"
            else:   # delete
                if sheets.Count == 1:
                    raise rt.InputError("마지막 남은 시트는 삭제할 수 없습니다.")
                prev = t.app.DisplayAlerts
                t.app.DisplayAlerts = False
                try:
                    ws.Delete()
                finally:
                    t.app.DisplayAlerts = prev
                msg = f"시트 '{existing}'을(를) 삭제했습니다 (되돌리기: office.py undo)"
        names = [str(sheets(i).Name) for i in range(1, sheets.Count + 1)]
        resolved = t.resolved()
    return rt.ok(TOOL, "sheet", msg, target_resolved=resolved, sheets=names)


def _same_cell(want: Any, got: Any) -> bool:
    """쓴 값과 다시 읽은 값이 같은지. 앞의 '(글자로 두기)는 값에 남지 않고, Excel이 날짜 등으로 바꾼 글자는 비어 있지만 않으면 인정."""
    if isinstance(want, str) and want.startswith("'"):
        want = want[1:]
    if want is None or want == "":
        return got is None or got == ""
    if isinstance(want, (int, float)) and not isinstance(want, bool):
        try:
            return abs(float(got) - float(want)) <= 1e-9 * max(1.0, abs(float(want)))
        except (TypeError, ValueError):
            return False
    if isinstance(got, str):
        return got == str(want)
    return got is not None and got != ""


def _verify_written(ws: Any, ops: List[Tuple[int, int, int, int, str, Any]]) -> List[str]:
    """R1: 쓴 칸을 다시 읽어 비교합니다. 수식은 첫 칸의 수식(대소문자·공백 무시), 값은 칸마다."""
    problems: List[str] = []
    norm = lambda f: re.sub(r"\s+", "", str(f)).upper()  # noqa: E731
    for a, b, c, d, kind, payload in ops:
        try:
            if kind == "formula":
                got = ws.Cells(a, b).Formula
                if norm(got) != norm(payload):
                    problems.append(f"{addr(a, b)} 수식이 {got!r} (쓴 수식 {payload!r})")
                continue
            got_rows = off.range_to_matrix(ws.Range(ws.Cells(a, b), ws.Cells(c, d)).Value)
            want_rows = payload if kind == "values" else [[payload] * (d - b + 1) for _ in range(c - a + 1)]
            bad = [addr(a + i, b + j) for i, row in enumerate(want_rows) for j, w in enumerate(row)
                   if not _same_cell(w, got_rows[i][j] if i < len(got_rows) and j < len(got_rows[i]) else None)]
            if bad:
                problems.append(f"다시 읽은 값이 다른 칸 {len(bad)}개: {', '.join(bad[:5])}")
        except Exception as e:
            problems.append(f"{addr(a, b, c, d)}를 다시 읽지 못함: {type(e).__name__}")
    return problems


def _is_number(text: str) -> bool:
    try:
        float(text.strip().replace(",", ""))
        return True
    except ValueError:
        return False


def _filled_cells(ws: Any, rng: Any) -> List[Tuple[str, Any]]:
    """범위에서 값이 있는 칸 [(주소, 값)] (문서 순서)."""
    r0, c0 = int(rng.Row), int(rng.Column)
    out = []
    for i, row in enumerate(off.range_to_matrix(rng.Value)):
        for j, v in enumerate(row):
            if v is not None and str(v) != "":
                out.append((addr(r0 + i, c0 + j), v))
    return out


def _axis(args) -> str:
    """--axis 생략 시 주소로 판단: '5:7', '5'처럼 행 번호만 있으면 rows (cols로 하면 시트의 모든 열이 숨겨짐)."""
    if args.axis:
        return args.axis
    cells = split_sheet(args.range or "")[1].replace("$", "")
    return "rows" if re.fullmatch(r"\d+(?::\d+)?", cells) else "cols"


def _inside(cell: Tuple[int, int], box: Tuple[int, int, int, int]) -> bool:
    return box[0] <= cell[0] <= box[2] and box[1] <= cell[1] <= box[3]


def _short(v: Any, n: int = 15) -> str:
    s = str(out_value(v)).replace(",", " ")
    return s if len(s) <= n else s[:n] + "..."


def cmd_range(args) -> Dict[str, Any]:
    a = args.action
    if a in ("merge", "unmerge", "width", "height", "hide", "unhide", "copy", "move", "clear") and not args.range:
        raise rt.InputError(f"range {a}에는 --range가 필요합니다." + (
            " 시트 전체를 지우려면 사용자에게 확인한 뒤 excel.py info로 사용 범위를 보고 그 범위를 --range로 지정하세요."
            if a == "clear" else ""))
    with write_session(args, "range") as t:
        ws = t.sheet(args.sheet)
        rng = target_range(t, ws, args.range) if args.range or a not in ("freeze", "unfreeze") else None
        if rng is not None:
            ws = rng.Worksheet
        detail = ""
        if a == "autofit":
            rng.Columns.AutoFit()
            if _axis(args) == "rows":
                rng.Rows.AutoFit()
        elif a in ("width", "height"):
            if args.size is None:
                raise rt.InputError(f"range {a}에는 --size가 필요합니다 (예: --size 20).")
            if a == "width":
                rng.ColumnWidth = float(args.size)
            else:
                rng.RowHeight = float(args.size)
        elif a in ("merge", "unmerge"):
            if a == "merge" and not args.overwrite:
                filled = _filled_cells(ws, rng)
                if len(filled) > 1:          # 병합하면 왼쪽 위 값만 남음
                    raise rt.UserDecisionError(
                        f"병합하면 왼쪽 위 칸만 남고 값 {len(filled) - 1}개가 지워집니다: "
                        + ", ".join(f"{c}={_short(v)}" for c, v in filled[1:6]) + (" ..." if len(filled) > 6 else ""),
                        "병합하지 않았습니다. 지워질 값을 사용자에게 보여 주고, 사용자가 지워도 된다고 하면 --overwrite를 붙여 다시 실행하세요.")
            prev = t.app.DisplayAlerts
            t.app.DisplayAlerts = False        # '왼쪽 위 값만 유지' 확인창 방지
            try:
                rng.Merge() if a == "merge" else rng.UnMerge()
            finally:
                t.app.DisplayAlerts = prev
        elif a in ("freeze", "unfreeze"):
            ws.Activate()
            win = t.app.ActiveWindow
            win.FreezePanes = False
            if a == "freeze":
                cell = args.range or "A2"           # 기본: 1행(머리글) 고정
                win.ScrollRow = 1
                win.ScrollColumn = 1
                ws.Range(split_sheet(cell)[1].split(":")[0]).Select()
                win.FreezePanes = True
                detail = f"{cell} 위/왼쪽 고정"
        elif a in ("hide", "unhide"):
            target = rng.EntireRow if _axis(args) == "rows" else rng.EntireColumn
            target.Hidden = a == "hide"
        elif a == "clear":
            rng = rng if rng is not None else ws.UsedRange
            {"contents": rng.ClearContents, "formats": rng.ClearFormats}.get(args.type, rng.Clear)()
            detail = args.type
        elif a in ("copy", "move"):
            if not args.dest:
                raise rt.InputError(f"range {a}에는 --dest가 필요합니다 (예: --dest Sheet2!A1).")
            dsheet, dcell = split_sheet(args.dest)
            dws = t.sheet(dsheet) if dsheet else ws
            dest = dws.Range(dcell)
            if not args.overwrite:
                dr, dc = parse_cell(dcell.replace("$", "").split(":")[0])
                nr, nc = int(rng.Rows.Count), int(rng.Columns.Count)
                block = dws.Range(dws.Cells(dr, dc), dws.Cells(dr + nr - 1, dc + nc - 1))
                src = (int(rng.Row), int(rng.Column), int(rng.Row) + nr - 1, int(rng.Column) + nc - 1)
                same_sheet = str(dws.Name) == str(ws.Name)
                filled = [(c, v) for c, v in _filled_cells(dws, block)
                          if not (same_sheet and _inside(parse_cell(c), src))]   # 원본과 겹치는 칸은 옮겨지는 값
                if filled:
                    raise rt.UserDecisionError(
                        f"대상 {dws.Name}!{addr(dr, dc, dr + nr - 1, dc + nc - 1)}에 값이 있는 칸 {len(filled)}개가 덮어써집니다: "
                        + ", ".join(f"{c}={_short(v)}" for c, v in filled[:5]) + (" ..." if len(filled) > 5 else ""),
                        "옮기지 않았습니다. 덮어쓸 칸을 사용자에게 보여 주고, 사용자가 덮어써도 된다고 하면 --overwrite를 붙여 "
                        "다시 실행하세요. 아니면 비어 있는 다른 --dest를 쓰세요.")
            if a == "copy":
                rng.Copy(Destination=dest)
            else:
                rng.Cut(Destination=dest)           # 겹쳐도 데이터 손실 없음, 수식 참조 갱신
            ws = dws
            detail = f"-> {dws.Name}!{dcell}"
        elif a == "border":
            _apply_border(rng if rng is not None else ws.UsedRange, args.weight, box=args.box)
            detail = args.weight + (" (외곽만)" if args.box else "")
        t.set_focus(sheet=ws)
        info = range_info(rng) if rng is not None else {"sheet": str(ws.Name)}
        resolved = t.resolved()
    where = f"{info['sheet']}!{info.get('range', '')}".rstrip("!")
    return rt.ok(TOOL, "range", f"{a} 완료: {where} {detail}".strip(), target_resolved=resolved, **info)


# ==============================================================================
# 9. chart
# ==============================================================================

def _chart_objects(ws: Any) -> List[Any]:
    objs = ws.ChartObjects()
    return [objs.Item(i) for i in range(1, objs.Count + 1)]


def _chart_title(obj: Any) -> str:
    try:
        return str(obj.Chart.ChartTitle.Text) if obj.Chart.HasTitle else ""
    except Exception:
        return ""


def find_chart(t: off.OfficeTarget, name: Optional[str], sheet: Optional[str] = None) -> Tuple[Any, Any]:
    """차트 이름/제목(대소문자 무시) 또는 번호로 차트를 찾습니다. 시트를 생략하면 모든 시트에서 찾습니다. (시트, 차트개체)"""
    wb = t.doc
    sheets = [t.sheet(sheet)] if sheet else [wb.Worksheets(i) for i in range(1, wb.Worksheets.Count + 1)]
    all_names: List[str] = []
    for ws in sheets:
        for obj in _chart_objects(ws):
            labels = [str(obj.Name), _chart_title(obj)]
            all_names += [x for x in labels if x]
            if name is None or any(x.strip().lower() == str(name).strip().lower() for x in labels if x):
                return ws, obj
    if name is not None and str(name).isdigit():
        ws = sheets[0]
        objs = _chart_objects(ws)
        if 1 <= int(name) <= len(objs):
            return ws, objs[int(name) - 1]
    if not all_names:
        raise rt.InputError("차트가 없습니다. excel.py chart --range ... 로 먼저 만드세요.")
    raise rt.not_found("차트", str(name), sorted(set(all_names)))


def cmd_chart(args) -> Dict[str, Any]:
    if args.action == "list":
        with off.open_target("excel", args.target, readonly=True) as t:
            wb = t.doc
            charts = [{"sheet": str(wb.Worksheets(i).Name), "name": str(o.Name), "title": _chart_title(o)}
                      for i in range(1, wb.Worksheets.Count + 1) for o in _chart_objects(wb.Worksheets(i))]
            resolved = t.resolved()
        return rt.ok(TOOL, "chart", f"차트 {len(charts)}개", target_resolved=resolved, charts=charts)

    if args.action == "delete":
        if not args.name and not args.all:
            raise rt.InputError("삭제할 차트를 --name(이름 또는 제목)으로 지정하거나, 시트의 모든 차트를 지우려면 --all을 쓰세요.")
        with write_session(args, "chart") as t:
            if args.all:
                ws = t.sheet(args.sheet)
                n = len(_chart_objects(ws))
                if n:
                    ws.ChartObjects().Delete()
                msg = f"시트 '{ws.Name}'의 차트 {n}개를 삭제했습니다"
            else:
                ws, obj = find_chart(t, args.name, args.sheet)
                msg = f"차트 '{obj.Name}'을(를) 삭제했습니다"
                obj.Delete()
            t.set_focus(sheet=ws)
            resolved = t.resolved()
        return rt.ok(TOOL, "chart", msg, target_resolved=resolved)

    # create
    theme = off.get_theme(args.theme)
    with write_session(args, "chart") as t:
        ws = t.sheet(args.sheet)
        rng = target_range(t, ws, args.range)
        ws = rng.Worksheet
        if args.range is None and int(rng.Rows.Count) * int(rng.Columns.Count) <= 1:
            raise rt.InputError("차트로 만들 데이터가 없습니다. --range로 데이터 범위를 지정하세요.")
        values = off.range_to_matrix(rng.Value)
        chart_type = args.type
        # 첫 열이 숫자(x값)인 꺾은선은 분산형(선)으로 그려야 x축이 숫자 축이 됨 (예: 사인 그래프)
        if chart_type == "line" and len(values) > 2 and len(values[0]) >= 2:
            first_col = [r[0] for r in values[1:] if r[0] is not None]
            if first_col and all(isinstance(v, (int, float)) and not isinstance(v, bool) for v in first_col):
                chart_type = "scatter_lines"
        if args.at:
            anchor = ws.Range(args.at)
        else:
            # 기본 위치: 데이터 범위 오른쪽, 단 시트의 다른 데이터도 가리지 않도록 사용 범위 오른쪽 끝 다음 열
            last_col = int(rng.Column) + int(rng.Columns.Count) - 1
            try:
                used = ws.UsedRange
                last_col = max(last_col, int(used.Column) + int(used.Columns.Count) - 1)
            except Exception:
                pass
            anchor = ws.Cells(int(rng.Row), last_col + 2)
        obj = ws.ChartObjects().Add(anchor.Left, anchor.Top, args.width, args.height)
        chart = obj.Chart
        # 종류 -> 데이터 -> 종류: 데이터 앞에 정해야 분산형이 첫 열을 x값으로 쓰고(뒤에만 정하면 x가 계열이 됨),
        # 데이터 뒤에 다시 정해야 세부 종류가 유지됨 (Excel 2007은 SetSourceData 때 분산형을 선+표식으로 바꿈)
        chart.ChartType = XL_CHART[chart_type]
        chart.SetSourceData(rng)
        chart.ChartType = XL_CHART[chart_type]
        if args.title:
            chart.HasTitle = True
            chart.ChartTitle.Text = args.title
        series = chart.SeriesCollection()
        if chart_type == "combo" and series.Count >= 2:
            series(1).ChartType = 51
            series(2).ChartType = 65
            series(2).AxisGroup = 2
        if args.labels:
            chart.ApplyDataLabels()
        if args.legend == "none" or (args.legend is None and series.Count == 1):
            chart.HasLegend = False
        else:
            chart.HasLegend = True
            chart.Legend.Position = XL_LEGEND[args.legend or "bottom"]
        if theme:
            colors = [off.parse_color(theme[k]) for k in ("primary", "secondary", "accent")]
            for i in range(1, min(series.Count, 3) + 1):
                s = series(i)
                if chart_type in ("line", "line_markers", "scatter_lines", "scatter") or (chart_type == "combo" and i == 2):
                    s.Format.Line.ForeColor.RGB = colors[i - 1]
                else:
                    s.Format.Fill.ForeColor.RGB = colors[i - 1]
        existing = [str(o.Name) for o in _chart_objects(ws) if o.Name != obj.Name]
        wanted = (args.name or args.title or "").strip()
        if wanted:
            final = wanted
            k = 2
            while rt.find_name(final, existing):
                final = f"{wanted} {k}"
                k += 1
            obj.Name = final
        chart_name = str(obj.Name)
        t.set_focus(sheet=ws)
        info = range_info(rng)
        resolved = t.resolved()
        tgt = t.path if t.source in ("file", "new") else args.target
    return rt.ok(TOOL, "chart", f"차트 '{chart_name}' 생성 ({chart_type}, 데이터 {info['sheet']}!{info['range']})",
                 "VERIFIED: 차트 생성 완료. 워드/PPT로 옮기려면 next_commands를 그대로 실행하세요.",
                 target_resolved=resolved, chart_name=chart_name, chart_type=chart_type, data_range=info,
                 next_commands=[rt.next_command("office.py", ["transfer", "--from", "excel", "--from-target", tgt,
                                                              "--chart", chart_name, "--to", "word,ppt"],
                                                "이 차트를 열린 워드와 PPT에 그림으로 넣기 (한쪽만이면 --to word 또는 --to ppt)")])


# ==============================================================================
# 10. 진입점
# ==============================================================================

def build_cli():
    parser, sub = rt.build_parser(TOOL, "엑셀 작업")

    def common(p, write: bool = False):
        p.add_argument("--target", "-t", default="active", help="active(기본) | 열린 통합문서 이름 | 파일 경로")
        p.add_argument("--sheet", "-s", help="시트 이름 (생략 시 활성 시트)")
        if write:
            p.add_argument("--no-focus", action="store_true", help="작업 후 화면을 고친 시트로 옮기지 않음")

    def fmt_args(p):
        p.add_argument("--bg", help="배경색 (yellow, #FFF2CC, 노랑 등)")
        p.add_argument("--color", help="글자색")
        p.add_argument("--bold", action="store_true", default=None, help="굵게")
        p.add_argument("--italic", action="store_true", default=None, help="기울임")
        p.add_argument("--size", type=float, help="글자 크기(pt)")
        p.add_argument("--align", choices=list(XL_ALIGN), help="가로 정렬")
        p.add_argument("--number-format", help="숫자 서식 (예: #,##0  0.0%  yyyy-mm-dd)")
        p.add_argument("--border", choices=list(XL_WEIGHT) + ["none"], help="테두리 (격자)")
        p.add_argument("--wrap", action="store_true", default=None, help="자동 줄바꿈")

    p = sub.add_parser("read", help="내용 읽기")
    common(p)
    p.add_argument("--range", "-r", help="읽을 범위 (생략 시 사용 범위 전체)")
    p.add_argument("--search", help="이 글자가 들어간 셀 찾기")
    p.add_argument("--max-rows", type=int, default=rt.DEFAULT_MAX_ROWS, help=f"표시할 최대 행 수 (기본 {rt.DEFAULT_MAX_ROWS})")
    p.add_argument("--save", nargs="?", const=True, help="CSV로 저장 (경로 생략 시 작업 폴더)")

    p = sub.add_parser("info", help="통합문서 요약")
    common(p)
    p.add_argument("--all", action="store_true", help="열린 통합문서 전체 목록")
    p.add_argument("--detail", action="store_true", help="차트 종류/계열 수/위치와 배치 경고(데이터를 가리는 차트)")

    p = sub.add_parser("write", help="값/표/수식 쓰기")
    common(p, write=True)
    p.add_argument("--range", "-r", help="시작 셀 또는 범위 (기본 A1). 수식은 범위 전체에 채움")
    p.add_argument("--value", help="값 하나 (범위를 주면 범위 전체에)")
    p.add_argument("--values", help='JSON 2차원 배열 (예: [["이름","값"],["가",1]])')
    p.add_argument("--file", "-f", help="표 파일 (.csv, .md 표, .txt)")
    p.add_argument("--formula", nargs="+", help="수식. 1개면 범위 전체, 여러 개면 열마다 (상대참조 자동 조정)")
    p.add_argument("--header", help="첫 행에 쓸 머리글 (쉼표 구분). 데이터/수식은 그 아래부터")
    p.add_argument("--append", action="store_true", help="기존 데이터 아래에 이어 쓰기 (같은 머리글은 생략)")
    p.add_argument("--overwrite", action="store_true", help="값이 있는 칸을 덮어쓰기 (사용자가 확인했을 때만)")
    p.add_argument("--new-sheet", action="store_true",
                   help="--sheet(또는 --range '이름'!A1)의 시트를 새로 만들어 쓰기. 없으면 없는 시트는 오류")
    p.add_argument("--no-autofit", action="store_true", help="열 너비 자동 맞춤 안 함")

    p = sub.add_parser("replace", help="찾아 바꾸기")
    common(p, write=True)
    p.add_argument("--find", help="찾을 글자")
    p.add_argument("--replace", help="바꿀 글자 (생략 시 서식만 적용)")
    p.add_argument("--map", help='여러 개 한 번에: JSON 객체 또는 JSON 파일 (예: {"초안":"최종"})')
    p.add_argument("--range", "-r", help="범위 (생략 시 사용 범위)")
    p.add_argument("--whole", action="store_true", help="셀 전체가 일치할 때만")
    p.add_argument("--all-sheets", action="store_true", help="모든 시트")
    fmt_args(p)

    p = sub.add_parser("style", help="서식")
    common(p, write=True)
    p.add_argument("--range", "-r", help="범위 (생략 시 사용 범위)")
    p.add_argument("--theme", help=f"표 테마: {', '.join(off.THEMES)}")
    p.add_argument("--no-zebra", action="store_true", help="테마 적용 시 줄무늬 없음")
    p.add_argument("--highlight", help="조건부 강조: '>=100:yellow', '=완료:lightgreen', '포함:지연:pink'")
    p.add_argument("--clear", action="store_true", help="기존 서식 지우기")
    fmt_args(p)

    p = sub.add_parser("export", help="CSV/XLSX/PDF로 저장")
    common(p)
    p.add_argument("--save", help="저장 경로 (.csv/.xlsx/.pdf, 생략 시 작업 폴더에 CSV)")
    p.add_argument("--range", "-r", help="범위 (생략 시 사용 범위)")
    p.add_argument("--columns", "-c", help="CSV로 뽑을 열만 (예: A,C,E)")
    p.add_argument("--encoding", default="utf-8-sig", choices=["utf-8-sig", "cp949", "utf-8"], help="CSV 인코딩")

    p = sub.add_parser("sheet", help="시트 관리")
    common(p, write=True)
    p.add_argument("--action", "-a", required=True, choices=["list", "add", "rename", "copy", "delete", "select"])
    p.add_argument("--name", "-n", help="대상 시트 이름")
    p.add_argument("--new-name", help="새 이름 (rename/copy)")

    p = sub.add_parser("range", help="범위 조작")
    common(p, write=True)
    p.add_argument("--action", "-a", required=True,
                   choices=["autofit", "width", "height", "merge", "unmerge", "freeze", "unfreeze",
                            "hide", "unhide", "clear", "copy", "move", "border"])
    p.add_argument("--range", "-r", help="대상 범위 (freeze는 기준 셀, 기본 A2)")
    p.add_argument("--dest", help="copy/move 목적지 (예: B2, 'Sheet2'!A1)")
    p.add_argument("--size", type=float, help="width/height 값")
    p.add_argument("--axis", choices=["rows", "cols"], help="hide/unhide/autofit 대상 (생략: '5:7'처럼 행 주소면 rows, 아니면 cols)")
    p.add_argument("--type", choices=["all", "contents", "formats"], default="all", help="clear 종류")
    p.add_argument("--weight", choices=list(XL_WEIGHT) + ["none"], default="thin", help="border 굵기")
    p.add_argument("--box", action="store_true", help="border: 외곽선만")
    p.add_argument("--overwrite", action="store_true",
                   help="merge: 왼쪽 위 밖의 값을 지우고 병합 / copy·move: 값이 있는 대상 칸을 덮어쓰기 (사용자가 원할 때만)")

    p = sub.add_parser("chart", help="차트")
    common(p, write=True)
    p.add_argument("--action", "-a", default="create", choices=["create", "list", "delete"])
    p.add_argument("--range", "-r", help="데이터 범위 (첫 행 머리글, 생략 시 사용 범위)")
    p.add_argument("--type", default="column", choices=list(XL_CHART), help="차트 종류 (기본 column)")
    p.add_argument("--title", help="차트 제목")
    p.add_argument("--name", help="차트 이름 (생략 시 제목). delete 대상 지정에도 사용")
    p.add_argument("--at", help="차트 왼쪽 위 셀 (생략 시 데이터 오른쪽)")
    p.add_argument("--width", type=float, default=480, help="너비(pt)")
    p.add_argument("--height", type=float, default=288, help="높이(pt)")
    p.add_argument("--labels", action="store_true", help="값 표시")
    p.add_argument("--legend", choices=list(XL_LEGEND) + ["none"], help="범례 위치 (계열 1개면 기본 없음)")
    p.add_argument("--theme", help=f"색 테마: {', '.join(off.THEMES)}")
    p.add_argument("--all", action="store_true", help="delete: 시트의 모든 차트")
    return parser


HANDLERS = {"read": cmd_read, "info": cmd_info, "write": cmd_write, "replace": cmd_replace, "style": cmd_style,
            "export": cmd_export, "sheet": cmd_sheet, "range": cmd_range, "chart": cmd_chart}


def main(argv: Optional[Sequence[str]] = None) -> None:
    args = build_cli().parse_args(argv)
    rt.run_cli(TOOL, args.command, lambda: HANDLERS[args.command](args))


if __name__ == "__main__":
    main()
