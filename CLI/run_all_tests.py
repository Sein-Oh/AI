"""CLI 도구 통합 테스트 (Windows 실기용, 사람이 실행)

사용법 (C:\\CLI 처럼 ml-kit, office-cli 폴더가 있는 곳에서):
    python run_all_tests.py                      # 환경 진단 + 오프라인 테스트 + 실기 테스트 전부
    python run_all_tests.py --auto               # office-cli 눈 확인 질문 없이 (자리 비울 때)
    python run_all_tests.py --only diag,pytest   # 일부 단계만 (diag, pytest, ml, office)
    python run_all_tests.py --label 회사PC --outlook

결과는 이 파일 옆의 test_result_<라벨>_<시각>.txt 한 개에 모두 모입니다.
    - 단계가 끝날 때마다 바로 기록하므로, 중간에 멈추거나 Ctrl+C로 끊어도 그때까지의 내용이 남습니다.
    - 이 파일 하나를 그대로 Claude에게 전달하세요 (편집하지 마세요).

단계:
    diag    환경 진단: Python, 라이브러리, 문서 보안(DRM) 영향 점검 (python-docx/pptx 기본 템플릿, 새로 쓴 파일)
    pytest  각 도구의 오프라인 테스트 (python -m pytest)
    ml      ml-kit 실기 테스트 (tools/run_win_tests.py, 질문 없음)
    office  office-cli 실기 테스트 (tools/run_win_tests.py, 화면에 질문이 나오면 답하세요. --auto면 질문 없음)
    끝에 각 도구의 오류 로그(<도구>/data/logs/errors.txt, 실제 사용 중 생긴 오류)도 함께 붙입니다.

이 스크립트는 사람이 실행하는 테스트 도구입니다. 에이전트용 출력 계약(stdout 1줄 JSON) 대상이 아닙니다.
"""
import argparse
import datetime as dt
import glob
import locale
import os
import platform
import shutil
import subprocess
import sys
import tempfile
import time
import traceback
import zipfile

HERE = os.path.dirname(os.path.abspath(__file__))
TOOLS = ("ml-kit", "office-cli")
STEPS = ("diag", "pytest", "ml", "office")
PACKAGES = ("numpy", "pandas", "scikit-learn", "scipy", "joblib", "openpyxl", "python-docx", "python-pptx",
            "matplotlib", "pymupdf", "olefile", "pywin32", "pytest")
IMPORTS = {"numpy": "numpy", "pandas": "pandas", "scikit-learn": "sklearn", "scipy": "scipy", "joblib": "joblib",
           "openpyxl": "openpyxl", "python-docx": "docx", "python-pptx": "pptx", "matplotlib": "matplotlib",
           "pymupdf": "fitz", "olefile": "olefile", "pywin32": "win32com.client", "pytest": "pytest"}


class Result:
    """결과 파일. 섹션마다 바로 덧붙여 저장합니다."""

    def __init__(self, path):
        self.path = path
        self.summary = []          # 단계별 한 줄 요약 (화면과 파일 끝에 출력)
        with open(path, "w", encoding="utf-8") as f:
            f.write(f"# CLI 통합 테스트 결과 (생성 {dt.datetime.now():%Y-%m-%d %H:%M:%S})\n")
            f.write("# 이 파일 전체를 그대로 Claude에게 전달하세요. 편집하지 마세요.\n")
            f.write("# 단계별 요약은 파일 맨 끝 <요약>에 있습니다.\n\n")

    def section(self, name, text):
        with open(self.path, "a", encoding="utf-8") as f:
            f.write(f"<{name}>\n{text.rstrip()}\n\n")


def say(msg):
    print(f"[STATUS] {msg}", flush=True)


def child_env():
    return dict(os.environ, PYTHONIOENCODING="utf-8", PYTHONUTF8="1")


def run_capture(args, cwd, timeout):
    """하위 프로세스를 실행해 (종료코드, 출력 전체, 걸린 초, Ctrl+C 여부)를 돌려줍니다. 출력은 화면에도 그대로 보여 줍니다."""
    start = time.time()
    lines = []
    interrupted = False
    proc = subprocess.Popen(args, cwd=cwd, env=child_env(), stdout=subprocess.PIPE, stderr=subprocess.STDOUT)
    try:
        for raw in proc.stdout:
            line = raw.decode("utf-8", errors="replace").rstrip("\r\n")
            print(line, flush=True)
            lines.append(line)
            if time.time() - start > timeout:
                proc.kill()
                lines.append(f"[ERROR] 시간 초과 ({timeout}초)로 중단했습니다.")
                break
        proc.wait()
    except KeyboardInterrupt:
        interrupted = True
        proc.kill()
        lines.append("[ERROR] 사용자가 중단했습니다 (Ctrl+C).")
    return proc.returncode, "\n".join(lines), round(time.time() - start, 1), interrupted


# ---------------------------------------------------------------------------
# 1. 환경 진단
# ---------------------------------------------------------------------------
def _version(dist):
    try:
        from importlib import metadata
        return metadata.version(dist)
    except Exception:
        return "없음"


def _zip_info(path):
    if not os.path.exists(path):
        return "파일 없음"
    with open(path, "rb") as f:
        head = f.read(16)
    return f"크기 {os.path.getsize(path)}바이트, zip 형식={zipfile.is_zipfile(path)}, 앞 16바이트={head.hex(' ')}"


def diag_docx(lines):
    try:
        import docx
    except Exception as e:
        lines.append(f"python-docx import 실패: {e}")
        return
    tdir = os.path.join(os.path.dirname(docx.__file__), "templates")
    default = os.path.join(tdir, "default.docx")
    lines.append(f"python-docx 기본 템플릿 {default}: {_zip_info(default)} (정상: 38116바이트, zip=True)")
    lines.append(f"python-docx 폴더 템플릿 default-docx-template 있음={os.path.isdir(os.path.join(tdir, 'default-docx-template'))}")
    try:
        d = docx.Document()
        names = [s.name for s in d.styles]
        lines.append(f"docx.Document(): 스타일 {len(names)}개 (정상 164), Caption 있음={'Caption' in names}, "
                     f"본문 단락 {len(d.paragraphs)}개 (정상 0), 앞 단락={[p.text[:60] for p in d.paragraphs][:3]}")
    except Exception as e:
        lines.append(f"docx.Document() 실패: {type(e).__name__}: {e}")


def diag_pptx(lines):
    try:
        import pptx
    except Exception as e:
        lines.append(f"python-pptx import 실패: {e}")
        return
    default = os.path.join(os.path.dirname(pptx.__file__), "templates", "default.pptx")
    lines.append(f"python-pptx 기본 템플릿 {default}: {_zip_info(default)} (정상: 34030바이트, zip=True, 50 4b 03 04로 시작)")
    try:
        prs = pptx.Presentation()
        lines.append(f"pptx.Presentation(): 정상 (레이아웃 {len(prs.slide_layouts)}개)")
    except Exception as e:
        lines.append(f"pptx.Presentation() 실패: {type(e).__name__}: {e}")


def diag_new_files(lines):
    """도구가 새로 쓴 Office 파일도 문서 보안(DRM)이 바꾸는지: 쓰고 잠시 뒤 다시 읽어 zip 형식인지 봅니다."""
    work = tempfile.mkdtemp(prefix="cli_diag_")
    made = []
    try:
        import openpyxl
        p = os.path.join(work, "진단.xlsx")
        wb = openpyxl.Workbook()
        wb.active["A1"] = "진단"
        wb.save(p)
        made.append(p)
    except Exception as e:
        lines.append(f"xlsx 쓰기 실패: {type(e).__name__}: {e}")
    try:
        import docx
        tdir = os.path.join(os.path.dirname(docx.__file__), "templates", "default-docx-template")
        d = docx.Document(tdir) if os.path.isdir(tdir) else docx.Document()
        d.add_paragraph("진단")
        p = os.path.join(work, "진단.docx")
        d.save(p)
        made.append(p)
    except Exception as e:
        lines.append(f"docx 쓰기 실패: {type(e).__name__}: {e}")
    p = os.path.join(work, "진단.txt")
    with open(p, "w", encoding="utf-8") as f:
        f.write("진단")
    made.append(p)
    time.sleep(3)                                  # 보안 프로그램이 파일을 바꿀 시간
    for p in made:
        lines.append(f"새로 쓴 {os.path.basename(p)}: {_zip_info(p) if not p.endswith('.txt') else f'크기 {os.path.getsize(p)}바이트 (정상 6)'}")
    try:
        import openpyxl
        xlsx = [p for p in made if p.endswith(".xlsx")]
        if xlsx:
            openpyxl.load_workbook(xlsx[0], read_only=True).close()
        lines.append("새로 쓴 xlsx 다시 열기: 정상")
    except Exception as e:
        lines.append(f"새로 쓴 xlsx 다시 열기 실패: {type(e).__name__}: {e}")
    shutil.rmtree(work, ignore_errors=True)


def diag_system(lines):
    """Office 버전, 문서 보안(DRM)으로 보이는 프로세스, 설치된 패키지 전체 (Windows에서만 일부)."""
    if os.name == "nt":
        try:
            import winreg
            key = winreg.OpenKey(winreg.HKEY_LOCAL_MACHINE, r"SOFTWARE\Microsoft\Office\ClickToRun\Configuration")
            for name in ("VersionToReport", "ProductReleaseIds", "Platform"):
                try:
                    lines.append(f"Office {name}: {winreg.QueryValueEx(key, name)[0]}")
                except OSError:
                    pass
        except OSError:
            lines.append("Office: ClickToRun 정보 없음 (MSI 설치이거나 Office 없음)")
        try:
            out = subprocess.run(["tasklist", "/FO", "CSV", "/NH"], capture_output=True, timeout=60).stdout
            names = sorted({row.split('","')[0].strip('"') for row in out.decode("mbcs", errors="replace").splitlines() if row})
            keys = ("drm", "fasoo", "softcamp", "markany", "nasca", "document", "secu", "guard", "protect", "dlp", "fsc", "scsk")
            hits = [n for n in names if any(k in n.lower() for k in keys)]
            lines.append(f"보안/DRM으로 보이는 프로세스 (이름에 drm, fasoo, softcamp, markany, secu, guard 등): {hits or '없음'}")
        except Exception as e:
            lines.append(f"프로세스 목록 실패: {e}")
    try:
        out = subprocess.run([sys.executable, "-m", "pip", "list", "--format=freeze", "--disable-pip-version-check"],
                             capture_output=True, timeout=120, env=child_env()).stdout.decode("utf-8", errors="replace")
        pkgs = [x for x in out.splitlines() if x.strip()]
        lines.append(f"설치된 패키지 {len(pkgs)}개: " + ", ".join(pkgs))
    except Exception as e:
        lines.append(f"pip list 실패: {e}")


def step_diag(res, args):
    lines = [
        f"OS: {platform.platform()} ({platform.version()})",
        f"Python: {sys.version.split()[0]} ({sys.executable})",
        f"인코딩: 파일시스템 {sys.getfilesystemencoding()}, 기본 {locale.getpreferredencoding(False)}",
        f"CLI 폴더: {HERE} (도구: {', '.join(t for t in TOOLS if os.path.isdir(os.path.join(HERE, t)))})",
        "",
        "[라이브러리 버전 / import]",
    ]
    for dist in PACKAGES:
        try:
            __import__(IMPORTS[dist])
            ok = "import 정상"
        except Exception as e:
            ok = f"import 실패: {type(e).__name__}: {e}"
        lines.append(f"  {dist} {_version(dist)} - {ok}")
    for tool in TOOLS:
        skill = os.path.join(HERE, tool, "SKILL.md")
        ver = "?"
        if os.path.exists(skill):
            for line in open(skill, encoding="utf-8"):
                if line.strip().startswith("version:"):
                    ver = line.split(":", 1)[1].strip().strip('"')
                    break
        lines.append(f"  도구 {tool} {ver}")
    lines += ["", "[문서 보안(DRM) 영향 점검]"]
    for fn in (diag_docx, diag_pptx, diag_new_files, diag_system):
        try:
            fn(lines)
        except Exception as e:
            lines.append(f"{fn.__name__} 오류: {type(e).__name__}: {e}")
    text = "\n".join(lines)
    print(text, flush=True)
    res.section("환경 진단", text)
    broken = [x for x in lines if ("zip 형식=False" in x or "실패" in x or "Caption 있음=False" in x)]
    res.summary.append(f"diag: 이상 {len(broken)}건" + (f" ({broken[0][:80]} ...)" if broken else " (없음)"))


# ---------------------------------------------------------------------------
# 2. 오프라인 테스트
# ---------------------------------------------------------------------------
def step_pytest(res, args):
    for tool in TOOLS:
        folder = os.path.join(HERE, tool)
        if not os.path.isdir(folder):
            continue
        say(f"{tool} 오프라인 테스트 시작 (몇 분 걸릴 수 있습니다)")
        code, out, sec, stop = run_capture([sys.executable, "-m", "pytest", "-q", "-p", "no:cacheprovider", "-rfE"],
                                           folder, args.timeout)
        res.section(f"pytest_{tool}", f"cwd={folder} exit={code} {sec}초\n{out}")
        tail = [x for x in out.splitlines() if x.strip()]
        res.summary.append(f"pytest {tool}: exit={code}, {tail[-1].strip('= ') if tail else '출력 없음'}")
        if stop:
            raise KeyboardInterrupt


# ---------------------------------------------------------------------------
# 3. 실기 테스트
# ---------------------------------------------------------------------------
def _newer(pattern, since):
    files = [p for p in glob.glob(pattern) if os.path.getmtime(p) >= since - 1]
    return sorted(files, key=os.path.getmtime)


def _grep(path, prefix):
    """파일에서 prefix로 시작하는 첫 줄."""
    try:
        for line in open(path, encoding="utf-8", errors="replace"):
            if line.strip().startswith(prefix):
                return line.strip()
    except Exception:
        pass
    return f"{prefix} (못 찾음)"


def _attach(res, name, path):
    try:
        with open(path, encoding="utf-8", errors="replace") as f:
            res.section(name, f"파일: {path}\n{f.read()}")
    except Exception as e:
        res.section(name, f"파일 {path} 읽기 실패: {e}")


def step_ml(res, args):
    folder = os.path.join(HERE, "ml-kit")
    start = time.time()
    cmd = [sys.executable, os.path.join("tools", "run_win_tests.py"), "--label", args.label]
    if args.ml_dir:
        cmd += ["--dir", args.ml_dir]
    say("ml-kit 실기 테스트 시작 (10~20분, 질문 없음). Office 창을 만지지 마세요.")
    code, out, sec, stop = run_capture(cmd, folder, args.timeout)
    res.section("ml-kit 실기 테스트 화면 출력", f"$ {' '.join(cmd)}\nexit={code} {sec}초\n{out[-20000:]}")
    logs = _newer(os.path.join(folder, "tools", "test_logs", "testlog_*.txt"), start)
    for p in logs:
        _attach(res, os.path.basename(p), p)
    res.summary.append(f"ml-kit 실기: exit={code}, " + (_grep(logs[-1], "결과:") if logs else "로그 파일 없음"))
    test_dir = args.ml_dir or r"C:\ml_test\ml 테스트"
    for p in _newer(os.path.join(test_dir, "report_*.txt"), start):
        _attach(res, os.path.basename(p), p)
    tool_log = os.path.join(test_dir, "_data", "logs", "errors.txt")
    if os.path.exists(tool_log) and os.path.getmtime(tool_log) >= start - 1:
        _attach(res, "ml-kit 실기 중 도구 오류 로그 (일부러 낸 오류 포함)", tool_log)
    if stop:
        raise KeyboardInterrupt


def step_office(res, args):
    folder = os.path.join(HERE, "office-cli")
    start = time.time()
    cmd = [sys.executable, os.path.join("tools", "run_win_tests.py")]
    if not args.outlook:
        cmd.append("--skip-outlook")
    if args.auto:
        cmd.append("--auto")
    if args.office_dir:
        cmd += ["--dir", args.office_dir]
    say("office-cli 실기 테스트 시작 (10~30분). 질문이 나오면 답하세요. 그 밖에는 Office 창을 만지지 마세요.")
    t0 = time.time()
    code, stop = None, False
    try:
        code = subprocess.call(cmd, cwd=folder, env=child_env())     # 질문에 답할 수 있게 화면을 그대로 씀
    except KeyboardInterrupt:
        code, stop = 130, True
    res.section("office-cli 실기 테스트 실행", f"$ {' '.join(cmd)}\nexit={code} {round(time.time() - t0, 1)}초"
                + (" (Ctrl+C로 중단)" if stop else ""))
    test_dir = args.office_dir or r"C:\office_test"
    errs = _newer(os.path.join(test_dir, "errors_*.txt"), start)
    for pat in ("errors_*.txt", "report_*.txt"):
        for p in _newer(os.path.join(test_dir, pat), start):
            _attach(res, os.path.basename(p), p)
    res.summary.append(f"office-cli 실기: exit={code}, " + (
        f"{_grep(errs[-1], '결과:')} / {_grep(errs[-1], '문제 항목:')}" if errs else "로그 파일 없음"))
    if stop:
        raise KeyboardInterrupt


def attach_tool_error_logs(res):
    """실제 사용 중 생긴 도구 오류 로그 (<도구>/data/logs/errors.txt). 실기 테스트용 로그와 별개입니다."""
    found = False
    for tool, env in (("ml-kit", "ML_KIT_DATA_DIR"), ("office-cli", "OFFICE_CLI_DATA_DIR")):
        data = os.environ.get(env) or os.path.join(HERE, tool, "data")
        for name in ("errors.old.txt", "errors.txt"):
            p = os.path.join(data, "logs", name)
            if os.path.exists(p):
                found = True
                _attach(res, f"{tool} 도구 오류 로그 {name}", p)
    if not found:
        res.section("도구 오류 로그", "없음 (<도구>/data/logs/errors.txt 파일이 없습니다)")


# ---------------------------------------------------------------------------
def main():
    p = argparse.ArgumentParser(description="CLI 도구 통합 테스트 (결과는 파일 1개)")
    p.add_argument("--only", help="실행할 단계 (쉼표로): diag,pytest,ml,office (기본: 전부)")
    p.add_argument("--label", default="PC", help="결과 파일 이름에 붙일 PC 라벨 (예: 회사PC)")
    p.add_argument("--auto", action="store_true", help="office-cli 눈 확인 질문 없이 자동 판정 항목만")
    p.add_argument("--outlook", action="store_true", help="office-cli Outlook(G그룹)도 시험 (기본: 건너뜀)")
    p.add_argument("--ml-dir", help=r"ml-kit 실기 테스트 폴더 (기본 C:\ml_test\ml 테스트)")
    p.add_argument("--office-dir", help=r"office-cli 실기 테스트 폴더 (기본 C:\office_test)")
    p.add_argument("--timeout", type=int, default=3600, help="단계별 시간 상한(초, 기본 3600)")
    args = p.parse_args()
    steps = [s.strip() for s in (args.only or ",".join(STEPS)).split(",") if s.strip()]
    bad = [s for s in steps if s not in STEPS]
    if bad:
        p.error(f"알 수 없는 단계: {', '.join(bad)} (가능: {', '.join(STEPS)})")

    label = "".join(c for c in args.label if c.isalnum() or c in "-_") or "PC"
    out = os.path.join(HERE, f"test_result_{label}_{dt.datetime.now():%Y%m%d_%H%M%S}.txt")
    res = Result(out)
    say(f"결과 파일: {out}")
    if any(s in steps for s in ("ml", "office")):
        print("[STATUS] 실기 테스트 중 Excel/Word/PowerPoint 창이 열리고 닫힙니다. "
              "열린 Office 문서는 모두 저장하고 닫은 뒤 시작하세요.", flush=True)
    runners = {"diag": step_diag, "pytest": step_pytest, "ml": step_ml, "office": step_office}
    try:
        for s in steps:
            t0 = time.time()
            try:
                runners[s](res, args)
            except KeyboardInterrupt:
                raise
            except Exception:
                res.section(f"{s} 단계 오류", traceback.format_exc())
                say(f"{s} 단계에서 오류가 났습니다. 다음 단계로 넘어갑니다.")
            say(f"{s} 단계 끝 ({round(time.time() - t0)}초)")
    except KeyboardInterrupt:
        res.section("중단", "사용자가 중단했습니다 (Ctrl+C). 위는 그때까지의 결과입니다.")
        say("중단했습니다.")
    finally:
        attach_tool_error_logs(res)
        summary = "\n".join(res.summary) or "실행된 단계 없음"
        res.section("요약", summary)
        print("\n" + "=" * 70 + "\n[STATUS] 요약\n" + summary + "\n" + "=" * 70, flush=True)
        say(f"끝. 이 파일 하나를 Claude에게 전달하세요: {out}")


if __name__ == "__main__":
    main()
