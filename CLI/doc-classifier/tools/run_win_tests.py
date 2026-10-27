"""doc-classifier Windows 실기 테스트 러너 (목록화, 분류, 확인, 보고서)

사용법 (CLI/doc-classifier 폴더에서 실행):
    python tools/run_win_tests.py                          # 전체 (실제 드라이브는 D:\\)
    python tools/run_win_tests.py --drive E:\\ --drive-limit 1800
    python tools/run_win_tests.py --only S1,S3 --files 200000

오프라인 pytest가 다루지 못하는 실제 NTFS 동작을 확인합니다.
    S1 대량 트리 속도/정확도        S2 파일 수가 늘어도 메모리 일정    S3 강제 종료 후 재개(중복 없음)
    S4 정션 고리 건너뜀             S5 260자 초과 경로                 S6 다른 프로세스가 독점 중인 파일도 목록에 포함
    S7 실제 드라이브 전체(시스템 폴더 건너뜀, 권한 오류 계속)         S8 시스템 드라이브(C:\\) 시간 상한 후 재개
    S10 대량 트리 분류 속도/메모리   S11 형식별 확인(peek)과 잠긴 파일, 원본 불변   S12 OneDrive 자리표시자를 열지 않음
    S13 엑셀 보고서(대량, 열려 있는 보고서)   S14 list 조회(Windows 경로 조건, 페이지)   S15 실제 드라이브(S7 작업) 분류->확인->보고서
    S9 모든 호출의 stdout 1줄 JSON (마지막)
작업 DB는 테스트 폴더의 jobs\\ 에 만들고(DOCCLS_DATA_DIR), 결과는 report_*.txt로 저장합니다.
"""
import argparse
import datetime as dt
import json
import os
import shutil
import sqlite3
import subprocess
import sys
import time
import traceback

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "tests"))          # dc_samples: 실제 형식 샘플 파일(HWP 포함) 만들기
PY = sys.executable
CONTRACT = []


class Ctx:
    def __init__(self, args):
        self.args = args
        self.dir = os.path.abspath(args.dir)
        self.jobs = os.path.join(self.dir, "jobs")
        self.results = []

    def path(self, *parts):
        return os.path.join(self.dir, *parts)


def log(msg):
    print(msg, flush=True)


def _parse(proc_out, proc_err, code, label):
    out = proc_out.decode("utf-8", errors="replace")
    lines = [ln for ln in out.splitlines() if ln.strip()]
    if len(lines) != 1:
        CONTRACT.append((label, f"stdout {len(lines)}줄"))
    data = None
    if lines:
        try:
            data = json.loads(lines[-1])
        except Exception:
            CONTRACT.append((label, "stdout 마지막 줄이 JSON이 아님"))
    return data, code, proc_err.decode("utf-8", errors="replace")[-1500:]


def env_for(ctx):
    return dict(os.environ, PYTHONIOENCODING="utf-8", DOCCLS_DATA_DIR=ctx.jobs)


def dc(ctx, *args, timeout=7200):
    cmd = [PY, os.path.join(ROOT, "doc_classifier.py")] + [str(a) for a in args]
    log("    $ python doc_classifier.py " + " ".join(str(a) for a in args))
    p = subprocess.run(cmd, cwd=ROOT, capture_output=True, timeout=timeout, env=env_for(ctx))
    return _parse(p.stdout, p.stderr, p.returncode, "doc_classifier " + (args[0] if args else ""))


def record(ctx, tid, title, ok, detail="", data=None, code=None, err=""):
    status = "SKIP" if ok is None else ("PASS" if ok else "FAIL")
    entry = {"id": tid, "title": title, "status": status, "detail": detail, "exit_code": code,
             "json": "" if data is None else json.dumps(data, ensure_ascii=False)[:500]}
    if not ok and err:
        entry["stderr_tail"] = err.strip()[-600:]
    ctx.results.append(entry)
    log(f"  => [{status}] {tid} {title} - {detail}")


def run_test(ctx, tid, title, fn):
    log(f"\n[{tid}] {title}")
    try:
        fn()
    except Exception as e:
        ctx.results.append({"id": tid, "title": title, "status": "ERROR", "detail": f"{type(e).__name__}: {e}",
                            "trace": traceback.format_exc()[-1000:]})
        log(f"  => [ERROR] {tid} {e}")


def ok_(d):
    return (d or {}).get("status") == "ok"


def make_tree(base, n_files, per_dir=100):
    """n_files개 빈 파일을 폴더당 per_dir개, 3단 깊이로 만듭니다. (만든 파일 수)"""
    names = ["보고서", "회의록", "계약서", "견적", "data", "도면", "사양서"]
    exts = [".hwp", ".docx", ".xlsx", ".pdf", ".pptx", ".dwg", ".txt"]
    made = 0
    d = 0
    while made < n_files:
        folder = os.path.join(base, f"부서{d // 100:02d}", f"프로젝트{(d // 10) % 10}", f"폴더{d % 10}_{d}")
        os.makedirs(folder, exist_ok=True)
        for i in range(min(per_dir, n_files - made)):
            k = (made + i) % len(exts)
            with open(os.path.join(folder, f"{names[k]}_{made + i:07d}{exts[k]}"), "wb"):
                pass
        made += min(per_dir, n_files - made)
        d += 1
    return made


def _mem_of(d):
    return f", 최대 메모리 {(d or {}).get('peak_memory_mb')}MB"


def db_check(ctx, job):
    con = sqlite3.connect(os.path.join(ctx.jobs, f"{job}.db"))
    total, distinct = con.execute("SELECT COUNT(*), COUNT(DISTINCT dir_id || '/' || name) FROM files").fetchone()
    con.close()
    return total, distinct


def test_all(ctx, only):
    a = ctx.args
    state = {}

    def s1():
        base = ctx.path("트리_대량")
        t0 = time.time()
        n = make_tree(base, a.files)
        log(f"    (트리 생성 {n:,}개 {time.time() - t0:.0f}초)")
        d, c, e = dc(ctx, "scan", "--root", base, "--job", "s1")
        tot = (d or {}).get("totals", {})
        run = (d or {}).get("this_run", {})
        total, distinct = db_check(ctx, "s1") if ok_(d) else (None, None)
        state["s1_mem"] = d.get("peak_memory_mb") if d else None
        ok = ok_(d) and d.get("complete") and tot.get("files") == n and total == distinct and d.get("priority") == "background"
        record(ctx, "S1", f"대량 트리 {n:,}개 목록화 (개수, 중복 없음, 백그라운드 우선순위)", ok,
               f"파일 {tot.get('files')}, {run.get('seconds')}초, 초당 {run.get('files_per_sec')}개, 최대 메모리 {d.get('peak_memory_mb') if d else None}MB, "
               f"우선순위 {d.get('priority') if d else None}", d, c, e)
        state["base"] = base
        state["n"] = n

    def s2():
        small = ctx.path("트리_소량")
        make_tree(small, max(1000, a.files // 10))
        d, c, e = dc(ctx, "scan", "--root", small, "--job", "s2")
        m_small, m_big = (d or {}).get("peak_memory_mb"), state.get("s1_mem")
        ok = ok_(d) and m_small and m_big and m_big < 200 and m_big < m_small * 2.5
        record(ctx, "S2", "파일 수 10배여도 메모리 일정 (200MB 미만, 2.5배 미만)", ok,
               f"소량 {m_small}MB -> 대량 {m_big}MB", d, c, e)

    def s3():
        base, n = state.get("base"), state.get("n")
        if not base:
            base = ctx.path("트리_대량")
            n = make_tree(base, a.files)
        cmd = [PY, os.path.join(ROOT, "doc_classifier.py"), "scan", "--root", base, "--job", "s3", "--throttle-ms", "3"]
        p = subprocess.Popen(cmd, cwd=ROOT, stdout=subprocess.PIPE, stderr=subprocess.PIPE, env=env_for(ctx))
        time.sleep(3)
        p.kill()                                      # 작업 관리자 '작업 끝내기'와 같은 강제 종료 (자신이 띄운 프로세스)
        p.communicate()
        partial, _ = db_check(ctx, "s3")
        d, c, e = dc(ctx, "scan", "--job", "s3", "--resume")
        total, distinct = db_check(ctx, "s3")
        ok = ok_(d) and d.get("complete") and d.get("resumed") and total == n and total == distinct and 0 < partial < n
        record(ctx, "S3", "목록화 중 강제 종료 -> 재개해 중복 없이 완료", ok,
               f"종료 시점 기록 {partial:,}개 -> 재개 후 {total:,}개 (고유 {distinct:,}, 기대 {n:,})", d, c, e)

    def s4():
        base = ctx.path("정션")
        os.makedirs(os.path.join(base, "a", "b"), exist_ok=True)
        open(os.path.join(base, "a", "b", "1.txt"), "w").close()
        import _winapi
        _winapi.CreateJunction(base, os.path.join(base, "a", "b", "위로"))
        d, c, e = dc(ctx, "scan", "--root", base, "--job", "s4", timeout=120)
        ok = ok_(d) and d["totals"]["files"] == 1 and d["skipped"]["by_reason"] == {"link": 1}
        record(ctx, "S4", "정션 고리(조상을 가리킴) 건너뜀", ok,
               f"파일 {(d or {}).get('totals', {}).get('files')}, 건너뜀 {(d or {}).get('skipped', {}).get('by_reason')}", d, c, e)

    def s5():
        base = ctx.path("긴경로")
        deep = base
        while len(deep) < 320:
            deep = os.path.join(deep, "가나다라마바사아자차카타파하" * 2)
        os.makedirs("\\\\?\\" + deep, exist_ok=True)
        open("\\\\?\\" + os.path.join(deep, "긴 경로 보고서.hwp"), "w").close()
        d, c, e = dc(ctx, "scan", "--root", base, "--job", "s5", timeout=120)
        ok = ok_(d) and d["totals"]["files"] == 1 and d["totals"]["dirs_error"] == 0
        record(ctx, "S5", f"260자 초과 경로({len(deep)}자) 목록화", ok,
               f"파일 {(d or {}).get('totals', {}).get('files')}, 오류 폴더 {(d or {}).get('totals', {}).get('dirs_error')}", d, c, e)

    def s6():
        base = ctx.path("잠긴파일")
        os.makedirs(base, exist_ok=True)
        target = os.path.join(base, "사용중_계약서.docx")
        open(target, "wb").write(b"x" * 1234)
        import win32file
        h = win32file.CreateFile(target, win32file.GENERIC_READ, 0, None, win32file.OPEN_EXISTING, 0, None)   # 공유 없음(독점)
        try:
            d, c, e = dc(ctx, "scan", "--root", base, "--job", "s6", timeout=120)
        finally:
            h.Close()
        ok = ok_(d) and d["totals"]["files"] == 1 and d["totals"]["bytes"] == 1234 and d["this_run"]["entry_errors"] == 0
        record(ctx, "S6", "다른 프로세스가 독점 중인 파일도 목록에 포함 (파일을 열지 않음)", ok,
               f"파일 {(d or {}).get('totals', {}).get('files')}, 크기 {(d or {}).get('totals', {}).get('bytes')}", d, c, e)

    def s7():
        drive = a.drive
        d, c, e = dc(ctx, "scan", "--root", drive, "--job", "s7", "--time-limit", a.drive_limit)
        tot = (d or {}).get("totals", {})
        run = (d or {}).get("this_run", {})
        ok = ok_(d) and (d.get("peak_memory_mb") or 0) < 300
        record(ctx, "S7", f"실제 드라이브 {drive} 목록화 (시간 상한 {a.drive_limit}초)", ok,
               f"완료={d.get('complete') if d else None}, 파일 {tot.get('files')}, {tot.get('size')}, 폴더 {tot.get('dirs_done')}, "
               f"{run.get('seconds')}초, 초당 {run.get('files_per_sec')}개, 메모리 {d.get('peak_memory_mb') if d else None}MB, "
               f"건너뜀 {(d or {}).get('skipped', {}).get('by_reason')}, 오류 {(d or {}).get('errors', {}).get('by_reason')}, "
               f"열지 않을 파일 {tot.get('no_read_files')}", d, c, e)

    def s8():
        d, c, e = dc(ctx, "scan", "--root", "C:\\", "--job", "s8", "--time-limit", "20")
        first_files = (d or {}).get("totals", {}).get("files")
        stopped = (d or {}).get("stopped_reason")
        d2, c2, e2 = dc(ctx, "scan", "--job", "s8", "--resume", "--time-limit", "20")
        skipped = (d2 or {}).get("skipped", {}).get("examples", [])
        sys_skipped = (d2 or {}).get("skipped", {}).get("by_reason", {}).get("system", 0)
        ok = ok_(d) and stopped == "time_limit" and ok_(d2) and (d2 or {}).get("totals", {}).get("files", 0) > first_files \
            and sys_skipped >= 3 and (d2 or {}).get("resumed")
        record(ctx, "S8", "C:\\ 20초 상한 -> 재개로 이어짐, 시스템 폴더(Windows, Program Files 등) 건너뜀", ok,
               f"1차 {first_files}개({stopped}) -> 2차 {(d2 or {}).get('totals', {}).get('files')}개, 시스템 건너뜀 {sys_skipped}개, "
               f"예: {[x['path'] for x in skipped[:4]]}", d2, c2, e2)

    def ensure_tree_job():
        """S10/S13/S14가 쓸 대량 트리 작업 (S1을 건너뛰었으면 작게 새로 만듦)."""
        if state.get("tree_job"):
            return state["tree_job"], state["tree_n"]
        if state.get("base"):
            job, n = "s1", state["n"]
        else:
            base = ctx.path("트리_분류")
            n = make_tree(base, min(a.files, 20000))
            dc(ctx, "scan", "--root", base, "--job", "s10")
            job = "s10"
        state["tree_job"], state["tree_n"] = job, n
        return job, n

    def rules_file():
        path = ctx.path("rules.json")
        if not os.path.exists(path):
            rules = {"계약서": ["계약", "협약"], "회의록": ["회의록", "회의"], "보고서": ["보고서", "보고"],
                     "견적/발주": ["견적", "발주", "단가"], "시험성적서": ["성적서"], "사양서": ["사양서"],
                     "도면": [".dwg", ".dxf"]}
            with open(path, "w", encoding="utf-8") as f:
                json.dump(rules, f, ensure_ascii=False)
        return path

    def s10():
        job, n = ensure_tree_job()
        d, c, e = dc(ctx, "classify", "--job", job, "--rules", rules_file())
        run = (d or {}).get("this_run", {})
        cats = {x["category"]: x["files"] for x in (d or {}).get("by_category", [])}
        total = sum(cats.values())
        mem = _mem_of(d)
        ok = ok_(d) and total == n and cats.get("미분류", 0) < n and ((d or {}).get("peak_memory_mb") or 0) < 300
        record(ctx, "S10", f"대량 트리 {n:,}개 파일명 분류 (개수, 속도)", ok,
               f"분류 {total:,}개, {run.get('seconds')}초, 초당 {run.get('files_per_sec')}개, 카테고리 {cats}, "
               f"중복 후보 {(d or {}).get('duplicates', {}).get('groups')}묶음{mem}", d, c, e)

    def s11():
        import dc_samples as ds
        base = ctx.path("확인")
        os.makedirs(base, exist_ok=True)
        j = lambda *p: os.path.join(base, *p)  # noqa: E731
        ds.make_docx(j("공급계약서.docx"), "물품 공급 계약서", "갑과 을은 계약한다.")
        ds.make_hwp(j("협약서.hwp"), "업무 협약 본문", preview="업무 협약서")
        ds.make_hwpx(j("회의록_0301.hwpx"), preview="주간 회의록")
        ds.make_pptx(j("회의자료.pptx"), "월간 회의")
        ds.make_xlsx(j("견적_사용중.xlsx"), "견적", ["품목", "단가"], ["볼트", 100])
        ds.make_pdf(j("시험성적서.pdf"), "시험 성적서")
        ds.make_doc(j("사양서.doc"), "제품 사양서 Specification")
        ds.make_eml(j("회의 안내.eml"), "회의 안내", "다음 주 회의")
        ds.make_docx(j("계약_독점중.docx"), "계약 본문")
        with open(j("계약_DRM.docx"), "wb") as f:
            f.write(b"<##DRM##>" + b"\x00" * 300)
        before = {n: os.stat(j(n)).st_mtime for n in os.listdir(base)}
        d, c, e = dc(ctx, "scan", "--root", base, "--job", "s11", timeout=120)
        dc(ctx, "classify", "--job", "s11", "--rules", rules_file(), timeout=120)
        import win32file
        # Excel처럼 다른 프로세스가 쓰기 중(읽기/쓰기 공유 허용)인 파일과, 공유 없이 독점 중인 파일
        shared = win32file.CreateFile(j("견적_사용중.xlsx"), win32file.GENERIC_READ | win32file.GENERIC_WRITE,
                                      win32file.FILE_SHARE_READ | win32file.FILE_SHARE_WRITE, None, win32file.OPEN_EXISTING, 0, None)
        excl = win32file.CreateFile(j("계약_독점중.docx"), win32file.GENERIC_READ, 0, None, win32file.OPEN_EXISTING, 0, None)
        try:
            t0 = time.time()
            d, c, e = dc(ctx, "peek", "--job", "s11", "--timeout", "15", timeout=600)
            secs = time.time() - t0
        finally:
            shared.Close()
            excl.Close()
        con = sqlite3.connect(os.path.join(ctx.jobs, "s11.db"))
        st = {n: (s, r) for n, s, r in con.execute("SELECT name, peek_status, peek_reason FROM files")}
        con.close()
        after = {n: os.stat(j(n)).st_mtime for n in os.listdir(base)}
        expect_ok = ["공급계약서.docx", "협약서.hwp", "회의록_0301.hwpx", "회의자료.pptx", "견적_사용중.xlsx", "시험성적서.pdf",
                     "사양서.doc", "회의 안내.eml"]
        bad = [f"{n}={st.get(n)}" for n in expect_ok if (st.get(n) or (None,))[0] != "confirmed"]
        drm_ok = (st.get("계약_DRM.docx") or (None,))[0] == "encrypted_or_drm"
        excl_ok = (st.get("계약_독점중.docx") or (None,))[0] == "unreadable"
        renamed = True
        try:                                   # peek가 끝난 뒤 파일을 잡고 있지 않은지 (이름 바꾸기 가능)
            os.rename(j("공급계약서.docx"), j("공급계약서_2.docx"))
            os.rename(j("공급계약서_2.docx"), j("공급계약서.docx"))
        except OSError:
            renamed = False
        ok = ok_(d) and not bad and drm_ok and excl_ok and before == after and renamed
        record(ctx, "S11", "형식별 확인(peek): docx/hwp/hwpx/pptx/xlsx(사용 중)/pdf/doc/eml, DRM, 독점 중 파일, 원본 불변", ok,
               f"{secs:.1f}초, 판정 {d.get('totals', {}).get('by_status') if d else None}, 틀림 {bad}, DRM판별 {drm_ok}, "
               f"독점중 {st.get('계약_독점중.docx')}, 수정시각 불변 {before == after}, 잠금 없음 {renamed}", d, c, e)

    def s12():
        od = os.environ.get("OneDrive") or os.environ.get("OneDriveCommercial") or os.environ.get("OneDriveConsumer")
        if not od or not os.path.isdir(od):
            record(ctx, "S12", "OneDrive 자리표시자를 열지 않음", None, "OneDrive 폴더 없음")
            return
        d, c, e = dc(ctx, "scan", "--root", od, "--job", "s12", "--include-cloud", "--time-limit", "120")
        cloud = (d or {}).get("totals", {}).get("no_read_files", 0)
        if not cloud:
            record(ctx, "S12", "OneDrive 자리표시자를 열지 않음", None,
                   f"자리표시자(온라인 전용) 파일이 없음 (파일 {(d or {}).get('totals', {}).get('files')}개). "
                   "탐색기에서 몇 개를 '공간 확보'로 바꾼 뒤 다시 실행하세요", d, c, e)
            return
        dc(ctx, "classify", "--job", "s12", "--rules", rules_file())
        con = sqlite3.connect(os.path.join(ctx.jobs, "s12.db"))
        targets = [os.path.join(p, n) for p, n in con.execute(
            "SELECT d.path, f.name FROM files f JOIN dirs d ON d.id = f.dir_id WHERE f.no_read = 1 LIMIT 200")]
        con.close()
        d, c, e = dc(ctx, "peek", "--job", "s12", "--per-category", "200", "--max-files", "2000", "--time-limit", "300")
        con = sqlite3.connect(os.path.join(ctx.jobs, "s12.db"))
        opened = con.execute("SELECT COUNT(*) FROM files WHERE no_read = 1 AND peek_status IS NOT NULL").fetchone()[0]
        con.close()
        still = 0
        for t in targets:
            try:
                attrs = os.stat(t, follow_symlinks=False).st_file_attributes
                still += bool(attrs & (0x400000 | 0x40000 | 0x1000))
            except OSError:
                pass
        ok = ok_(d) and opened == 0 and still == len(targets)
        record(ctx, "S12", "OneDrive 자리표시자를 열지 않음 (peek 표본 제외, 내려받기 없음)", ok,
               f"자리표시자 {cloud}개, peek가 연 자리표시자 {opened}개, 확인 후에도 자리표시자 {still}/{len(targets)}", d, c, e)

    def s13():
        job, n = ensure_tree_job()
        report = ctx.path(f"보고서_{job}.xlsx")
        t0 = time.time()
        d, c, e = dc(ctx, "summary", "--job", job, "--save", report)
        secs = time.time() - t0
        saved = (d or {}).get("report", {}).get("report_file")
        rows = None
        csvs = (d or {}).get("report", {}).get("file_list_files") or []
        if csvs:                                      # 10만 행 초과: 목록은 CSV (104만 행씩)
            rows = 0
            for p in csvs:
                with open(p, encoding="utf-8-sig") as f:
                    rows += sum(1 for _ in f) - 1
        elif saved and os.path.exists(saved):
            import openpyxl
            wb = openpyxl.load_workbook(saved, read_only=True)
            rows = sum(wb[s].max_row - 1 for s in wb.sheetnames if s.startswith("파일 목록"))
            wb.close()
        import win32file
        h = win32file.CreateFile(report, win32file.GENERIC_READ, 0, None, win32file.OPEN_EXISTING, 0, None)   # 엑셀이 연 상태처럼 잠금
        try:
            d2, c2, e2 = dc(ctx, "summary", "--job", job, "--save", report, "--file-list", "none")
        finally:
            h.Close()
        alt = (d2 or {}).get("report", {}).get("report_file")
        ok = ok_(d) and rows == n and ok_(d2) and alt and alt != report and os.path.exists(alt)
        record(ctx, "S13", f"엑셀 보고서 ({n:,}행), 열려 있는 보고서는 다른 이름으로 저장", ok,
               f"{secs:.1f}초, 파일 목록 {rows}행({'CSV %d개' % len(csvs) if csvs else '시트'}), 크기 {os.path.getsize(saved) / 1048576 if saved else 0:.1f}MB, "
               f"잠긴 경우 저장 위치 {alt}{_mem_of(d)}", d2, c2, e2)

    def s14():
        job, n = ensure_tree_job()
        con = sqlite3.connect(os.path.join(ctx.jobs, f"{job}.db"))
        folder = con.execute("SELECT path FROM dirs WHERE depth = 1 ORDER BY id LIMIT 1").fetchone()[0]
        per_dir = con.execute("SELECT d.path, COUNT(*) FROM files f JOIN dirs d ON d.id = f.dir_id GROUP BY d.id").fetchall()
        con.close()
        cnt = sum(k for p, k in per_dir if p == folder)
        sub = sum(k for p, k in per_dir if p.startswith(folder + os.sep))
        d, c, e = dc(ctx, "list", "--job", job, "--folder", folder, "--page-size", "50")
        d2, c2, e2 = dc(ctx, "list", "--job", job, "--category", "계약서", "--ext", "hwp", "--since", "2000", "--page", "2")
        nxt = (d or {}).get("next_commands", [{}])[0].get("command", "")
        ok = ok_(d) and d["total"] == cnt + sub and len(d["items"]) == min(50, cnt + sub) and "--page 2" in nxt and ok_(d2)
        record(ctx, "S14", "list: Windows 경로 폴더 조건(하위 포함), 카테고리+확장자+기간, 페이지", ok,
               f"폴더 {folder}: {d.get('total') if d else None}개 (기대 {cnt + sub}), 계약서 hwp {d2.get('total') if d2 else None}개", d2, c2, e2)

    def s15():
        if not os.path.exists(os.path.join(ctx.jobs, "s7.db")):
            record(ctx, "S15", "실제 드라이브 분류->확인->보고서", None, "S7(실제 드라이브 목록화)을 먼저 실행해야 함")
            return
        d1, c1, e1 = dc(ctx, "classify", "--job", "s7", "--rules", rules_file())
        t0 = time.time()
        d2, c2, e2 = dc(ctx, "peek", "--job", "s7", "--per-category", "15", "--max-files", "150", "--time-limit", str(a.peek_limit))
        peek_secs = time.time() - t0
        d3, c3, e3 = dc(ctx, "summary", "--job", "s7", "--save", ctx.path("보고서_실제드라이브.xlsx"), "--file-list", "none")
        tot = (d2 or {}).get("totals", {})
        ok = ok_(d1) and ok_(d2) and ok_(d3)
        record(ctx, "S15", f"실제 드라이브 {a.drive} 분류 -> 표본 확인 -> 보고서", ok,
               f"분류 {(d1 or {}).get('this_run', {}).get('files')}개 {(d1 or {}).get('this_run', {}).get('seconds')}초; "
               f"확인 {(d2 or {}).get('this_run', {}).get('files')}개 {peek_secs:.0f}초, 판정 {tot.get('by_status')}, "
               f"확인율 {tot.get('confirm_rate')}, 시간초과 재시작 {(d2 or {}).get('this_run', {}).get('worker_restarts')}; "
               f"못 읽은 이유 {[r['reason'] for r in tot.get('not_read_reasons', [])[:3]]}", d2, c2, e2)

    def s9():
        bad = [f"{label}: {problem}" for label, problem in CONTRACT]
        record(ctx, "S9", "모든 호출의 stdout 1줄 JSON", not bad, "위반 없음" if not bad else "; ".join(bad[:10]))

    tests = [("S1", "대량 트리", s1), ("S2", "메모리 일정", s2), ("S3", "강제 종료 후 재개", s3), ("S4", "정션 고리", s4),
             ("S5", "긴 경로", s5), ("S6", "독점 중인 파일", s6), ("S7", "실제 드라이브", s7), ("S8", "C: 상한과 재개", s8),
             ("S10", "분류", s10), ("S11", "형식별 확인", s11), ("S12", "OneDrive 자리표시자", s12), ("S13", "엑셀 보고서", s13),
             ("S14", "list 조회", s14), ("S15", "실제 드라이브 전체 흐름", s15),
             ("S9", "출력 계약", s9)]
    for tid, title, fn in tests:
        if only and tid not in only and tid != "S9":
            continue
        run_test(ctx, tid, title, fn)


def write_report(ctx):
    stamp = dt.datetime.now().strftime("%Y%m%d_%H%M%S")
    path = ctx.path(f"report_{stamp}.txt")
    counts = {}
    for r in ctx.results:
        counts[r["status"]] = counts.get(r["status"], 0) + 1
    lines = [f"doc-classifier Windows 테스트 보고서 ({stamp})", f"Python {sys.version.split()[0]} / 테스트 폴더 {ctx.dir}",
             "요약: " + ", ".join(f"{k}={v}" for k, v in sorted(counts.items())), ""]
    for r in ctx.results:
        lines.append(f"[{r['status']}] {r['id']} {r['title']}")
        for key, label in (("detail", "내용"), ("exit_code", "종료코드"), ("json", "JSON"), ("stderr_tail", "stderr(끝부분)"),
                           ("trace", "trace")):
            if r.get(key) not in (None, ""):
                lines.append(f"    {label}: " + str(r[key]).replace("\n", "\n        "))
        lines.append("")
    with open(path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))
    return path, counts


def main():
    parser = argparse.ArgumentParser(description="doc-classifier Windows 실기 테스트 러너")
    parser.add_argument("--dir", default=r"C:\dc_test", help="테스트 폴더 (실행 때마다 새로 만듦)")
    parser.add_argument("--only", help="실행할 항목 (예: S1,S3)")
    parser.add_argument("--files", type=int, default=100000, help="S1 대량 트리 파일 수 (기본 10만)")
    parser.add_argument("--drive", default="D:\\", help="S7 실제 드라이브 (기본 D:\\)")
    parser.add_argument("--drive-limit", type=int, default=1800, help="S7 시간 상한(초)")
    parser.add_argument("--peek-limit", type=int, default=600, help="S15 표본 확인 시간 상한(초)")
    args = parser.parse_args()
    ctx = Ctx(args)
    shutil.rmtree(ctx.dir, ignore_errors=True)
    os.makedirs(ctx.jobs, exist_ok=True)
    only = {x.strip().upper() for x in args.only.split(",")} if args.only else None
    log("=" * 70)
    log(f" doc-classifier Windows 테스트 / 테스트 폴더: {ctx.dir}")
    log("=" * 70)
    test_all(ctx, only)
    path, counts = write_report(ctx)
    log("\n" + "=" * 70)
    log(" 결과 요약: " + ", ".join(f"{k}={v}" for k, v in sorted(counts.items())))
    for r in ctx.results:
        if r["status"] != "PASS":
            log(f"   [{r['status']}] {r['id']} {r['title']} {r.get('detail', '')}")
    log(f" 보고서: {path}")
    log("=" * 70)


if __name__ == "__main__":
    main()
