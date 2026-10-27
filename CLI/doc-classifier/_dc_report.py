"""summary / list / 엑셀 보고서.

- summary: 형식별·분류별 개수와 용량, 상위 폴더 분포, 수정 시기, 확인(peek) 결과, 중복 후보, 미분류 단어.
  JSON은 LLM 컨텍스트에 맞게 상위 N개로 자르고(truncated), 전체는 엑셀 보고서로.
- 엑셀: openpyxl 쓰기 전용 모드(메모리 일정). 파일 목록이 시트 한도(1,048,576행)를 넘으면 시트를 나누고,
  너무 크면(기본 300만 행 초과) 목록만 CSV로 따로 저장합니다.
- list: 조건(분류, 확장자, 형식, 기간, 폴더, 이름) 조회, 페이지 단위.
"""
import csv
import datetime as dt
import json
import os
import re
import time
from typing import Any, Dict, List, Optional, Sequence, Tuple

import _dc_classify as cl
import _dc_peek as pk
import _dc_rules as ru
import _dc_runtime as rt
import _dc_store as st

SHEET_ROWS = 1048575            # 머리글 1행을 뺀 시트당 최대 행
CSV_THRESHOLD = 3000000          # 이보다 많으면 파일 목록은 CSV로
AGE_BUCKETS = [("1년 이내", 1), ("1~3년", 3), ("3~5년", 5), ("5~10년", 10), ("10년 이상", None)]


# ==============================================================================
# 1. 집계
# ==============================================================================

def age_counts(con: Any, now: Optional[float] = None) -> List[Dict[str, Any]]:
    now = now or time.time()
    year = 365.25 * 86400
    parts: List[str] = []
    for label, yrs in AGE_BUCKETS:
        cond = f"mtime > {now - yrs * year}" if yrs else "1"
        parts.append(f"WHEN {cond} THEN {len(parts)}")
    sql = (f"SELECT CASE {' '.join(parts)} END AS b, COUNT(*), COALESCE(SUM(size), 0) FROM files "
           "WHERE mtime IS NOT NULL GROUP BY b ORDER BY b")
    rows = {b: (n, s) for b, n, s in con.execute(sql)}
    total = sum(n for n, _ in rows.values()) or 1
    return [{"age": label, "files": rows.get(i, (0, 0))[0], "size": rt.human_size(rows.get(i, (0, 0))[1]),
             "share": round(rows.get(i, (0, 0))[0] * 100 / total, 1)} for i, (label, _) in enumerate(AGE_BUCKETS)]


def top_folders(con: Any, limit: int = 10, category: Optional[str] = None) -> List[Dict[str, Any]]:
    """대상 폴더 바로 아래 폴더별 파일 수와 용량 (대상 폴더에 바로 있는 파일은 '(대상 폴더)')."""
    where, params = ("WHERE f.category = ?", [category]) if category else ("", [])
    sql = (f"SELECT d.top_id, COUNT(*), COALESCE(SUM(f.size), 0) FROM files f JOIN dirs d ON d.id = f.dir_id {where} "
           "GROUP BY d.top_id ORDER BY 2 DESC LIMIT ?")
    rows = con.execute(sql, (*params, limit)).fetchall()
    out = []
    for top_id, n, b in rows:
        r = con.execute("SELECT path, depth FROM dirs WHERE id=?", (top_id,)).fetchone() if top_id is not None else None
        folder = (r[0] if r[1] == 1 else "(대상 폴더)") if r else "(알 수 없음)"
        out.append({"folder": folder, "files": n, "size": rt.human_size(b), "bytes": b})
    return out


def dense_folders(con: Any, category: str, limit: int = 5) -> List[Dict[str, Any]]:
    """그 카테고리 파일이 가장 많이 모인 폴더 (바로 들어 있는 폴더 기준)."""
    rows = con.execute("SELECT d.path, COUNT(*), COALESCE(SUM(f.size), 0) FROM files f JOIN dirs d ON d.id = f.dir_id "
                       "WHERE f.category = ? GROUP BY f.dir_id ORDER BY 2 DESC LIMIT ?", (category, limit)).fetchall()
    return [{"folder": p, "files": n, "size": rt.human_size(b)} for p, n, b in rows]


def rules_from_meta(meta: Dict[str, Any]) -> ru.Rules:
    return ru.parse_rules(meta.get("rules_text") or "{}", (meta.get("rules") or {}).get("path", ""))


def peek_overview(con: Any, rules: ru.Rules) -> Dict[str, Any]:
    by_status = {s: n for s, n in con.execute("SELECT peek_status, COUNT(*) FROM files WHERE peek_status IS NOT NULL "
                                               "GROUP BY peek_status")}
    stats = pk.category_stats(con, rules)
    judged = sum(by_status.get(k, 0) for k in pk.JUDGED)
    reasons = [{"status": s, "reason": r, "files": n} for s, r, n in con.execute(
        "SELECT peek_status, peek_reason, COUNT(*) FROM files WHERE peek_status NOT IN ('confirmed', 'changed', 'no_match') "
        "GROUP BY peek_status, peek_reason ORDER BY 3 DESC LIMIT 10")]
    return {"files": sum(by_status.values()), "by_status": by_status,
            "confirm_rate": round(by_status.get("confirmed", 0) / judged, 2) if judged else None,
            "by_category": stats, "not_read_reasons": reasons}


def rate_text(rate: Optional[float]) -> str:
    return f"{round(rate * 100)}%" if rate is not None else "-(판정한 표본 없음)"


def changed_examples(con: Any, limit: int = 10) -> List[Dict[str, Any]]:
    rows = con.execute("SELECT d.path, f.name, f.category, f.peek_category, f.peek_excerpt FROM files f "
                       "JOIN dirs d ON d.id = f.dir_id WHERE f.peek_status='changed' ORDER BY f.category, f.id LIMIT ?",
                       (limit,)).fetchall()
    return [{"path": st.full_path(p, n), "category": c, "suggested": s, "excerpt": (e or "")[:80]} for p, n, c, s, e in rows]


def recommendations(summary: Dict[str, Any]) -> List[str]:
    out: List[str] = []
    cats = {c["category"]: c for c in summary["by_category"]}
    total = summary["totals"]["files"] or 1
    un = cats.get(ru.UNCLASSIFIED, {}).get("files", 0)
    if un * 100 / total >= 30:
        words = ", ".join(w["word"] for w in summary.get("unclassified_words", [])[:8])
        out.append(f"미분류가 {un * 100 / total:.0f}%입니다. 미분류 파일명에 자주 나오는 단어({words})를 규칙 키워드로 추가할지 검토하세요.")
    zero = [c["category"] for c in summary["by_category"] if c["files"] == 0]
    if zero:
        out.append(f"맞는 파일이 하나도 없는 카테고리: {', '.join(zero[:10])}. 키워드를 확인하세요.")
    low = [c["category"] for c in summary["by_category"]
           if c["category"] != ru.UNCLASSIFIED and c["files"] >= 20 and c.get("low", 0) * 100 / c["files"] >= 30]
    if low:
        out.append(f"신뢰도 low 비율이 높은 카테고리: {', '.join(low)}. 여러 카테고리에 겹치는 키워드가 있는지 확인하세요.")
    peek = summary.get("peek") or {}
    weak = [s for s in peek.get("by_category", []) if s["category"] != ru.UNCLASSIFIED and s.get("confirm_rate") is not None
            and s["judged"] >= 3 and s["confirm_rate"] < 0.5]
    for s in weak[:5]:
        sugg = s.get("suggested") or {}
        tail = f" 내용상으로는 '{max(sugg, key=sugg.get)}'가 많습니다." if sugg else ""
        out.append(f"'{s['category']}' 표본 확인율이 {s['confirm_rate'] * 100:.0f}%로 낮습니다. 규칙 보완 대상입니다.{tail}")
    found = [s for s in peek.get("by_category", []) if s["category"] == ru.UNCLASSIFIED and s.get("changed")]
    if found:
        sugg = found[0].get("suggested") or {}
        out.append(f"미분류 표본 {found[0]['changed']}개는 내용상 다른 카테고리로 보입니다({', '.join(f'{k} {v}' for k, v in sugg.items())}). "
                   "그 파일명의 단어를 규칙에 추가하면 분류가 늘어납니다.")
    dup = summary.get("duplicates") or {}
    if dup.get("wasted_bytes", 0) >= 1024 ** 3:
        out.append(f"이름과 크기가 같은 중복 후보가 {dup['groups']:,}묶음, 정리하면 약 {dup['wasted_size']}를 줄일 수 있습니다(내용은 확인 필요).")
    if summary.get("scan", {}).get("complete") is False:
        out.append("목록화가 끝나지 않았습니다. scan --resume으로 마친 뒤 classify를 다시 실행하세요.")
    if summary.get("classify", {}).get("not_classified_files"):
        out.append(f"분류 이후 새로 목록화된 파일 {summary['classify']['not_classified_files']:,}개가 있습니다. classify를 다시 실행하세요.")
    return out


def build_summary(con: Any, job: str, meta: Dict[str, Any], top: int = 10) -> Dict[str, Any]:
    rules = rules_from_meta(meta)
    counts = st.scan_counts(con)
    by_cat = cl.category_counts(con, rules.names)
    for c in by_cat:
        if c["files"] and c["category"] != ru.UNCLASSIFIED:
            c["top_folders"] = dense_folders(con, c["category"], 3)
    words = meta.get("unclassified_words") or []
    dups = dict(meta.get("duplicates") or {})
    dup_examples = dups.pop("examples", [])
    s: Dict[str, Any] = {
        "job": job, "db_file": st.db_path(job), "root": meta.get("root"),
        "scan": {"complete": meta.get("scan_status") == "complete", "status": meta.get("scan_status"),
                 "dirs_skipped": counts["dirs_skipped"], "dirs_error": counts["dirs_error"],
                 "skipped_by_reason": st.reason_counts(con, "skipped"), "errors_by_reason": st.reason_counts(con, "error")},
        "classify": {"rules_file": (meta.get("rules") or {}).get("path"), "classified_at": meta.get("classified_at"),
                     "not_classified_files": st.unclassified_count(con)},
        "totals": {"files": counts["files"], "bytes": counts["bytes"], "size": counts["size"],
                   "no_read_files": counts["no_read_files"]},
        "by_category": by_cat,
        "by_kind": cl.kind_counts(con),
        "top_folders": top_folders(con, top),
        "age": age_counts(con),
        "duplicates": dict(dups, examples=dup_examples[:3]),
        "unclassified_words": words[:15],
        "keyword_hits": meta.get("keyword_hits") or {},
    }
    if meta.get("peek_files"):
        s["peek"] = peek_overview(con, rules)
        s["peek"]["changed_examples"] = changed_examples(con, 5)
    s["recommendations"] = recommendations(s)
    s["truncated"] = {"top_folders": top, "unclassified_words": 15, "duplicates_examples": 3,
                      "note": "전체는 summary --save 엑셀 보고서에 있습니다."}
    return s


# ==============================================================================
# 2. 엑셀 보고서
# ==============================================================================

def _excel_tools():
    try:
        import openpyxl
        from openpyxl.cell import WriteOnlyCell
        from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
        from openpyxl.styles import Font, PatternFill
    except ImportError:
        raise rt.InputError("엑셀 보고서에는 openpyxl이 필요합니다 (python -m pip install openpyxl).")
    return openpyxl, WriteOnlyCell, ILLEGAL_CHARACTERS_RE, Font, PatternFill


class Book:
    def __init__(self) -> None:
        self.openpyxl, self.Cell, self.illegal, Font, Fill = _excel_tools()
        self.wb = self.openpyxl.Workbook(write_only=True)
        self.bold = Font(bold=True)
        self.head_fill = Fill("solid", fgColor="DDEBF7")

    def sheet(self, title: str, widths: Sequence[int]):
        ws = self.wb.create_sheet(title)
        for i, w in enumerate(widths):
            ws.column_dimensions[self.openpyxl.utils.get_column_letter(i + 1)].width = w
        ws.freeze_panes = "A2"
        return ws

    def value(self, ws: Any, v: Any) -> Any:
        if isinstance(v, str):
            v = self.illegal.sub("", v)
            if v[:1] in ("=", "+", "-", "@") and len(v) > 1:      # 파일명이 수식으로 해석되지 않게 문자열로 고정
                c = self.Cell(ws, value=v)
                c.data_type = "s"
                return c
        return v

    def head(self, ws: Any, cols: Sequence[str]) -> None:
        cells = []
        for c in cols:
            cell = self.Cell(ws, value=c)
            cell.font = self.bold
            cell.fill = self.head_fill
            cells.append(cell)
        ws.append(cells)

    def row(self, ws: Any, values: Sequence[Any]) -> None:
        ws.append([self.value(ws, v) for v in values])

    def title_row(self, ws: Any, text: str) -> None:
        cell = self.Cell(ws, value=text)
        cell.font = self.bold
        ws.append([cell])


def _ts(t: Optional[float]) -> Optional[dt.datetime]:
    if t is None:
        return None
    try:
        return dt.datetime.fromtimestamp(t).replace(microsecond=0)
    except (OverflowError, OSError, ValueError):
        return None


FILE_COLS = ["경로", "파일명", "확장자", "형식", "카테고리", "신뢰도", "근거", "크기(바이트)", "수정일", "만든 날짜",
             "이름 속 날짜", "버전", "표시", "확인 결과", "열지 않음"]
FILE_SQL = ("SELECT d.path, f.name, f.ext, f.kind, f.category, f.confidence, f.matches, f.size, f.mtime, f.ctime, "
            "f.name_date, f.name_version, f.name_flags, f.peek_status, f.no_read, f.flags "
            "FROM files f JOIN dirs d ON d.id = f.dir_id ORDER BY f.dir_id, f.name")


def _file_values(r: Tuple[Any, ...]) -> List[Any]:
    dpath, name, ext, kind, cat, conf, matches, size, mtime, ctime, ndate, nver, nflags, peek, no_read, flags = r
    try:
        why = ru.matches_text(json.loads(matches)) if matches else ""
    except ValueError:
        why = matches
    return [st.full_path(dpath, name), name, ext, kind, cat, conf, why, size, _ts(mtime), _ts(ctime), ndate, nver, nflags,
            pk.STATUS_TEXT.get(peek, peek) if peek else None, (flags or "열지 않음") if no_read else None]


def write_report(con: Any, job: str, meta: Dict[str, Any], summary: Dict[str, Any], path: str,
                 file_list: str = "auto") -> Dict[str, Any]:
    total = summary["totals"]["files"]
    mode = file_list
    if mode == "auto":
        mode = "xlsx" if total <= CSV_THRESHOLD else "csv"
    rules = rules_from_meta(meta)
    out: Dict[str, Any] = {"file_list": mode}
    t0 = time.time()

    def write(tmp: str) -> None:
        b = Book()
        _sheet_summary(b, summary)
        _sheet_categories(b, con, summary)
        _sheet_peek(b, con, rules)
        _sheet_words(b, summary, meta)
        _sheet_duplicates(b, meta)
        _sheet_skipped(b, con)
        if mode == "xlsx":
            out["file_list_sheets"] = _sheet_files(b, con, total)
        b.wb.save(tmp)

    rt.log("INFO", f"엑셀 보고서 작성 중 (파일 {total:,}개, 목록 {mode})")
    out["report_file"] = rt.save_with_lock_fallback(path, write)
    if mode == "csv":
        csv_path = os.path.splitext(out["report_file"])[0] + "_파일목록.csv"
        out["file_list_csv"] = rt.save_with_lock_fallback(csv_path, lambda tmp: _write_csv(con, tmp, total))
    out["seconds"] = round(time.time() - t0, 1)
    return out


def _sheet_summary(b: Book, s: Dict[str, Any]) -> None:
    ws = b.sheet("요약", [26, 14, 10, 14, 10, 10, 10, 10, 10, 12])
    b.head(ws, ["항목", "값"])
    tot = s["totals"]
    info = [("작업", s["job"]), ("대상 폴더", s["root"]), ("목록화 상태", s["scan"]["status"]),
            ("파일 수", tot["files"]), ("전체 용량", tot["size"]), ("열지 않을 파일(클라우드/링크)", tot["no_read_files"]),
            ("건너뛴 폴더", s["scan"]["dirs_skipped"]), ("읽지 못한 폴더", s["scan"]["dirs_error"]),
            ("분류 규칙 파일", s["classify"]["rules_file"]), ("분류 시각", s["classify"]["classified_at"]),
            ("확인(peek) 표본 수", (s.get("peek") or {}).get("files", 0)),
            ("중복 후보", f"{s['duplicates'].get('groups', 0):,}묶음, 약 {s['duplicates'].get('wasted_size', '0B')}"),
            ("보고서 작성", dt.datetime.now().strftime("%Y-%m-%d %H:%M"))]
    for k, v in info:
        b.row(ws, [k, v])
    ws.append([])
    b.title_row(ws, "카테고리별")
    b.head(ws, ["카테고리", "파일 수", "비율(%)", "용량", "high", "medium", "low", "표본", "맞음", "확인율(%)"])
    peek = {p["category"]: p for p in (s.get("peek") or {}).get("by_category", [])}
    for c in s["by_category"]:
        p = peek.get(c["category"], {})
        rate = p.get("confirm_rate")
        b.row(ws, [c["category"], c["files"], c["share"], c["size"], c["high"], c["medium"], c["low"],
                   p.get("sampled"), p.get("confirmed"), round(rate * 100) if rate is not None else None])
    ws.append([])
    b.title_row(ws, "형식별")
    b.head(ws, ["형식", "파일 수", "", "용량"])
    for k in s["by_kind"]:
        b.row(ws, [k["kind"], k["files"], None, k["size"]])
    ws.append([])
    b.title_row(ws, "수정 시기")
    b.head(ws, ["수정 시기", "파일 수", "비율(%)", "용량"])
    for a in s["age"]:
        b.row(ws, [a["age"], a["files"], a["share"], a["size"]])
    if s.get("recommendations"):
        ws.append([])
        b.title_row(ws, "권장 사항")
        for r in s["recommendations"]:
            b.row(ws, [r])


def _sheet_categories(b: Book, con: Any, s: Dict[str, Any]) -> None:
    ws = b.sheet("카테고리별", [18, 14, 70, 12, 12])
    b.head(ws, ["카테고리", "구분", "폴더", "파일 수", "용량"])
    for c in s["by_category"]:
        if not c["files"]:
            continue
        for f in top_folders(con, 30, c["category"]):
            b.row(ws, [c["category"], "상위 폴더", f["folder"], f["files"], f["size"]])
        for f in dense_folders(con, c["category"], 20):
            b.row(ws, [c["category"], "많이 모인 폴더", f["folder"], f["files"], f["size"]])


def _sheet_peek(b: Book, con: Any, rules: ru.Rules) -> None:
    ws = b.sheet("확인 결과", [16, 14, 16, 70, 60, 40, 10])
    b.head(ws, ["카테고리", "판정", "제안 카테고리", "파일 경로", "발췌(앞 200자)", "이유", "읽은 시간(ms)"])
    order = {n: i for i, n in enumerate(rules.names + [ru.UNCLASSIFIED])}
    rows = con.execute("SELECT f.category, f.peek_status, f.peek_category, d.path, f.name, f.peek_excerpt, f.peek_reason, "
                       "f.peek_ms FROM files f JOIN dirs d ON d.id = f.dir_id WHERE f.peek_status IS NOT NULL").fetchall()
    rows.sort(key=lambda r: (order.get(r[0], 999), r[1] != "changed", r[1], r[3], r[4]))
    for cat, status, sugg, dpath, name, exc, reason, ms in rows:
        b.row(ws, [cat, pk.STATUS_TEXT.get(status, status), sugg, st.full_path(dpath, name), exc, reason, ms])


def _sheet_words(b: Book, s: Dict[str, Any], meta: Dict[str, Any]) -> None:
    ws = b.sheet("미분류 단어", [24, 12, 4, 18, 24, 12])
    b.head(ws, ["미분류 파일명의 단어", "파일 수", "", "카테고리", "맞은 키워드", "파일 수"])
    words = meta.get("unclassified_words") or []
    hits = [(cat, kw, n) for cat, kws in (meta.get("keyword_hits") or {}).items() for kw, n in kws.items()]
    for i in range(max(len(words), len(hits))):
        w = words[i] if i < len(words) else {}
        h = hits[i] if i < len(hits) else (None, None, None)
        b.row(ws, [w.get("word"), w.get("files"), None, h[0], h[1], h[2]])


def _sheet_duplicates(b: Book, meta: Dict[str, Any]) -> None:
    ws = b.sheet("중복 후보", [36, 14, 8, 14, 90])
    b.head(ws, ["파일명", "크기(바이트)", "개수", "줄일 수 있는 용량", "경로 (최대 5개)"])
    for e in (meta.get("duplicates") or {}).get("examples", []):
        b.row(ws, [e["name"], e["size"], e["copies"], rt.human_size(e["wasted"]), "\n".join(e["paths"])])


def _sheet_skipped(b: Book, con: Any) -> None:
    ws = b.sheet("오류·건너뜀", [10, 50, 90])
    b.head(ws, ["상태", "이유", "폴더"])
    for status, reason, path in con.execute("SELECT status, reason, path FROM dirs WHERE status IN ('skipped', 'error') "
                                            "ORDER BY status, id LIMIT 100000"):
        b.row(ws, ["건너뜀" if status == "skipped" else "오류", reason, path])


def _sheet_files(b: Book, con: Any, total: int) -> int:
    widths = [70, 30, 8, 10, 14, 8, 40, 14, 18, 18, 12, 8, 10, 14, 10]
    sheets = 0
    ws = None
    n = 0
    last = time.time()
    for r in con.execute(FILE_SQL):
        if ws is None or n % SHEET_ROWS == 0:
            sheets += 1
            ws = b.sheet("파일 목록" if sheets == 1 else f"파일 목록 {sheets}", widths)
            b.head(ws, FILE_COLS)
        b.row(ws, _file_values(r))
        n += 1
        if time.time() - last >= 2:
            last = time.time()
            rt.log("PROGRESS", f"엑셀 파일 목록 {n:,}/{total:,}행")
    if ws is None:
        ws = b.sheet("파일 목록", widths)
        b.head(ws, FILE_COLS)
        sheets = 1
    return sheets


def _write_csv(con: Any, path: str, total: int) -> None:
    last = time.time()
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(FILE_COLS)
        for i, r in enumerate(con.execute(FILE_SQL), 1):
            vals = _file_values(r)
            w.writerow(["" if v is None else (v.strftime("%Y-%m-%d %H:%M:%S") if isinstance(v, dt.datetime) else v) for v in vals])
            if time.time() - last >= 2:
                last = time.time()
                rt.log("PROGRESS", f"CSV 파일 목록 {i:,}/{total:,}행")


# ==============================================================================
# 3. list
# ==============================================================================

_SIZE_RE = re.compile(r"^\s*(\d+(?:\.\d+)?)\s*(b|kb|mb|gb|tb)?\s*$", re.I)


def parse_size(text: str) -> int:
    m = _SIZE_RE.match(str(text))
    if not m:
        raise rt.InputError(f"크기 '{text}'를 알 수 없습니다. 예: 500KB, 100MB, 2GB")
    mult = {"b": 1, "kb": 1024, "mb": 1024 ** 2, "gb": 1024 ** 3, "tb": 1024 ** 4}[(m.group(2) or "b").lower()]
    return int(float(m.group(1)) * mult)


def parse_date(text: str, end: bool = False) -> float:
    """'2025', '2025-03', '2025-03-15'(구분자 . / - 가능) -> 시각. end=True면 그 기간의 끝(다음 날/달/해 시작)."""
    m = re.fullmatch(r"\s*(\d{4})(?:[-./](\d{1,2})(?:[-./](\d{1,2}))?)?\s*", str(text))
    if not m:
        raise rt.InputError(f"날짜 '{text}'를 알 수 없습니다. 예: 2025, 2025-03, 2025-03-15")
    y, mo, d = int(m.group(1)), m.group(2), m.group(3)
    try:
        if d:
            t = dt.datetime(y, int(mo), int(d))
            t = t + dt.timedelta(days=1) if end else t
        elif mo:
            t = dt.datetime(y, int(mo), 1)
            if end:
                t = dt.datetime(y + (int(mo) == 12), int(mo) % 12 + 1, 1)
        else:
            t = dt.datetime(y + (1 if end else 0), 1, 1)
    except ValueError:
        raise rt.InputError(f"날짜 '{text}'가 올바르지 않습니다.")
    return t.timestamp()


def _like(s: str) -> str:
    return s.replace("\\", "\\\\").replace("%", "\\%").replace("_", "\\_")


def list_query(con: Any, args: Any, known_categories: List[str]) -> Tuple[str, List[Any], Dict[str, Any]]:
    where: List[str] = []
    params: List[Any] = []
    shown: Dict[str, Any] = {}
    if args.category:
        cats = [c.strip() for c in args.category.split(",") if c.strip()]
        for c in cats:
            if c not in known_categories:
                raise rt.not_found("카테고리", c, known_categories)
        where.append(f"f.category IN ({','.join('?' * len(cats))})")
        params += cats
        shown["category"] = cats
    if args.ext:
        exts = ["." + e.strip().lower().lstrip(".") for e in args.ext.split(",") if e.strip()]
        where.append(f"f.ext IN ({','.join('?' * len(exts))})")
        params += exts
        shown["ext"] = exts
    if args.kind:
        kinds = [k.strip() for k in args.kind.split(",") if k.strip()]
        for k in kinds:
            if k not in ru.KIND_ORDER:
                raise rt.not_found("형식", k, ru.KIND_ORDER)
        where.append(f"f.kind IN ({','.join('?' * len(kinds))})")
        params += kinds
        shown["kind"] = kinds
    if args.confidence:
        where.append("f.confidence = ?")
        params.append(args.confidence)
        shown["confidence"] = args.confidence
    if args.peek:
        if args.peek == "any":
            where.append("f.peek_status IS NOT NULL")
        else:
            where.append("f.peek_status = ?")
            params.append(args.peek)
        shown["peek"] = args.peek
    if args.since:
        where.append("f.mtime >= ?")
        params.append(parse_date(args.since))
        shown["since"] = args.since
    if args.until:
        where.append("f.mtime < ?")
        params.append(parse_date(args.until, end=True))
        shown["until"] = args.until
    if args.folder:
        folder = args.folder.strip().rstrip("\\/")
        if re.match(r"^([A-Za-z]:|\\\\|/)", folder):           # 절대 경로: 그 폴더와 하위 폴더
            where.append("(d.path = ? OR d.path LIKE ? ESCAPE '\\' OR d.path LIKE ? ESCAPE '\\')")
            params += [folder, _like(folder) + "\\\\%", _like(folder) + "/%"]
        else:                                                    # 폴더 이름 일부
            where.append("d.path LIKE ? ESCAPE '\\'")
            params.append("%" + _like(folder) + "%")
        shown["folder"] = args.folder
    if args.name:
        where.append("f.name LIKE ? ESCAPE '\\'")
        params.append("%" + _like(args.name) + "%")
        shown["name"] = args.name
    if args.min_size:
        where.append("f.size >= ?")
        params.append(parse_size(args.min_size))
        shown["min_size"] = args.min_size
    if args.flag:
        where.append("(',' || COALESCE(f.name_flags, '') || ',') LIKE ?")
        params.append(f"%,{args.flag},%")
        shown["flag"] = args.flag
    return ("WHERE " + " AND ".join(where)) if where else "", params, shown


SORTS = {"path": "d.path, f.name", "size": "f.size", "mtime": "f.mtime", "name": "f.name"}


def list_files(con: Any, args: Any, known_categories: List[str]) -> Dict[str, Any]:
    where, params, shown = list_query(con, args, known_categories)
    total, size = con.execute(f"SELECT COUNT(*), COALESCE(SUM(f.size), 0) FROM files f JOIN dirs d ON d.id = f.dir_id {where}",
                              params).fetchone()
    page_size = args.page_size
    pages = max(1, -(-total // page_size))
    page = args.page
    if page > pages and total:
        raise rt.InputError(f"페이지 {page}은(는) 없습니다. 전체 {pages}쪽입니다. [추천] --page {pages}")
    order = SORTS[args.sort] + (" DESC" if args.desc else "")
    rows = con.execute(f"SELECT d.path, f.name, f.size, f.mtime, f.kind, f.category, f.confidence, f.peek_status, f.no_read "
                       f"FROM files f JOIN dirs d ON d.id = f.dir_id {where} ORDER BY {order}, f.id LIMIT ? OFFSET ?",
                       (*params, page_size, (page - 1) * page_size)).fetchall()
    items = []
    for dpath, name, sz, mtime, kind, cat, conf, peek, no_read in rows:
        t = _ts(mtime)
        item = {"path": st.full_path(dpath, name), "size": rt.human_size(sz), "mtime": t.strftime("%Y-%m-%d") if t else None,
                "kind": kind, "category": cat, "confidence": conf}
        if peek:
            item["peek"] = peek
        if no_read:
            item["no_read"] = True
        items.append(item)
    return {"total": total, "total_size": rt.human_size(size), "page": page, "pages": pages, "page_size": page_size,
            "filters": shown, "sort": args.sort + (" desc" if args.desc else ""), "items": items}
