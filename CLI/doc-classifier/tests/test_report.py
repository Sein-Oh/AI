"""summary / list / 엑셀 보고서 테스트."""
import json
import os
import time

import pytest

import _dc_report as rp
import _dc_runtime as rt
import dc_samples as ds
from dc_tree import make_tree

RULES = {"계약서": ["계약"], "회의록": ["회의록"], "보고서": ["보고"], "도면": [".dwg"]}


@pytest.fixture
def job(tmp_path, run):
    root = tmp_path / "자료"
    spec = {f"계약/{y}/계약서_{y}_{i}.hwp": 100 * (i + 1) for y in (2019, 2025) for i in range(3)}
    spec.update({"보고/결과보고.pdf": 40, "보고/=수식처럼.txt": 3, "보고/회의록_0301.docx": 30, "도면/A.dwg": 9,
                 "기타/메모.txt": 5, "기타/메모 (1).txt": 5, "기타/사진.jpg": 7})
    make_tree(str(root), spec)
    os.utime(root / "계약" / "2019" / "계약서_2019_0.hwp", (time.time(), time.time() - 6 * 365 * 86400))
    ds.make_docx(str(root / "보고" / "보고_내용.docx"), "회의록 정리")
    run("scan", "--root", str(root), "--job", "r")
    p = tmp_path / "rules.json"
    p.write_text(json.dumps(RULES, ensure_ascii=False), encoding="utf-8")
    code, out = run("classify", "--job", "r", "--rules", str(p))
    assert code == 0
    return root


def test_summary_json_is_compact_and_has_counts(job, run):
    code, out = run("summary", "--job", "r")
    assert code == 0 and out["agent_directive"].startswith("VERIFIED:")
    by = {c["category"]: c for c in out["by_category"]}
    assert by["계약서"]["files"] == 6 and by["계약서"]["top_folders"][0]["files"] == 3
    assert out["totals"]["files"] == 14 and sum(c["files"] for c in out["by_category"]) == 14
    assert {k["kind"] for k in out["by_kind"]} >= {"한글", "PDF", "도면", "이미지", "문서"}
    assert out["top_folders"][0]["folder"].endswith("계약") and out["top_folders"][0]["files"] == 6
    ages = {a["age"]: a["files"] for a in out["age"]}
    assert ages["5~10년"] == 1 and sum(ages.values()) == 14
    assert "peek" not in out and any("peek" in c["command"] for c in out["next_commands"])
    assert any(c["needs_input"] and "<카테고리>" in c["command"] for c in out["next_commands"])
    assert "--save" in out["next_commands"][0]["command"]
    assert out["truncated"]["note"] and "report" not in out
    assert len(json.dumps(out, ensure_ascii=False)) < 20000                      # LLM 컨텍스트에 맞는 크기


def test_summary_save_writes_all_sheets(job, run, data_dir):
    run("peek", "--job", "r")
    code, out = run("summary", "--job", "r", "--save")
    assert code == 0 and "보고서:" in out["preview"]
    path = out["report"]["report_file"]
    assert path == str(data_dir / "r_보고서.xlsx") and os.path.exists(path) and out["report"]["file_list"] == "xlsx"
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True)
    assert wb.sheetnames == ["요약", "카테고리별", "확인 결과", "미분류 단어", "중복 후보", "오류·건너뜀", "파일 목록"]
    rows = list(wb["파일 목록"].iter_rows(values_only=True))
    assert rows[0][0] == "경로" and len([r for r in rows[1:] if r[0]]) == 14
    names = {r[1]: r for r in rows[1:] if r[0]}
    assert names["=수식처럼.txt"][1] == "=수식처럼.txt"                         # 수식으로 바뀌지 않음
    assert names["계약서_2025_0.hwp"][4] == "계약서" and "계약@파일명" in names["계약서_2025_0.hwp"][6]
    assert names["메모 (1).txt"][12] == "copy"
    peek_rows = [r for r in wb["확인 결과"].iter_rows(values_only=True)][1:]
    assert any(r[1] == "다른 분류 제안" and r[2] == "회의록" for r in peek_rows if r[0])
    summary = [tuple(r) + (None,) for r in wb["요약"].iter_rows(values_only=True) if r and r[0]]
    assert ("파일 수", 14) in [(r[0], r[1]) for r in summary]
    assert any(r[0] == "계약서" and r[1] == 6 for r in summary)
    wb.close()


def test_save_to_explicit_path_and_bad_extension(job, run, tmp_path):
    target = tmp_path / "보고서" / "결과.xlsx"
    code, out = run("summary", "--job", "r", "--save", str(target))
    assert code == 0 and out["report"]["report_file"] == str(target) and target.exists()
    code, out = run("summary", "--job", "r", "--save", str(tmp_path / "결과.csv"))
    assert code == 1 and "[추천]" in out["message"] and out["message"].endswith(".xlsx")


def test_file_list_splits_sheets_and_switches_to_csv(job, run, tmp_path, monkeypatch):
    monkeypatch.setattr(rp, "SHEET_ROWS", 5)
    code, out = run("summary", "--job", "r", "--save", str(tmp_path / "a.xlsx"))
    assert out["report"]["file_list_sheets"] == 3                                 # 14행 -> 5, 5, 4
    import openpyxl
    wb = openpyxl.load_workbook(str(tmp_path / "a.xlsx"), read_only=True)
    assert wb.sheetnames[-3:] == ["파일 목록", "파일 목록 2", "파일 목록 3"]
    wb.close()
    monkeypatch.setattr(rp, "CSV_THRESHOLD", 10)
    code, out = run("summary", "--job", "r", "--save", str(tmp_path / "b.xlsx"))
    csv_path = out["report"]["file_list_csv"]
    assert out["report"]["file_list"] == "csv" and csv_path.endswith("b_파일목록.csv")
    with open(csv_path, encoding="utf-8-sig") as f:
        lines = f.read().splitlines()
    assert lines[0].startswith("경로,파일명") and len(lines) == 15
    assert out["report"]["file_list_files"] == [csv_path]
    import openpyxl
    wb = openpyxl.load_workbook(str(tmp_path / "b.xlsx"), read_only=True)
    assert "파일 목록" not in wb.sheetnames                                    # 엑셀은 요약 시트만
    note = [r for r in wb["요약"].iter_rows(values_only=True) if r and r[0] == "파일 목록"][0][1]
    assert "b_파일목록.csv" in note
    wb.close()
    code, out = run("summary", "--job", "r", "--save", str(tmp_path / "c.xlsx"), "--file-list", "none")
    assert out["report"]["file_list"] == "none" and "file_list_csv" not in out["report"]


def test_csv_file_list_is_split_for_excel_and_stale_parts_removed(job, run, tmp_path, monkeypatch):
    monkeypatch.setattr(rp, "CSV_THRESHOLD", 10)
    monkeypatch.setattr(rp, "CSV_ROWS", 5)
    code, out = run("summary", "--job", "r", "--save", str(tmp_path / "d.xlsx"))
    files = out["report"]["file_list_files"]
    assert [os.path.basename(f) for f in files] == ["d_파일목록_1.csv", "d_파일목록_2.csv", "d_파일목록_3.csv"]
    counts = []
    paths = set()
    for f in files:
        with open(f, encoding="utf-8-sig") as fh:
            lines = fh.read().splitlines()
        assert lines[0].startswith("경로,파일명")                              # 파일마다 머리글
        counts.append(len(lines) - 1)
        paths.update(ln.split(",")[0] for ln in lines[1:])
    assert counts == [5, 5, 4] and len(paths) == 14                           # 빠짐·중복 없이 이어짐
    assert out["report"]["file_list_csv"] == files[0]
    monkeypatch.setattr(rp, "CSV_ROWS", 10)                                    # --overwrite로 다시 저장해 2개로 줄면 3번째는 지움
    code, out = run("summary", "--job", "r", "--save", str(tmp_path / "d.xlsx"), "--overwrite")
    assert len(out["report"]["file_list_files"]) == 2 and not (tmp_path / "d_파일목록_3.csv").exists()
    assert out["applied"] == ["덮어쓰기: d.xlsx", "덮어쓰기: d_파일목록_1.csv", "덮어쓰기: d_파일목록_2.csv",
                              "삭제: d_파일목록_3.csv (같은 이름의 이전 목록 CSV)"]
    monkeypatch.setattr(rp, "CSV_ROWS", 100)                                   # 1개면 번호 없는 이름
    code, out = run("summary", "--job", "r", "--save", str(tmp_path / "d.xlsx"), "--overwrite")
    assert sorted(p for p in os.listdir(tmp_path) if p.startswith("d_파일목록")) == ["d_파일목록.csv"]


def _snapshot(folder):
    """보고서와 목록 CSV 파일의 (크기, 수정 시각)."""
    return {p: (os.path.getsize(os.path.join(folder, p)), os.path.getmtime(os.path.join(folder, p)))
            for p in os.listdir(folder) if p.endswith((".xlsx", ".csv"))}


def test_existing_report_is_not_overwritten_without_flag(job, run, tmp_path, monkeypatch, capsys):
    """AGENTS.md 복구 금지(덮어쓰기, 삭제): 기본은 시각을 붙인 새 이름, 이전 파일은 그대로 (CLI/README.md R4)."""
    monkeypatch.setattr(rp, "CSV_THRESHOLD", 10)
    monkeypatch.setattr(rp, "CSV_ROWS", 5)
    run("summary", "--job", "r", "--save", str(tmp_path / "e.xlsx"))           # e.xlsx + CSV 3개
    before = _snapshot(tmp_path)
    monkeypatch.setattr(rp, "CSV_ROWS", 10)
    code, out = run("summary", "--job", "r", "--save", str(tmp_path / "e.xlsx"))
    saved = out["report"]["report_file"]
    assert code == 0 and os.path.basename(saved).startswith("e_") and saved.endswith(".xlsx")
    assert all(os.path.basename(f).startswith(os.path.basename(saved)[:-5] + "_파일목록_") for f in out["report"]["file_list_files"])
    assert out["applied"] == [] and "이미 있음" in out["recovered"][0] and "--overwrite" in out["recovered"][0]
    after = _snapshot(tmp_path)
    assert {p: after[p] for p in before} == before                            # 이전 보고서와 CSV 3개는 그대로
    code, out = run("summary", "--job", "r", "--save", saved)                  # 시각 붙은 이름을 다시 저장해도 시각을 두 번 붙이지 않음
    import re
    assert re.fullmatch(r"e_\d{8}_\d{6}(_\d+)?\.xlsx", os.path.basename(out["report"]["report_file"]))
    assert out["report"]["report_file"] != saved


def test_summary_dry_run_changes_nothing_and_lists_plan(job, run, tmp_path, monkeypatch):
    monkeypatch.setattr(rp, "CSV_THRESHOLD", 10)
    monkeypatch.setattr(rp, "CSV_ROWS", 5)
    target = str(tmp_path / "f.xlsx")
    code, out = run("summary", "--job", "r", "--save", target, "--dry-run")
    assert code == 0 and out["dry_run"] is True and _snapshot(tmp_path) == {}       # 아무 파일도 만들지 않음
    assert out["planned"] == ["새 파일: f.xlsx", "새 파일: f_파일목록_1.csv", "새 파일: f_파일목록_2.csv", "새 파일: f_파일목록_3.csv"]
    assert out["agent_directive"].startswith("VERIFIED: 미리보기만") and "--dry-run" not in out["next_commands"][0]["command"]
    run("summary", "--job", "r", "--save", target)
    monkeypatch.setattr(rp, "CSV_ROWS", 10)
    before = _snapshot(tmp_path)
    code, out = run("summary", "--job", "r", "--save", target, "--overwrite", "--dry-run")
    assert out["planned"] == ["덮어쓰기: f.xlsx", "덮어쓰기: f_파일목록_1.csv", "덮어쓰기: f_파일목록_2.csv",
                              "삭제: f_파일목록_3.csv (같은 이름의 이전 목록 CSV)"]
    assert _snapshot(tmp_path) == before
    assert out["next_commands"][0]["command"].endswith("--overwrite")
    code, out = run("summary", "--job", "r", "--save", target, "--dry-run")    # --overwrite 없으면 새 이름 계획
    assert out["planned"][0].startswith("새 파일: f_") and not any(p.startswith(("덮어쓰기", "삭제")) for p in out["planned"])
    code, out = run("summary", "--job", "r", "--dry-run")
    assert code == 1 and "--save" in out["message"] and "[추천]" in out["message"]


def test_default_threshold_moves_large_lists_to_csv():
    assert rp.CSV_THRESHOLD == 100000 and rp.CSV_ROWS == 1048575
    assert len(rp.csv_file_paths("x", 2100000)) == 3 and rp.csv_file_paths("x", 5) == ["x_파일목록.csv"]


def test_locked_report_is_saved_under_new_name(tmp_path, monkeypatch):
    target = tmp_path / "잠김.xlsx"
    target.write_bytes(b"old")
    real = os.replace

    def locked(src, dst):
        if str(dst) == str(target):
            raise PermissionError(13, "다른 프로세스가 사용 중")
        return real(src, dst)
    monkeypatch.setattr(rt.os, "replace", locked)
    rt.reset_run_state()
    saved = rt.save_with_lock_fallback(str(target), lambda p: open(p, "wb").write(b"new"))
    assert rt._RECOVERED == [f"잠김.xlsx 잠김(열려 있음) -> {os.path.basename(saved)}로 저장"]          # R3 recovered
    rt.reset_run_state()
    assert saved != str(target) and os.path.basename(saved).startswith("잠김_") and open(saved, "rb").read() == b"new"
    assert target.read_bytes() == b"old" and not [p for p in os.listdir(tmp_path) if ".saving" in p]


def test_list_filters_and_pages(job, run):
    code, out = run("list", "--job", "r", "--category", "계약서", "--page-size", "4")
    assert code == 0 and out["total"] == 6 and out["pages"] == 2 and len(out["items"]) == 4
    nxt = out["next_commands"][0]["command"]
    assert nxt.endswith("--page 2") and "--category 계약서" in nxt
    reader = out["next_commands"][1]
    assert reader["needs_input"] and os.path.join("office-cli", "reader.py") in reader["command"] and " read --file " in reader["command"]
    code, out = run("list", "--job", "r", "--category", "계약서", "--page-size", "4", "--page", "2")
    assert len(out["items"]) == 2 and not any("다음 페이지" in c["purpose"] for c in out["next_commands"])
    assert run("list", "--job", "r", "--ext", "hwp,.DWG")[1]["total"] == 7
    assert run("list", "--job", "r", "--kind", "도면")[1]["total"] == 1
    assert run("list", "--job", "r", "--folder", "2019")[1]["total"] == 3
    root = str(job)
    assert run("list", "--job", "r", "--folder", os.path.join(root, "계약"))[1]["total"] == 6
    assert run("list", "--job", "r", "--folder", os.path.join(root, "계"))[1]["total"] == 0     # 앞부분만 같은 다른 폴더 아님
    assert run("list", "--job", "r", "--name", "메모")[1]["total"] == 2
    assert run("list", "--job", "r", "--flag", "copy")[1]["total"] == 1
    assert run("list", "--job", "r", "--min-size", "0.25KB")[1]["total"] == 3            # 300, 300, docx
    assert run("list", "--job", "r", "--until", "2020")[1]["total"] == 1
    assert run("list", "--job", "r", "--since", "2021-01")[1]["total"] == 13
    code, out = run("list", "--job", "r", "--category", "미분류", "--sort", "size", "--desc")
    assert out["items"][0]["path"].endswith("사진.jpg")
    assert run("list", "--job", "r", "--name", "50%_")[1]["total"] == 0                   # LIKE 특수문자는 글자 그대로


def test_list_errors_suggest(job, run):
    code, out = run("list", "--job", "r", "--category", "계약")
    assert code == 1 and "[추천] 계약서" in out["message"]
    code, out = run("list", "--job", "r", "--kind", "문서류")
    assert code == 1 and "[추천] 문서" in out["message"]
    code, out = run("list", "--job", "r", "--since", "작년")
    assert code == 1 and "2025-03" in out["message"]
    code, out = run("list", "--job", "r", "--page", "9")
    assert code == 1 and "[추천] --page 1" in out["message"]
    assert run("list", "--job", "r", "--page-size", "0")[0] == 1
    assert run("list", "--job", "r", "--min-size", "큼")[0] == 1


def test_parse_date_ranges():
    import datetime as dt
    assert rp.parse_date("2025-03") == dt.datetime(2025, 3, 1).timestamp()
    assert rp.parse_date("2025-12", end=True) == dt.datetime(2026, 1, 1).timestamp()
    assert rp.parse_date("2025.3.15", end=True) == dt.datetime(2025, 3, 16).timestamp()
    assert rp.parse_size("1.5MB") == int(1.5 * 1024 * 1024) and rp.parse_size("10") == 10
